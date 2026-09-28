"""C07/G18 two-set regulatory reporting engine.

Built-in RP01/RP02 schemas are procurement baselines. They are deliberately NOT marked as
an officially verified target-platform template. ACCEPTED/REJECTED cannot be written until
a real receipt adapter with source/signature verification is implemented and enabled.
"""
from __future__ import annotations

from datetime import datetime
from hashlib import sha256
from io import BytesIO
from secrets import token_hex

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from sqlalchemy import func, select

from app.config import settings
from app.core.exceptions import AppException, not_found
from app.models import (
    InternshipApplication,
    InternshipAuditTrail,
    InternshipBatch,
    InternshipInsurance,
    InternshipRecord,
    InternshipRegulatoryTask,
    InternshipRegulatoryTaskRow,
    InternshipRegulatoryTemplateVersion,
    InternshipSpecialFiling,
    Major,
    SchoolClass,
    StudentProfile,
)
from app.services.db_service import _as_id, _iso, _tid, session

REPORT_CODES = ("RP01", "RP02")
TASK_STATUSES = (
    "GENERATED",
    "VALIDATED",
    "EXPORTED",
    "SUBMITTED_EXTERNAL",
    "RECEIPT_PENDING",
    "ACCEPTED",
    "REJECTED",
)

RP01_FIELDS = [
    ("majorCode", "专业国标代码", True),
    ("studentNo", "学号", True),
    ("studentName", "姓名", True),
    ("enrollYear", "入学年份", True),
    ("educationYears", "学制", True),
    ("gender", "性别", True),
    ("organizationForm", "实习组织形式", True),
    ("companyName", "实习企业名称", True),
    ("companyCreditCode", "统一社会信用代码", True),
    ("companyNature", "企业性质", False),
    ("workCountry", "实习国家（地区）", True),
    ("workProvince", "实习省份", False),
    ("workCity", "实习地市", False),
    ("internshipStartDate", "实习开始时间", True),
    ("internshipEndDate", "实习结束时间", True),
    ("advisorEmployeeNo", "指导教师工号", True),
    ("insured", "是否购买保险", True),
    ("agreementSigned", "是否签订三方协议", True),
    ("missingDocumentExplanation", "未购买保险或未签协议说明", False),
    ("majorMatch", "专业对口情况", True),
    ("highRisk", "是否高风险实习", True),
    ("highRiskType", "高风险类型", False),
    ("nightOrOvertime", "是否加班夜班", True),
    ("holidayInternship", "是否休息日或法定节假日实习", True),
    ("filingStatus", "备案情况", True),
]

RP02_FIELDS = [
    ("majorName", "专业名称", True),
    ("majorCode", "专业代码", True),
    ("educationYears", "学制", True),
    ("className", "班级名称", True),
    ("startTerm", "起始学期", True),
    ("positionName", "实习岗位", True),
    ("majorMatch", "是否对口", True),
    ("internshipType", "实习类型", True),
    ("internshipArrangement", "实习安排", True),
    ("internshipForm", "实习形式", True),
    ("durationDays", "实习时长（天）", True),
    ("internshipStartDate", "实习开始时间", True),
    ("internshipEndDate", "实习结束时间", True),
    ("companyName", "实习单位名称", True),
    ("workCountry", "国家（地区）", True),
    ("workProvince", "所在省份", False),
    ("workCity", "所在地市", False),
    ("advisorName", "校内指导教师姓名", True),
    ("advisorPhone", "校内指导教师电话", False),
    ("enterpriseMentorName", "企业指导师傅姓名", False),
    ("enterpriseMentorPhone", "企业指导师傅电话", False),
    ("remuneration", "生均报酬", False),
    ("supervisionPhone", "校级督导咨询电话", False),
    ("insuranceType", "保险险种名称", False),
    ("policyNo", "保单号", False),
    ("insuranceFunder", "保险费出资方名称", False),
    ("insuranceBuyer", "保险购买方名称", False),
    ("agreementSigned", "是否签订三方协议", True),
    ("crossProvince", "是否跨省实习", False),
    ("overseas", "是否赴国（境）外实习", True),
]

_ENUMS = {
    "gender": ["男", "女", "未知"],
    "insured": ["是", "否"],
    "agreementSigned": ["是", "否"],
    "majorMatch": ["是", "否", "未知"],
    "highRisk": ["是", "否"],
    "nightOrOvertime": ["是", "否", "未知"],
    "holidayInternship": ["是", "否", "未知"],
    "filingStatus": ["已备案", "无需备案", "待备案"],
    "crossProvince": ["是", "否", "未知"],
    "overseas": ["是", "否"],
}

_BASELINES = {
    "RP01": {
        "name": "教育部/职业学校学生实习监管上报（益阳采购字段基线）",
        "fields": RP01_FIELDS,
        "crossRules": ["MISSING_DOC_EXPLANATION", "HIGH_RISK_TYPE", "DATE_ORDER"],
    },
    "RP02": {
        "name": "全国高等职业学校人才培养工作状态数据上报（益阳采购字段基线）",
        "fields": RP02_FIELDS,
        "crossRules": ["OVERSEAS_COUNTRY", "DATE_ORDER"],
    },
}

# Procurement requires every regulatory field to expose its business source and unit.
# These descriptions document the real authority used by _load_facts; unavailable facts
# are stated explicitly instead of silently inventing a default value.
_FIELD_METADATA = {
    "majorCode": ("Major.code", "—"),
    "majorName": ("Major.major_name", "—"),
    "studentNo": ("StudentProfile.student_no", "—"),
    "studentName": ("StudentProfile.real_name", "—"),
    "enrollYear": ("StudentProfile.enroll_date.year；无入学日期时回退 StudentProfile.grade", "年"),
    "educationYears": ("Major.education_years", "年"),
    "gender": ("StudentProfile.gender", "—"),
    "organizationForm": ("InternshipApplication.internship_mode / application_type", "—"),
    "companyName": ("InternshipApplication.company_name；缺失时回退 InternshipRecord.enterprise_name", "—"),
    "companyCreditCode": ("InternshipApplication.company_credit_code", "—"),
    "companyNature": ("InternshipApplication.company_nature", "—"),
    "workCountry": ("InternshipApplication.work_country", "—"),
    "workProvince": ("InternshipApplication.work_province", "—"),
    "workCity": ("InternshipApplication.work_city", "—"),
    "internshipStartDate": ("InternshipApplication.internship_start_date；缺失时回退 InternshipRecord.intern_start_date", "日期"),
    "internshipEndDate": ("InternshipApplication.internship_end_date；缺失时回退 InternshipRecord.intern_end_date", "日期"),
    "advisorEmployeeNo": ("当前 Standalone 正式事实尚未接入指导教师工号；必填时保持校验失败", "—"),
    "insured": ("InternshipInsurance.status=VERIFIED", "—"),
    "agreementSigned": ("InternshipRecord.agreement_info 状态解析", "—"),
    "missingDocumentExplanation": ("由保险与三方协议正式事实派生", "—"),
    "majorMatch": ("InternshipApplication.major_match", "—"),
    "highRisk": ("InternshipSpecialFiling(HIGH_RISK) / InternshipRecord.risk_level", "—"),
    "highRiskType": ("InternshipSpecialFiling.filing_type", "—"),
    "nightOrOvertime": ("当前正式业务事实未接入夜班/加班结论，保持“未知”", "—"),
    "holidayInternship": ("当前正式业务事实未接入节假日实习结论，保持“未知”", "—"),
    "filingStatus": ("InternshipSpecialFiling 审核状态派生", "—"),
    "className": ("SchoolClass.class_name", "—"),
    "startTerm": ("当前正式业务事实尚未接入实习起始学期；必填时保持校验失败", "—"),
    "positionName": ("InternshipApplication.position_name；缺失时回退 InternshipRecord.position_name", "—"),
    "internshipType": ("InternshipRecord.destination_type", "—"),
    "internshipArrangement": ("InternshipApplication.internship_mode / application_type", "—"),
    "internshipForm": ("InternshipApplication.internship_mode / application_type", "—"),
    "durationDays": ("由实习开始/结束日期计算（含首尾日）", "天"),
    "advisorName": ("InternshipRecord.advisor_name", "—"),
    "advisorPhone": ("当前正式业务事实尚未接入指导教师联系电话；缺失时保持空值", "—"),
    "enterpriseMentorName": ("InternshipApplication.enterprise_mentor_name；缺失时回退 InternshipRecord.enterprise_mentor_name", "—"),
    "enterpriseMentorPhone": ("InternshipApplication.enterprise_mentor_phone", "—"),
    "remuneration": ("InternshipApplication.agreed_salary", "元/月"),
    "supervisionPhone": ("当前正式业务事实尚未接入校级督导咨询电话；缺失时保持空值", "—"),
    "insuranceType": ("InternshipInsurance.coverage_type", "—"),
    "policyNo": ("InternshipInsurance.policy_no", "—"),
    "insuranceFunder": ("当前正式业务事实尚未接入保险费出资方；缺失时保持空值", "—"),
    "insuranceBuyer": ("当前正式业务事实尚未接入保险购买方；缺失时保持空值", "—"),
    "crossProvince": ("当前缺少学校所在地省份权威事实，不能可靠推导，保持“未知”", "—"),
    "overseas": ("InternshipSpecialFiling(OVERSEAS) / InternshipApplication.work_country 派生", "—"),
}


def _field_schema(code: str) -> list[dict]:
    result = []
    for key, label, required in _BASELINES[code]["fields"]:
        source, unit = _FIELD_METADATA.get(key, ("未配置正式业务来源", "—"))
        result.append({
            "key": key,
            "label": label,
            "required": required,
            "source": source,
            "unit": unit,
        })
    return result


def baseline_definition(code: str) -> dict:
    code = str(code or "").upper()
    if code not in REPORT_CODES:
        raise AppException("VALIDATION_ERROR", "上报类型仅支持 RP01/RP02")
    base = _BASELINES[code]
    return {
        "reportCode": code,
        "templateName": base["name"],
        "sourceLabel": "YIYANG_PROCUREMENT_BASELINE_2026",
        "sourceReference": "益阳职业技术学院岗位实习采购需求；目标平台当年正式模板需现场复核",
        "officialVerified": False,
        "fields": _field_schema(code),
        "enums": {key: value for key, value in _ENUMS.items()
                  if any(f[0] == key for f in base["fields"])},
        "crossRules": list(base["crossRules"]),
    }


def _blank(value) -> bool:
    return value is None or (isinstance(value, str) and not value.strip())


def validate_row(report_code: str, row: dict, template: dict | None = None) -> list[dict]:
    definition = template or baseline_definition(report_code)
    errors: list[dict] = []
    for field in definition.get("fields") or []:
        key = str(field.get("key") or "")
        label = str(field.get("label") or key)
        value = row.get(key)
        if field.get("required") and _blank(value):
            errors.append({"field": key, "label": label, "code": "REQUIRED",
                           "message": f"{label}不能为空"})
            continue
        allowed = (definition.get("enums") or {}).get(key)
        if allowed and not _blank(value) and str(value) not in {str(v) for v in allowed}:
            errors.append({"field": key, "label": label, "code": "ENUM",
                           "message": f"{label}值不在允许枚举中"})

    rules = set(definition.get("crossRules") or [])
    if "MISSING_DOC_EXPLANATION" in rules:
        if (row.get("insured") == "否" or row.get("agreementSigned") == "否") and _blank(
                row.get("missingDocumentExplanation")):
            errors.append({"field": "missingDocumentExplanation", "label": "说明",
                           "code": "CROSS_FIELD",
                           "message": "未购买保险或未签订三方协议时必须填写说明"})
    if "HIGH_RISK_TYPE" in rules and row.get("highRisk") == "是" and _blank(row.get("highRiskType")):
        errors.append({"field": "highRiskType", "label": "高风险类型", "code": "CROSS_FIELD",
                       "message": "高风险实习必须填写高风险类型"})
    if "OVERSEAS_COUNTRY" in rules and row.get("overseas") == "是" and _blank(row.get("workCountry")):
        errors.append({"field": "workCountry", "label": "国家（地区）", "code": "CROSS_FIELD",
                       "message": "赴国（境）外实习必须填写国家（地区）"})
    if "DATE_ORDER" in rules:
        start = str(row.get("internshipStartDate") or "")
        end = str(row.get("internshipEndDate") or "")
        if start and end and end < start:
            errors.append({"field": "internshipEndDate", "label": "实习结束时间",
                           "code": "CROSS_FIELD", "message": "实习结束时间不得早于开始时间"})
    return errors


def build_workbook(report_code: str, rows: list[dict], template: dict | None = None) -> bytes:
    definition = template or baseline_definition(report_code)
    fields = definition.get("fields") or []
    wb = Workbook()
    ws = wb.active
    ws.title = report_code
    for col, field in enumerate(fields, 1):
        cell = ws.cell(1, col, str(field.get("label") or field.get("key") or ""))
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
        ws.column_dimensions[cell.column_letter].width = min(28, max(12, len(str(cell.value)) * 2 + 2))
    for row_no, row in enumerate(rows, 2):
        for col, field in enumerate(fields, 1):
            key = str(field.get("key") or "")
            value = row.get(key)
            cell = ws.cell(row_no, col, "" if value is None else value)
            if key in {"studentNo", "companyCreditCode", "policyNo", "advisorEmployeeNo"}:
                cell.number_format = "@"
                if value is not None:
                    cell.value = str(value)
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def build_error_workbook(report_code: str, rows: list[tuple[int, dict, list[dict]]]) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = f"{report_code}-错误行"
    ws.append(["原始行号", "学号", "姓名", "错误字段", "错误代码", "错误说明"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
    for row_no, payload, errors in rows:
        for error in errors:
            ws.append([
                row_no,
                str(payload.get("studentNo") or ""),
                str(payload.get("studentName") or ""),
                error.get("label") or error.get("field") or "",
                error.get("code") or "",
                error.get("message") or "",
            ])
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def _yes_no(value, *, unknown=False):
    if value is True:
        return "是"
    if value is False:
        return "否"
    return "未知" if unknown else ""


def _date(value):
    if value is None:
        return ""
    return value.date().isoformat() if hasattr(value, "date") else str(value)[:10]


def _duration_days(start, end):
    if not start or not end:
        return ""
    a = start.date() if hasattr(start, "date") else start
    b = end.date() if hasattr(end, "date") else end
    try:
        return max(0, (b - a).days + 1)
    except Exception:
        return ""


def _agreement_signed(record):
    value = str(record.agreement_info or "").upper()
    if any(key in value for key in ("SIGNED", "EFFECTIVE", "VALID", "已签", "生效")):
        return "是"
    return "否" if value else ""


def _history(status: str, actor: str, note: str = "") -> dict:
    return {"status": status, "at": datetime.utcnow().isoformat(timespec="seconds") + "Z",
            "actor": actor, "note": note}


def _actor(user=None) -> str:
    return str((user or {}).get("realName") or "系统")


def _audit(db, task_id: int, action: str, detail: dict, user=None):
    db.add(InternshipAuditTrail(
        tenant_id=_tid(), target_id=task_id, target_type="REGULATORY_REPORT",
        action=action, operator_name=_actor(user), detail_json=detail,
        occurred_at=datetime.utcnow(),
    ))


def _template_view(row) -> dict:
    return {
        "id": str(row.id),
        "reportCode": row.report_code,
        "versionNo": int(row.version_no),
        "templateName": row.template_name,
        "sourceLabel": row.source_label,
        "sourceReference": row.source_reference or "",
        "officialVerified": bool(row.official_verified),
        "fields": list(row.field_schema_json or []),
        "enums": dict(row.enum_schema_json or {}),
        "crossRules": list(row.cross_rule_json or []),
        "status": row.status,
        "externalReadiness": "READY" if row.official_verified else "BLOCKED_EXTERNAL",
    }


def _task_view(row) -> dict:
    return {
        "id": str(row.id),
        "taskNo": row.task_no,
        "reportCode": row.report_code,
        "templateVersionId": str(row.template_version_id),
        "batchId": str(row.batch_id),
        "status": row.status,
        "statusHistory": list(row.status_history_json or []),
        "totalRows": int(row.total_rows or 0),
        "validRows": int(row.valid_rows or 0),
        "errorRows": int(row.error_rows or 0),
        "outputFilename": row.output_filename or "",
        "outputSha256": row.output_sha256 or "",
        "externalSubmissionRef": row.external_submission_ref or "",
        "receiptCode": row.receipt_code or "",
        "receiptMessage": row.receipt_message or "",
        "generatedAt": _iso(row.generated_at) or "",
        "validatedAt": _iso(row.validated_at) or "",
        "exportedAt": _iso(row.exported_at) or "",
        "submittedAt": _iso(row.submitted_at) or "",
        "receiptAt": _iso(row.receipt_at) or "",
        "externalReadiness": (
            "READY" if bool(getattr(settings, "REGULATORY_RECEIPT_ADAPTER_ENABLED", False))
            else "BLOCKED_EXTERNAL"
        ),
    }


def _ensure_baseline(db, code: str):
    row = db.scalar(select(InternshipRegulatoryTemplateVersion).where(
        InternshipRegulatoryTemplateVersion.tenant_id == _tid(),
        InternshipRegulatoryTemplateVersion.report_code == code,
        InternshipRegulatoryTemplateVersion.status == "ACTIVE",
        InternshipRegulatoryTemplateVersion.is_deleted.is_(False),
    ).order_by(InternshipRegulatoryTemplateVersion.version_no.desc()))
    if row:
        fields = []
        changed = False
        for raw in list(row.field_schema_json or []):
            field = dict(raw or {})
            key = str(field.get("key") or "")
            source, unit = _FIELD_METADATA.get(key, ("未配置正式业务来源", "—"))
            if not str(field.get("source") or "").strip():
                field["source"] = source
                changed = True
            if not str(field.get("unit") or "").strip():
                field["unit"] = unit
                changed = True
            fields.append(field)
        if changed:
            row.field_schema_json = fields
        return row

    definition = baseline_definition(code)
    row = InternshipRegulatoryTemplateVersion(
        tenant_id=_tid(), report_code=code, version_no=1,
        template_name=definition["templateName"], source_label=definition["sourceLabel"],
        source_reference=definition["sourceReference"], official_verified=False,
        field_schema_json=definition["fields"], enum_schema_json=definition["enums"],
        cross_rule_json=definition["crossRules"], mapping_schema_json={},
        status="ACTIVE", effective_at=datetime.utcnow(),
        change_reason="C07/G18 首版采购字段基线；正式目标平台模板待外部联调确认",
    )
    db.add(row)
    db.flush()
    return row


def list_templates(report_code: str | None = None):
    code = str(report_code or "").upper()
    if code and code not in REPORT_CODES:
        raise AppException("VALIDATION_ERROR", "上报类型仅支持 RP01/RP02")
    with session() as db:
        if code:
            _ensure_baseline(db, code)
        else:
            for item in REPORT_CODES:
                _ensure_baseline(db, item)
        db.commit()
        query = select(InternshipRegulatoryTemplateVersion).where(
            InternshipRegulatoryTemplateVersion.tenant_id == _tid(),
            InternshipRegulatoryTemplateVersion.is_deleted.is_(False),
        )
        if code:
            query = query.where(InternshipRegulatoryTemplateVersion.report_code == code)
        rows = db.scalars(query.order_by(
            InternshipRegulatoryTemplateVersion.report_code,
            InternshipRegulatoryTemplateVersion.version_no.desc(),
        )).all()
        return [_template_view(row) for row in rows]


def create_template_version(report_code: str, body: dict, user=None):
    code = str(report_code or "").upper()
    baseline_definition(code)
    payload = body or {}
    fields = payload.get("fields")
    if not isinstance(fields, list) or not fields:
        raise AppException("VALIDATION_ERROR", "模板字段不能为空")
    keys = []
    normalized_fields = []
    for field in fields:
        key = str((field or {}).get("key") or "").strip()
        label = str((field or {}).get("label") or "").strip()
        if not key or not label or key in keys:
            raise AppException("VALIDATION_ERROR", "模板字段 key/label 必须完整且不能重复")
        if key not in _FIELD_METADATA:
            raise AppException(
                "VALIDATION_ERROR",
                f"字段 {key} 尚未建立 Standalone 正式业务映射，不能直接启用为上报模板字段",
            )
        source = str((field or {}).get("source") or _FIELD_METADATA[key][0]).strip()
        unit = str((field or {}).get("unit") or _FIELD_METADATA[key][1]).strip() or "—"
        if not source:
            raise AppException("VALIDATION_ERROR", f"字段 {key} 必须填写真实业务来源")
        keys.append(key)
        normalized_fields.append({
            "key": key,
            "label": label[:120],
            "required": bool((field or {}).get("required")),
            "source": source[:500],
            "unit": unit[:80],
        })
    source_reference = str(payload.get("sourceReference") or "").strip()
    change_reason = str(payload.get("changeReason") or "").strip()
    if len(source_reference) < 5:
        raise AppException("VALIDATION_ERROR", "新模板版本必须填写可追溯的模板来源/文件版本")
    if len(change_reason) < 2:
        raise AppException("VALIDATION_ERROR", "新模板版本必须填写变更原因")

    with session() as db:
        current = _ensure_baseline(db, code)
        version = int(db.scalar(select(func.max(InternshipRegulatoryTemplateVersion.version_no)).where(
            InternshipRegulatoryTemplateVersion.tenant_id == _tid(),
            InternshipRegulatoryTemplateVersion.report_code == code,
        )) or 0) + 1
        for old in db.scalars(select(InternshipRegulatoryTemplateVersion).where(
            InternshipRegulatoryTemplateVersion.tenant_id == _tid(),
            InternshipRegulatoryTemplateVersion.report_code == code,
            InternshipRegulatoryTemplateVersion.status == "ACTIVE",
            InternshipRegulatoryTemplateVersion.is_deleted.is_(False),
        )).all():
            old.status = "RETIRED"
        row = InternshipRegulatoryTemplateVersion(
            tenant_id=_tid(), report_code=code, version_no=version,
            template_name=str(payload.get("templateName") or current.template_name)[:200],
            source_label=str(payload.get("sourceLabel") or "SCHOOL_CONFIRMED_TEMPLATE")[:120],
            source_reference=source_reference[:500],
            official_verified=False,
            field_schema_json=normalized_fields,
            enum_schema_json=dict(payload.get("enums") or {}),
            cross_rule_json=list(payload.get("crossRules") or []),
            mapping_schema_json=dict(payload.get("mapping") or {}),
            status="ACTIVE", effective_at=datetime.utcnow(),
            change_reason=change_reason[:500],
        )
        db.add(row)
        db.flush()
        _audit(db, row.id, "REGULATORY_TEMPLATE_VERSION_CREATE", {
            "reportCode": code, "versionNo": version, "officialVerified": False,
            "note": "人工录入模板不得自行宣称目标平台已官方核验",
        }, user)
        db.commit()
        return _template_view(row)


def build_field_dictionary_workbook(definition: dict) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "字段来源清单"
    ws.append(["序号", "字段Key", "上报列名", "必填", "单位", "业务来源", "枚举"])
    for cell in ws[1]:
        cell.font = Font(bold=True)
        cell.alignment = Alignment(horizontal="center", vertical="center")
    enums = definition.get("enums") or {}
    for index, field in enumerate(definition.get("fields") or [], 1):
        key = str(field.get("key") or "")
        ws.append([
            index,
            key,
            str(field.get("label") or ""),
            "是" if field.get("required") else "否",
            str(field.get("unit") or "—"),
            str(field.get("source") or ""),
            " / ".join(str(x) for x in (enums.get(key) or [])),
        ])
    widths = [8, 28, 28, 10, 16, 64, 36]
    for index, width in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(1, index).column_letter].width = width
    ws.freeze_panes = "A2"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def field_dictionary_file(report_code: str, user=None) -> tuple[bytes, str]:
    code = str(report_code or "").upper()
    baseline_definition(code)
    with session() as db:
        template = _ensure_baseline(db, code)
        definition = {
            "fields": list(template.field_schema_json or []),
            "enums": dict(template.enum_schema_json or {}),
        }
        payload = build_field_dictionary_workbook(definition)
        filename = f"{code}-V{int(template.version_no)}-field-dictionary.xlsx"
        _audit(db, int(template.id), "REGULATORY_FIELD_DICTIONARY_EXPORT", {
            "reportCode": code,
            "templateVersionId": str(template.id),
            "versionNo": int(template.version_no),
            "fieldCount": len(definition["fields"]),
            "filename": filename,
            "sha256": sha256(payload).hexdigest(),
        }, user)
        db.commit()
        return payload, filename


def _approved_application_map(db, record_ids: list[int]) -> dict[int, InternshipApplication]:
    if not record_ids:
        return {}
    rows = db.scalars(select(InternshipApplication).where(
        InternshipApplication.tenant_id == _tid(),
        InternshipApplication.record_id.in_(record_ids),
        InternshipApplication.status == "APPROVED",
        InternshipApplication.is_deleted.is_(False),
    )).all()
    out = {}
    for row in rows:
        rid = int(row.record_id or 0)
        old = out.get(rid)
        current_key = (row.reviewed_at or row.submitted_at or row.created_at, row.id)
        old_key = ((old.reviewed_at or old.submitted_at or old.created_at), old.id) if old else None
        if not old or current_key > old_key:
            out[rid] = row
    return out


def _load_facts(db, batch_id: int) -> list[dict]:
    records = db.scalars(select(InternshipRecord).where(
        InternshipRecord.tenant_id == _tid(),
        InternshipRecord.batch_id == batch_id,
        InternshipRecord.is_deleted.is_(False),
    ).order_by(InternshipRecord.id)).all()
    if not records:
        return []

    student_ids = [int(r.student_id) for r in records]
    students = {int(x.id): x for x in db.scalars(select(StudentProfile).where(
        StudentProfile.tenant_id == _tid(),
        StudentProfile.id.in_(student_ids),
        StudentProfile.is_deleted.is_(False),
    )).all()}
    major_ids = {int(x.major_id) for x in students.values() if x.major_id}
    class_ids = {int(x.class_id) for x in students.values() if x.class_id}
    majors = ({int(x.id): x for x in db.scalars(select(Major).where(
        Major.tenant_id == _tid(), Major.id.in_(major_ids), Major.is_deleted.is_(False),
    )).all()} if major_ids else {})
    classes = ({int(x.id): x for x in db.scalars(select(SchoolClass).where(
        SchoolClass.tenant_id == _tid(), SchoolClass.id.in_(class_ids),
        SchoolClass.is_deleted.is_(False),
    )).all()} if class_ids else {})
    record_ids = [int(r.id) for r in records]
    apps = _approved_application_map(db, record_ids)
    insurances = {int(x.internship_id): x for x in db.scalars(select(InternshipInsurance).where(
        InternshipInsurance.tenant_id == _tid(),
        InternshipInsurance.internship_id.in_(record_ids),
        InternshipInsurance.is_deleted.is_(False),
    )).all()}
    filings: dict[int, list] = {}
    for filing in db.scalars(select(InternshipSpecialFiling).where(
        InternshipSpecialFiling.tenant_id == _tid(),
        InternshipSpecialFiling.internship_id.in_(record_ids),
        InternshipSpecialFiling.is_deleted.is_(False),
    )).all():
        filings.setdefault(int(filing.internship_id), []).append(filing)

    facts = []
    for record in records:
        student = students.get(int(record.student_id))
        if not student:
            continue
        app = apps.get(int(record.id))
        insurance = insurances.get(int(record.id))
        major = majors.get(int(student.major_id)) if student.major_id else None
        clazz = classes.get(int(student.class_id)) if student.class_id else None
        approved_filings = [x for x in filings.get(int(record.id), []) if x.status == "APPROVED"]
        high = next((x for x in approved_filings if x.filing_type == "HIGH_RISK"), None)
        overseas_filing = next((x for x in approved_filings if x.filing_type == "OVERSEAS"), None)

        insured = "是" if insurance and insurance.status == "VERIFIED" else ("否" if insurance else "")
        agreement = _agreement_signed(record)
        country = str(getattr(app, "work_country", "") or "") if app else ""
        overseas_yes = bool(overseas_filing or (
            country and country.upper() not in {"中国", "中国大陆", "CHINA", "CN"}
        ))
        start = (getattr(app, "internship_start_date", None) if app else None) or record.intern_start_date
        end = (getattr(app, "internship_end_date", None) if app else None) or record.intern_end_date
        major_match = getattr(app, "major_match", None) if app else None
        organization = str(
            (getattr(app, "internship_mode", "") or getattr(app, "application_type", ""))
            if app else ""
        )
        missing_doc = []
        if insured == "否":
            missing_doc.append("未取得已核验保险")
        if agreement == "否":
            missing_doc.append("三方协议当前状态未确认签署")

        facts.append({
            "_studentId": int(student.id),
            "_internshipId": int(record.id),
            "majorCode": str(major.code or "") if major else "",
            "majorName": str(major.major_name or "") if major else "",
            "studentNo": str(student.student_no or ""),
            "studentName": str(student.real_name or ""),
            "enrollYear": str(student.enroll_date.year if student.enroll_date else student.grade or ""),
            "educationYears": str(major.education_years or "") if major else "",
            "gender": str(student.gender or "未知"),
            "organizationForm": organization,
            "companyName": str((getattr(app, "company_name", "") if app else "") or record.enterprise_name or ""),
            "companyCreditCode": str(getattr(app, "company_credit_code", "") or "") if app else "",
            "companyNature": str(getattr(app, "company_nature", "") or "") if app else "",
            "workCountry": country,
            "workProvince": str(getattr(app, "work_province", "") or "") if app else "",
            "workCity": str(getattr(app, "work_city", "") or "") if app else "",
            "internshipStartDate": _date(start),
            "internshipEndDate": _date(end),
            "advisorEmployeeNo": "",
            "insured": insured,
            "agreementSigned": agreement,
            "missingDocumentExplanation": "；".join(missing_doc),
            "majorMatch": _yes_no(major_match, unknown=True),
            "highRisk": "是" if (high or record.risk_level == "HIGH") else "否",
            "highRiskType": str(high.filing_type if high else ""),
            "nightOrOvertime": "未知",
            "holidayInternship": "未知",
            "filingStatus": (
                "已备案" if approved_filings else
                "待备案" if (high or overseas_yes) else
                "无需备案"
            ),
            "className": str(clazz.class_name or "") if clazz else "",
            "startTerm": "",
            "positionName": str((getattr(app, "position_name", "") if app else "") or record.position_name or ""),
            "internshipType": str(record.destination_type or ""),
            "internshipArrangement": organization,
            "internshipForm": organization,
            "durationDays": _duration_days(start, end),
            "advisorName": str(record.advisor_name or ""),
            "advisorPhone": "",
            "enterpriseMentorName": str(
                (getattr(app, "enterprise_mentor_name", "") if app else "")
                or record.enterprise_mentor_name or ""
            ),
            "enterpriseMentorPhone": str(getattr(app, "enterprise_mentor_phone", "") or "") if app else "",
            "remuneration": str(getattr(app, "agreed_salary", "") or "") if app else "",
            "supervisionPhone": "",
            "insuranceType": str(insurance.coverage_type or "") if insurance else "",
            "policyNo": str(insurance.policy_no or "") if insurance else "",
            "insuranceFunder": "",
            "insuranceBuyer": "",
            "crossProvince": "未知",
            "overseas": "是" if overseas_yes else "否",
        })
    return facts


def create_task(body: dict, user=None):
    payload = body or {}
    code = str(payload.get("reportCode") or "").upper()
    baseline_definition(code)
    try:
        batch_id = int(payload.get("batchId"))
    except (TypeError, ValueError):
        raise AppException("VALIDATION_ERROR", "请选择实习批次") from None

    with session() as db:
        batch = db.scalar(select(InternshipBatch).where(
            InternshipBatch.id == batch_id,
            InternshipBatch.tenant_id == _tid(),
            InternshipBatch.is_deleted.is_(False),
        ))
        if not batch:
            raise not_found("实习批次不存在")
        template = _ensure_baseline(db, code)
        facts = _load_facts(db, batch_id)
        now = datetime.utcnow()
        task = InternshipRegulatoryTask(
            tenant_id=_tid(),
            task_no=f"{code}-{now.strftime('%Y%m%d%H%M%S')}-{token_hex(3).upper()}",
            report_code=code, template_version_id=template.id, batch_id=batch_id,
            status="GENERATED",
            status_history_json=[_history("GENERATED", _actor(user), "冻结当前正式业务事实")],
            total_rows=len(facts), valid_rows=0, error_rows=0,
            created_by_name=_actor(user), generated_at=now,
        )
        db.add(task)
        db.flush()
        for index, fact in enumerate(facts, 1):
            snapshot = dict(fact)
            student_id = snapshot.pop("_studentId", None)
            internship_id = snapshot.pop("_internshipId", None)
            db.add(InternshipRegulatoryTaskRow(
                tenant_id=_tid(), task_id=task.id, row_no=index,
                student_id=student_id, internship_id=internship_id,
                payload_json=snapshot, validation_errors_json=[], is_valid=False,
            ))
        _audit(db, task.id, "REGULATORY_TASK_GENERATE", {
            "reportCode": code, "batchId": str(batch_id), "totalRows": len(facts),
            "templateVersionId": str(template.id),
            "officialVerified": bool(template.official_verified),
        }, user)
        db.commit()
        return _task_view(task)


def list_tasks(report_code: str | None = None, batch_id: int | None = None):
    code = str(report_code or "").upper()
    if code and code not in REPORT_CODES:
        raise AppException("VALIDATION_ERROR", "上报类型仅支持 RP01/RP02")
    with session() as db:
        query = select(InternshipRegulatoryTask).where(
            InternshipRegulatoryTask.tenant_id == _tid(),
            InternshipRegulatoryTask.is_deleted.is_(False),
        )
        if code:
            query = query.where(InternshipRegulatoryTask.report_code == code)
        if batch_id:
            query = query.where(InternshipRegulatoryTask.batch_id == int(batch_id))
        rows = db.scalars(query.order_by(InternshipRegulatoryTask.id.desc()).limit(200)).all()
        return [_task_view(row) for row in rows]


def _task(db, task_id: int, *, lock=False):
    query = select(InternshipRegulatoryTask).where(
        InternshipRegulatoryTask.id == _as_id(task_id),
        InternshipRegulatoryTask.tenant_id == _tid(),
        InternshipRegulatoryTask.is_deleted.is_(False),
    )
    row = db.scalar(query.with_for_update() if lock else query)
    if not row:
        raise not_found("上报任务不存在")
    return row


def task_detail(task_id: int):
    with session() as db:
        task = _task(db, task_id)
        template = db.get(InternshipRegulatoryTemplateVersion, task.template_version_id)
        result = _task_view(task)
        result["template"] = _template_view(template) if template else None
        return result


def list_task_rows(
    task_id: int, *, page: int = 1, page_size: int = 50, error_only: bool = False,
) -> dict:
    page = max(1, int(page or 1))
    page_size = min(200, max(1, int(page_size or 50)))
    with session() as db:
        task = _task(db, task_id)
        base = select(InternshipRegulatoryTaskRow).where(
            InternshipRegulatoryTaskRow.tenant_id == _tid(),
            InternshipRegulatoryTaskRow.task_id == task.id,
            InternshipRegulatoryTaskRow.is_deleted.is_(False),
        )
        if error_only:
            if task.validated_at is None:
                return {
                    "items": [], "total": 0, "page": page, "pageSize": page_size,
                    "errorOnly": True,
                }
            base = base.where(
                func.json_length(InternshipRegulatoryTaskRow.validation_errors_json) > 0
            )
        total = int(db.scalar(select(func.count()).select_from(base.subquery())) or 0)
        rows = db.scalars(
            base.order_by(InternshipRegulatoryTaskRow.row_no)
            .offset((page - 1) * page_size)
            .limit(page_size)
        ).all()
        return {
            "items": [
                {
                    "id": str(row.id),
                    "rowNo": int(row.row_no),
                    "studentId": str(row.student_id) if row.student_id else "",
                    "internshipId": str(row.internship_id) if row.internship_id else "",
                    "payload": dict(row.payload_json or {}),
                    "errors": list(row.validation_errors_json or []),
                    "isValid": bool(row.is_valid),
                    "validationState": (
                        "PENDING" if task.validated_at is None
                        else ("FAILED" if row.validation_errors_json else "PASSED")
                    ),
                }
                for row in rows
            ],
            "total": total,
            "page": page,
            "pageSize": page_size,
            "errorOnly": bool(error_only),
        }


def validate_task(task_id: int, user=None):
    with session() as db:
        task = _task(db, task_id, lock=True)
        if task.status not in {"GENERATED", "VALIDATED"}:
            raise AppException("DATA_CONFLICT", "当前任务状态不可重新校验")
        template = db.get(InternshipRegulatoryTemplateVersion, task.template_version_id)
        if not template:
            raise not_found("上报模板版本不存在")
        definition = {
            "fields": list(template.field_schema_json or []),
            "enums": dict(template.enum_schema_json or {}),
            "crossRules": list(template.cross_rule_json or []),
        }
        rows = db.scalars(select(InternshipRegulatoryTaskRow).where(
            InternshipRegulatoryTaskRow.tenant_id == _tid(),
            InternshipRegulatoryTaskRow.task_id == task.id,
            InternshipRegulatoryTaskRow.is_deleted.is_(False),
        ).order_by(InternshipRegulatoryTaskRow.row_no)).all()
        valid = 0
        error_rows = 0
        for row in rows:
            found = validate_row(task.report_code, dict(row.payload_json or {}), definition)
            row.validation_errors_json = found
            row.is_valid = not found
            if found:
                error_rows += 1
            else:
                valid += 1
        task.valid_rows = valid
        task.error_rows = error_rows
        task.validated_at = datetime.utcnow()
        if error_rows == 0:
            task.status = "VALIDATED"
            task.status_history_json = list(task.status_history_json or []) + [
                _history("VALIDATED", _actor(user), f"{valid} 行全部通过")
            ]
        _audit(db, task.id, "REGULATORY_TASK_VALIDATE", {
            "totalRows": len(rows), "validRows": valid, "errorRows": error_rows,
            "validationPassed": error_rows == 0,
        }, user)
        db.commit()
        result = _task_view(task)
        result["validationPassed"] = error_rows == 0
        return result


def _rows(db, task_id: int):
    return db.scalars(select(InternshipRegulatoryTaskRow).where(
        InternshipRegulatoryTaskRow.tenant_id == _tid(),
        InternshipRegulatoryTaskRow.task_id == task_id,
        InternshipRegulatoryTaskRow.is_deleted.is_(False),
    ).order_by(InternshipRegulatoryTaskRow.row_no)).all()


def error_file(task_id: int, user=None):
    with session() as db:
        task = _task(db, task_id, lock=True)
        rows = [
            (row.row_no, dict(row.payload_json or {}), list(row.validation_errors_json or []))
            for row in _rows(db, task.id) if row.validation_errors_json
        ]
        payload = build_error_workbook(task.report_code, rows)
        filename = f"{task.task_no}-errors.xlsx"
        task.error_filename = filename
        task.error_sha256 = sha256(payload).hexdigest()
        _audit(db, task.id, "REGULATORY_ERROR_XLSX_EXPORT", {
            "filename": filename, "sha256": task.error_sha256, "errorRows": len(rows),
        }, user)
        db.commit()
        return payload, filename


def export_file(task_id: int, user=None):
    with session() as db:
        task = _task(db, task_id, lock=True)
        if task.status not in {"VALIDATED", "EXPORTED"}:
            raise AppException("DATA_CONFLICT", "只有全部校验通过的任务才能生成正式上报文件")
        template = db.get(InternshipRegulatoryTemplateVersion, task.template_version_id)
        if not template:
            raise not_found("上报模板版本不存在")
        definition = {
            "fields": list(template.field_schema_json or []),
            "enums": dict(template.enum_schema_json or {}),
            "crossRules": list(template.cross_rule_json or []),
        }
        rows = _rows(db, task.id)
        if any(not row.is_valid for row in rows):
            raise AppException("DATA_CONFLICT", "任务中仍存在未通过校验的数据行")
        payload = build_workbook(
            task.report_code, [dict(row.payload_json or {}) for row in rows], definition
        )
        filename = f"{task.task_no}.xlsx"
        digest = sha256(payload).hexdigest()
        task.status = "EXPORTED"
        task.output_filename = filename
        task.output_sha256 = digest
        task.output_size = len(payload)
        task.exported_at = datetime.utcnow()
        history = list(task.status_history_json or [])
        if not history or history[-1].get("status") != "EXPORTED":
            history.append(_history("EXPORTED", _actor(user), filename))
        task.status_history_json = history
        _audit(db, task.id, "REGULATORY_FORMAL_XLSX_EXPORT", {
            "filename": filename, "sha256": digest, "bytes": len(payload),
            "totalRows": len(rows), "templateVersionId": str(task.template_version_id),
        }, user)
        db.commit()
        return payload, filename


def mark_submitted(task_id: int, body: dict, user=None):
    reference = str((body or {}).get("externalSubmissionRef") or "").strip()
    if not reference:
        raise AppException("VALIDATION_ERROR", "请填写真实外部提交凭据/任务号；禁止用空值冒充已提交")
    with session() as db:
        task = _task(db, task_id, lock=True)
        if task.status != "EXPORTED":
            raise AppException("DATA_CONFLICT", "必须先导出正式上报文件，才能登记外部提交")
        history = list(task.status_history_json or [])
        history.append(_history("SUBMITTED_EXTERNAL", _actor(user), reference))
        history.append(_history("RECEIPT_PENDING", _actor(user), "尚未取得目标平台真实回执"))
        task.status = "RECEIPT_PENDING"
        task.status_history_json = history
        task.external_submission_ref = reference
        task.submitted_at = datetime.utcnow()
        _audit(db, task.id, "REGULATORY_EXTERNAL_SUBMISSION_RECORDED", {
            "externalSubmissionRef": reference,
            "finalStatus": "RECEIPT_PENDING",
            "acceptedClaimed": False,
        }, user)
        db.commit()
        return _task_view(task)


def record_external_receipt(task_id: int, body: dict, user=None):
    if not bool(getattr(settings, "REGULATORY_RECEIPT_ADAPTER_ENABLED", False)):
        raise AppException(
            "BLOCKED_EXTERNAL",
            "未配置并验收目标监管平台真实回执适配器；禁止人工把任务标记为 ACCEPTED/REJECTED",
            http_status=409,
        )
    status = str((body or {}).get("status") or "").upper()
    receipt_code = str((body or {}).get("receiptCode") or "").strip()
    if status not in {"ACCEPTED", "REJECTED"} or not receipt_code:
        raise AppException("VALIDATION_ERROR", "真实回执必须包含 ACCEPTED/REJECTED 与回执编号")
    raise AppException(
        "BLOCKED_EXTERNAL",
        "回执适配器开关已开启但签名/来源校验实现尚未装配，继续阻断状态落库",
        http_status=409,
    )
