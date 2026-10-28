"""师生开户唯一批量命令：持久化预检、全批次事务、幂等确认和限时加密回执。"""
import hashlib
import io
import json
import re
import secrets
from datetime import datetime, timedelta
from openpyxl import load_workbook
from sqlalchemy import select, or_
from app.core.exceptions import AppException, no_permission, check_version
from app.core.field_crypto import encrypt_field, decrypt_field
from app.core.permissions import enforce_permission, has_permission
from app.core.security import hash_password
from app.models import (User, Role, UserRole, College, Major, SchoolClass, StudentProfile,
                        StaffProfile, StudentAccountLink, ExcelImportJob, SchoolSetting)
from app.services import system_management as sm
from app.services.system_authorization import audit, role_permissions
from app.services.student_account_link_service import bind_in_session
from app.services.xlsx_util import read_xlsx, build_template_xlsx, build_ledger_xlsx
from app.db.session import get_sessionmaker

HEADERS = {'学号或工号': 'number', '姓名': 'name', '院系代码': 'collegeCode', '院系名称': 'collegeName',
           '专业代码': 'majorCode', '专业名称': 'majorName', '班级代码': 'classCode', '班级名称': 'className', '年级': 'grade',
           '部门代码':'departmentCode','部门名称':'departmentName'}
_NUMBER = re.compile(r'^[\w.-]{1,50}$')
ORG_HEADERS={label:key for label,key in HEADERS.items() if key not in ('number','name')}


def import_permission(kind,action):
    return f'system.org.{action}' if kind=='organizations' else f'system.identity-import.{action}'


def template(user, kind):
    enforce_permission(user, import_permission(kind,'view'))
    if kind=='organizations':
        return build_template_xlsx(list(ORG_HEADERS),sample=['COL001','请替换为院系','MAJ001','请替换为专业','CLS001','请替换为班级','2026','',''],
            notes=['组织代码按文本填写；可只填院系、院系与专业、完整班级链或校内部门。','已有代码必须与原名称一致；改名、停用从组织维护办理。不会创建人员或账号。'])
    sample = ['0001001', '请替换为姓名', 'COL001', '请替换为院系', 'MAJ001', '请替换为专业', 'CLS001', '请替换为班级', '2024']
    if kind == 'staff': sample[4:] = [''] * 5
    sample += ['', '']
    return build_template_xlsx(list(HEADERS), sample=sample, notes=[
        '学号、工号和所有组织代码必须使用文本格式，保留前导零。示例行请删除或替换。',
        '新账号首次登录须修改随机临时密码；更新资料不会重置密码或覆盖角色。',
        '同一代码只对应同一组织；同名组织用不同代码区分。',
        '教师必填工号、姓名及院系或校内部门；学生须填院系、专业、班级及年级。'])


def _text_identifiers(content):
    wb = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        ws = wb['导入模板'] if '导入模板' in wb.sheetnames else wb.worksheets[0]
        it = ws.iter_rows(); header = next(it, [])
        columns = {i for i, c in enumerate(header) if str(c.value or '').rstrip(' *') in
                   {'学号或工号', '院系代码', '专业代码', '班级代码', '部门代码'}}
        for n, cells in enumerate(it, 2):
            for i in columns:
                if i < len(cells) and cells[i].value is not None and cells[i].data_type != 's':
                    raise AppException('VALIDATION_ERROR', f'第{n}行编号必须为文本，不能是数字或公式；请核对前导零后重传')
    finally:
        wb.close()


def _evaluate(db, user, kind, rows):
    """Codes are exact keys. Name collisions never choose the first candidate."""
    errors, entries, new_orgs = [], [], {}
    dictionary=db.scalar(sm.tenant_query(SchoolSetting,user).where(SchoolSetting.setting_key=='dictionary'))
    aliases={r['alias']:r['code'] for r in (dictionary.value_json or {}).get('aliases',[])} if dictionary else {}
    unique_alias_codes=sm.unique_organization_codes(db,user,aliases.values())
    catalogs={}
    for org_kind,prefix in [('colleges','college'),('majors','major'),('classes','class'),('departments','department')]:
        model,_,code_col,_=sm.ORG[org_kind];catalogs[org_kind]={}
        for item in db.scalars(sm.tenant_query(model,user).where(getattr(model,code_col).in_({r.get(prefix+'Code') for r in rows if r.get(prefix+'Code')}))):
            catalogs[org_kind].setdefault(getattr(item,code_col),[]).append(item)
    person_model=StudentProfile if kind=='students' else StaffProfile
    number_col=person_model.student_no if kind=='students' else person_model.staff_no
    numbers={r.get('number') for r in rows}
    people={getattr(p,'student_no' if kind=='students' else 'staff_no'):p for p in db.scalars(select(person_model).where(
        person_model.tenant_id==sm.ids(user)[0],number_col.in_(numbers)))}
    accounts={a.login_name:a for a in db.scalars(select(User).where(User.tenant_id==sm.ids(user)[0],User.login_name.in_(numbers)))}
    scoped = not sm.school_wide(user)
    scopes = sm.scope_ids(db,user) if scoped else {}
    if scoped:
        people_query = sm.scoped_students(db,user) if kind=='students' else sm.scoped_staff(db,user)
        allowed_people = set(db.scalars(people_query.with_only_columns(person_model.id).where(number_col.in_(numbers))))
        allowed_accounts = set(db.scalars(sm.scoped_accounts(db,user).with_only_columns(User.id).where(User.login_name.in_(numbers))))
    opening_role = None
    opening_permissions = []
    if kind != 'organizations':
        opening_role = db.scalar(sm.tenant_query(Role,user).where(Role.role_code==('STUDENT' if kind=='students' else 'INTERN_MENTOR'),Role.status=='ACTIVE'))
        opening_permissions = role_permissions(db,opening_role) if opening_role else []
    bindings={};by_account={};staff_by_account={}
    if kind=='students':
        for link in db.scalars(sm.tenant_query(StudentAccountLink,user).where(StudentAccountLink.link_status=='ACTIVE',or_(
            StudentAccountLink.student_id.in_([p.id for p in people.values()]),StudentAccountLink.user_id.in_([a.id for a in accounts.values()])))):
            bindings[link.student_id]=link;by_account.setdefault(link.user_id,[]).append(link)
    elif kind=='staff':
        # Include archived profiles: their account association is not free to reuse.
        for owner in db.scalars(select(StaffProfile).where(StaffProfile.tenant_id==sm.ids(user)[0],
                StaffProfile.user_id.in_([a.id for a in accounts.values()]))):
            staff_by_account.setdefault(owner.user_id,[]).append(owner)
    seen = set()
    for index, data in enumerate(rows, 2):
        problems, refs, facts = [], {}, []
        if scoped:
            colleges = catalogs['colleges'].get(data.get('collegeCode'),[])
            majors = catalogs['majors'].get(data.get('majorCode'),[])
            classes = catalogs['classes'].get(data.get('classCode'),[])
            person,account = people.get(data.get('number')),accounts.get(data.get('number'))
            within_scope = len(colleges)==1 and (colleges[0].id in scopes['COLLEGE'] or (kind!='staff' and
                ((len(majors)==1 and majors[0].id in scopes['MAJOR']) or (len(classes)==1 and classes[0].id in scopes['CLASS']))))
            if (not within_scope or data.get('departmentCode') or data.get('departmentName') or
                len(majors)>1 or (majors and majors[0].college_id!=colleges[0].id) or len(classes)>1 or
                (classes and (not majors or classes[0].major_id!=majors[0].id)) or
                (person and person.id not in allowed_people) or (account and account.id not in allowed_accounts)):
                message='该行组织、人员或账号不在当前负责范围内，请交由校级管理员核对'
                entries.append({'row':index,'number':data.get('number') or data.get('collegeCode'),'name':data.get('name') or data.get('collegeName'),
                    'operation':'需要核对','accountAction':'不予开通','refs':{},'facts':[],'data':data,'errors':[message]})
                errors.append({'row':index,'message':message})
                continue
        department_only=kind in ('staff','organizations') and not data.get('collegeCode') and bool(data.get('departmentCode'))
        required = ['number','name'] + (['departmentCode','departmentName'] if department_only else ['collegeCode','collegeName']) + ([] if kind == 'staff' else ['majorCode', 'majorName', 'classCode', 'className', 'grade'])
        if kind=='organizations':
            required=[]
            for prefix in ('college','major','class','department'):
                if data.get(prefix+'Code') or data.get(prefix+'Name'): required += [prefix+'Code',prefix+'Name']
            if data.get('classCode'):required+=['majorCode','collegeCode']
            if data.get('majorCode'):required+=['collegeCode']
            if not required:problems.append('至少填写一个组织代码及名称')
        for field in required:
            if not data.get(field): problems.append(f'{next(k for k,v in HEADERS.items() if v == field)}不能为空')
        for field, value in data.items():
            if len(value) > (20 if field == 'grade' else 100): problems.append('字段内容过长')
        if kind!='organizations' and not _NUMBER.fullmatch(data.get('number', '')): problems.append('学号或工号只能包含文字、数字、下划线、点和短横线，最多50字')
        if kind!='organizations' and data.get('number') in seen: problems.append('文件内学号或工号重复')
        seen.add(data.get('number'))
        chain = [('colleges', 'college'), ('majors', 'major'), ('classes', 'class')][:1 if kind == 'staff' else 3]
        if department_only: chain=[]
        if kind=='organizations':chain=[(k,p) for k,p in chain if data.get(p+'Code')]
        if kind in ('staff','organizations') and data.get('departmentCode'): chain.append(('departments','department'))
        parent_id = None
        for org_kind, prefix in chain:
            model, name_col, code_col, parent_col = sm.ORG[org_kind]
            code, name = data.get(prefix+'Code', ''), data.get(prefix+'Name', '')
            if not _NUMBER.fullmatch(code):
                problems.append('组织代码只能包含文字、数字、下划线、点和短横线，最多50字');continue
            candidates = catalogs[org_kind].get(code,[])
            if len(candidates) > 1:
                problems.append(f'{name}代码存在多个候选，须先在组织架构核对'); continue
            if candidates:
                item = candidates[0]
                if getattr(item, name_col) != name:
                    if aliases.get(name)!=code: problems.append(f'{code}名称与现有组织不一致')
                    elif code not in unique_alias_codes: problems.append(f'{name}别名对应组织已不唯一或不可用，请先核对基础字典和组织架构')
                if item.status != 'ACTIVE': problems.append(f'{name}已停用')
                if parent_col and getattr(item, parent_col) != parent_id: problems.append(f'{name}的上级组织不一致')
                refs[prefix] = str(item.id); facts.append([org_kind, str(item.id), item.version])
                parent_id = item.id
            else:
                if not name or (parent_col and not refs.get('college' if org_kind=='majors' else 'major')): continue
                key = org_kind + ':' + code
                spec = {'kind': org_kind, 'code': code, 'name': name,
                        'parentCode': data.get('collegeCode' if org_kind == 'majors' else 'majorCode', '') if parent_col else '',
                        'grade': data.get('grade','') if org_kind=='classes' else ''}
                if key in new_orgs and new_orgs[key] != spec: problems.append(f'{code}在文件中对应多个名称或上级')
                new_orgs[key] = spec; refs[prefix] = key; parent_id = None
        if kind=='organizations':
            entries.append({'row':index,'number':data.get('classCode') or data.get('majorCode') or data.get('collegeCode') or data.get('departmentCode'),
                'name':data.get('className') or data.get('majorName') or data.get('collegeName') or data.get('departmentName'),
                'operation':'新增或复用组织','accountAction':'不创建账号','refs':refs,'facts':facts,'data':data,'errors':problems})
            errors.extend({'row':index,'message':message} for message in problems);continue
        person=people.get(data.get('number'));account=accounts.get(data.get('number'))
        if not account:
            if not opening_role:
                problems.append('开户默认角色未安装或已停用，请联系学校系统管理员')
            elif any(not (has_permission(user,code) or
                (code.startswith('internship.') and has_permission(user,'system.role.grant-business')) or
                (code.startswith('system.') and has_permission(user,'system.role.grant-admin'))) for code in opening_permissions):
                problems.append('新账号默认身份的权限超过当前授权能力，请交由学校系统管理员开通')
        if person:
            facts.append(['person', str(person.id), person.version])
            if person.is_deleted: problems.append('该编号属于已归档人员，请先核对原档案')
            if person.status!='ACTIVE': problems.append('人员已停用，请先通过资料维护核对，导入不能自动启用')
            if person.real_name != data.get('name'): problems.append('已有同编号人员姓名不同，请到人员资料核对，导入不能覆盖身份')
            # Existing affiliations require the personnel impact review, never a bulk overwrite.
            affiliations = ('college', 'major', 'class') if kind == 'students' else ('college', 'department')
            if any(str(getattr(person, prefix + '_id') or '') != str(refs.get(prefix) or '') for prefix in affiliations):
                problems.append('已有人员的组织归属发生变化，请先到人员资料核对在办事项并确认调整，再重新导入')
            if kind == 'students':
                bound = bindings.get(person.id)
                if bound and (not account or bound.user_id!=account.id): problems.append('学生已绑定其他登录账号，请先核对原绑定')
            elif person.user_id and (not account or person.user_id!=account.id):
                problems.append('教职工已绑定其他登录账号，请先到教职工资料核对原绑定；导入不能替换账号')
        if account:
            facts.append(['account', str(account.id), account.version, account.credential_version])
            expected_type = 'STUDENT' if kind == 'students' else 'TEACHER'
            if account.is_deleted or account.status != 'ACTIVE' or account.user_type != expected_type or account.real_name != data.get('name'):
                problems.append('已有同名登录账号身份、姓名或状态不匹配，请先核对')
            if kind == 'students':
                links=by_account.get(account.id,[])
                if any(not person or link.student_id != person.id or link.user_id != account.id for link in links): problems.append('学生账号绑定冲突')
                facts.extend(['binding', str(link.id), link.version] for link in links)
            else:
                owners=staff_by_account.get(account.id,[])
                if any(not person or owner.id!=person.id for owner in owners):
                    problems.append('该登录账号已绑定其他教职工资料，请先核对；不能重复绑定')
                facts.extend(['staffBinding',str(owner.id),owner.version] for owner in owners)
        entries.append({'row': index, 'number': data.get('number'), 'name': data.get('name'),
            'operation': '更新资料' if person else '新增人员',
            'accountAction': '需核对，暂不处理' if problems else ('保留已有账号' if account else '开通账号'),
            'personId': str(person.id) if person else None, 'accountId': str(account.id) if account else None,
            'refs': refs, 'facts': facts, 'data': data, 'errors': problems})
        errors.extend({'row': index, 'message': message} for message in problems)
    return {'kind': kind, 'entries': entries, 'newOrganizations': list(new_orgs.values()), 'errors': errors,
            'roleCode': user['currentRoleCode'], 'permissionVersion': user.get('permissionVersion'), 'total': len(rows),
            'openingRoleVersion': [str(opening_role.id),opening_role.version,sorted(opening_permissions)] if opening_role else None}


def _digest(value):
    return hashlib.sha256(json.dumps(value, sort_keys=True, ensure_ascii=False, separators=(',', ':')).encode()).hexdigest()


def preview(user, kind, filename, content):
    rows = read_xlsx(content, ORG_HEADERS if kind=='organizations' else HEADERS); _text_identifiers(content)
    if not rows: raise AppException('VALIDATION_ERROR', '表格没有可导入的数据')
    with sm.transaction(user, import_permission(kind,'manage')) as db:
        result = _evaluate(db, user, kind, rows)
        job = ExcelImportJob(tenant_id=sm.ids(user)[0], module_key='system', biz_type='IDENTITY_'+kind.upper(),
            file_name=filename[:255], file_sha256=hashlib.sha256(content).hexdigest(), template_version='identity-v1',
            status='FAILED' if result['errors'] else 'VALIDATED', total_rows=len(rows),
            invalid_rows=len({x['row'] for x in result['errors']}), valid_rows=len(rows)-len({x['row'] for x in result['errors']}),
            operator_id=sm.ids(user)[1], operator_name=user.get('realName'), started_at=datetime.utcnow(),
            expires_at=datetime.utcnow()+timedelta(minutes=30), preview_json=result, dry_run_sha256=_digest(result))
        db.add(job); db.flush()
        audit(db, user, 'IDENTITY_PREVIEW', 'import', job.id, {'rows': len(rows), 'invalid': job.invalid_rows, 'sha256': job.file_sha256})
        return public_preview(job)


def public_preview(job, page=1, size=100):
    result = job.preview_json or {}
    entries=result.get('entries', [])[(page-1)*size:page*size]
    refs={ref for entry in entries for ref in entry.get('refs',{}).values()}
    proposals=[r for r in result.get('newOrganizations',[]) if r.get('name') and _NUMBER.fullmatch(r.get('code','')) and
        (r.get('kind') not in ('majors','classes') or _NUMBER.fullmatch(r.get('parentCode','')))]
    return {'id': str(job.id), 'kind': result.get('kind'), 'version': job.version, 'status': job.status,
        'statusLabel': sm.STATUS_LABELS.get(job.status, '待核对'), 'fileName': job.file_name,
        'total': job.total_rows, 'invalidRows': job.invalid_rows, 'expiresAt': job.expires_at,
        'entries': [{k:v for k,v in item.items() if k not in {'facts','data'}} for item in entries],
        'newOrganizations': [r for r in proposals if r['kind']+':'+r['code'] in refs],
        'newOrganizationCount':len(proposals), 'errors': result.get('errors', [])[:100],
        'message':job.remark or '',
        'page': page, 'pageSize': size, 'receiptAvailable': bool(job.receipt_encrypted and job.expires_at > datetime.utcnow())}


def owned_job(db, user, identifier, lock=True):
    job = sm.row(db, ExcelImportJob, identifier, user, lock)
    if job.module_key != 'system' or not job.biz_type.startswith('IDENTITY_') or job.operator_id != sm.ids(user)[1]:
        raise no_permission('仅本批次上传人可确认或领取开户回执')
    snapshot = job.preview_json or {}
    if snapshot.get('roleCode') != user['currentRoleCode'] or snapshot.get('permissionVersion') != user.get('permissionVersion'):
        raise no_permission('当前身份或授权已变化，请使用原身份重新预检；旧批次内容与凭据不可继续领取')
    if not sm.school_wide(user) and snapshot.get('kind')!='organizations':
        student=snapshot.get('kind')=='students'
        people=sm.scoped_students(db,user) if student else sm.scoped_staff(db,user)
        number=StudentProfile.student_no if student else StaffProfile.staff_no
        for query,column,field in ((people,number,'personId'),(sm.scoped_accounts(db,user),User.login_name,'accountId')):
            numbers={entry['number'] for entry in snapshot.get('entries',[]) if entry.get(field) or job.status=='IMPORTED'}
            visible=set(db.scalars(query.with_only_columns(column).where(column.in_(numbers))))
            if numbers!=visible:
                raise no_permission('本批次人员或账号已移出当前负责范围，不能读取原明细、错误行或开户回执')
    return job


def get_preview(user, identifier, page=1):
    with get_sessionmaker()() as db:
        job=owned_job(db,user,identifier,False);enforce_permission(user,import_permission(job.preview_json['kind'],'view'))
        return public_preview(job, page)


def error_rows(user, identifier):
    kind=get_preview(user,identifier)['kind']
    with sm.transaction(user, import_permission(kind,'view')) as db:
        job=owned_job(db,user,identifier)
        headers=ORG_HEADERS if kind=='organizations' else HEADERS
        entries=[entry for entry in (job.preview_json or {}).get('entries',[]) if entry['errors']]
        audit(db,user,'IMPORT_ERRORS_DOWNLOAD','import',job.id,{'rows':len(entries)})
        return build_ledger_xlsx('错误行',list(headers)+['表格行','错误原因'],[
            [entry['data'].get(key,'') for key in headers.values()]+[entry['row'],'；'.join(entry['errors'])]
            for entry in entries])


def enqueue(user, identifier, version):
    preview_data=get_preview(user,identifier)
    if preview_data['total']<=25: return confirm(user,identifier,version)
    with sm.transaction(user,import_permission(preview_data['kind'],'manage')) as db:
        job=owned_job(db,user,identifier)
        if job.status in ('QUEUED','IMPORTED'): return public_preview(job)
        check_version(job.version,version)
        if job.status not in ('VALIDATED','FAILED') or job.invalid_rows or job.expires_at<=datetime.utcnow():
            raise AppException('PREVIEW_EXPIRED','请修正错误或重新预检后再确认',http_status=409)
        current=_evaluate(db,user,preview_data['kind'],[r['data'] for r in job.preview_json['entries']])
        if _digest(current)!=job.dry_run_sha256:
            raise AppException('DATA_CONFLICT','人员、组织或授权已变化，请重新预检后核对',http_status=409)
        job.status='QUEUED';job.version+=1;job.remark=None;job.failed_rows=0
        job.data_scope_snapshot={'roleCode':user['currentRoleCode']}
        audit(db,user,'IDENTITY_QUEUE','import',job.id,{'rows':job.total_rows})
        return public_preview(job)


def confirm(user, identifier, version, queued=False):
    kind=get_preview(user,identifier)['kind']
    with sm.transaction(user, import_permission(kind,'manage')) as db:
        job = owned_job(db, user, identifier)
        if job.status == 'IMPORTED': return public_preview(job)
        check_version(job.version, version)
        if job.status != ('QUEUED' if queued else 'VALIDATED') or job.invalid_rows or job.expires_at <= datetime.utcnow():
            raise AppException('PREVIEW_EXPIRED', '预检未通过或已过期，请修正后重新上传', http_status=409)
        preview_data = job.preview_json
        current = _evaluate(db, user, preview_data['kind'], [entry['data'] for entry in preview_data['entries']])
        if _digest(current) != job.dry_run_sha256:
            raise AppException('DATA_CONFLICT', '人员、组织或授权已变化，请重新预检后核对', http_status=409)
        org_ids = {}
        for kind in ('colleges','majors','classes','departments'):
            for spec in (s for s in current['newOrganizations'] if s['kind'] == kind):
                model, name, code, parent = sm.ORG[kind]
                item = model(tenant_id=sm.ids(user)[0], status='ACTIVE', **{name: spec['name'], code: spec['code']})
                if kind=='classes':item.grade=spec.get('grade','')
                if parent:
                    parent_kind = 'colleges' if kind == 'majors' else 'majors'
                    parent_key = parent_kind+':'+spec['parentCode']
                    if parent_key in org_ids: parent_id = org_ids[parent_key]
                    else:
                        parent_model, _, parent_code, _ = sm.ORG[parent_kind]
                        parent_id = db.scalar(sm.tenant_query(parent_model, user).where(getattr(parent_model, parent_code) == spec['parentCode'])).id
                    setattr(item, parent, parent_id)
                db.add(item); db.flush(); org_ids[kind+':'+spec['code']] = item.id
        if current['kind']=='organizations':
            job.status='IMPORTED';job.success_rows=job.total_rows;job.confirm_at=job.finished_at=datetime.utcnow();job.version+=1
            audit(db,user,'ORGANIZATION_IMPORT','import',job.id,{'rows':job.total_rows,'created':len(org_ids),'sha256':job.file_sha256})
            return public_preview(job)
        role_code = 'STUDENT' if current['kind'] == 'students' else 'INTERN_MENTOR'
        role = db.scalar(sm.tenant_query(Role, user).where(Role.role_code == role_code, Role.status == 'ACTIVE'))
        if not role and any(not entry['accountId'] for entry in current['entries']):
            raise AppException('ROLE_NOT_READY', '初始角色未安装或已停用，无法开通账号', http_status=409)
        receipt = []
        for entry in current['entries']:
            data = entry['data']; password = None
            if entry['accountId']: account = sm.row(db, User, entry['accountId'], user)
            else:
                password = 'Aa9!' + secrets.token_urlsafe(15)
                account = User(tenant_id=sm.ids(user)[0], login_name=data['number'], real_name=data['name'],
                    user_type='STUDENT' if current['kind']=='students' else 'TEACHER', status='ACTIVE',
                    password_hash=hash_password(password), must_change_password=True, credential_version=0)
                db.add(account); db.flush()
                db.add(UserRole(tenant_id=sm.ids(user)[0], user_id=account.id, role_id=role.id, status='ACTIVE'))
            model = StudentProfile if current['kind']=='students' else StaffProfile
            person = sm.row(db, model, entry['personId'], user) if entry['personId'] else model(tenant_id=sm.ids(user)[0])
            person.real_name = data['name']; person.status = 'ACTIVE'; person.version = (person.version or 0) + 1
            def ref(key):
                value = entry['refs'].get(key); return (int(value) if value.isdigit() else org_ids[value]) if value else None
            person.college_id = ref('college')
            if current['kind']=='students':
                person.student_no, person.major_id, person.class_id, person.grade = data['number'], ref('major'), ref('class'), data['grade']
            else: person.staff_no, person.user_id, person.department_id = data['number'], account.id, ref('department')
            db.add(person); db.flush()
            if current['kind']=='students':
                bind_in_session(db, tenant_id=sm.ids(user)[0], student_id=person.id, user_id=account.id,
                    source='IDENTITY_IMPORT', login_name=account.login_name, student_no=person.student_no)
            receipt.append([data['number'], data['name'], account.login_name, password or '已有账号，密码保持不变'])
        job.receipt_encrypted = encrypt_field(json.dumps(receipt, ensure_ascii=False))
        job.status, job.success_rows, job.confirm_at, job.finished_at = 'IMPORTED', len(receipt), datetime.utcnow(), datetime.utcnow()
        job.expires_at = datetime.utcnow()+timedelta(hours=24); job.version += 1
        audit(db, user, 'IDENTITY_CONFIRM', 'import', job.id, {'rows': len(receipt), 'sha256': job.file_sha256})
        return public_preview(job)


def receipt(user, identifier):
    with sm.transaction(user, 'system.identity-import.manage') as db:
        job = owned_job(db, user, identifier)
        if not job.receipt_encrypted or job.expires_at <= datetime.utcnow():
            raise AppException('RECEIPT_EXPIRED', '开户回执已过期，后续需要由授权人员受控重置', http_status=410)
        values = json.loads(decrypt_field(job.receipt_encrypted, allow_legacy_plaintext=False))
        audit(db, user, 'IDENTITY_RECEIPT_DOWNLOAD', 'import', job.id, {'rows': len(values)})
        return build_ledger_xlsx('开户回执', ['编号','姓名','登录账号','临时密码'], values,
                                watermark='仅供学校管理员受控交付；首次登录必须修改密码，请勿公开转发')
