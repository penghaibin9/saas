"""Normal-role payroll writes and same-filter exports on a named isolated MySQL target."""
import base64
from concurrent.futures import ThreadPoolExecutor
from threading import Barrier
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import get_sessionmaker
from app.models import College, Major, SchoolClass, StudentProfile, InternshipRecord, Tenant, User, UserRole, Role, TeacherStudentScope, InternshipAuditTrail
from test_student_support_story_mysql import seed_support, context, result
from test_yiyang_plan_export_options_mysql import login
from test_yiyang_write_flows_mysql import png_bytes

ROOT = '/api/v1/internship/payroll'
STUDENT = '/api/v1/mobile/internship/context/payroll'
TEACHER = '/api/v1/mobile/teacher/internship/context/payroll'


def seed_payroll():
    story = seed_support()
    tid = int(story['tenantId'])
    with get_sessionmaker()() as db:
        db.get(Tenant, tid).school_name = '工资全链合成验收学校'
        college = College(tenant_id=tid, college_name='工资验收院系', code='PAY', status='ACTIVE')
        db.add(college); db.flush()
        major = Major(tenant_id=tid, college_id=college.id, major_name='软件技术', status='ACTIVE')
        db.add(major); db.flush()
        group = SchoolClass(tenant_id=tid, major_id=major.id, class_name='工资验收班', status='ACTIVE')
        db.add(group); db.flush()
        student = db.get(StudentProfile, int(story['studentId']))
        student.real_name, student.student_no = '=工资%验收', '000000000000000000123'
        student.college_id, student.major_id, student.class_id = college.id, major.id, group.id
        record = db.get(InternshipRecord, int(story['recordId']))
        record.enterprise_name, record.position_name = '工资验收合成企业', '软件实习生'
        for code, name in [('STUDENT','工资验收学生'),('INTERN_MENTOR','工资验收导师'),('SCHOOL_ADMIN','工资验收学校管理员')]:
            db.get(User, story['users'][code]).real_name = name
        for code, role_code in [('OTHER_MENTOR','INTERN_MENTOR'),('COLLEGE_ADMIN','COLLEGE_ADMIN')]:
            user = User(tenant_id=tid, login_name=code.lower(), real_name='工资验收'+('其他导师' if code=='OTHER_MENTOR' else '院系管理员'),
                password_hash=db.get(User, story['users']['INTERN_MENTOR']).password_hash,
                user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
            db.add(user); db.flush()
            role = db.scalar(select(Role).where(Role.tenant_id==tid, Role.role_code==role_code))
            db.add(UserRole(tenant_id=tid, user_id=user.id, role_id=role.id, status='ACTIVE'))
            story['users'][code] = user.id
            if code=='COLLEGE_ADMIN':
                db.add(TeacherStudentScope(tenant_id=tid, user_id=user.id, teacher_key=code.lower(),
                    role_code=code, scope_type='COLLEGE', ref_id=college.id, status='ACTIVE'))
        db.commit()
        story['collegeId'] = str(college.id)
    return story


def upload_photo(client, story):
    from app.services.file_scan_service import process_next_scan_job
    meta = result(client.post('/api/v1/files', params={'bizType':'TEMP_PRIVATE'},
        files={'file':('工资验收凭证.png', png_bytes(), 'image/png')}))
    if not meta['readyForBusiness']:
        process_next_scan_job(tenant_id=int(story['tenantId']), file_id=int(meta['fileId']))
        meta = result(client.get('/api/v1/files/'+meta['fileId']))
    assert meta['readyForBusiness'], meta.get('scanStatus')
    return meta


def submit(client, story, month='2026-10', amount='1234.56', **extra):
    photo = upload_photo(client, story)
    return result(client.post(STUDENT, json={**context(story), 'payMonth':month, 'actualAmount':amount,
        'currency':'CNY', 'paidOn':'2026-10-14', 'evidenceFileId':photo['fileId'], **extra}))


def sheet_rows(response):
    payload = result(response)
    sheet = load_workbook(BytesIO(base64.b64decode(payload['contentBase64']))).active
    headers = [cell.value for cell in sheet[2]]
    rows = [dict(zip(headers,row)) for row in sheet.iter_rows(min_row=3, values_only=True)]
    assert len(rows)==payload['rowCount']
    return rows


@pytest.fixture
def story(record_property):
    story = seed_payroll()
    record_property('tenantId',story['tenantId'])
    yield story


def test_student_return_correct_zero_approve_same_statement_and_history(story):
    with login(story,'STUDENT') as student, login(story,'INTERN_MENTOR') as teacher, login(story) as school:
        first = submit(student, story)
        queue = result(teacher.get(TEACHER,params={'batchId':story['batchId'],'status':'SUBMITTED'}))
        assert queue['items'][0]['id']==first['id']
        assert teacher.get('/api/v1/files/download/'+first['current']['evidenceFileId']).content==png_bytes()
        returned = result(teacher.post(TEACHER+'/versions/'+first['current']['id']+'/review',json={
            'action':'RETURN','comment':'请更正为实际到账金额','expectedVersion':first['current']['version']}))
        photo = upload_photo(student,story)
        body = {**context(story),'payMonth':'2026-10','actualAmount':'0','currency':'CNY',
            'evidenceFileId':photo['fileId'],'correctionReason':'本月实际尚未收到工资','expectedVersion':first['version']}
        stale = student.post(STUDENT,json=body)
        assert stale.status_code==409 and stale.json()['bizCode']=='DATA_CONFLICT'
        body['expectedVersion']=returned['version']
        corrected = result(student.post(STUDENT,json=body))
        assert corrected['id']==first['id'] and corrected['revisionCount']==2
        assert corrected['current']['actualAmount']==0
        approved = result(teacher.post(TEACHER+'/versions/'+corrected['current']['id']+'/review',json={
            'action':'APPROVE','comment':'核对实际到账为零','expectedVersion':corrected['current']['version']}))
        assert result(student.get(STUDENT,params=context(story)))[0]['current']['status']=='APPROVED'
        assert len(result(school.get(ROOT+'/'+first['id']))['history'])==2
        rows = sheet_rows(school.post(ROOT+'/export',params={'batchId':story['batchId'],'payMonth':'2026-10','status':'APPROVED'}))
        assert len(rows)==1 and rows[0]['工资单编号']==first['id'] and rows[0]['实发金额']==0
        assert rows[0]['学号']=='000000000000000000123' and rows[0]['姓名']=="'=工资%验收"
        assert rows[0]['院系']=='工资验收院系' and rows[0]['班级']=='工资验收班'
        assert rows[0]['版本']==2 and rows[0]['审核人']=='工资验收导师'
        assert rows[0]['币种']=='人民币'
        duplicate = teacher.post(TEACHER+'/versions/'+corrected['current']['id']+'/review',json={
            'action':'APPROVE','expectedVersion':corrected['current']['version']})
        assert duplicate.status_code==409
        assert approved['current']['status']=='APPROVED'

@pytest.mark.parametrize('currency,label',[('USD','美元'),('AUD','AUD（其他币种）')])
def test_export_currency_names_preserve_original_contract_and_zero_amount(story,currency,label):
    with login(story,'STUDENT') as student,login(story) as school:
        item=submit(student,story,currency=currency,actualAmount='0')
        assert item['current']['currency']==currency
        rows=sheet_rows(school.post(ROOT+'/export',params={'batchId':story['batchId'],'payMonth':'2026-10'}))
        assert len(rows)==1 and rows[0]['币种']==label and rows[0]['实发金额']==0


def test_month_filters_pagination_literal_keyword_scope_export_and_audit(story):
    with login(story,'STUDENT') as student:
        ids = [submit(student,story,month)['id'] for month in ('2026-09','2026-10','2026-11')]
    for role in ['SCHOOL_ADMIN','INTERN_MENTOR','COLLEGE_ADMIN']:
        with login(story,role) as reader:
            filters = {'batchId':story['batchId'],'keyword':'%','collegeId':story['collegeId']}
            found = [result(reader.get(ROOT,params={**filters,'page':p,'pageSize':1}))['items'][0]['id'] for p in (1,2,3)]
            assert set(found)==set(ids)
            assert {row['工资单编号'] for row in sheet_rows(reader.post(ROOT+'/export',params=filters))}==set(found)
            one = sheet_rows(reader.post(ROOT+'/export',params={**filters,'payMonth':'2026-10'}))
            assert len(one)==1 and one[0]['工资月份']=='2026-10'
            assert not sheet_rows(reader.post(ROOT+'/export',params={**filters,'keyword':'不存在'}))
    with get_sessionmaker()() as db:
        audits=db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id==int(story['tenantId']),
            InternshipAuditTrail.target_type=='EXPORT')).all()
        assert any(x.detail_json['ledger']=='月度工资台账' and x.detail_json['totalRows']==3 for x in audits)


def test_other_teacher_student_and_other_school_cannot_read_or_export(story):
    with login(story,'STUDENT') as student:
        row=submit(student,story)
        assert student.get(ROOT,params={'batchId':story['batchId']}).status_code==403
        assert student.post(ROOT+'/export',params={'batchId':story['batchId']}).status_code==403
    other=seed_payroll()
    for who,role in [(story,'OTHER_MENTOR'),(other,'SCHOOL_ADMIN')]:
        with login(who,role) as client:
            assert result(client.get(ROOT,params={'batchId':story['batchId']}))['total']==0
            assert not sheet_rows(client.post(ROOT+'/export',params={'batchId':story['batchId']}))
            assert client.get(ROOT+'/'+row['id']).status_code in (403,404)
            assert client.get('/api/v1/files/download/'+row['current']['evidenceFileId']).status_code in (403,404)
            assert client.post(TEACHER+'/versions/'+row['current']['id']+'/review',json={
                'action':'APPROVE','expectedVersion':row['current']['version']}).status_code in (403,404)


@pytest.mark.parametrize('amount',['NaN','Infinity','-1','10000000000'])
def test_invalid_amount_never_creates_statement(story,amount):
    with login(story,'STUDENT') as student:
        response=student.post(STUDENT,json={**context(story),'payMonth':'2026-10','actualAmount':amount})
        assert response.status_code==400 and response.json()['bizCode']=='VALIDATION_ERROR'
        assert result(student.get(STUDENT,params=context(story)))==[]


def test_two_independent_reviewers_only_one_audit_and_no_overwrite(story):
    with login(story,'STUDENT') as student: row=submit(student,story)
    def review(role):
        with login(story,role) as client:
            return client.post(TEACHER+'/versions/'+row['current']['id']+'/review',json={
                'action':'APPROVE','expectedVersion':row['current']['version']}).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(review,['INTERN_MENTOR','SCHOOL_ADMIN']))==[200,409]
    with get_sessionmaker()() as db:
        audits=db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id==int(story['tenantId']),
            InternshipAuditTrail.target_type=='PAYROLL',InternshipAuditTrail.action=='PAYROLL_APPROVE')).all()
        assert len(audits)==1


def test_concurrent_corrections_with_earlier_read_snapshot_do_not_replace_pending_version(story,monkeypatch):
    from app.models import InternshipPayrollStatement
    from app.modules.internship.services import internship_payroll_service as payrolls
    with login(story,'STUDENT') as student, login(story,'INTERN_MENTOR') as teacher:
        first=submit(student,story)
        returned=result(teacher.post(TEACHER+'/versions/'+first['current']['id']+'/review',json={
            'action':'RETURN','expectedVersion':first['current']['version'],'comment':'请更正本月工资金额'}))
        photos=[upload_photo(student,story)['fileId'] for _ in range(2)]
    original=payrolls.require_explicit_context
    barrier=Barrier(2)
    def earlier_snapshot(db,user,payload,*,for_write):
        db.scalar(select(InternshipPayrollStatement.id).where(InternshipPayrollStatement.id==int(first['id'])))
        barrier.wait(timeout=15)
        return original(db,user,payload,for_write=for_write)
    monkeypatch.setattr(payrolls,'require_explicit_context',earlier_snapshot)
    def correct(index):
        with login(story,'STUDENT') as student:
            return student.post(STUDENT,json={**context(story),'payMonth':'2026-10','actualAmount':str(index),
                'evidenceFileId':photos[index],'expectedVersion':returned['version'],'correctionReason':'核对实际工资金额后更正'}).status_code
    with ThreadPoolExecutor(max_workers=2) as pool:
        assert sorted(pool.map(correct,[0,1]))==[200,409]
    with login(story) as school:
        row=result(school.get(ROOT+'/'+first['id']))
        assert row['revisionCount']==2 and row['current']['status']=='SUBMITTED' and len(row['history'])==2
