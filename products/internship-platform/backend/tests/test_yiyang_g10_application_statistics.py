"""Yiyang G10: internship application coverage/review statistics must use full scoped data."""

from __future__ import annotations

import pytest

TID = 1000000000000000001
USER = {
    "tenantId": str(TID),
    "userId": "90001",
    "realName": "G10验收管理员",
    "userType": "SCHOOL_ADMIN",
    "currentRoleCode": "SCHOOL_ADMIN",
}


@pytest.fixture()
def standalone_db(tmp_path):
    import app.db.session as db_session
    from app.config import settings
    from app.models import (
        College,
        EmpCompany,
        InternshipApplication,
        InternshipBatch,
        InternshipPosition,
        InternshipRecord,
        Major,
        SchoolClass,
        StudentProfile,
    )

    old_url = settings.DATABASE_URL
    settings.DATABASE_URL = f"sqlite+pysqlite:///{(tmp_path / 'yiyang-g10.db').as_posix()}"
    db_session._engine = None
    db_session._factory = None

    engine = db_session.get_engine()
    for table in (
        College.__table__,
        Major.__table__,
        SchoolClass.__table__,
        EmpCompany.__table__,
        InternshipBatch.__table__,
        StudentProfile.__table__,
        InternshipRecord.__table__,
        InternshipPosition.__table__,
        InternshipApplication.__table__,
    ):
        table.create(bind=engine, checkfirst=True)

    try:
        yield
    finally:
        engine.dispose()
        settings.DATABASE_URL = old_url
        db_session._engine = None
        db_session._factory = None


def _seed_g10():
    from app.db.session import get_sessionmaker
    from app.models import (
        College, InternshipApplication, InternshipBatch, InternshipRecord,
        Major, SchoolClass, StudentProfile,
    )

    db = get_sessionmaker()()
    try:
        batch = InternshipBatch(
            tenant_id=TID,
            batch_name="益阳G10统计验收批次",
            batch_no="YIYANG-G10",
            planned_count=65,
            status="RUNNING",
        )
        db.add(batch)
        db.flush()

        college = College(
            tenant_id=TID, college_name="信息工程学院", status="ACTIVE")
        db.add(college)
        db.flush()
        major = Major(
            tenant_id=TID, college_id=college.id,
            major_name="软件技术", status="ACTIVE")
        db.add(major)
        db.flush()
        school_class = SchoolClass(
            tenant_id=TID, major_id=major.id,
            class_name="软件技术2401", grade="2024",
            status="ACTIVE", class_status="NORMAL")
        db.add(school_class)
        db.flush()

        for index in range(65):
            student = StudentProfile(
                tenant_id=TID,
                student_no=f"G10-{index + 1:03d}",
                real_name=f"G10学生{index + 1:03d}",
                current_stage="INTERNSHIP",
                student_status="NORMAL",
                status="ACTIVE",
                college_id=college.id,
                major_id=major.id,
                class_id=school_class.id,
            )
            db.add(student)
            db.flush()

            record = InternshipRecord(
                tenant_id=TID,
                student_id=student.id,
                batch_id=batch.id,
                advisor_name="G10指导老师",
                eligibility_status="QUALIFIED",
                destination_type="NONE",
                status="PREPARING",
                risk_level="NONE",
            )
            db.add(record)
            db.flush()

            if index < 45:
                status = "PENDING_REVIEW" if index < 30 else ("APPROVED" if index < 40 else "REJECTED")
                db.add(InternshipApplication(
                    tenant_id=TID,
                    record_id=record.id,
                    student_id=student.id,
                    batch_id=batch.id,
                    campaign_id=None,
                    application_type="SELF_ARRANGED",
                    volunteer_no=0,
                    company_name=f"G10企业{index + 1:03d}",
                    position_name="测试岗位",
                    work_address="湖南省益阳市赫山区",
                    internship_department="研发部",
                    position_category="技术类",
                    agreed_salary=3500,
                    company_industry="软件和信息技术服务业",
                    enterprise_mentor_name="企业导师张老师",
                    contact_name="企业联系人",
                    contact_phone="13800138000",
                    status=status,
                ))
        db.commit()
        return batch.id
    finally:
        db.close()


def test_g10_full_scope_counts_and_drilldown_are_not_page_length(standalone_db):
    from app.core.context import set_current_user, set_tenant
    from app.modules.internship.services import internship_application_service as svc

    set_tenant({"tenantId": str(TID)})
    set_current_user(USER)
    batch_id = _seed_g10()

    pending_items, pending_total = svc.list_applications(
        1, 20, status="PENDING_REVIEW", batch_id=batch_id, user=USER
    )
    assert pending_total == 30
    assert len(pending_items) == 20

    summary = svc.application_summary(batch_id=batch_id, user=USER)
    assert summary["totalStudents"] == 65
    assert summary["filledStudents"] == 45
    assert summary["unfilledStudents"] == 20
    assert summary["filledRate"] == 69.2
    assert summary["pendingReviewApplications"] == 30
    assert summary["reviewedApplications"] == 15
    assert summary["approvedApplications"] == 10
    assert summary["rejectedApplications"] == 5

    reviewed_items, reviewed_total = svc.list_applications(
        1, 20, status="REVIEWED", batch_id=batch_id, user=USER
    )
    assert reviewed_total == 15
    assert len(reviewed_items) == 15

    unfilled_items, unfilled_total = svc.list_application_students(
        1, 100, "UNFILLED", batch_id=batch_id, user=USER
    )
    assert unfilled_total == 20
    assert len(unfilled_items) == 20
    assert all(item["applicationState"] == "UNFILLED" for item in unfilled_items)

    filled_items, filled_total = svc.list_application_students(
        1, 20, "FILLED", batch_id=batch_id, user=USER
    )
    assert filled_total == 45
    assert len(filled_items) == 20
    assert all(item["applicationState"] == "FILLED" for item in filled_items)


def test_g10_ap05_export_contains_procurement_fields(standalone_db):
    from io import BytesIO

    from openpyxl import load_workbook

    from app.core.context import set_current_user, set_tenant
    from app.modules.internship.services import internship_application_service as svc

    set_tenant({"tenantId": str(TID)})
    set_current_user(USER)
    batch_id = _seed_g10()

    payload, filename, row_count = svc.export_applications(
        status="REVIEWED", batch_id=batch_id, user=USER)
    assert filename.endswith(".xlsx")
    assert row_count == 15

    workbook = load_workbook(BytesIO(payload), read_only=True, data_only=True)
    sheet = workbook.active
    rows = list(sheet.iter_rows(values_only=True))
    assert rows
    required_headers = (
        "申请时间", "学号", "姓名", "院系", "实习单位/免实习",
        "实习岗位/免实习去向", "实习单位地址", "所属科室",
        "职位类别", "实习薪资(元/月)", "所属行业", "校内指导老师",
        "企业老师", "状态", "免实习原因", "免实习佐证",
    )
    header_index = next(
        (
            index for index, row in enumerate(rows)
            if set(required_headers) <= {str(value or "") for value in row}
        ),
        None,
    )
    assert header_index is not None, rows[:5]
    headers = [str(value or "") for value in rows[header_index]]
    for required in required_headers:
        assert required in headers

    data = [
        dict(zip(headers, row))
        for row in rows[header_index + 1:]
        if any(value is not None for value in row)
    ]
    assert data
    first = data[0]
    assert first["院系"] == "信息工程学院"
    assert first["所属行业"] == "软件和信息技术服务业"
    assert first["企业老师"] == "企业导师张老师"
