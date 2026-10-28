"""Package 5 normal-role chain on new synthetic schools; existing rows are retained."""
from datetime import datetime
from base64 import b64decode
from io import BytesIO
import os
import pytest
from sqlalchemy import select
from openpyxl import Workbook,load_workbook
from app.db.session import get_sessionmaker
from app.models import (College,Major,SchoolClass,StaffProfile,StudentProfile,StudentAccountLink,
 InternshipBatchPlan,InternshipRecord,InternshipBatch,EmpCompany,InternshipEnterpriseContact,
 InternshipEnterpriseCollegeScope,User,Tenant,Role,UserRole,TeacherStudentScope,InternshipPlanAck)
from test_yiyang_plan_export_options_mysql import seed_workspace,login,PASSWORD

ROOT='/api/v1/internship'
PORTAL='/api/v1/portal/internship/context'

def ok(response):
    assert response.status_code==200 and response.json()['code']==0,response.text
    return response.json()['data']

def seed_plan_story():
    ws=seed_workspace(2);tid=int(ws['tenantId'])
    with get_sessionmaker()() as db:
        from app.services.system_authorization import bootstrap_school
        bootstrap_school(db,tid)
        db.get(Tenant,tid).school_name='计划岗位分配合成验收学校'
        college=College(tenant_id=tid,college_name='合成信息学院');other=College(tenant_id=tid,college_name='合成机电学院')
        db.add_all([college,other]);db.flush()
        major=Major(tenant_id=tid,college_id=college.id,major_name='软件技术');db.add(major);db.flush()
        group=SchoolClass(tenant_id=tid,major_id=major.id,class_name='计划验收一班',grade='2024');db.add(group);db.flush()
        db.add(StaffProfile(tenant_id=tid,user_id=ws['users']['INTERN_MENTOR'],staff_no='000000000000000051',real_name='计划验收导师',college_id=college.id,status='ACTIVE'))
        db.get(User,ws['users']['INTERN_MENTOR']).real_name='计划验收导师'
        for code,base in [('COLLEGE_ADMIN','COLLEGE_ADMIN'),('OTHER_MENTOR','INTERN_MENTOR')]:
            actor=User(tenant_id=tid,login_name=code.lower(),real_name='计划验收院系管理员' if code=='COLLEGE_ADMIN' else '外院验收导师',user_type='TEACHER',status='ACTIVE',credential_version=0,must_change_password=False,password_hash=db.get(User,ws['users']['INTERN_MENTOR']).password_hash)
            db.add(actor);db.flush();ws['users'][code]=actor.id
            role=db.scalar(select(Role).where(Role.tenant_id==tid,Role.role_code==base))
            db.add(UserRole(tenant_id=tid,user_id=actor.id,role_id=role.id,status='ACTIVE'))
        db.add(StaffProfile(tenant_id=tid,user_id=ws['users']['OTHER_MENTOR'],staff_no='000000000000000052',real_name='外院验收导师',college_id=other.id,status='ACTIVE'))
        db.add(TeacherStudentScope(tenant_id=tid,user_id=ws['users']['COLLEGE_ADMIN'],teacher_key=str(ws['users']['COLLEGE_ADMIN']),role_code='COLLEGE_ADMIN',scope_type='COLLEGE',ref_id=college.id,status='ACTIVE'))
        for entry in ws['batches']:
            plan=db.get(InternshipBatchPlan,int(entry['planId']));plan.status='DRAFT';plan.published_at=None
            plan.tasks_json=[{'sortOrder':1,'name':'岗位实践记录','requirement':'说明本次岗位实践与安全要求','deadline':'2026-09-30'}]
            batch=db.get(InternshipBatch,int(entry['id']));batch.rules_config={'weeklyReport':{'requiredCount':0,'minWordCount':0},'checkin':{'requiredDays':0,'requireDaily':False}}
        rec=db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id==int(ws['batches'][0]['id'])))
        stu=db.get(StudentProfile,rec.student_id);stu.student_no='000000000000000005001';stu.real_name='计划验收学生';stu.college_id=college.id;stu.major_id=major.id;stu.class_id=group.id
        from test_yiyang_gap09_minimum_age import encrypted_student
        stu.id_card_encrypted=encrypted_student(datetime(2005,1,1).date()).id_card_encrypted
        rec.advisor_user_id=None;rec.advisor_name=None;rec.status='READY'
        other_rec=db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id==int(ws['batches'][1]['id'])))
        db.get(StudentProfile,other_rec.student_id).college_id=other.id
        db.add(StudentAccountLink(tenant_id=tid,student_id=stu.id,user_id=ws['users']['STUDENT'],link_status='ACTIVE',source='IDENTITY_IMPORT',bound_student_no=stu.student_no,bound_login_name='student'))
        db.commit();ws.update(batchId=str(rec.batch_id),recordId=str(rec.id),studentId=str(stu.id),studentNo=stu.student_no,collegeName=college.college_name,collegeId=str(college.id))
    return ws

@pytest.fixture
def story():
    if os.environ.get('GAP09_MYSQL_ACCEPTANCE')!='1':pytest.skip('requires isolated MySQL opt-in')
    return seed_plan_story()

def company_row(ws):
    return {'name':'合成合作企业有限公司','creditCode':'91430000'+ws['tenantCode'][-10:].upper(),
      'nature':'民营企业','industry':'软件','region':'湖南省','city':'益阳市','contactPerson':'合成联系人',
      'contactPhone':'13800000000','contactEmail':'contact@example.com','collegeNames':ws['collegeName']}

def publish(client,ws,index=0):
    batch=ws['batches'][index]['id'];view=ok(client.get(ROOT+'/plans/batch/'+batch))
    saved=ok(client.put(ROOT+'/plans/batch/'+batch,json={**view,'expectedVersion':view['version']}))
    return ok(client.post(ROOT+'/plans/batch/'+batch+'/publish',json={'expectedVersion':saved['version']}))

def context(ws):return {'batchId':ws['batchId'],'internshipId':ws['recordId']}

def test_enterprise_template_and_real_upload_keep_complete_procurement_fields(story):
    from app.modules.internship.services import internship_enterprise_service as svc
    row=company_row(story);spec=svc.build_import_spec()
    with login(story) as school:
        template=school.get(ROOT+'/enterprises/import/template');assert template.status_code==200
        wb=load_workbook(BytesIO(template.content));assert {'企业性质','城市','联系人邮箱','适用学院'}.issubset({cell.value for cell in wb['导入模板'][1]})
        file=Workbook();sheet=file.active;sheet.title='导入模板';sheet.append(spec.headers);sheet.append(spec.row_values(row));out=BytesIO();file.save(out)
        pre=ok(school.post(ROOT+'/enterprises/import/xlsx',files={'file':('企业导入.xlsx',out.getvalue(),'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}))
        assert pre['passed'] and pre['rows'][0]['city']=='益阳市'
        assert ok(school.post(ROOT+'/enterprises/import/confirm',json={'rows':pre['rows']}))['created']==1
        company=ok(school.get(ROOT+'/enterprises',params={'keyword':row['name']}))['items'][0]
        detail=ok(school.get(ROOT+'/enterprises/'+company['id']))
        assert (detail['nature'],detail['region'],detail['city'])==('民营企业','湖南省','益阳市')
        assert detail['collegeScopeIds']==[story['collegeId']] and detail['coopStatus']=='PENDING'
        assert detail['contacts'][0]['email']=='contact@example.com'
        exported=ok(school.post(ROOT+'/enterprises/export',params={'keyword':row['name']}))
        sheet=load_workbook(BytesIO(b64decode(exported['contentBase64']))).active
        assert {'企业性质','城市','联系人邮箱','适用学院'}.issubset({c.value for c in sheet[2]})
        assert 'contact@example.com' in [c.value for c in sheet[3]]
    with get_sessionmaker()() as db:
        contact=db.scalar(select(InternshipEnterpriseContact).where(InternshipEnterpriseContact.company_id==int(company['id'])))
        assert contact.phone_encrypted and contact.phone_encrypted!='13800000000'

@pytest.mark.parametrize('change',[{'contactEmail':'bad-email'},{'collegeNames':'不存在学院'},{'contactEmail':'a@example.com','contactPerson':''},{'nature':'长'*51}])
def test_enterprise_invalid_rows_reject_confirm_without_partial_write(story,change):
    row={**company_row(story),**change}
    with login(story) as school:
        pre=ok(school.post(ROOT+'/enterprises/import/dry-run',json={'rows':[row]}));assert not pre['passed'] and pre['invalidRows']==1
        assert school.post(ROOT+'/enterprises/import/confirm',json={'rows':[company_row(story),row]}).json()['code']!=0
        errors=ok(school.post(ROOT+'/enterprises/import/errors-xlsx',json={'rows':[row],'errors':pre['errors']}))
        assert b64decode(errors['contentBase64']).startswith(b'PK')
    with get_sessionmaker()() as db:
        assert not db.scalar(select(EmpCompany.id).where(EmpCompany.tenant_id==int(story['tenantId'])))

def test_published_plan_assignment_ack_task_return_and_resubmit_share_one_record(story):
    with login(story) as school,login(story,'INTERN_MENTOR') as teacher,login(story,'STUDENT') as student:
        plan=publish(school,story)
        rec=ok(school.get(ROOT+'/intern-students/'+story['recordId']))
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/advisor',json={'advisorUserId':str(story['users']['INTERN_MENTOR']),'expectedVersion':rec['version']}))
        detail=ok(student.get(PORTAL+'/plan',params={**context(story),'planId':plan['id']}))
        assert detail['content']==plan['content'] and detail['version']==plan['version']
        ack={**context(story),'planId':plan['id'],'planVersion':plan['version'],'expectedVersion':detail['ackVersion']}
        ok(student.post(PORTAL+'/plan/acknowledge',json=ack));assert student.post(PORTAL+'/plan/acknowledge',json=ack).json()['code']!=0
        tasks=ok(student.get(PORTAL+'/plan/tasks',params={**context(story),'planId':plan['id']}))
        body={**context(story),'planId':plan['id'],'expectedVersion':tasks['tasks'][0]['progressVersion'],'studentNote':'完成岗位实践并核对企业岗位劳动安全要求。'}
        submitted=ok(student.post(PORTAL+'/plan/tasks/1/submit',json=body))
        returned=ok(teacher.post(ROOT+'/plan-task-progress/'+submitted['id']+'/review',json={'action':'REJECT','comment':'请补充实践成果和安全核对过程','expectedVersion':submitted['version']}))
        assert student.post(PORTAL+'/plan/tasks/1/submit',json=body).json()['code']!=0
        body.update(expectedVersion=returned['version'],studentNote=body['studentNote']+'已按导师意见补充成果记录。')
        current=ok(student.post(PORTAL+'/plan/tasks/1/submit',json=body));assert current['id']==submitted['id']
        ok(teacher.post(ROOT+'/plan-task-progress/'+current['id']+'/review',json={'action':'APPROVE','expectedVersion':current['version']}))
        assert ok(student.get(PORTAL+'/plan/tasks',params={**context(story),'planId':plan['id']}))['summary']['approved']==1

def test_multiple_plans_have_independent_confirmation_and_removal_guards(story):
    with login(story) as school,login(story,'STUDENT') as student:
        primary=publish(school,story);extra=publish(school,story,1)
        assigned=ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/plan-assignments',json={'planIds':[extra['id']]}))
        assert assigned['added']==1
        rows=ok(student.get(PORTAL+'/plans',params=context(story)))['items'];assert {x['planId'] for x in rows}=={primary['id'],extra['id']}
        detail=ok(student.get(PORTAL+'/plan',params={**context(story),'planId':extra['id']}))
        ok(student.post(PORTAL+'/plan/acknowledge',json={**context(story),'planId':extra['id'],'expectedVersion':detail['ackVersion'],'planVersion':detail['version']}))
        primary_detail=ok(student.get(PORTAL+'/plan',params={**context(story),'planId':primary['id']}));assert primary_detail['ackStatus']=='PENDING'
        item=assigned['items'][0]
        assert school.post(ROOT+'/intern-students/'+story['recordId']+'/plan-assignments/'+item['id']+'/remove',json={'reason':'合成移除测试','expectedVersion':item['version']}).json()['code']!=0
        for suffix in ['pdf','xlsx']:
            result=ok(school.post(ROOT+'/plans/bulk-export.'+suffix,json={'batchIds':[x['id'] for x in story['batches']]}))
            data=b64decode(result['contentBase64']);assert data.startswith(b'%PDF' if suffix=='pdf' else b'PK')

def test_import_and_plan_write_reject_student_and_foreign_school(story):
    foreign=seed_plan_story()
    with login(story,'STUDENT') as student,login(foreign) as outsider:
        assert student.post(ROOT+'/enterprises/import/confirm',json={'rows':[company_row(story)]}).status_code==403
        assert student.post(ROOT+'/plans/batch/'+story['batchId']+'/publish',json={'expectedVersion':2}).status_code==403
        assert outsider.get(ROOT+'/plans/batch/'+story['batchId']).json()['code']!=0
        assert outsider.post(ROOT+'/intern-students/'+story['recordId']+'/plan-assignments',json={'planIds':[story['batches'][0]['planId']]}).json()['code']!=0

def test_college_advisor_picker_and_direct_import_writes_use_same_scope(story):
    with login(story,'COLLEGE_ADMIN') as college:
        advisors=ok(college.get(ROOT+'/intern-students/advisors'))
        assert [row['id'] for row in advisors]==[str(story['users']['INTERN_MENTOR'])]
        rec=ok(college.get(ROOT+'/intern-students/'+story['recordId']))
        denied=college.post(ROOT+'/intern-students/'+story['recordId']+'/advisor',json={'advisorUserId':str(story['users']['OTHER_MENTOR']),'expectedVersion':rec['version']})
        assert denied.json()['code']!=0
        pre=ok(college.post(ROOT+'/intern-students/allocation-import/dry-run',json={'batchId':story['batchId'],'rows':[{'studentNo':story['studentNo'],'teacherNo':'000000000000000052','advisorName':'外院验收导师'}]}))
        assert pre['invalidRows']==1
        assert college.post(ROOT+'/plans/batch/'+story['batches'][1]['id']+'/publish',json={'expectedVersion':2}).status_code==403

def test_allocation_preview_version_conflict_preserves_newer_assignment(story):
    endpoint=ROOT+'/intern-students/allocation-import/'
    row={'studentNo':story['studentNo'],'teacherNo':'000000000000000051','advisorName':'计划验收导师'}
    with login(story) as school:
        pre=ok(school.post(endpoint+'dry-run',json={'batchId':story['batchId'],'rows':[row]}));assert pre['validRows']==1
        assert isinstance(pre['rows'][0]['expectedVersion'],int)
        rec=ok(school.get(ROOT+'/intern-students/'+story['recordId']));assert not rec['advisorUserId']
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/advisor',json={'advisorUserId':str(story['users']['OTHER_MENTOR']),'expectedVersion':rec['version']}))
        assert school.post(endpoint+'confirm',json={'batchId':story['batchId'],'rows':pre['rows']}).json()['code']!=0
        assert ok(school.get(ROOT+'/intern-students/'+story['recordId']))['advisorUserId']==str(story['users']['OTHER_MENTOR'])
        assert school.post(endpoint+'confirm',json={'batchId':story['batchId'],'rows':[row]}).json()['code']!=0
        fresh=ok(school.post(endpoint+'dry-run',json={'batchId':story['batchId'],'rows':[row]}))
        assert ok(school.post(endpoint+'confirm',json={'batchId':story['batchId'],'rows':fresh['rows']}))['advisorAssigned']==1
        assert school.post(endpoint+'confirm',json={'batchId':story['batchId'],'rows':fresh['rows']}).json()['code']!=0

def test_combined_advisor_position_import_changes_both_and_preserves_capacity(story):
    row=company_row(story)
    with login(story) as school:
        ok(school.post(ROOT+'/enterprises/import/confirm',json={'rows':[row]}))
        company=ok(school.get(ROOT+'/enterprises'))['items'][0]
        ok(school.post(ROOT+'/enterprises/'+company['id']+'/review',json={'action':'APPROVE','expectedVersion':company['version'],'comment':'合成企业校内前置'}))
        with get_sessionmaker()() as db:batch_no=db.get(InternshipBatch,int(story['batchId'])).batch_no
        position={'templateVersion':'POSITION_IMPORT_V2','batchNo':batch_no,'company':row['creditCode'],'title':'软件实践岗位','workContent':'完成岗位实践和安全教育任务',
          'dailyHours':'8','weeklyHours':'40','nightShift':'否','overtimeAllowed':'否','restDaysPerWeek':'2','remunerationType':'MONTHLY','remunerationAmount':'3000','remunerationCycle':'MONTHLY',
          'accommodationProvided':'是','mealProvided':'是','hazardousFlag':'否','headcount':'2','workAddress':'合成实践地址'}
        ok(school.post(ROOT+'/positions/import/confirm',json={'templateVersion':'POSITION_IMPORT_V2','rows':[position]}))
        pos=ok(school.get(ROOT+'/positions'))['items'][0]
        pos=ok(school.post(ROOT+'/positions/'+pos['id']+'/status',json={'action':'SUBMIT','expectedVersion':pos['version']}))
        pos=ok(school.post(ROOT+'/positions/'+pos['id']+'/status',json={'action':'PUBLISH','expectedVersion':pos['version']}))
        allocation={'studentNo':story['studentNo'],'teacherNo':'000000000000000051','advisorName':'计划验收导师','companyCreditCode':row['creditCode'],'positionName':position['title']}
        pre=ok(school.post(ROOT+'/intern-students/allocation-import/dry-run',json={'batchId':story['batchId'],'rows':[allocation]}));assert pre['validRows']==1
        receipt=ok(school.post(ROOT+'/intern-students/allocation-import/confirm',json={'batchId':story['batchId'],'rows':pre['rows']}))
        assert receipt['advisorAssigned']==1 and receipt['positionAssigned']==1
        detail=ok(school.get(ROOT+'/intern-students/'+story['recordId']))
        assert detail['positionId']==pos['id'] and detail['advisorUserId']==str(story['users']['INTERN_MENTOR'])
        assert ok(school.get(ROOT+'/positions/'+pos['id']))['allocatedCount']==1

def test_related_teacher_reads_shared_and_additional_plans_without_other_students(story):
    with get_sessionmaker()() as db:
        other=db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id==int(story['batches'][1]['id'])))
        db.add(InternshipRecord(tenant_id=int(story['tenantId']),batch_id=int(story['batchId']),student_id=other.student_id,
            advisor_user_id=story['users']['OTHER_MENTOR'],status='READY'))
        db.commit()
    with login(story) as school,login(story,'INTERN_MENTOR') as teacher,login(story,'COLLEGE_ADMIN') as college:
        rec=ok(school.get(ROOT+'/intern-students/'+story['recordId']))
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/advisor',json={'advisorUserId':str(story['users']['INTERN_MENTOR']),'expectedVersion':rec['version']}))
        # Sharing the plan body does not grant an institute authority to publish for other colleges.
        assert college.post(ROOT+'/plans/batch/'+story['batchId']+'/publish',json={'expectedVersion':2}).status_code==403
        primary=publish(school,story);extra=publish(school,story,1)
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/plan-assignments',json={'planIds':[extra['id']]}))
        for entry in story['batches']:
            assert ok(teacher.get(ROOT+'/plans/batch/'+entry['id']))['id']==entry['planId']
            rows=ok(teacher.get(ROOT+'/plan-task-progress',params={'batchId':entry['id']}))['items']
            assert rows and {row['internId'] for row in rows}=={story['recordId']}
        assert {x['id'] for x in ok(teacher.get(ROOT+'/plans/export-options'))['items']}=={x['id'] for x in story['batches']}

def test_published_plan_attachment_requires_membership_binding_scan_and_current_role(story):
    from app.services.file_scan_service import process_next_scan_job
    from test_yiyang_write_flows_mysql import png_bytes
    from app.models.file import FileBinding
    with login(story) as school,login(story,'STUDENT') as student,login(story,'INTERN_MENTOR') as teacher,login(story,'OTHER_MENTOR') as unrelated:
        uploaded=ok(school.post('/api/v1/files',params={'bizType':'TEMP_PRIVATE'},files={'file':('计划安排.png',png_bytes(),'image/png')}))
        fid=uploaded['fileId']
        if not uploaded['readyForBusiness']:
            assert process_next_scan_job(tenant_id=int(story['tenantId']),file_id=int(fid))
            assert ok(school.get('/api/v1/files/'+fid))['readyForBusiness']
        assert student.get('/api/v1/files/download/'+fid).status_code==404
        plan=ok(school.get(ROOT+'/plans/batch/'+story['batchId']))
        plan=ok(school.put(ROOT+'/plans/batch/'+story['batchId'],json={**plan,'attachmentFileIds':[fid],'expectedVersion':plan['version']}))
        assert student.get('/api/v1/files/download/'+fid).status_code==404
        rec=ok(school.get(ROOT+'/intern-students/'+story['recordId']))
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/advisor',json={'advisorUserId':str(story['users']['INTERN_MENTOR']),'expectedVersion':rec['version']}))
        ok(school.post(ROOT+'/plans/batch/'+story['batchId']+'/publish',json={'expectedVersion':plan['version']}))
        detail=ok(student.get(PORTAL+'/plan',params=context(story)));assert [f['fileId'] for f in detail['attachments']]==[fid]
        for client in [student,teacher]:
            assert client.get('/api/v1/files/download/'+fid).content==png_bytes()
        assert unrelated.get('/api/v1/files/download/'+fid).status_code==404
        with get_sessionmaker()() as db:
            binding=db.scalar(select(FileBinding).where(FileBinding.file_id==int(fid),FileBinding.tenant_id==int(story['tenantId'])))
            binding.is_current=False;db.commit()
        for client in [student,teacher,school]:
            assert client.get('/api/v1/files/download/'+fid).status_code==404

@pytest.mark.parametrize('entry',['create','import'])
def test_joining_after_publication_receives_primary_plan_ack_and_tasks_atomically(story,entry):
    batch=story['batches'][1]['id']
    with login(story) as school,login(story,'STUDENT') as student:
        plan=publish(school,story,1)
        if entry=='create':
            record=ok(school.post(ROOT+'/intern-students',json={'studentId':story['studentId'],'batchId':batch,'advisorUserId':str(story['users']['INTERN_MENTOR'])}))
        else:
            imported=ok(school.post(ROOT+'/intern-students/import/confirm',json={'batchId':batch,'rows':[{'studentNo':story['studentNo'],'teacherNo':'000000000000000051','advisorName':'计划验收导师'}]}))
            assert imported['created']==1
            record=next(row for row in ok(school.get(ROOT+'/intern-students',params={'batchId':batch}))['items'] if row['studentId']==story['studentId'])
        ctx={'batchId':batch,'internshipId':record['id']}
        detail=ok(student.get(PORTAL+'/plan',params=ctx));assert detail['ackId'] and detail['isPrimary']
        assert len(ok(student.get(PORTAL+'/plan/tasks',params=ctx))['tasks'])==1
        ok(student.post(PORTAL+'/plan/acknowledge',json={**ctx,'planId':plan['id'],'planVersion':detail['version'],'expectedVersion':detail['ackVersion']}))

def test_participant_remove_and_rejoin_preserve_same_record_and_unstarted_plan(story):
    batch=story['batches'][1]['id'];endpoint=ROOT+'/batches/'+batch+'/participants'
    with login(story) as school,login(story,'STUDENT') as student,login(story,'INTERN_MENTOR') as teacher:
        publish(school,story,1)
        assert ok(school.post(endpoint+'/add',json={'studentIds':[story['studentId']],'reason':'合成漏选补录'}))['added']==1
        member=next(row for row in ok(school.get(endpoint))['items'] if row['studentId']==story['studentId'])
        ctx={'batchId':batch,'internshipId':member['internshipId']}
        rec=ok(school.get(ROOT+'/intern-students/'+member['internshipId']))
        ok(school.post(ROOT+'/intern-students/'+rec['id']+'/advisor',json={'advisorUserId':str(story['users']['INTERN_MENTOR']),'expectedVersion':rec['version']}))
        assert teacher.get(ROOT+'/intern-students/'+rec['id']).status_code==200
        first=ok(student.get(PORTAL+'/plan/tasks',params=ctx));assert first['ackId'] and len(first['tasks'])==1
        ok(school.post(endpoint+'/'+member['id']+'/remove',json={'reason':'合成名单调整','version':member['version']}))
        assert student.get(PORTAL+'/plan',params=ctx).json()['code']!=0
        assert teacher.get(ROOT+'/intern-students/'+rec['id']).json()['code']!=0
        summary=ok(school.get(endpoint+'/summary'));assert summary['removedCount']==1
        history=ok(school.get(endpoint,params={'includeRemoved':True}))['items']
        assert next(row for row in history if row['id']==member['id'])['status']=='REMOVED'
        assert ok(school.post(endpoint+'/add',json={'studentIds':[story['studentId']],'reason':'合成恢复原名单'}))['added']==1
        restored=next(row for row in ok(school.get(endpoint))['items'] if row['studentId']==story['studentId'])
        assert restored['id']==member['id'] and restored['internshipId']==member['internshipId']
        assert ok(teacher.get(ROOT+'/intern-students/'+rec['id']))['advisorUserId']==str(story['users']['INTERN_MENTOR'])
        again=ok(student.get(PORTAL+'/plan/tasks',params=ctx));assert again['ackId']==first['ackId'] and again['tasks'][0]['progressId']==first['tasks'][0]['progressId']

def test_started_plan_task_prevents_removing_student_from_active_roster(story):
    batch=story['batches'][1]['id'];endpoint=ROOT+'/batches/'+batch+'/participants'
    with login(story) as school,login(story,'STUDENT') as student:
        plan=publish(school,story,1)
        ok(school.post(endpoint+'/add',json={'studentIds':[story['studentId']]}))
        member=next(row for row in ok(school.get(endpoint))['items'] if row['studentId']==story['studentId'])
        ctx={'batchId':batch,'internshipId':member['internshipId']}
        detail=ok(student.get(PORTAL+'/plan',params=ctx))
        ok(student.post(PORTAL+'/plan/acknowledge',json={**ctx,'planId':plan['id'],'planVersion':detail['version'],'expectedVersion':detail['ackVersion']}))
        tasks=ok(student.get(PORTAL+'/plan/tasks',params=ctx))
        ok(student.post(PORTAL+'/plan/tasks/1/submit',json={**ctx,'planId':plan['id'],'expectedVersion':tasks['tasks'][0]['progressVersion'],'studentNote':'合成新增学生完成本计划任务，保留原办理记录。'}))
        assert school.post(endpoint+'/'+member['id']+'/remove',json={'reason':'不得隐去在办任务','version':member['version']}).json()['code']!=0
        assert next(row for row in ok(school.get(endpoint))['items'] if row['id']==member['id'])['status']=='ACTIVE'

def test_concurrent_publication_and_student_join_never_lose_plan_delivery(story):
    from concurrent.futures import ThreadPoolExecutor
    from threading import Barrier
    batch=story['batches'][1]['id'];barrier=Barrier(2)
    with login(story) as school:
        def publish_now():
            barrier.wait()
            return ok(school.post(ROOT+'/plans/batch/'+batch+'/publish',json={'expectedVersion':2}))
        def join_now():
            barrier.wait()
            return ok(school.post(ROOT+'/intern-students',json={'studentId':story['studentId'],'batchId':batch}))
        with ThreadPoolExecutor(max_workers=2) as pool:
            publication=pool.submit(publish_now);joining=pool.submit(join_now)
            plan=publication.result();record=joining.result()
        assignments=ok(school.get(ROOT+'/intern-students/'+record['id']+'/plan-assignments'))['items']
        assert len(assignments)==1 and assignments[0]['planId']==plan['id'] and assignments[0]['taskCount']==1
        with get_sessionmaker()() as db:
            acks=db.scalars(select(InternshipPlanAck).where(InternshipPlanAck.tenant_id==int(story['tenantId']),InternshipPlanAck.internship_id==int(record['id']))).all()
            assert len(acks)==1 and acks[0].status=='PENDING'

def test_freezing_a_cohort_after_plan_publication_delivers_the_same_plan(story):
    batch=story['batches'][1]['id']
    with get_sessionmaker()() as db:
        db.get(InternshipBatch,int(batch)).status='DRAFT';db.commit()
    with login(story) as school,login(story,'STUDENT') as student:
        plan=publish(school,story,1)
        frozen=ok(school.post(ROOT+'/batches/'+batch+'/participants/freeze',json={'rule':{'studentIds':[story['studentId']]}}))
        assert frozen['batchStatus']=='RUNNING'
        members=ok(school.get(ROOT+'/batches/'+batch+'/participants'))['items']
        member=next(row for row in members if row['studentId']==story['studentId'])
        ctx={'batchId':batch,'internshipId':member['internshipId']}
        detail=ok(student.get(PORTAL+'/plan',params=ctx))
        assert detail['id']==plan['id'] and detail['ackId'] and detail['isPrimary']

@pytest.mark.parametrize('kind,label',[
    ('POST','顶岗实习（岗位实习）'),('COGNITIVE','认知实习'),('FOLLOW_POST','跟岗实习'),
    ('APPRENTICESHIP','学徒制'),('COMPREHENSIVE','综合实训'),('OTHER','其他'),
])
def test_plan_categories_preserve_complete_fields_in_publication_and_both_exports(story,kind,label):
    from pypdf import PdfReader
    with login(story) as school:
        view=ok(school.get(ROOT+'/plans/batch/'+story['batchId']))
        saved=ok(school.put(ROOT+'/plans/batch/'+story['batchId'],json={**view,'internshipType':kind,'expectedVersion':view['version']}))
        plan=ok(school.post(ROOT+'/plans/batch/'+story['batchId']+'/publish',json={'expectedVersion':saved['version']}))
        assert plan['internshipType']==kind and plan['internshipTypeLabel']==label
        pdf=ok(school.post(ROOT+'/plans/bulk-export.pdf',json={'batchIds':[story['batchId']]}))
        text='\n'.join(page.extract_text() for page in PdfReader(BytesIO(b64decode(pdf['contentBase64']))).pages)
        for expected in [label,plan['planNo'],plan['objectives'],plan['responsibleName'],'考核内容','计划签到天数']:
            assert expected in text
        xlsx=ok(school.post(ROOT+'/plans/bulk-export.xlsx',json={'batchIds':[story['batchId']]}))
        book=load_workbook(BytesIO(b64decode(xlsx['contentBase64'])))
        cells=[c.value for sheet in book for row in sheet for c in row]
        assert label in cells and plan['planNo'] in cells and plan['responsibleName'] in cells

def test_task_database_pagination_filters_and_definition_use_bounded_queries(story):
    from app.models import InternshipPlanTaskProgress
    from app.db.session import get_engine
    from sqlalchemy import event
    with get_sessionmaker()() as db:
        for i in range(25):
            student=StudentProfile(tenant_id=int(story['tenantId']),student_no=f'TASK-{i:02}',real_name=f'分页任务学生{i:02}',
                college_id=int(story['collegeId']),status='ACTIVE',student_status='NORMAL',current_stage='ENROLLED')
            db.add(student);db.flush();db.add(InternshipRecord(tenant_id=int(story['tenantId']),student_id=student.id,
                batch_id=int(story['batchId']),advisor_user_id=story['users']['INTERN_MENTOR'],status='PREPARING'))
        db.commit()
    with login(story) as school,login(story,'INTERN_MENTOR') as teacher,login(story,'COLLEGE_ADMIN') as college:
        plan=publish(school,story)
        with get_sessionmaker()() as db:
            other=db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id==int(story['batches'][1]['id'])))
            db.add(InternshipPlanTaskProgress(tenant_id=int(story['tenantId']),plan_id=int(plan['id']),internship_id=int(story['recordId']),
                student_id=other.student_id,task_sort_order=2,task_name='错误学生关联不得展示',status='NOT_STARTED'));db.commit()
        pages=[ok(school.get(ROOT+'/plan-task-progress',params={'batchId':story['batchId'],'page':i,'pageSize':10})) for i in [1,2,3]]
        assert [len(p['items']) for p in pages]==[10,10,6] and all(p['total']==26 for p in pages)
        assert len({r['id'] for p in pages for r in p['items']})==26
        for p in pages:
            assert all(r['taskRequirement']=='说明本次岗位实践与安全要求' and r['taskDeadline'] for r in p['items'])
        assert ok(teacher.get(ROOT+'/plan-task-progress',params={'batchId':story['batchId']}))['total']==25
        assert ok(college.get(ROOT+'/plan-task-progress',params={'batchId':story['batchId']}))['total']==26
        assert ok(school.get(ROOT+'/plan-task-progress',params={'batchId':story['batchId'],'keyword':'%'}))['total']==0
        assert ok(school.get(ROOT+'/plans/batch/'+story['batchId']+'/task-summary'))['studentCount']==26
        selects=[]
        def counted(_conn,_cursor,statement,_parameters,_context,_many):
            if statement.lstrip().upper().startswith('SELECT') and any(t in statement for t in ['t_internship_plan_task_progress','t_internship_record','t_student_profile','t_internship_batch_plan']):selects.append(statement)
        engine=get_engine();event.listen(engine,'before_cursor_execute',counted)
        try:
            page=ok(school.get(ROOT+'/plan-task-progress',params={'batchId':story['batchId'],'pageSize':50}))
            assert len(page['items'])==26 and len(selects)<=4
        finally:event.remove(engine,'before_cursor_execute',counted)

def test_removed_additional_plan_disappears_from_scoped_task_list_and_summary(story):
    with login(story) as school,login(story,'INTERN_MENTOR') as teacher:
        rec=ok(school.get(ROOT+'/intern-students/'+story['recordId']))
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/advisor',json={'advisorUserId':str(story['users']['INTERN_MENTOR']),'expectedVersion':rec['version']}))
        publish(school,story);extra=publish(school,story,1)
        assignment=ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/plan-assignments',json={'planIds':[extra['id']]}))['items'][0]
        batch=story['batches'][1]['id']
        assert ok(teacher.get(ROOT+'/plan-task-progress',params={'batchId':batch}))['total']==1
        ok(school.post(ROOT+'/intern-students/'+story['recordId']+'/plan-assignments/'+assignment['id']+'/remove',json={'reason':'合成方案调整','expectedVersion':assignment['version']}))
        assert ok(teacher.get(ROOT+'/plan-task-progress',params={'batchId':batch}))['total']==0
        assert ok(teacher.get(ROOT+'/plans/batch/'+batch+'/task-summary'))['studentCount']==0
