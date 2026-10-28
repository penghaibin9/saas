from types import SimpleNamespace
from contextlib import contextmanager
from datetime import date, datetime

import pytest

from app.core.exceptions import AppException
from app.modules.internship.services import internship_report_quality_service as svc


def test_g15_default_rules_and_minimum_words():
    assert svc.DEFAULT_RULES["weeklyMinWords"] == 30
    assert svc.DEFAULT_RULES["monthlyMinWords"] == 100
    assert svc.DEFAULT_RULES["summaryMinWords"] == 300
    assert svc.minimum_words(svc.DEFAULT_RULES, "DAILY") == 30
    assert svc.minimum_words(svc.DEFAULT_RULES, "MONTHLY") == 100
    assert svc.minimum_words(svc.DEFAULT_RULES, "SUMMARY") == 300


def test_g15_attachment_validation_accepts_images_and_videos(monkeypatch):
    metas = {
        "img": {
            "fileId": "img", "fileName": "现场.jpg", "mimeType": "image/jpeg",
            "sizeBytes": 123, "sha256": "a" * 64,
        },
        "vid": {
            "fileId": "vid", "fileName": "过程.mp4", "mimeType": "video/mp4",
            "sizeBytes": 456, "sha256": "b" * 64,
        },
    }
    monkeypatch.setattr(svc.file_service, "get_file_meta", lambda fid: metas.get(fid))
    ids, views = svc.validate_attachments(["img", "vid", "img"], svc.DEFAULT_RULES)
    assert ids == ["img", "vid"]
    assert [v["kind"] for v in views] == ["IMAGE", "VIDEO"]
    assert views[0]["sha256"] == "a" * 64


def test_g15_attachment_limits_are_server_enforced(monkeypatch):
    monkeypatch.setattr(
        svc.file_service,
        "get_file_meta",
        lambda fid: {
            "fileId": fid, "fileName": f"{fid}.jpg", "mimeType": "image/jpeg",
            "sizeBytes": 1, "sha256": "c" * 64,
        },
    )
    with pytest.raises(AppException):
        svc.validate_attachments(
            ["1", "2"],
            {**svc.DEFAULT_RULES, "maxImages": 1},
        )


def test_g15_zero_attachment_limits_keep_documents_but_reject_media(monkeypatch):
    metas = {
        "img": {"mimeType": "image/jpeg"},
        "vid": {"mimeType": "video/mp4"},
        "doc": {"mimeType": "application/pdf", "ext": "pdf"},
    }
    monkeypatch.setattr(svc.file_service, "get_file_meta", metas.get)
    rules = {**svc.DEFAULT_RULES, "maxImages": 0, "maxVideos": 0}
    assert svc.validate_attachments([], rules) == ([], [])
    assert svc.validate_attachments(["doc"], rules)[0] == ["doc"]
    for fid, message in (("img", "图片最多 0 张"), ("vid", "视频最多 0 个")):
        with pytest.raises(AppException, match=message):
            svc.validate_attachments([fid], rules)


def test_g15_batch_rules_preserve_explicit_zero_limits(monkeypatch):
    monkeypatch.setattr(svc, "_tid", lambda: 11)
    row = SimpleNamespace(weekly_min_words=0, plan_task_min_words=10,
                          monthly_min_words=100, summary_min_words=300,
                          max_images=0, max_videos=0)
    class Db:
        def scalar(self, _): return row
        def get(self, *_): return SimpleNamespace(rules_config=config)
    for config in ({}, {"weeklyReport": {"minWordCount": 0},
                       "processReport": {"maxImages": 0, "maxVideos": 0}}):
        rules = svc.rules_for_batch(Db(), 18)
        assert (rules["weeklyMinWords"], rules["maxImages"], rules["maxVideos"]) == (0, 0, 0)


def test_g15_weekly_submit_uses_zero_minimum_but_requires_both_body_fields(monkeypatch):
    from app.modules.internship.services import internship_student_report_context_service as reports
    class ReachedReportLookup(Exception): pass
    class Db:
        def scalar(self, _): raise ReachedReportLookup
    @contextmanager
    def fake_session(): yield Db()
    monkeypatch.setattr(reports, "session", fake_session)
    monkeypatch.setattr(reports, "_tid", lambda: 11)
    monkeypatch.setattr(reports, "require_explicit_context",
                        lambda *_a, **_k: (SimpleNamespace(id=18, batch_id=18), None, 18))
    monkeypatch.setattr(svc, "rules_for_batch", lambda *_: {**svc.DEFAULT_RULES, "weeklyMinWords": 0})
    body = {"weekNo": 1, "workContent": "工", "harvestContent": "获"}
    with pytest.raises(ReachedReportLookup):
        reports.submit_weekly(_user(), body)
    for field in ("workContent", "harvestContent"):
        with pytest.raises(AppException, match="不能为空"):
            reports.submit_weekly(_user(), {**body, field: ""})


class _FakeDb:
    def scalar(self, _query):
        return None

    def add(self, value):
        self.value = value

    def flush(self):
        if getattr(self, "value", None) is not None and getattr(self.value, "id", None) is None:
            self.value.id = 99


def _row(report_type="MONTHLY"):
    return SimpleNamespace(id=10, report_type=report_type)


def _user():
    return {"userId": "7", "realName": "指导教师"}


def test_g15_approve_requires_five_level_rating(monkeypatch):
    monkeypatch.setattr(svc, "_tid", lambda: 1)
    monkeypatch.setattr(
        svc,
        "latest_process_snapshot",
        lambda db, report_id: SimpleNamespace(id=88),
    )
    with pytest.raises(AppException) as exc:
        svc.record_process_review(
            _FakeDb(), row=_row(), action="APPROVE", comment="通过",
            user=_user(), rating_level=None, summary_score=None,
        )
    assert "五级评价" in str(exc.value)


def test_g15_summary_requires_0_to_100_score(monkeypatch):
    monkeypatch.setattr(
        svc,
        "latest_process_snapshot",
        lambda db, report_id: SimpleNamespace(id=88),
    )
    with pytest.raises(AppException):
        svc.record_process_review(
            _FakeDb(), row=_row("SUMMARY"), action="APPROVE", comment="通过",
            user=_user(), rating_level=5, summary_score=101,
        )
    with pytest.raises(AppException):
        svc.record_process_review(
            _FakeDb(), row=_row("SUMMARY"), action="APPROVE", comment="通过",
            user=_user(), rating_level=5, summary_score=None,
        )


def test_g15_return_can_be_unrated_but_is_version_bound(monkeypatch):
    monkeypatch.setattr(svc, "_tid", lambda: 1)
    monkeypatch.setattr(
        svc,
        "latest_process_snapshot",
        lambda db, report_id: SimpleNamespace(id=88),
    )
    db = _FakeDb()
    review = svc.record_process_review(
        db, row=_row("SUMMARY"), action="RETURN", comment="请补充现场过程",
        user=_user(), rating_level=None, summary_score=None,
    )
    assert review.report_version_id == 88
    assert review.action == "RETURN"
    assert review.rating_level is None
    assert review.summary_score is None


def test_grade_rejects_fractional_empty_and_invalid_values_and_summary_accepts_zero(monkeypatch):
    for value in (True, 0, 6, 2.5, "2.5", " ", "nan", "inf"):
        with pytest.raises(AppException):
            svc._rating(value)
    assert svc._rating("5") == 5
    assert svc._rating(None) is None
    monkeypatch.setattr(svc, "_tid", lambda: 1)
    monkeypatch.setattr(svc, "latest_process_snapshot", lambda *_: SimpleNamespace(id=88))
    for value in (None, "", True, "nan", "inf", -1, 101):
        with pytest.raises(AppException):
            svc.record_process_review(_FakeDb(), row=_row("SUMMARY"), action="APPROVE",
                comment="合成评分", user=_user(), rating_level=None, summary_score=value)
    review = svc.record_process_review(_FakeDb(), row=_row("SUMMARY"), action="APPROVE",
        comment="合成评分", user=_user(), rating_level=None, summary_score=0)
    assert review.summary_score == 0 and review.rating_level is None


def test_submission_dates_are_inclusive_local_days_and_reject_reversed_range():
    from sqlalchemy import select
    from app.models import InternshipProcessReport
    query = svc.filter_submission_dates(select(InternshipProcessReport), InternshipProcessReport,
        date(2026, 10, 3), date(2026, 10, 4))
    values = list(query.compile().params.values())
    assert values == [datetime(2026, 10, 2, 16), datetime(2026, 10, 4, 16)]
    assert svc.format_report_time("2026-10-06T16:05:00Z") == "2026-10-07 00:05"
    with pytest.raises(AppException):
        svc.filter_submission_dates(query, InternshipProcessReport, "2026-10-05", "2026-10-03")


def test_latest_review_query_requires_the_current_submission_snapshot(monkeypatch):
    from sqlalchemy.dialects import mysql
    captured = []
    class Db:
        def scalars(self, query):
            captured.append(query.compile(dialect=mysql.dialect()))
            return SimpleNamespace(all=lambda: [])
    monkeypatch.setattr(svc, "_tid", lambda: 11)
    assert svc.latest_review_map(Db(), "PROCESS", [10]) == {}
    sql = str(captured[0]).lower()
    assert "max(" in sql and "version_no" in sql
    assert "report_version_id" in sql and "join" in sql
    assert "tenant_id" in sql and "report_kind" in sql


def test_obligation_totals_and_page_use_database_queries_not_roster_slicing(monkeypatch):
    from sqlalchemy.dialects import mysql
    from app.modules.internship.services import internship_batch_context as batches, internship_scope as scope
    captured = []
    record = SimpleNamespace(id=18, advisor_name="合成导师", enterprise_name="并行B-报告", position_name="合成岗位")
    student = SimpleNamespace(id=18, real_name="合成学生", student_no="B001")
    class Result:
        def one(self): return (1, 0, 3, 2, 1, 4, 0, 0, 2, 0, 0, 2)
        def all(self): return [(record, student, 1, 0, 2, 1, 0, 0, 0, 0)]
    class Db:
        def scalar(self, query): captured.append(query.compile(dialect=mysql.dialect())); return 2
        def execute(self, query): captured.append(query.compile(dialect=mysql.dialect())); return Result()
    @contextmanager
    def fake_session(): yield Db()
    monkeypatch.setattr(svc, "session", fake_session)
    monkeypatch.setattr(svc, "_tid", lambda: 11)
    monkeypatch.setattr(batches, "resolve_batch", lambda *_a, **_k: SimpleNamespace(id=18))
    monkeypatch.setattr(scope, "apply_internship_record_scope", lambda query, _: query)
    monkeypatch.setattr(svc, "rules_for_batch", lambda *_: {**svc.DEFAULT_RULES,
        "dailyRequiredCount": 2, "weeklyRequiredCount": 3, "monthlyRequiredCount": 1, "summaryRequiredCount": 1})
    data = svc.report_obligations(_user(), 18, page=2, page_size=1, missing_only=True, keyword="B%")
    assert data["total"] == 2 and len(data["items"]) == 1
    assert data["summary"]["daily"] == {"required": 4, "submitted": 1, "approved": 0, "missing": 3, "configured": True}
    assert data["items"][0]["weekly"]["submitted"] == 2
    assert data["items"][0]["missingTotal"] == 4
    sql = str(captured[-1]).lower()
    assert "limit" in sql and "group by" in sql and "tenant_id" in sql
    assert any("PENDING_REVIEW" in str(value) for value in captured[-1].params.values())
    assert 1 in captured[-1].params.values()


def test_missing_reminder_commits_a_real_event_with_audit_and_refuses_complete_or_repeated(monkeypatch):
    from app.modules.internship.services import internship_batch_context as batches, internship_scope as scope
    from app.services import message_event_outbox_service as outbox, student_account_link_service as links
    record = SimpleNamespace(id=18, student_id=18)
    student = SimpleNamespace(id=18, tenant_id=11, is_deleted=False, student_no="B001")
    class Db:
        def __init__(self, submitted=1, recent=None): self.responses=iter((record, submitted, recent)); self.saved=[]; self.committed=False
        def scalar(self, _): return next(self.responses)
        def get(self, *_): return student
        def add(self, row): self.saved.append(row)
        def commit(self): self.committed=True
    db = Db()
    @contextmanager
    def fake_session(): yield db
    emitted=[]; deliveries=[]
    monkeypatch.setattr(svc, "session", fake_session); monkeypatch.setattr(svc, "_tid", lambda: 11)
    monkeypatch.setattr(batches, "resolve_batch", lambda *_a, **_k: SimpleNamespace(id=18, batch_name="并行B-报告独立批次"))
    monkeypatch.setattr(scope, "apply_internship_record_scope", lambda query, _: query)
    monkeypatch.setattr(svc, "rules_for_batch", lambda *_: {**svc.DEFAULT_RULES, "dailyRequiredCount":2})
    monkeypatch.setattr(links, "resolve_user_id_for_student", lambda *_a, **_k: 18)
    monkeypatch.setattr(outbox, "emit_receiver_notice", lambda _db, **kw: emitted.append(kw) or SimpleNamespace(id=71))
    monkeypatch.setattr(outbox, "try_process_pending_outbox",
                        lambda **kw: deliveries.append((db.committed, kw["outbox_ids"])))
    result = svc.remind_report_obligation(_user(), 18, 18, "DAILY")
    assert result["queued"] and result["missing"] == 1 and db.committed
    assert emitted[0]["event_code"] == "INTERNSHIP.REPORT_REMIND"
    assert "并行B-报告独立批次" in emitted[0]["content"] and "编号18" in emitted[0]["content"]
    assert db.saved[0].detail_json["eventId"] == "71"
    assert deliveries == [(True, [71])]
    for submitted, recent in ((2,None),(1,9)):
        db=Db(submitted, recent)
        with pytest.raises(AppException): svc.remind_report_obligation(_user(), 18, 18, "DAILY")
        assert not db.committed and not db.saved
    assert len(emitted)==1
    assert len(deliveries)==1


def test_report_document_contains_only_the_current_submission_grade_and_real_pdf_bytes(monkeypatch):
    import base64
    from app.modules.internship.services import internship_student_report_context_service as student_reports
    from app.services import pdf_util
    row=SimpleNamespace(id=42,report_type="SUMMARY",period_key="FINAL",content="并行B-报告总结正文",
        status="PENDING_REVIEW",review_comment="原版本退回意见",submitted_at=datetime(2026,10,3))
    record=SimpleNamespace(id=18,enterprise_name="合成企业",position_name="合成岗位")
    student=SimpleNamespace(real_name="合成学生",student_no="B001")
    snap=SimpleNamespace(version_no=2,attachment_meta_json=[])
    class Db:
        def __init__(self): self.saved=[]; self.commits=0
        def scalar(self, _): return row
        def add(self, audit): self.saved.append(audit)
        def commit(self): self.commits+=1
    db=Db()
    @contextmanager
    def fake_session(): yield db
    monkeypatch.setattr(student_reports,"session",fake_session);monkeypatch.setattr(student_reports,"_tid",lambda:11)
    monkeypatch.setattr(student_reports,"require_explicit_context",lambda *_a,**_k:(record,student,18))
    monkeypatch.setattr(svc,"latest_process_snapshot",lambda *_:snap)
    review={}
    monkeypatch.setattr(svc,"latest_review_map",lambda *_:review)
    bodies=[]; titles=[]; builder=pdf_util.build_text_pdf
    def build(title,body,**kw): titles.append(title);bodies.append(body);return builder(title,body,**kw)
    monkeypatch.setattr(pdf_util,"build_text_pdf",build)
    def export():return student_reports.export_report_pdf(_user(),report_kind="PROCESS",report_id=42,batch_id=18,internship_id=18)
    pending=export();assert "原版本退回意见" not in bodies[-1]
    row.status="APPROVED";review[42]={"summaryScore":0,"comment":"本版评分","reviewerName":"合成导师"}
    approved=export();assert "0分（百分制）" in bodies[-1] and "本版评分" in bodies[-1]
    assert base64.b64decode(approved["contentBase64"]).startswith(b"%PDF")
    assert approved["reportId"]==pending["reportId"]=="42" and approved["reportVersion"]==2
    assert titles[-1]=="实习总结" and "FINAL" not in approved["filename"]
    assert db.commits==2 and db.saved[-1].action=="EXPORT_PDF"


def test_process_list_returns_current_snapshot_number_and_zero_score(monkeypatch):
    from app.modules.internship.services import internship_process_report_service as process_reports
    from app.modules.internship.services import internship_batch_context as batches, internship_scope as scope
    class Db:
        def scalar(self, _):return 1
        def execute(self, _):return SimpleNamespace(all=lambda:[(SimpleNamespace(id=1),None,None,None)])
    @contextmanager
    def fake_session():yield Db()
    monkeypatch.setattr(process_reports,"session",fake_session);monkeypatch.setattr(process_reports,"_tid",lambda:11)
    monkeypatch.setattr(batches,"resolve_batch",lambda *_:SimpleNamespace(id=18))
    monkeypatch.setattr(scope,"apply_internship_record_scope",lambda query,_:query)
    monkeypatch.setattr(process_reports,"_row",lambda *_a,**_k:{"id":"1","version":4})
    monkeypatch.setattr(svc,"latest_snapshot_map",lambda *_:{1:SimpleNamespace(version_no=2)})
    monkeypatch.setattr(svc,"latest_review_map",lambda *_:{1:{"summaryScore":0,"reportVersionId":"2"}})
    rows,total=process_reports.list_reports(1,20,batch_id=18,user=_user())
    assert total==1 and rows[0]["reportVersion"]==2 and rows[0]["review"]["summaryScore"]==0


def test_process_export_creates_real_xlsx_with_current_submission_and_grade(monkeypatch):
    import base64
    from io import BytesIO
    from openpyxl import load_workbook
    from app.modules.internship.services import internship_process_report_service as process_reports, internship_export_util as exports
    item={"id":"1","studentNo":"000B001","studentName":"合成学生","enterpriseName":"并行B-报告",
        "reportType":"SUMMARY","reportTypeLabel":"实习总结","periodKey":"FINAL","reportVersion":2,"wordCount":301,
        "submitAt":"2026-10-06T16:05:00Z","statusLabel":"已通过","review":{"summaryScore":0,"ratingLevel":None,
        "reviewerName":"合成导师","reviewedAt":"2026-10-06T16:10:00Z","comment":"本次提交版评分"}}
    monkeypatch.setattr(exports,"load_export_rows",lambda *_a,**_k:([item],1))
    monkeypatch.setattr(exports,"write_export_audit",lambda *_a,**_k:None)
    @contextmanager
    def fake_session():yield SimpleNamespace(commit=lambda:None)
    monkeypatch.setattr(process_reports,"session",fake_session)
    packed=process_reports.export_reports(batch_id=18,user=_user())
    content=base64.b64decode(packed["contentBase64"])
    assert content.startswith(b"PK")
    sheet=load_workbook(BytesIO(content),data_only=True).active
    values=list(sheet.values)
    assert any("提交版本" in row for row in values)
    row=next(row for row in values if "000B001" in row)
    assert "FINAL" not in row and "实习总结" in row
    assert row[5]==2 and row[7]=="2026-10-07 00:05" and row[9] is None and row[10]==0 and row[-1]=="本次提交版评分"


def test_weekly_export_uses_submission_version_instead_of_lock_version(monkeypatch):
    import base64
    from io import BytesIO
    from openpyxl import load_workbook
    from app.modules.internship.services import internship_service as weekly, internship_export_util as exports
    item={"id":"3","studentName":"并行B-报告","className":"合成班级","enterpriseName":"合成企业",
        "week":"第 2 周","submitAt":"2026-10-06T16:38:00Z","version":3,"reportVersion":"v2",
        "wordCount":130,"riskFlag":"","statusLabel":"已通过",
        "review":{"ratingLevel":3,"reviewerName":"合成教师","reviewedAt":"2026-10-06T16:39:00Z","comment":"第2版评价"}}
    monkeypatch.setattr(exports,"load_export_rows",lambda *_a,**_k:([item],1))
    monkeypatch.setattr(weekly,"_op_name",lambda:"合成管理员")
    packed=weekly.export_weekly_reports(batch_id=18,user=_user())
    rows=list(load_workbook(BytesIO(base64.b64decode(packed["contentBase64"])),data_only=True).active.values)
    row=next(row for row in rows if "并行B-报告" in row)
    assert row[5]==2 and row[9]==3 and row[-1]=="第2版评价"
