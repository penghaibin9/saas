"""Normal-role statistics on new synthetic schools; historical facts are read-only prerequisites."""
from datetime import datetime,date,timedelta
from base64 import b64decode
from io import BytesIO
import pytest
from sqlalchemy import select
from openpyxl import load_workbook
from app.db.session import get_sessionmaker
from app.models import (Tenant,User,UserRole,Role,StaffProfile,College,Major,SchoolClass,StudentProfile,TeacherStudentScope,
 InternshipRecord,InternshipBatch,InternshipApplication,InternshipBatchPlan,InternshipPlanAssignment,
 WeeklyReport,InternshipReportVersion,InternshipReportReview,InternshipProcessReport,InternshipCheckin,
 InternshipGuidance,InternshipVisit,InternshipPayrollStatement,InternshipPayrollVersion,InternshipFinalScore,InternshipChangeRequest)
from test_student_support_story_mysql import seed_support,result
from test_yiyang_plan_export_options_mysql import login,PASSWORD

ROOT='/api/v1/internship/stats'
def seed_stats_story():
    s=seed_support();tid=int(s['tenantId']);now=datetime(2026,9,9,9)
    with get_sessionmaker()() as db:
        db.get(Tenant,tid).school_name='学校统计与教师绩效合成学校'
        db.get(User,s['users']['SCHOOL_ADMIN']).real_name='统计验收学校管理员'
        batch=db.get(InternshipBatch,int(s['batchId']));batch.batch_name='学校统计与教师绩效合成批次'
        batch.start_date=datetime(2026,9,1);batch.end_date=datetime(2026,9,30)
        batch.rules_config={'weeklyReport':{'requiredCount':2,'reviewSlaHours':48},'processReport':{'monthlyRequiredCount':1,'summaryRequiredCount':1,'reviewSlaHours':48}}
        for code,role_code in [('OTHER_MENTOR','INTERN_MENTOR'),('COUNSELOR','COUNSELOR'),('COLLEGE_ADMIN','COLLEGE_ADMIN')]:
            role=db.scalar(select(Role).where(Role.tenant_id==tid,Role.role_code==role_code))
            actor=User(tenant_id=tid,login_name=code.lower(),real_name='同名统计教师',user_type='TEACHER',
                password_hash=db.get(User,s['users']['INTERN_MENTOR']).password_hash,status='ACTIVE',credential_version=0,must_change_password=False)
            db.add(actor);db.flush();db.add(UserRole(tenant_id=tid,user_id=actor.id,role_id=role.id,status='ACTIVE'));s['users'][code]=actor.id
        teacher=db.get(User,s['users']['INTERN_MENTOR']);teacher.real_name='同名统计教师'
        colleges=[];classes=[]
        for i in range(2):
            college=College(tenant_id=tid,college_name=['合成信息学院','合成机电学院'][i]);db.add(college);db.flush();colleges.append(college)
            major=Major(tenant_id=tid,college_id=college.id,major_name=['合成软件技术','合成机电技术'][i]);db.add(major);db.flush()
            group=SchoolClass(tenant_id=tid,major_id=major.id,class_name=['合成软件一班','合成机电一班'][i],grade='2024',head_teacher_id=s['users']['COUNSELOR'])
            db.add(group);db.flush();classes.append(group)
        for code,no in [('INTERN_MENTOR','000000000000000091'),('OTHER_MENTOR','000000000000000092'),('COUNSELOR','000000000000000093')]:
            db.add(StaffProfile(tenant_id=tid,user_id=s['users'][code],staff_no=no,real_name='同名统计教师',college_id=colleges[0].id,status='ACTIVE'))
        for code,kind,ref in [('COLLEGE_ADMIN','COLLEGE',colleges[0].id),('COUNSELOR','CLASS',classes[0].id)]:
            db.add(TeacherStudentScope(tenant_id=tid,user_id=s['users'][code],teacher_key=str(s['users'][code]),role_code=code,scope_type=kind,ref_id=ref,status='ACTIVE'))
        records=[]
        for i in range(3):
            if i==0:
                rec=db.get(InternshipRecord,int(s['recordId']));student=db.get(StudentProfile,int(s['studentId']))
            else:
                student=StudentProfile(tenant_id=tid,student_no='STATS-'+str(i),real_name='统计样例'+str(i),status='ACTIVE',student_status='NORMAL')
                db.add(student);db.flush();rec=InternshipRecord(tenant_id=tid,student_id=student.id,batch_id=batch.id);db.add(rec);db.flush()
            student.real_name=['统计甲','统计乙','统计丙'][i];student.class_id=classes[0 if i<2 else 1].id;student.grade='2024'
            if i==0:student.student_no='000000000000000000321'
            rec.advisor_user_id=s['users']['INTERN_MENTOR' if i<2 else 'OTHER_MENTOR'];rec.advisor_name='同名统计教师'
            rec.status=['ONBOARD','ASSESSING','ARCHIVED'][i];rec.destination_type='ASSIGNED'
            rec.intern_start_date=batch.start_date;rec.intern_end_date=batch.end_date
            rec.enterprise_name='合成企业'+str(i);rec.enterprise_id=900+i
            records.append(rec)
            db.add(InternshipApplication(tenant_id=tid,record_id=rec.id,student_id=student.id,batch_id=batch.id,
                application_type='SELF_ARRANGED',major_match=[True,False,None][i],status='APPROVED',reviewed_at=now))
            statement=InternshipPayrollStatement(tenant_id=tid,internship_id=rec.id,student_id=student.id,batch_id=batch.id,pay_month='2026-09',revision_count=1)
            db.add(statement);db.flush();pay=InternshipPayrollVersion(tenant_id=tid,statement_id=statement.id,revision_no=1,
                actual_amount=['1234.56','0','100'][i],currency=['CNY','CNY','USD'][i],status='APPROVED',is_current=True)
            db.add(pay);db.flush();statement.current_version_id=pay.id
            if i:db.add(InternshipFinalScore(tenant_id=tid,internship_id=rec.id,student_id=student.id,batch_id=batch.id,status='PUBLISHED',total_score=[0,0,96.4][i]))
        def report(rec,week,state,author,revision=1):
            row=WeeklyReport(tenant_id=tid,internship_id=rec.id,week_number=week,status=state,report_version=revision,
                work_content='合成工作记录用于统计核对。'*4,submitted_at=now,reviewed_at=now+timedelta(hours=1),reviewed_by_name='同名统计教师')
            db.add(row);db.flush()
            for v in range(1,revision+1):
                snap=InternshipReportVersion(tenant_id=tid,report_kind='WEEKLY',report_id=row.id,version_no=v,
                    internship_id=rec.id,student_id=rec.student_id,content_json={'workContent':row.work_content},submitted_at=now,word_count=80)
                db.add(snap);db.flush()
                if v<revision or state=='APPROVED':db.add(InternshipReportReview(tenant_id=tid,report_kind='WEEKLY',report_id=row.id,
                    report_version_id=snap.id,action='RETURN' if v<revision else 'APPROVE',reviewer_user_id=str(author),reviewer_name='同名统计教师',reviewed_at=now+timedelta(hours=1)))
            return row.id
        pending=report(records[0],1,'PENDING_REVIEW',teacher.id,2)
        report(records[0],2,'APPROVED',teacher.id)
        report(records[1],1,'APPROVED',s['users']['OTHER_MENTOR'])
        db.add(WeeklyReport(tenant_id=tid,internship_id=records[2].id,week_number=1,status='DRAFT'))
        for day,state in [('2026-09-09','NORMAL'),('2026-09-10','MOCK_LOCATION'),('2026-08-31','NORMAL')]:
            db.add(InternshipCheckin(tenant_id=tid,internship_id=records[0].id,checkin_date=day,result=state))
        for index in [0,0,1]:db.add(InternshipGuidance(tenant_id=tid,internship_id=records[index].id,student_id=records[index].student_id,advisor_name='同名统计教师',guidance_date=date(2026,9,9),status='NORMAL',content='合成指导记录'))
        for index in [0,2]:db.add(InternshipVisit(tenant_id=tid,internship_id=records[index].id,student_id=records[index].student_id,visit_at=datetime(2026,9,9),monthly_report='合成巡访记录'))
        db.add(InternshipChangeRequest(tenant_id=tid,internship_id=records[0].id,student_id=records[0].student_id,change_type='CHANGE_POSITION',reason='合成换岗事实',status='APPROVED'))
        for i in range(2):
            owner=batch
            if i:
                owner=InternshipBatch(tenant_id=tid,batch_name='追加方案批次',batch_no=s['tenantCode']+'-OTHER',status='RUNNING')
                db.add(owner);db.flush()
            plan=InternshipBatchPlan(tenant_id=tid,batch_id=owner.id,title='合成计划'+str(i),status='PUBLISHED')
            db.add(plan);db.flush();db.add(InternshipPlanAssignment(tenant_id=tid,internship_id=records[0].id,
                student_id=records[0].student_id,plan_id=plan.id,plan_batch_id=owner.id,status='ACTIVE'))
        db.commit();s.update(recordIds=[str(x.id) for x in records],collegeName=colleges[0].college_name,pendingReportId=str(pending))
    return s

@pytest.fixture
def story(record_property):
    s=seed_stats_story();record_property('tenantId',s['tenantId']);return s

def read(client,s,path,**params):return result(client.get(ROOT+path,params={'batchId':s['batchId'],**params}))
def metric(data,key):return next(x for x in data['metrics'] if x['key']==key)
def workbook(response):return list(load_workbook(BytesIO(b64decode(result(response)['contentBase64']))).active.values)

def test_current_submission_only_stable_reviewer_identity_and_attendance(story):
    with login(story) as school:
        rows=read(school,story,'/process-analytics')['items'];a=next(x for x in rows if x['studentNo'].startswith('000'))
        assert a['actualCheckins']==1
        teachers=read(school,story,'/process-analytics',groupBy='ADVISOR')['items']
        own=next(x for x in teachers if x['employeeNo']=='000000000000000091')
        assert all(x['college']=='合成信息学院' for x in teachers)
        assert own['studentCount']==2 and own['weeklyReviewed']==1 and own['weeklyReviewRate']==33.3
        heads=read(school,story,'/process-analytics',groupBy='HOMEROOM')['items'];assert sum(x['weeklyReviewed'] for x in heads)==2

def test_every_metric_matches_paginated_contributions_and_export(story):
    with login(story) as school:
        data=read(school,story,'/overview')
        assert metric(data,'guidanceCoverRate')['rate']==66.7 and metric(data,'visitCoverRate')['rate']==66.7
        assert metric(data,'weeklyReviewRate')['rate']==66.7
        for item in data['metrics']:
            if not item.get('drilldownAvailable'):continue
            for subset in ['numerator','denominator']:
                page=read(school,story,'/metrics/'+item['key']+'/drilldown',subset=subset,pageSize=2)
                assert page['totalValue']==item[subset]
                rows=page['items']
                if page['total']>2:rows+=read(school,story,'/metrics/'+item['key']+'/drilldown',subset=subset,pageSize=2,page=2)['items']
                assert sum(row['contribution'] for row in rows)==item[subset]
        values=workbook(school.post(ROOT+'/metrics/reportTaskCompletionRate/drilldown/export',params={'batchId':story['batchId'],'subset':'denominator'}))
        assert any(row[1]=='000000000000000000321' and row[5]==2 for row in values if len(row)>5)

@pytest.mark.parametrize('group',['STUDENT','ADVISOR','HOMEROOM','COLLEGE','MAJOR','CLASS'])
def test_org_filter_custom_columns_and_real_excel_share_the_same_population(story,group):
    with login(story) as school:
        data=read(school,story,'/process-analytics',groupBy=group,college=story['collegeName'])
        assert all(row.get('college')==story['collegeName'] for row in data['items'])
        cols=[x['key'] for x in data['availableColumns'][:4]]
        data=read(school,story,'/process-analytics',groupBy=group,college=story['collegeName'],columns=','.join(cols))
        values=workbook(school.post(ROOT+'/process-analytics/export',params={'batchId':story['batchId'],'groupBy':group,'college':story['collegeName'],'columns':','.join(cols)}))
        assert all(any(tuple(row[:4])==tuple(item[key] if item[key] is not None else None for key in cols) for row in values) for item in data['items'])

def test_quality_wages_keep_zero_currency_and_applied_org_scope(story):
    with get_sessionmaker()() as db:
        rec=db.get(InternshipRecord,int(story['recordIds'][1]))
        db.add(InternshipChangeRequest(tenant_id=int(story['tenantId']),internship_id=rec.id,student_id=rec.student_id,change_type='RESUME',reason='合成恢复事实不算转企转岗',status='APPROVED'));db.commit()
    with login(story) as school:
        data=read(school,story,'/procurement-overview',college=story['collegeName'])
        quality=next(x for x in data['groups'] if x['key']=='quality')
        assert quality['wageByCurrency']==[{'currency':'CNY','count':2,'averageActualAmount':617.28,'totalActualAmount':1234.56}]
        assert quality['ratios'][0]['rate']==50
        summary=read(school,story,'/process-analytics',college=story['collegeName'],groupBy='COLLEGE')
        assert summary['items'][0]['stabilityRate']==50
        values=workbook(school.post(ROOT+'/procurement-overview/export',params={'batchId':story['batchId'],'college':story['collegeName']}))
        assert any(row[1]=='实际工资合计' and row[2]==2 and row[4]==1234.56 for row in values if len(row)>4)

def test_normal_role_scope_exports_and_cross_school_denials(story):
    with login(story,'INTERN_MENTOR') as teacher,login(story,'COUNSELOR') as head,login(story,'COLLEGE_ADMIN') as college,login(story,'STUDENT') as student:
        for actor in [teacher,head,college]:
            data=read(actor,story,'/process-analytics');assert data['total']==2
            assert actor.post(ROOT+'/process-analytics/export',params={'batchId':story['batchId']}).status_code==200
        assert student.get(ROOT+'/overview',params={'batchId':story['batchId']}).status_code==403
        assert student.post(ROOT+'/process-analytics/export',params={'batchId':story['batchId']}).status_code==403
        with get_sessionmaker()() as db:
            scope=db.scalar(select(TeacherStudentScope).where(TeacherStudentScope.tenant_id==int(story['tenantId']),TeacherStudentScope.role_code=='COLLEGE_ADMIN'));scope.status='REVOKED';db.commit()
        assert read(college,story,'/process-analytics')['total']==0
    foreign=seed_support()
    with login(foreign) as outsider:
        assert outsider.get(ROOT+'/overview',params={'batchId':story['batchId']}).status_code==404
    with get_sessionmaker()() as db:
        foreign_org=College(tenant_id=int(foreign['tenantId']),college_name='外校合成组织不应被当前学校读取')
        db.add(foreign_org);db.flush();student=db.get(StudentProfile,int(story['studentId']))
        student.college_id=foreign_org.id;student.major_id=None;student.class_id=None;db.commit()
    with login(story) as school:
        for path in ['/overview','/dimensions','/procurement-overview','/process-analytics','/trends']:
            response=school.get(ROOT+path,params={'batchId':story['batchId']})
            assert response.status_code==200 and foreign_org.college_name not in response.text

def test_zero_review_sla_does_not_revert_to_default_and_multi_plan_does_not_duplicate(story):
    from app.modules.internship.services.internship_process_statistics_service import _as_date
    assert _as_date(datetime(2026,9,7,16,30))==date(2026,9,8)
    with get_sessionmaker()() as db:
        batch=db.get(InternshipBatch,int(story['batchId']));rules=dict(batch.rules_config);rules['weeklyReport']={**rules['weeklyReport'],'reviewSlaHours':0};batch.rules_config=rules;db.commit()
    with login(story) as school:
        data=read(school,story,'/process-analytics',groupBy='ADVISOR')
        assert sum(row['studentCount'] for row in data['items'])==3
        assert all(row['weeklyTimelyReviewRate'] is None for row in data['items'])

def test_export_audit_failure_returns_no_file_and_wrong_payroll_owner_is_excluded(story,monkeypatch):
    from app.modules.internship.services import internship_export_util
    with login(story) as school:
        with get_sessionmaker()() as db:
            statements=db.scalars(select(InternshipPayrollStatement).where(InternshipPayrollStatement.tenant_id==int(story['tenantId'])).order_by(InternshipPayrollStatement.id)).all()
            statements[0].current_version_id=statements[2].current_version_id;db.commit()
        quality=next(x for x in read(school,story,'/procurement-overview')['groups'] if x['key']=='quality')
        assert quality['wageRecordCount']==2 and quality['wageByCurrency'][0]['totalActualAmount']==0
        def fail(*args,**kwargs):raise RuntimeError('synthetic required audit failure')
        monkeypatch.setattr(internship_export_util,'write_export_audit',fail)
        with pytest.raises(RuntimeError,match='synthetic required audit failure'):
            school.post(ROOT+'/process-analytics/export',params={'batchId':story['batchId']})

def test_export_exceeds_screen_page_keeps_long_numbers_zero_scores_and_formula_safety(story):
    tid=int(story['tenantId'])
    with get_sessionmaker()() as db:
        extra=[StudentProfile(tenant_id=tid,student_no='ADDITIONAL-'+str(i),real_name='=SUM(1,2)' if i==0 else '合成分页学生'+str(i),status='ACTIVE') for i in range(501)]
        db.add_all(extra);db.flush()
        db.add_all([InternshipRecord(tenant_id=tid,student_id=student.id,batch_id=int(story['batchId']),advisor_user_id=story['users']['INTERN_MENTOR'],
            status='ONBOARD',intern_start_date=datetime(2026,9,1),intern_end_date=datetime(2026,9,30)) for student in extra]);db.commit()
    with login(story) as school:
        page=read(school,story,'/process-analytics',pageSize=500)
        assert page['total']==504 and len(page['items'])==500
        response=school.post(ROOT+'/process-analytics/export',params={'batchId':story['batchId'],'columns':'studentNo,studentName,totalScore'})
        data=result(response);assert data['rowCount']==504
        assert '全部时间' in data['filename']
        sheet=load_workbook(BytesIO(b64decode(data['contentBase64']))).active
        assert ' · ALL · ' not in str(list(sheet.values))
        header=next([cell.value for cell in row] for row in sheet if '学号' in [cell.value for cell in row])
        number,name,score=[header.index(label) for label in ['学号','姓名','成绩分数']]
        row=next(row for row in sheet if row[number].value=='000000000000000000321');assert row[number].data_type=='s'
        assert any(row[number].value=='STATS-1' and row[score].value==0 for row in sheet)
        formula=next(row for row in sheet if row[number].value=='ADDITIONAL-0');assert formula[name].data_type!='f'

def test_trends_share_direct_org_links_and_school_local_month_boundaries(story,monkeypatch):
    from app.core import timeutil
    monkeypatch.setattr(timeutil,'local_now',lambda:datetime(2026,10,14))
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(story['recordId']));student=db.get(StudentProfile,record.student_id)
        group=db.get(SchoolClass,student.class_id);major=db.get(Major,group.major_id)
        student.college_id=major.college_id;student.major_id=major.id;student.class_id=None
        report=db.get(WeeklyReport,int(story['pendingReportId']));report.submitted_at=datetime(2026,8,31,16,30);db.commit()
    with login(story) as school:
        data=read(school,story,'/trends',college=story['collegeName'])
        reports=next(row for row in data['series'] if row['key']=='reports')['points']
        assert next(row['value'] for row in reports if row['month']=='2026-09')==3
        assert next(row['value'] for row in reports if row['month']=='2026-08')==0
        assert read(school,story,'/overview',college=story['collegeName'])['counters'][0]['value']==2

def test_legacy_approved_without_review_time_keeps_progress_without_fake_timeliness(story):
    with get_sessionmaker()() as db:
        legacy=db.scalar(select(WeeklyReport).where(WeeklyReport.internship_id==int(story['recordIds'][2])))
        legacy.status='APPROVED';legacy.reviewed_at=None;legacy.submitted_at=None;db.commit()
    with login(story) as school:
        overall=read(school,story,'/overview');assert metric(overall,'weeklyReviewRate')['numerator']==3
        data=read(school,story,'/process-analytics',groupBy='HOMEROOM')
        assert sum(row['weeklyReviewed'] for row in data['items'])==3
        legacy_class=next(row for row in data['items'] if row['className']=='合成机电一班')
        assert legacy_class['weeklyTimelyReviewRate'] is None
