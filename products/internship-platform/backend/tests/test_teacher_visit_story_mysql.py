"""正常账号巡访执行、原记录及附件回读；仅隔离自动测试库。"""
from datetime import date
from concurrent.futures import ThreadPoolExecutor
from uuid import uuid4

import pytest

from sqlalchemy import func, select
from sqlalchemy.exc import SQLAlchemyError

from app.db.session import get_sessionmaker
from app.models import InternshipVisitPlan, InternshipRecord, InternshipVisit, StudentProfile, User, InternshipAuditTrail
from app.models.file import FileBinding, FileObject
from test_yiyang_write_flows_mysql import flow, login, upload, png_bytes, seed_flow, PASSWORD

TARGETS = '/api/v1/teacher-mobile/internship/visit-targets'
VISITS = '/api/v1/teacher-mobile/internship/visits/'


def file_paths(file_id):
    return ('/api/v1/files/' + file_id, '/api/v1/files/preview/' + file_id,
            '/api/v1/files/download/' + file_id, '/api/v1/files/' + file_id + '/url')


def denied_file(client, file_id):
    for path in file_paths(file_id):
        response = client.get(path)
        assert response.status_code == 404 and not response.json().get('data'), response.text


def ok(response):
    assert response.status_code == 200 and response.json()['code'] == 0, response.text
    return response.json()['data']


def plan_fixture(data, **fields):
    # 合成测试前置计划，不代替真实页面的下发或巡访办理。
    with get_sessionmaker()() as db:
        rec = db.get(InternshipRecord, int(data['recordId']))
        student = db.get(StudentProfile, rec.student_id)
        owner = db.get(User, data['users']['INTERN_MENTOR'])
        plan = InternshipVisitPlan(tenant_id=rec.tenant_id, batch_id=rec.batch_id,
            owner_name=owner.real_name, enterprise_name=rec.enterprise_name,
            student_scope=student.student_no, plan_date=date.today().isoformat(),
            objective='合成联调巡访执行', status='PUBLISHED', plan_type='VISIT')
        for name, value in fields.items():
            setattr(plan, name, value)
        db.add(plan); db.commit()
        return str(plan.id)


def target(client, data, plan_id):
    result = ok(client.get(TARGETS, params={'batchId': data['batchId'], 'planId': plan_id,
                                          'internshipId': data['recordId']}))
    return next(s for p in result['plans'] for s in p['students'])


def body(data, plan_id, current, **fields):
    return {'batchId': data['batchId'], 'planId': plan_id,
        'expectedPlanVersion': current['expectedPlanVersion'], 'expectedVersion': current['expectedVersion'],
        'visitType': 'PHONE', 'contactPerson': '合成企业导师', 'workStatus': '合成学生在岗状态正常',
        'enterpriseFeedback': '合成沟通反馈岗位适应良好',
        'facts': '合成巡访联调，核对岗位任务及安全要求，不构成真实巡访。', **fields}


def test_normal_teacher_reads_current_batch_plan_and_saves_bound_evidence(flow):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        result = ok(teacher.get(TARGETS, params={'batchId': flow['batchId']}))
        plan = next(p for p in result['plans'] if p['id'] == plan_id)
        student = next(s for s in plan['students'] if s['internshipId'] == flow['recordId'])
        assert student['resolvable'] and not student['visited']
        proof = upload(teacher, name='visit-proof.png', biz='INTERNSHIP_VISIT')
        saved = ok(teacher.post(VISITS + flow['recordId'], json={
            'batchId': flow['batchId'], 'planId': plan_id, 'expectedPlanVersion': student['expectedPlanVersion'],
            'visitType': 'PHONE', 'contactPerson': '合成企业导师',
            'workStatus': '合成学生在岗状态正常', 'enterpriseFeedback': '合成沟通反馈岗位适应良好',
            'facts': '合成巡访联调，核对岗位任务及安全要求，不构成真实巡访。',
            'fileIds': [proof['fileId']], 'expectedVersion': student['expectedVersion']}))
        detail = ok(school.get('/api/v1/internship/visits/' + saved['visitId']))
        assert detail['internId'] == flow['recordId'] and '合成巡访联调' in detail['monthlyReport']
        assert teacher.get('/api/v1/files/download/' + proof['fileId']).content == png_bytes()
        assert school.get('/api/v1/files/download/' + proof['fileId']).status_code == 200
        reread = ok(teacher.get(TARGETS, params={'batchId': flow['batchId']}))
        current = next(s for p in reread['plans'] if p['id'] == plan_id for s in p['students'])
        assert current['visited'] and current['expectedVersion'] == saved['expectedVersion']
    with get_sessionmaker()() as db:
        visit = db.get(InternshipVisit, int(saved['visitId']))
        assert visit.student_id == int(flow['studentId']) and visit.file_id == proof['fileId']
        bound = db.scalars(select(FileBinding).where(FileBinding.tenant_id == int(flow['tenantId']),
            FileBinding.file_id == int(proof['fileId']), FileBinding.status == 'ACTIVE',
            FileBinding.is_current.is_(True), FileBinding.is_deleted.is_(False))).all()
        assert len(bound) == 1 and bound[0].biz_type == 'INTERNSHIP_VISIT' and bound[0].biz_id == saved['visitId']
        assert bound[0].student_id == int(flow['studentId']) and bound[0].batch_id == flow['batchId']
        assert bound[0].scope_json['internshipId'] == flow['recordId']
        file = db.get(FileObject, int(proof['fileId']))
        assert file.biz_type == 'INTERNSHIP_VISIT' and file.biz_id == saved['visitId']
        assert db.scalar(select(InternshipAuditTrail.id).where(
            InternshipAuditTrail.tenant_id == int(flow['tenantId']),
            InternshipAuditTrail.target_type == 'VISIT',
            InternshipAuditTrail.target_id == visit.id,
            InternshipAuditTrail.action == 'CREATE_MOBILE_EVIDENCE'))


@pytest.mark.parametrize('change', ['record_version', 'plan_version', 'cancelled', 'student_scope', 'enterprise'])
def test_opened_plan_and_record_cannot_be_silently_replaced(flow, change):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher:
        current = target(teacher, flow, plan_id)
        with get_sessionmaker()() as db:
            plan = db.get(InternshipVisitPlan, int(plan_id))
            rec = db.get(InternshipRecord, int(flow['recordId']))
            if change == 'record_version': rec.version += 1
            elif change == 'plan_version': plan.version += 1
            elif change == 'cancelled': plan.status = 'CANCELLED'
            elif change == 'student_scope': plan.student_scope = '不存在的测试学号'
            elif change == 'enterprise': rec.enterprise_name = '另一个合成企业'
            db.commit()
        response = teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current))
        assert response.status_code == (403 if change == 'student_scope' else 409), response.text
    with get_sessionmaker()() as db:
        assert not db.scalar(select(InternshipVisit.id).where(InternshipVisit.tenant_id == int(flow['tenantId'])))


def test_batch_and_schema_are_required_and_student_cannot_register(flow):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow, 'STUDENT') as student:
        current = target(teacher, flow, plan_id)
        assert teacher.get(TARGETS).status_code == 422
        for field in ['batchId', 'expectedVersion', 'expectedPlanVersion']:
            payload = body(flow, plan_id, current); payload.pop(field)
            assert teacher.post(VISITS + flow['recordId'], json=payload).status_code == 422
        for field in ['expectedVersion', 'expectedPlanVersion']:
            assert teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, **{field: True})).status_code == 422
        assert teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, location={'latitude': 1})).status_code == 422
        wrong = teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, batchId=flow['batches'][1]['id']))
        assert wrong.status_code == 400 and wrong.json()['bizCode'] == 'DATA_CONFLICT'
        assert teacher.get(TARGETS, params={'batchId': flow['batches'][1]['id']}).status_code == 403
        assert student.get(TARGETS, params={'batchId': flow['batchId']}).status_code == 403
        assert student.post(VISITS + flow['recordId'], json=body(flow, plan_id, current)).status_code == 403


def test_other_plan_is_not_marked_visited_and_same_version_is_not_replayed(flow):
    one, two = plan_fixture(flow), plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher:
        current = target(teacher, flow, one)
        saved = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, one, current)))
        assert target(teacher, flow, one)['visited'] is True
        second = target(teacher, flow, two)
        assert second['visited'] is False and second['expectedVersion'] == saved['expectedVersion']
        assert teacher.post(VISITS + flow['recordId'], json=body(flow, one, current)).status_code == 409


@pytest.mark.parametrize('brackets',[('（','）'),('(',')')])
def test_bare_name_ambiguity_is_denied_but_exact_student_number_survives(flow,brackets):
    with get_sessionmaker()() as db:
        original = db.get(StudentProfile, int(flow['studentId']))
        rec = db.get(InternshipRecord, int(flow['recordId']))
        duplicate = StudentProfile(tenant_id=rec.tenant_id, real_name=original.real_name,
            student_no='DUP-' + uuid4().hex[:12], current_stage='ENROLLED', student_status='NORMAL', status='ACTIVE')
        db.add(duplicate); db.flush()
        db.add(InternshipRecord(tenant_id=rec.tenant_id, student_id=duplicate.id, batch_id=rec.batch_id,
            advisor_user_id=flow['users']['SCHOOL_ADMIN'], enterprise_name=rec.enterprise_name, status='ONBOARD'))
        name, number = original.real_name, original.student_no
        db.commit()
    ambiguous = plan_fixture(flow, student_scope=name)
    exact = plan_fixture(flow, student_scope=f'{name}{brackets[0]}{number}{brackets[1]}')
    with login(flow, 'INTERN_MENTOR') as teacher:
        result = ok(teacher.get(TARGETS, params={'batchId': flow['batchId'], 'planId': ambiguous}))
        assert result['plans'] == []
        current = target(teacher, flow, exact)
        denied = teacher.post(VISITS + flow['recordId'], json=body(flow, ambiguous, current))
        assert denied.status_code == 403
        assert current['studentNo'] == number


def test_file_owner_readiness_and_binding_are_enforced_in_same_transaction(flow):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        current = target(teacher, flow, plan_id)
        other = upload(school, biz='INTERNSHIP_VISIT')
        denied = teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, fileIds=[other['fileId']]))
        assert denied.status_code == 404
        pending = upload(teacher, biz='INTERNSHIP_VISIT')
        with get_sessionmaker()() as db:
            file = db.get(FileObject, int(pending['fileId']))
            file.status, file.scan_status = 'QUARANTINED', 'PENDING'; db.commit()
        denied = teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, fileIds=[pending['fileId']]))
        assert denied.status_code == 404
        assert target(teacher, flow, plan_id)['expectedVersion'] == current['expectedVersion']
    with get_sessionmaker()() as db:
        assert not db.scalar(select(InternshipVisit.id).where(InternshipVisit.tenant_id == int(flow['tenantId'])))
        assert not db.scalar(select(FileBinding.id).where(FileBinding.file_id == int(pending['fileId'])))


@pytest.mark.parametrize('change', ['advisor', 'binding'])
def test_bound_photo_uses_current_scope_even_with_legacy_file_purpose(flow, change):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school, login(flow, 'STUDENT') as student:
        current = target(teacher, flow, plan_id)
        photo = upload(teacher, biz='INTERNSHIP_VISIT')
        saved = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, fileIds=[photo['fileId']])))
        denied_file(student, photo['fileId'])
        for reader in (teacher, school):
            for path in file_paths(photo['fileId']):
                assert reader.get(path).status_code == 200
        with get_sessionmaker()() as db:
            file = db.get(FileObject, int(photo['fileId'])); file.biz_type = 'ATTACHMENT'
            if change == 'advisor': db.get(InternshipRecord, int(flow['recordId'])).advisor_user_id = flow['users']['SCHOOL_ADMIN']
            else:
                binding = db.scalar(select(FileBinding).where(FileBinding.file_id == file.id)); binding.status = 'INVALIDATED'
            db.commit()
        denied_file(teacher, photo['fileId'])
        if change != 'binding': assert school.get('/api/v1/files/download/' + photo['fileId']).status_code == 200


def test_bound_photo_denies_missing_visit_permission_other_teacher_and_foreign_school(flow):
    from app.core.security import hash_password
    from app.models import Role, TeacherStudentScope, UserRole
    foreign = seed_flow()
    with get_sessionmaker()() as db:
        tenant = int(flow['tenantId'])
        leader_role = db.scalar(select(Role).where(Role.tenant_id==tenant,Role.role_code=='PROFESSIONAL_LEADER'))
        mentor_role = db.scalar(select(Role).where(Role.tenant_id == tenant, Role.role_code == 'INTERN_MENTOR'))
        leader = User(tenant_id=tenant, login_name='professional_leader', real_name='合成专业负责人', password_hash=hash_password(PASSWORD),
                      user_type='PROFESSIONAL_LEADER', status='ACTIVE', credential_version=0, must_change_password=False)
        other = User(tenant_id=tenant, login_name='other_mentor', real_name='合成非指导教师', password_hash=hash_password(PASSWORD),
                     user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
        db.add_all([leader, other]); db.flush()
        db.add_all([UserRole(tenant_id=tenant, user_id=leader.id, role_id=leader_role.id, status='ACTIVE'),
                    UserRole(tenant_id=tenant, user_id=other.id, role_id=mentor_role.id, status='ACTIVE'),
                    TeacherStudentScope(tenant_id=tenant, user_id=leader.id, teacher_key='db-' + str(leader.id), role_code='PROFESSIONAL_LEADER',
                                        scope_type='STUDENT', ref_id=int(flow['studentId']), ref_value='S0', status='ACTIVE')]); db.commit()
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow, 'PROFESSIONAL_LEADER') as leader, login(flow, 'OTHER_MENTOR') as other, login(foreign) as outsider:
        current = target(teacher, flow, plan_id)
        proof = upload(teacher, biz='INTERNSHIP_VISIT')
        ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, fileIds=[proof['fileId']])))
        assert leader.get('/api/v1/internship/students', params={'batchId': flow['batchId']}).status_code == 200
        assert leader.get('/api/v1/internship/visits', params={'batchId': flow['batchId']}).status_code == 403
        for reader in (leader, other, outsider):
            denied_file(reader, proof['fileId'])


@pytest.mark.parametrize('surface', ['school', 'mobile'])
def test_binding_risk_todo_audit_and_record_version_roll_back_after_late_failure(flow, monkeypatch, surface):
    from app.models import RiskRecord, UnifiedTodo
    from app.modules.internship.services import internship_todo_helper as ix_todo
    plan_id = plan_fixture(flow)
    models = (InternshipVisit, FileBinding, InternshipAuditTrail, RiskRecord, UnifiedTodo)
    def snapshot():
        with get_sessionmaker()() as db:
            rec = db.get(InternshipRecord, int(flow['recordId']))
            return ([db.scalar(select(func.count()).select_from(model).where(model.tenant_id == int(flow['tenantId']))) for model in models],
                    rec.version, rec.risk_level)
    called = []
    original = ix_todo.push_visit_rectify_todo
    def fail_after_formal_todo(db, visit, rec):
        original(db, visit, rec); db.flush(); called.append(str(visit.id))
        assert db.scalar(select(FileBinding.id).where(FileBinding.file_id == int(visit.file_id), FileBinding.biz_type == 'INTERNSHIP_VISIT'))
        raise SQLAlchemyError('synthetic late transaction failure')
    with login(flow, 'INTERN_MENTOR') as teacher:
        current = target(teacher, flow, plan_id)
        proof = upload(teacher, biz='INTERNSHIP_VISIT')
        before = snapshot()
        monkeypatch.setattr(ix_todo, 'push_visit_rectify_todo', fail_after_formal_todo)
        try:
            response = teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, fileIds=[proof['fileId']],
                title='合成晚失败巡访', location=visit_location(),
                needFollow=True, needRisk=True, riskLevel='HIGH', riskReason='合成晚失败风险验证')) if surface == 'mobile' else teacher.post(
                '/api/v1/internship/visits', json={'internshipId': flow['recordId'], 'fileId': proof['fileId'], 'safetyIssue': '合成整改事项',
                'title': '合成晚失败巡访', 'location': visit_location()})
        except SQLAlchemyError:
            pass
        else:
            assert response.status_code >= 500 or response.json().get('code') != 0
        assert called and snapshot() == before
        with get_sessionmaker()() as db:
            file = db.get(FileObject, int(proof['fileId']))
            assert file.biz_type == 'TEMP_PRIVATE' and file.biz_id is None
        assert teacher.get('/api/v1/files/download/' + proof['fileId']).content == png_bytes()


def test_concurrent_same_opened_version_commits_one_visit(flow):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as first, login(flow, 'INTERN_MENTOR') as second:
        current = target(first, flow, plan_id)
        with ThreadPoolExecutor(max_workers=2) as pool:
            responses = list(pool.map(lambda client: client.post(VISITS + flow['recordId'], json=body(flow, plan_id, current)), [first, second]))
        assert sorted(r.status_code for r in responses) == [200, 409]
    with get_sessionmaker()() as db:
        visits = db.scalars(select(InternshipVisit).where(InternshipVisit.tenant_id == int(flow['tenantId']))).all()
        assert len(visits) == 1


def test_school_creates_exact_student_plan_and_teacher_executes_that_plan(flow):
    with login(flow) as school, login(flow, 'INTERN_MENTOR') as teacher:
        plan = ok(school.post('/api/v1/internship/visit-plans', json={
            'batchId': flow['batchId'], 'recordIds': [flow['recordId']], 'objective': '合成学校正常下发巡访',
            'planDate': date.today().isoformat(), 'studentScope': '客户端伪造学号', 'collaborators': '客户端伪造教师'}))
        assert plan['status'] == 'DRAFT' and plan['studentScope'] != '客户端伪造学号'
        published = ok(school.post('/api/v1/internship/visit-plans/' + plan['id'] + '/transition',
            json={'action': 'PUBLISH', 'expectedVersion': plan['version']}))
        current = target(teacher, flow, plan['id'])
        assert current['expectedPlanVersion'] == published['version']
        saved = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan['id'], current)))
        assert ok(school.get('/api/v1/internship/visits/' + saved['visitId']))['internId'] == flow['recordId']
        denied = school.post('/api/v1/internship/visit-plans', json={
            'batchId': flow['batchId'], 'recordIds': [flow['recordId'], '9223372036854775807'], 'objective': '不能创建缺失学生计划'})
        assert denied.status_code == 403


def test_pc_visit_uses_same_file_binding_and_retains_legacy_record_readability(flow):
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        file = upload(teacher, biz='INTERNSHIP_VISIT')
        visit = ok(teacher.post('/api/v1/internship/visits', json={
            'internshipId': flow['recordId'], 'method': 'VIDEO', 'monthlyReport': '合成学校巡访记录', 'fileId': file['fileId']}))
        detail = ok(school.get('/api/v1/internship/visits/' + visit['id']))
        assert detail['attachment']['fileId'] == file['fileId'] and detail['methodLabel'] == '视频'
        with get_sessionmaker()() as db:
            binding = db.scalar(select(FileBinding).where(FileBinding.file_id == int(file['fileId'])))
            binding.status = 'INVALIDATED'; db.commit()
        detail = ok(school.get('/api/v1/internship/visits/' + visit['id']))
        assert detail['monthlyReport'] == '合成学校巡访记录' and detail['attachment'] is None and detail['attachmentError']
        assert school.get('/api/v1/files/download/' + file['fileId']).status_code == 404


def test_followup_todo_opens_exact_visit_and_completion_is_versioned(flow):
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher:
        current = target(teacher, flow, plan_id)
        visit = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current,
            needFollow=True, needRisk=True, riskLevel='HIGH', riskReason='合成检查发现需要跟进的风险')))
        todos = ok(teacher.get('/api/v1/teacher-mobile/internship/workbench/todos'))['items']
        todo = next(t for t in todos if t['sourceBizId'] == visit['visitId'] and t['todoType'] == 'INTERN_VISIT_RECTIFY')
        assert todo['actionPath'].endswith('?tab=visit&visitId=' + visit['visitId'])
        detail = ok(teacher.get('/api/v1/internship/visits/' + visit['visitId']))
        uri = '/api/v1/internship/visits/' + visit['visitId'] + '/rectify'
        assert teacher.post(uri, json={'status': 'DONE', 'note': '合成整改已完成', 'batchId': flow['batchId']}).status_code == 400
        follow = ok(teacher.post(uri, json={'batchId': flow['batchId'], 'expectedVersion': detail['version'],
            'status': 'PENDING', 'note': '合成跟进中，仍需继续整改'}))
        assert teacher.post(uri, json={'batchId': flow['batchId'], 'expectedVersion': detail['version'],
            'status': 'DONE', 'note': '不能覆盖新版整改'}).status_code == 409
        done = ok(teacher.post(uri, json={'batchId': flow['batchId'], 'expectedVersion': follow['version'],
            'status': 'DONE', 'note': '合成整改已完成并核对'}))
        assert done['rectifyStatus'] == 'DONE'
        assert not any(t['todoType'] == 'INTERN_VISIT_RECTIFY' and t['sourceBizId'] == visit['visitId']
            for t in ok(teacher.get('/api/v1/teacher-mobile/internship/workbench/todos'))['items'])
    from app.models import RiskRecord, UnifiedTodo
    with get_sessionmaker()() as db:
        risk = db.get(RiskRecord, int(visit['riskId']))
        assert risk.source_id == int(visit['visitId']) and risk.risk_level == 'HIGH'
        assert db.get(InternshipRecord, int(flow['recordId'])).risk_level == 'HIGH'
        assert db.scalar(select(UnifiedTodo.status).where(UnifiedTodo.id == int(todo['todoId']))) == 'DONE'


def test_plan_and_student_pages_keep_totals_and_exact_second_page_target(flow):
    with get_sessionmaker()() as db:
        rec = db.get(InternshipRecord, int(flow['recordId'])); original = db.get(StudentProfile, rec.student_id)
        numbers = [original.student_no]
        for index in range(20):
            student = StudentProfile(tenant_id=rec.tenant_id, student_no=f'PAGE-{index:03}', real_name=f'合成巡访学生{index}',
                current_stage='ENROLLED', student_status='NORMAL', status='ACTIVE')
            db.add(student); db.flush(); numbers.append(student.student_no)
            db.add(InternshipRecord(tenant_id=rec.tenant_id, student_id=student.id, batch_id=rec.batch_id,
                advisor_user_id=rec.advisor_user_id, enterprise_name=rec.enterprise_name, status='ONBOARD'))
        db.commit()
    ids = [plan_fixture(flow, student_scope=','.join(numbers)) for _ in range(21)]
    with login(flow, 'INTERN_MENTOR') as teacher:
        first = ok(teacher.get(TARGETS, params={'batchId': flow['batchId']}))
        second = ok(teacher.get(TARGETS, params={'batchId': flow['batchId'], 'page': 2}))
        assert len(first['plans']) == 20 and first['total'] == 21 and first['hasMore']
        assert len(second['plans']) == 1 and not second['hasMore']
        assert {p['id'] for p in first['plans'] + second['plans']} == set(ids)
        page1 = first['plans'][0]
        page2 = ok(teacher.get(TARGETS, params={'batchId': flow['batchId'], 'planId': page1['id'], 'studentPage': 2}))['plans'][0]
        assert page1['studentTotal'] == 21 and len(page1['students']) == 20 and page1['studentHasMore']
        assert len(page2['students']) == 1 and not page2['studentHasMore']
        exact = ok(teacher.get(TARGETS, params={'batchId': flow['batchId'], 'planId': page1['id'],
            'internshipId': page2['students'][0]['internshipId']}))['plans'][0]
        assert exact['students'][0]['internshipId'] == page2['students'][0]['internshipId']


def visit_location():
    # 合成坐标，仅检查合同与权限，不构成实地定位证明。
    return {'consent': True, 'latitude': 28.1234567, 'longitude': 112.7654321,
            'accuracyM': 12.5, 'coordinateSystem': 'WGS84', 'capturedAt': '2026-10-06T02:00:00Z'}


def test_visit_title_encrypted_location_and_audited_disclosure(flow):
    import base64, json
    from io import BytesIO
    from openpyxl import load_workbook
    from app.core.field_crypto import decrypt_field
    from app.models import AuditOutbox
    from app.modules.internship.services import internship_evidence_package_service as packages
    plan_id, capture = plan_fixture(flow), visit_location()
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        current = target(teacher, flow, plan_id)
        saved = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current,
            title='  合成巡访标题  ', location=capture)))
        uri = '/api/v1/internship/visits/' + saved['visitId']
        detail = ok(teacher.get(uri))
        listing = ok(school.get('/api/v1/internship/visits', params={'batchId': flow['batchId']}))
        assert detail['title'] == '合成巡访标题' and detail['hasLocation'] and detail['canViewLocation']
        assert next(x for x in listing['items'] if x['id'] == saved['visitId'])['title'] == detail['title']
        with get_sessionmaker()() as db:
            visit = db.get(InternshipVisit, int(saved['visitId']))
            cipher = visit.location_encrypted
            assert visit.visit_title == detail['title'] and cipher.startswith('gAAAA')
            assert json.loads(decrypt_field(cipher, allow_legacy_plaintext=False)) == capture
            safe = packages._safe_row(visit)
            assert safe['has_location'] and 'location_encrypted' not in safe
        ordinary = json.dumps([detail, listing, safe], ensure_ascii=False)
        assert cipher not in ordinary and 'latitude' not in ordinary and 'longitude' not in ordinary
        exported = ok(school.post('/api/v1/internship/visits/export', params={'batchId': flow['batchId']}))
        assert exported['rowCount'] == 1
        book = load_workbook(BytesIO(base64.b64decode(exported['contentBase64'])), read_only=True, data_only=True)
        try:
            cells = [value for row in book.active.iter_rows(values_only=True) for value in row if value is not None]
        finally:
            book.close()
        assert all(v in cells for v in ('巡访标题', '合成巡访标题', '位置登记', '已登记'))
        text = json.dumps(cells, ensure_ascii=False, default=str)
        assert cipher not in text and all(str(capture[k]) not in text for k in ('latitude', 'longitude'))
        disclosed = ok(teacher.post(uri + '/location/disclose', json={'batchId': flow['batchId'],
            'expectedVersion': detail['version'], 'reason': '核对本次合成巡访登记位置'}))
        assert disclosed == {'visitId': saved['visitId'], 'batchId': flow['batchId'], 'location': capture}
        assert 'location' not in ok(school.get(uri))
    with get_sessionmaker()() as db:
        trails = db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id == int(flow['tenantId']),
            InternshipAuditTrail.target_id == int(saved['visitId']), InternshipAuditTrail.action == 'DISCLOSE_LOCATION')).all()
        events = db.scalars(select(AuditOutbox).where(AuditOutbox.tenant_id == int(flow['tenantId']),
            AuditOutbox.event_type == 'INTERNSHIP_DISCLOSE_LOCATION')).all()
        assert len(trails) == len(events) == 1
        assert trails[0].detail_json['reason'] == '核对本次合成巡访登记位置'
        assert trails[0].detail_json['actorUserId'] == 'db-' + str(flow['users']['INTERN_MENTOR'])
        assert trails[0].detail_json['actorRole'] == 'INTERN_MENTOR'
        assert events[0].payload_json['targetId'] == saved['visitId']
        audit = json.dumps([trails[0].detail_json, events[0].payload_json], ensure_ascii=False)
        assert 'latitude' not in audit and 'longitude' not in audit and cipher not in audit


def test_visit_location_denials_context_and_legacy_record(flow):
    from app.core.security import hash_password
    from app.models import Role, TeacherStudentScope, UserRole
    foreign = seed_flow()
    with get_sessionmaker()() as db:
        tenant = int(flow['tenantId']); stu = db.get(StudentProfile, int(flow['studentId']))
        mentor = db.scalar(select(Role).where(Role.tenant_id == tenant, Role.role_code == 'INTERN_MENTOR'))
        role = db.scalar(select(Role).where(Role.tenant_id==tenant,Role.role_code=='COUNSELOR'))
        counselor = User(tenant_id=tenant, login_name='counselor', real_name='合成辅导员', password_hash=hash_password(PASSWORD),
            user_type='COUNSELOR', status='ACTIVE', credential_version=0, must_change_password=False)
        other = User(tenant_id=tenant, login_name='other_mentor', real_name='合成非指导教师', password_hash=hash_password(PASSWORD),
            user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
        db.add_all([counselor, other]); db.flush()
        db.add_all([UserRole(tenant_id=tenant, user_id=counselor.id, role_id=role.id, status='ACTIVE'),
            UserRole(tenant_id=tenant, user_id=other.id, role_id=mentor.id, status='ACTIVE'),
            TeacherStudentScope(tenant_id=tenant, user_id=counselor.id, teacher_key='db-' + str(counselor.id), role_code='COUNSELOR',
                scope_type='STUDENT', ref_id=stu.id, ref_value=stu.student_no, status='ACTIVE')]); db.commit()
    plan_id = plan_fixture(flow)
    with login(flow, 'INTERN_MENTOR') as teacher:
        current = target(teacher, flow, plan_id)
        for invalid_title in (True, 123, {'x': '不能作为标题'}):
            response = teacher.post('/api/v1/internship/visits', json={'internshipId': flow['recordId'], 'title': invalid_title})
            assert response.status_code == 400 and response.json()['bizCode'] == 'VALIDATION_ERROR'
            assert teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current, title=invalid_title)).status_code == 422
        for key, value in (('consent', False), ('latitude', 91.0), ('accuracyM', -1.0), ('capturedAt', '2026-10-06T02:00:00')):
            assert teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current,
                location={**visit_location(), key: value})).status_code == 422
        saved = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, current,
            title='合成权限巡访', location=visit_location(), needFollow=True)))
        uri = '/api/v1/internship/visits/' + saved['visitId']; detail = ok(teacher.get(uri))
        payload = {'batchId': flow['batchId'], 'expectedVersion': detail['version'], 'reason': '核对本次合成巡访登记位置'}
        with login(flow, 'COUNSELOR') as counselor, login(flow, 'OTHER_MENTOR') as other, login(flow, 'STUDENT') as student, login(foreign) as outsider:
            assert ok(counselor.get(uri))['hasLocation']
            for reader in (counselor, other, student):
                assert reader.post(uri + '/location/disclose', json=payload).status_code == 403
            assert outsider.post(uri + '/location/disclose', json=payload).status_code == 404
        wrong = teacher.post(uri + '/location/disclose', json={**payload, 'batchId': flow['batches'][1]['id']})
        assert wrong.status_code == 400 and wrong.json()['bizCode'] == 'DATA_CONFLICT'
        ok(teacher.post(uri + '/rectify', json={'batchId': flow['batchId'], 'expectedVersion': detail['version'],
            'status': 'PENDING', 'note': '合成跟进仍需继续核对'}))
        assert teacher.post(uri + '/location/disclose', json=payload).status_code == 409
        legacy = ok(teacher.post('/api/v1/internship/visits', json={'internshipId': flow['recordId'], 'monthlyReport': '合成旧入口无定位记录'}))
        old = ok(teacher.get('/api/v1/internship/visits/' + legacy['id']))
        assert old['title'] == '' and old['hasLocation'] is False
        assert teacher.post('/api/v1/internship/visits/' + legacy['id'] + '/location/disclose',
            json={**payload, 'expectedVersion': old['version']}).status_code == 404
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(InternshipAuditTrail).where(
            InternshipAuditTrail.tenant_id == int(flow['tenantId']), InternshipAuditTrail.action == 'DISCLOSE_LOCATION')) == 0


def test_visit_location_disclosure_rolls_back_after_required_audit_failure(flow, monkeypatch):
    from app.models import AuditOutbox
    from app.modules.internship.services import internship_visit_service as service
    plan_id = plan_fixture(flow); called = []
    with login(flow, 'INTERN_MENTOR') as teacher:
        saved = ok(teacher.post(VISITS + flow['recordId'], json=body(flow, plan_id, target(teacher, flow, plan_id),
            title='合成审计失败巡访', location=visit_location())))
        uri = '/api/v1/internship/visits/' + saved['visitId']; detail = ok(teacher.get(uri))
        def snapshot():
            with get_sessionmaker()() as db:
                visit = db.get(InternshipVisit, int(saved['visitId']))
                return (visit.version, visit.location_encrypted, *[db.scalar(select(func.count()).select_from(m).where(
                    m.tenant_id == int(flow['tenantId']))) for m in (InternshipAuditTrail, AuditOutbox)])
        before, original = snapshot(), service._trail
        def fail_after_audit_flush(db, vid, action, detail=None, operator='系统'):
            original(db, vid, action, detail, operator); db.flush(); called.append(action)
            assert db.scalar(select(InternshipAuditTrail.id).where(InternshipAuditTrail.target_id == int(vid),
                InternshipAuditTrail.action == 'DISCLOSE_LOCATION'))
            raise SQLAlchemyError('synthetic required disclosure audit failure')
        with monkeypatch.context() as patch:
            patch.setattr(service, '_trail', fail_after_audit_flush)
            try:
                response = teacher.post(uri + '/location/disclose', json={'batchId': flow['batchId'],
                    'expectedVersion': detail['version'], 'reason': '合成必要审计失败检查'})
            except SQLAlchemyError:
                pass
            else:
                assert response.status_code >= 500 and 'latitude' not in response.text and 'longitude' not in response.text
        assert called == ['DISCLOSE_LOCATION'] and snapshot() == before


def test_visit_enterprise_literal_filter_keeps_plan_and_student_pagination(flow):
    with get_sessionmaker()() as db:
        rec = db.get(InternshipRecord, int(flow['recordId'])); stu = db.get(StudentProfile, rec.student_id)
        rec.enterprise_name = '合成%_公司'; numbers = [stu.student_no]; ids = []
        for index, company, owner in ((1, '合成%_公司', flow['users']['INTERN_MENTOR']),
            (2, '合成普通公司', flow['users']['INTERN_MENTOR']), (3, '合成%_公司', flow['users']['SCHOOL_ADMIN'])):
            person = StudentProfile(tenant_id=rec.tenant_id, student_no=f'FILTER-{index}', real_name=f'合成筛选学生{index}',
                current_stage='ENROLLED', student_status='NORMAL', status='ACTIVE')
            db.add(person); db.flush(); numbers.append(person.student_no)
            row = InternshipRecord(tenant_id=rec.tenant_id, student_id=person.id, batch_id=rec.batch_id, advisor_user_id=owner,
                advisor_name=db.get(User, owner).real_name, enterprise_name=company, status='ONBOARD')
            db.add(row); db.flush(); ids.append(str(row.id))
        db.commit()
    matching = plan_fixture(flow, enterprise_name=None, student_scope=','.join(numbers))
    empty = plan_fixture(flow, enterprise_name=None, student_scope=numbers[2])
    with login(flow, 'INTERN_MENTOR') as teacher:
        params = {'batchId': flow['batchId'], 'keyword': '%_', 'pageSize': 1, 'studentPageSize': 1}
        first = ok(teacher.get(TARGETS, params=params)); second = ok(teacher.get(TARGETS, params={**params, 'page': 2}))
        assert first['total'] == second['total'] == 2 and first['hasMore'] and not second['hasMore']
        assert first['plans'][0]['id'] == empty and first['plans'][0]['studentTotal'] == 0
        matched = second['plans'][0]
        assert matched['id'] == matching and matched['studentTotal'] == 2 and matched['studentHasMore']
        more = ok(teacher.get(TARGETS, params={**params, 'planId': matching, 'studentPage': 2}))['plans'][0]
        assert more['studentTotal'] == 2 and not more['studentHasMore']
        assert {s['internshipId'] for s in matched['students'] + more['students']} == {flow['recordId'], ids[0]}
        ok(teacher.post(VISITS + flow['recordId'], json=body(flow, matching, target(teacher, flow, matching), title='合成企业检索巡访')))
        for needle in ('%', '_', '%_'):
            found = ok(teacher.get(TARGETS, params={'batchId': flow['batchId'], 'planId': matching, 'keyword': needle, 'studentPageSize': 1}))
            assert found['total'] == 1
            assert found['plans'][0]['studentTotal'] == 2 and found['plans'][0]['visitedStudentCount'] == 1
            assert found['plans'][0]['students'][0]['enterpriseName'] == '合成%_公司'
