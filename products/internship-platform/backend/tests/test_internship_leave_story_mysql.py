"""请假原单跨端、材料版本和销假交接：普通账号与隔离 MySQL。"""
from concurrent.futures import ThreadPoolExecutor
from datetime import date, datetime
from threading import Event

import pytest
from sqlalchemy import select

from app.core.context import set_tenant
from app.core.security import hash_password
from app.db.session import get_sessionmaker
from app.models import InternshipAuditTrail, InternshipBatch, InternshipCheckin, InternshipLeave, InternshipRecord, RiskRecord, UnifiedTodo, Role, TeacherStudentScope, User, UserRole
from app.models.file import FileBinding
from test_yiyang_write_flows_mysql import flow, context, login, png_bytes, seed_flow, upload, PASSWORD

PC = '/api/v1/portal/internship/context/leaves'
MOBILE = '/api/v1/mobile/internship/context/leaves'
TEACHER = '/api/v1/mobile/teacher/internship/context/leaves'
LEAVE_KINDS = {'LEAVE_REVIEW', 'LEAVE_RETURN'}


def values(data, **extra):
    return {**context(data), 'leaveType': 'SICK', 'startDate': '2026-10-03',
            'endDate': '2026-10-05', 'reason': '合成测试：需就诊并补充证明', **extra}


def ok(response):
    assert response.status_code == 200 and response.json()['code'] == 0, response.text
    return response.json()['data']


def frozen(data, row, **extra):
    return {'batchId': data['batchId'], 'expectedVersion': row['version'],
            'evidenceFileId': row.get('evidenceFileId', ''), **extra}


def review(teacher, data, row, **extra):
    return teacher.post(TEACHER + '/' + row['id'] + '/review',
        params={'batchId': data['batchId']}, json=frozen(data, row, **extra))


def conflict(response):
    assert response.status_code in (400, 409) and response.json()['bizCode'] == 'DATA_CONFLICT', response.text


def persisted(data, leave_id):
    with get_sessionmaker()() as db:
        row = db.get(InternshipLeave, int(leave_id))
        trails = db.scalars(select(InternshipAuditTrail).where(
            InternshipAuditTrail.tenant_id == int(data['tenantId']),
            InternshipAuditTrail.target_type == 'LEAVE',
            InternshipAuditTrail.target_id == row.id).order_by(InternshipAuditTrail.id)).all()
        todos = db.scalars(select(UnifiedTodo).where(
            UnifiedTodo.tenant_id == int(data['tenantId']), UnifiedTodo.source_biz_id == row.id,
            UnifiedTodo.source_module == 'internship').order_by(UnifiedTodo.id)).all()
        return (row.status, row.version, row.file_id,
                [(t.action, t.detail_json) for t in trails], [(t.todo_type, t.status) for t in todos])


def dashboard_leaves(client, data, leave_id, kind=None):
    dashboard = ok(client.get('/api/v1/internship/dashboard', params={'batchId': data['batchId']}))
    items = [x for x in dashboard['workItems'] if x['kind'] in LEAVE_KINDS and x['objectId'] == leave_id]
    if kind is None:
        assert items == []
    else:
        assert len(items) == 1 and items[0]['kind'] == kind
        panel = 'pending' if kind == 'LEAVE_REVIEW' else 'return'
        assert items[0]['route'] == f'/admin/internship/leaves?panel={panel}&id={leave_id}&batchId={data["batchId"]}'
        assert items[0]['internshipId'] == data['recordId']
    return dashboard


def return_clock(monkeypatch):
    from zoneinfo import ZoneInfo
    from app.modules.internship.services import internship_student_leave_context_service as leaves
    from app.modules.internship.services import internship_student_checkin_service as checkins
    from app.modules.internship.services import internship_process_statistics_service as stats
    from app.modules.internship.services import internship_student_dashboard_service as dashboard
    class Clock(datetime):
        @classmethod
        def utcnow(cls):
            return datetime(2026, 10, 5, 16, 30)  # 学校当地已到 10 月 6 日，UTC 仍为前日。
    monkeypatch.setattr(leaves, 'datetime', Clock)
    local = datetime(2026, 10, 6, 0, 30, tzinfo=ZoneInfo('Asia/Shanghai'))
    monkeypatch.setattr(stats, 'local_now', lambda: local)
    monkeypatch.setattr(dashboard, 'local_now', lambda: local)
    monkeypatch.setattr(checkins, '_local_clock', lambda name=None: (
        ZoneInfo('Asia/Shanghai'), local, 480))


def approve_period(student, teacher, data, surface=PC, **extra):
    proof = upload(student, biz='INTERNSHIP_LEAVE')
    row = ok(student.post(surface, json=values(data, **{'startDate': '2026-10-05',
        'endDate': '2026-10-08', 'fileId': proof['fileId'], **extra})))
    row = ok(teacher.get('/api/v1/internship/leaves/' + row['id'], params={'batchId': data['batchId']}))
    assert teacher.get('/api/v1/files/download/' + proof['fileId']).status_code == 200
    ok(teacher.post(TEACHER + '/' + row['id'] + '/evidence-viewed', json=frozen(data, row)))
    return {**row, **ok(review(teacher, data, row, action='APPROVE', comment='合成证明已核验'))}


@pytest.mark.parametrize('surface', [PC, MOBILE])
def test_early_return_releases_day_and_preserves_calendar_and_quantity(flow, monkeypatch, surface):
    from app.modules.internship.services.internship_compliance_facts import material_quantity_facts
    return_clock(monkeypatch)
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        row = approve_period(student, teacher, flow, surface)
        initial = ok(student.get('/api/v1/mobile/internship/context/my', params={'batchId': flow['batchId']}))
        assert initial['todayCheckin']['done'] is False and initial['todayCheckin']['totalDays'] == 0
        returned = ok(student.post(surface + '/' + row['id'] + '/return', json={
            **context(flow), 'expectedVersion': row['version'], 'note': '合成测试：提前返岗即刻恢复出勤'}))
        credential = ok(student.post('/api/v1/mobile/internship/checkin/preflight',
            params={'batchId': flow['batchId'], 'timezoneName': 'Asia/Shanghai'}))
        checkin_body = {'idempotencyKey': 'early-return-day', 'timezoneName': 'Asia/Shanghai',
                        'lat': 28.2, 'lng': 112.9, 'accuracy': 20, 'checkinToken': credential['token']}
        signed = ok(student.post('/api/v1/mobile/internship/checkin', params={'batchId': flow['batchId']},
                                json=checkin_body))
        assert signed['date'] == '2026-10-06' and signed['result'] == 'RECORDED'
        replay = ok(student.post('/api/v1/mobile/internship/checkin', params={'batchId': flow['batchId']},
                                json=checkin_body))
        assert replay['id'] == signed['id'] and replay['idempotentReplay']
        calendar = ok(student.get('/api/v1/mobile/internship/checkin/calendar', params={
            'batchId': flow['batchId'], 'month': '2026-10', 'timezoneName': 'Asia/Shanghai'}))
        days = {x['date']: x for x in calendar['days']}
        assert days['2026-10-05']['status'] == 'LEAVE'
        assert days['2026-10-05']['time'] is None
        assert days['2026-10-06']['status'] == 'CHECKIN'
        assert days['2026-10-07']['status'] == days['2026-10-08']['status'] == 'FUTURE'
        stats = ok(school.get('/api/v1/internship/stats/process-analytics', params={
            'batchId': flow['batchId'], 'period': 'WEEK', 'anchor': '2026-10-06'}))['items'][0]
        assert stats['leaveDays'] == 1 and stats['actualCheckins'] == 1 and stats['expectedCheckins'] == 1
        dashboard = ok(student.get('/api/v1/mobile/internship/context/my', params={'batchId': flow['batchId']}))
        assert dashboard['todayCheckin']['done'] and dashboard['todayCheckin']['totalDays'] == 1
        for target in (PC, MOBILE):
            pending_list = ok(student.get(target, params=context(flow)))
            pending_rows = pending_list if isinstance(pending_list, list) else pending_list['items']
            assert next(item for item in pending_rows if item['id'] == row['id'])['returnAcknowledged'] is False
        ok(teacher.post(TEACHER + '/' + row['id'] + '/ack-return', json=frozen(flow, returned, note='已核实提前返岗')))
        for target in (PC, MOBILE):
            final = ok(student.get(target + '/' + row['id'], params=context(flow)))
            assert final['returnAcknowledged'] and final['startDate'] == '2026-10-05' and final['endDate'] == '2026-10-08'
            completed_list = ok(student.get(target, params=context(flow)))
            completed_rows = completed_list if isinstance(completed_list, list) else completed_list['items']
            assert next(item for item in completed_rows if item['id'] == row['id'])['returnAcknowledged'] is True
    set_tenant(int(flow['tenantId']))
    with get_sessionmaker()() as db:
        records = db.scalars(select(InternshipCheckin).where(InternshipCheckin.internship_id == int(flow['recordId']))).all()
        assert len(records) == 4 and len([x for x in records if x.checkin_date == '2026-10-06']) == 1
        assert {x.checkin_date for x in records if not x.is_deleted} == {'2026-10-05', '2026-10-06'}
        facts = material_quantity_facts(db, db.get(InternshipRecord, int(flow['recordId'])), db.get(InternshipBatch, int(flow['batchId'])))
        assert facts['leaveDaysDeducted'] == 1


def test_return_preserves_other_leave_real_attendance_and_reuses_reapproved_day(flow, monkeypatch):
    return_clock(monkeypatch)
    # 仅隔离自动测试的已有签到前置数据；请假审批、销假和再次签到均走正式账号命令。
    with get_sessionmaker()() as db:
        real = InternshipCheckin(tenant_id=int(flow['tenantId']), internship_id=int(flow['recordId']),
            checkin_date='2026-10-08', checkin_at=datetime(2026, 10, 8), result='NORMAL', note='保留原真实签到')
        db.add(real); db.commit(); real_id = real.id
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        first = approve_period(student, teacher, flow)
        overlapping = approve_period(student, teacher, flow, startDate='2026-10-07', endDate='2026-10-07')
        ok(student.post(PC + '/' + first['id'] + '/return', json={
            **context(flow), 'expectedVersion': first['version'], 'note': '提前返岗，后日另有有效请假'}))
        with get_sessionmaker()() as db:
            rows = {x.checkin_date: x for x in db.scalars(select(InternshipCheckin).where(
                InternshipCheckin.internship_id == int(flow['recordId']))).all()}
            assert rows['2026-10-06'].is_deleted and not rows['2026-10-07'].is_deleted
            assert rows['2026-10-08'].id == real_id and rows['2026-10-08'].result == 'NORMAL'
        again = approve_period(student, teacher, flow, startDate='2026-10-06', endDate='2026-10-06')
        conflict(student.post('/api/v1/mobile/internship/checkin', params={'batchId': flow['batchId']}, json={}))
        ok(student.post(PC + '/' + again['id'] + '/return', json={
            **context(flow), 'expectedVersion': again['version'], 'note': '本日请假亦已销假'}))
        ok(student.post('/api/v1/mobile/internship/checkin', params={'batchId': flow['batchId']},
                        json={'idempotencyKey': 'reapproved-day'}))
        calendar = ok(student.get('/api/v1/mobile/internship/checkin/calendar', params={
            'batchId': flow['batchId'], 'month': '2026-10'}))
        days = {x['date']: x for x in calendar['days']}
        assert days['2026-10-05']['status'] == days['2026-10-07']['status'] == 'LEAVE'
        assert days['2026-10-06']['status'] == days['2026-10-08']['status'] == 'CHECKIN'


def test_overdue_return_task_is_not_hidden_behind_four_older_returned_leaves(flow, monkeypatch):
    from app.modules.internship.services import internship_leave_service as service
    return_clock(monkeypatch)
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        for _ in range(4):
            row = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', startDate='2026-10-01', endDate='2026-10-01')))
            approved = ok(review(teacher, flow, row, action='APPROVE'))
            ok(student.post(PC + '/' + row['id'] + '/return', json={
                **context(flow), 'expectedVersion': approved['version'], 'note': '合成历史返岗待确认'}))
        urgent = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', startDate='2026-10-05', endDate='2026-10-05')))
        ok(review(teacher, flow, urgent, action='APPROVE'))
        set_tenant(int(flow['tenantId']))
        assert service.refresh_overdue(reference_date=date(2026, 10, 6), system=True)['markedOverdue'] == 1
        dashboard = dashboard_leaves(teacher, flow, urgent['id'], 'LEAVE_RETURN')
        returns = [x for x in dashboard['workItems'] if x['kind'] == 'LEAVE_RETURN']
        assert returns[0]['objectId'] == urgent['id']


def test_early_return_day_can_be_made_up_without_duplicate_date_row(flow, monkeypatch):
    return_clock(monkeypatch)
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        row = approve_period(student, teacher, flow)
        ok(student.post(PC + '/' + row['id'] + '/return', json={
            **context(flow), 'expectedVersion': row['version'], 'note': '已提前返岗，漏签由导师核实补录'}))
        makeup = ok(student.post('/api/v1/portal/internship/context/makeups', json={
            **context(flow), 'checkinDate': '2026-10-06', 'makeupType': 'MISSING', 'reason': '合成测试：提前返岗当天忘记签到'}))
        ok(school.post('/api/v1/internship/makeups/' + makeup['id'] + '/approve',
            json={'expectedVersion': makeup['version'], 'comment': '核实合成漏签事实'}))
        calendar = ok(student.get('/api/v1/mobile/internship/checkin/calendar', params={'batchId': flow['batchId'], 'month': '2026-10'}))
        assert next(x for x in calendar['days'] if x['date'] == '2026-10-06')['status'] == 'MAKEUP'
    with get_sessionmaker()() as db:
        rows = db.scalars(select(InternshipCheckin).where(InternshipCheckin.internship_id == int(flow['recordId']),
            InternshipCheckin.checkin_date == '2026-10-06')).all()
        assert len(rows) == 1 and not rows[0].is_deleted and rows[0].result == 'RECORDED'


def test_concurrent_overlapping_approval_cannot_be_erased_by_return(flow, monkeypatch):
    from app.modules.internship.services import internship_leave_service as service
    return_clock(monkeypatch)
    paused, release, review_started, review_locked = Event(), Event(), Event(), Event()
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        first = approve_period(student, teacher, flow)
        proof = upload(student, biz='INTERNSHIP_LEAVE')
        second = ok(student.post(PC, json=values(flow, startDate='2026-10-06', endDate='2026-10-08', fileId=proof['fileId'])))
        second = ok(teacher.get('/api/v1/internship/leaves/' + second['id'], params={'batchId': flow['batchId']}))
        assert teacher.get('/api/v1/files/download/' + proof['fileId']).status_code == 200
        ok(teacher.post(TEACHER + '/' + second['id'] + '/evidence-viewed', json=frozen(flow, second)))
        original = service._release_leave_checkins
        def paused_release(db, leave):
            if str(leave.id) == first['id']:
                paused.set()
                assert release.wait(15)
            return original(db, leave)
        monkeypatch.setattr(service, '_release_leave_checkins', paused_release)
        original_lock = service.lock_attendance_record
        def observed_lock(db, record_id):
            reviewing = paused.is_set() and not release.is_set()
            if reviewing:
                review_started.set()
            record = original_lock(db, record_id)
            if reviewing:
                review_locked.set()
            return record
        monkeypatch.setattr(service, 'lock_attendance_record', observed_lock)
        def approve_second():
            return review(teacher, flow, second, action='APPROVE', comment='核实第二份重叠请假')
        with ThreadPoolExecutor(max_workers=2) as pool:
            returning = pool.submit(student.post, PC + '/' + first['id'] + '/return', json={
                **context(flow), 'expectedVersion': first['version'], 'note': '并发测试：首单提前返岗'})
            try:
                assert paused.wait(15)
                approving = pool.submit(approve_second)
                assert review_started.wait(15)
                assert not review_locked.wait(0.2), '同一实习记录的审批应等待返岗事务释放记录锁'
            finally:
                release.set()
            ok(returning.result(timeout=20)); ok(approving.result(timeout=20))
        conflict(student.post('/api/v1/mobile/internship/checkin', params={'batchId': flow['batchId']}, json={}))
        with get_sessionmaker()() as db:
            rows = db.scalars(select(InternshipCheckin).where(InternshipCheckin.internship_id == int(flow['recordId']),
                InternshipCheckin.checkin_date >= '2026-10-06')).all()
            assert len(rows) == 3 and all(not x.is_deleted and x.result == 'LEAVE' for x in rows)


def test_mismatched_owned_record_is_rejected_before_second_record_lock(flow, monkeypatch):
    from app.modules.internship.services import internship_leave_service as service
    with get_sessionmaker()() as db:
        second = db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id == int(flow['batches'][1]['id'])))
        second.student_id = int(flow['studentId']); second.status = 'ONBOARD'
        second.intern_start_date, second.intern_end_date = datetime(2026, 9, 1), datetime(2026, 12, 30)
        db.get(InternshipBatch, second.batch_id).end_date = datetime(2026, 12, 30)
        db.commit(); second_id = str(second.id)
    with login(flow, 'STUDENT') as student:
        row = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', endDate='2026-10-03')))
        calls, original = [], service.lock_attendance_record
        def observed_lock(db, record_id):
            calls.append(str(record_id))
            return original(db, record_id)
        monkeypatch.setattr(service, 'lock_attendance_record', observed_lock)
        denied = student.post(PC + '/' + row['id'] + '/withdraw', json={
            'batchId': flow['batches'][1]['id'], 'internshipId': second_id, 'expectedVersion': row['version']})
        assert denied.status_code == 403 and calls == []
        assert persisted(flow, row['id'])[0] == 'PENDING'


@pytest.mark.parametrize('first_surface,second_surface', [(PC, MOBILE), (MOBILE, PC)])
def test_original_leave_return_resubmit_approve_and_return_ack(flow, first_surface, second_surface):
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        first = upload(student, biz='INTERNSHIP_LEAVE')
        assert teacher.get('/api/v1/files/download/' + first['fileId']).status_code == 404
        row = ok(student.post(first_surface, json=values(flow, fileId=first['fileId'])))
        dashboard_leaves(teacher, flow, row['id'], 'LEAVE_REVIEW')
        url = TEACHER + '/' + row['id']
        queued = ok(teacher.get(TEACHER, params={'batchId': flow['batchId']}))['items']
        pending = next(item for item in queued if item['id'] == row['id'])
        assert pending['batchId'] == flow['batchId']
        assert teacher.get('/api/v1/files/download/' + first['fileId']).content == png_bytes()
        ok(teacher.post(url + '/evidence-viewed', json=frozen(flow, pending)))
        rejected = ok(review(teacher, flow, pending, action='REJECT', comment='请补充清晰的本次就诊证明'))
        dashboard_leaves(teacher, flow, row['id'])
        original = ok(student.get(second_surface + '/' + row['id'], params=context(flow)))
        assert original['status'] == 'REJECTED' and original['version'] == rejected['version']
        second = upload(student, name='corrected-leave.png', biz='INTERNSHIP_LEAVE')
        updated = ok(student.post(second_surface + '/' + row['id'] + '/resubmit', json=values(
            flow, fileId=second['fileId'], expectedVersion=original['version'], reason='合成测试：按意见补充就诊说明')))
        assert updated['id'] == row['id']
        dashboard_leaves(teacher, flow, row['id'], 'LEAVE_REVIEW')
        current = ok(student.get(first_surface + '/' + row['id'], params=context(flow)))
        assert current['status'] == 'PENDING' and not current['evidenceViewed']
        assert current['submissionHistory'][0]['reason'] != current['reason']
        assert current['submissionHistory'][0]['evidenceFileId'] == first['fileId']
        assert current['submissionHistory'][0]['reviewComment'] == original['reviewComment']
        assert current['submissionHistory'][0]['reviewAt'] == original['reviewAt']
        before = persisted(flow, row['id'])
        conflict(teacher.post(url + '/evidence-viewed', json=frozen(flow, pending)))
        conflict(review(teacher, flow, current, action='APPROVE'))
        assert persisted(flow, row['id']) == before
        assert teacher.get('/api/v1/files/download/' + second['fileId']).content == png_bytes()
        viewed = ok(teacher.post(url + '/evidence-viewed', json=frozen(flow, current)))
        assert viewed['version'] == current['version'] and viewed['evidenceFileId'] == second['fileId']
        approved = ok(review(teacher, flow, current, action='APPROVE', comment='证明已核验'))
        dashboard_leaves(teacher, flow, row['id'])
        returned = ok(student.post(second_surface + '/' + row['id'] + '/return', json={
            **context(flow), 'expectedVersion': approved['version'], 'note': '合成测试：已返岗'}))
        dashboard_leaves(teacher, flow, row['id'], 'LEAVE_RETURN')
        todos = ok(teacher.get('/api/v1/teacher-mobile/internship/workbench/todos'))['items']
        todo = next(t for t in todos if t['todoType'] == 'INTERN_LEAVE_RETURN')
        assert todo['typeLabel'] == '返岗确认' and todo['group'] == 'LEAVE'
        assert todo['actionPath'].endswith('?tab=overdue&leaveId=' + row['id'])
        queue = ok(teacher.get(TEACHER + '/overdue', params={'batchId': flow['batchId']}))['list']
        assert [x['id'] for x in queue] == [row['id']]
        ack = ok(teacher.post(url + '/ack-return', json=frozen(flow, returned, note='已核实合成测试返岗')))
        dashboard_leaves(teacher, flow, row['id'])
        assert ok(teacher.get(TEACHER + '/overdue', params={'batchId': flow['batchId']}))['list'] == []
        for surface in (PC, MOBILE):
            final = ok(student.get(surface + '/' + row['id'], params=context(flow)))
            assert final['status'] == 'RETURNED' and final['version'] == ack['version']
            assert final['returnAcknowledged'] and len(final['submissionHistory']) == 2
            assert any(t['action'] == 'REVIEW_REJECT' for t in final['auditTrail'])
        conflict(teacher.post(url + '/ack-return', json=frozen(flow, ack, note='重复确认返岗')))
        with get_sessionmaker()() as db:
            assert len(db.scalars(select(InternshipLeave).where(InternshipLeave.internship_id == int(flow['recordId']))).all()) == 1
            bound = db.scalars(select(FileBinding).where(FileBinding.file_id.in_([int(first['fileId']), int(second['fileId'])]))).all()
            assert len(bound) == 2 and all(b.biz_type == 'INTERNSHIP_LEAVE' and b.biz_id == row['id']
                and b.student_id == int(flow['studentId']) and b.batch_id == flow['batchId']
                and b.scope_json['internshipId'] == flow['recordId'] for b in bound)
            checkins = db.scalars(select(InternshipCheckin).where(InternshipCheckin.internship_id == int(flow['recordId']))).all()
            assert len(checkins) == 3 and all(x.result == 'LEAVE' for x in checkins)
        for reader in (student, teacher):
            for proof in (first, second):
                assert reader.get('/api/v1/files/download/' + proof['fileId']).content == png_bytes()
        assert all(status == 'DONE' for _, status in persisted(flow, row['id'])[-1])


def test_leave_query_and_full_export_share_fields_scope_filters_and_durable_audit(flow, monkeypatch):
    from base64 import b64decode
    from io import BytesIO
    from openpyxl import load_workbook
    from app.models import StudentProfile
    from app.core.exceptions import AppException
    from app.modules.internship.services import internship_export_util as exports
    foreign = seed_flow()
    with login(flow, 'STUDENT') as student, login(flow) as school, login(flow, 'INTERN_MENTOR') as teacher, login(foreign) as outsider:
        proof = upload(student, name='synthetic-leave-proof.png', biz='INTERNSHIP_LEAVE')
        original = ok(student.post(PC, json=values(flow, fileId=proof['fileId'])))
        with get_sessionmaker()() as db:
            profile = db.get(StudentProfile, int(flow['studentId']))
            profile.student_no = '00090071992547409930'
            profile.real_name = '=合成查询学生'
            db.get(InternshipRecord, int(flow['recordId'])).enterprise_name = '+合成实习企业'
            db.get(InternshipLeave, int(original['id'])).reason = '=SUM(1,2)'
            # 仅准备隔离查询样本，不把这些状态计入界面办理验收。
            rows = [InternshipLeave(tenant_id=int(flow['tenantId']), internship_id=int(flow['recordId']),
                student_id=int(flow['studentId']), leave_type='PERSONAL', start_date='2026-10-10',
                end_date='2026-10-10', days=1, reason='合成查询样本', status='WITHDRAWN', version=0) for _ in range(24)]
            db.add_all(rows); db.flush()
            db.commit()
        params = {'batchId': flow['batchId'], 'keyword': '00090071992547409930'}
        first = ok(school.get('/api/v1/internship/leaves', params={**params, 'pageSize': 20}))
        second = ok(school.get('/api/v1/internship/leaves', params={**params, 'page': 2, 'pageSize': 20}))
        assert first['total'] == second['total'] == 25
        assert len(first['items']) == 20 and len(second['items']) == 5
        all_rows = first['items'] + second['items']
        row = next(item for item in all_rows if item['id'] == original['id'])
        assert row['enterpriseName'] == '+合成实习企业' and row['evidenceFileName'] == proof['fileName']
        exported = ok(school.post('/api/v1/internship/leaves/export', params=params))
        assert exported['rowCount'] == exported['totalRows'] == exported['exportedRows'] == 25
        assert exported['truncated'] is False and exported['safeExportMax'] == 10000
        sheet = load_workbook(BytesIO(b64decode(exported['contentBase64']))).active
        headers = [cell.value for cell in sheet[2]]
        required = {'姓名', '学号', '实习企业', '开始日期', '结束日期', '请假原因', '证明文件', '老师意见', '原请假单号'}
        assert required.issubset(headers)
        values_by_id = {str(row[-1].value): row for row in list(sheet.rows)[2:]}
        assert set(values_by_id) == {item['id'] for item in all_rows}
        cells = values_by_id[original['id']]
        assert cells[0].value == '00090071992547409930' and cells[0].data_type == 's'
        assert cells[1].value == "'=合成查询学生" and cells[2].value == "'+合成实习企业"
        assert cells[headers.index('请假原因')].value == "'=SUM(1,2)"
        assert cells[headers.index('证明文件')].value == proof['fileName']
        assert not any(cell.data_type == 'f' for row in sheet for cell in row)
        assert cells[headers.index('请假原因')].alignment.wrap_text
        assert sheet.column_dimensions['I'].width == 58 and sheet.column_dimensions['J'].width == 34
        assert sheet.freeze_panes == 'B3' and sheet.row_dimensions[cells[0].row].height >= 24
        assert teacher.post('/api/v1/internship/leaves/export', params=params).status_code == 403
        assert not ok(school.get('/api/v1/internship/leaves', params={'batchId': flow['batches'][1]['id']}))['items']
        assert outsider.post('/api/v1/internship/leaves/export', params=params).json()['code'] != 0
        filtered = ok(school.post('/api/v1/internship/leaves/export', params={**params, 'status': 'PENDING'}))
        assert filtered['rowCount'] == 0
        assert ok(school.get('/api/v1/internship/leaves', params={**params, 'status': 'PENDING'}))['total'] == 0
        assigned = ok(teacher.get('/api/v1/internship/leaves', params={**params, 'status': 'PENDING'}))
        assert assigned['total'] == 1 and assigned['items'][0]['id'] == original['id']
        with get_sessionmaker()() as db:
            audits = db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id == int(flow['tenantId']),
                InternshipAuditTrail.target_type == 'EXPORT').order_by(InternshipAuditTrail.id)).all()
            assert len(audits) == 2 and all(item.action == 'SENSITIVE_EXPORT' for item in audits)
            assert audits[0].detail_json['totalRows'] == 25
            assert audits[0].detail_json['actorUserId'] and audits[0].detail_json['actorRole'] == 'SCHOOL_ADMIN'
            assert audits[0].detail_json['filters'] == {**params, 'status': '', 'returnPending': False}
        def unavailable(*args, **kwargs):
            raise AppException('SERVICE_UNAVAILABLE', '合成测试：导出审计暂不可写', http_status=503)
        monkeypatch.setattr(exports, 'write_export_audit', unavailable)
        denied = school.post('/api/v1/internship/leaves/export', params=params)
        assert denied.status_code == 503 and not denied.json().get('data')


@pytest.mark.parametrize('mismatch', [None, 'tenant_id', 'record_id', 'batch_id', 'company_id', 'position_id'])
def test_leave_enterprise_uses_only_its_current_matching_placement(flow, mismatch):
    from base64 import b64decode
    from io import BytesIO
    from openpyxl import load_workbook
    from app.models.internship_placement_snapshot import InternshipPlacementSnapshot
    with login(flow, 'STUDENT') as student, login(flow) as school, login(flow, 'INTERN_MENTOR') as teacher:
        leave = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', endDate='2026-10-03')))
        with get_sessionmaker()() as db:
            record = db.get(InternshipRecord, int(flow['recordId']))
            # 合成投影样本，仅验证关系边界，不计作正式落岗或界面办理。
            record.enterprise_name = '旧岗位名称不应回退'
            record.enterprise_id, record.position_id = 90001, 90002
            fields = dict(tenant_id=record.tenant_id, record_id=record.id, batch_id=record.batch_id,
                          company_id=record.enterprise_id, position_id=record.position_id)
            if mismatch:
                fields[mismatch] += 1000000
            snapshot = InternshipPlacementSnapshot(**fields, placement_seq=1, snapshot_version=1,
                company_name='当前正式落岗企业', position_title='合成岗位', position_version=0,
                snapshot_json={}, snapshot_hash='0' * 64, snapshot_sha256='0' * 64,
                placement_at=datetime.utcnow(), captured_at=datetime.utcnow())
            db.add(snapshot); db.flush()
            record.current_placement_snapshot_id = snapshot.id
            db.commit()
        expected = '' if mismatch else '当前正式落岗企业'
        params = {'batchId': flow['batchId']}
        listed = ok(school.get('/api/v1/internship/leaves', params=params))['items']
        assert next(row for row in listed if row['id'] == leave['id'])['enterpriseName'] == expected
        assert ok(school.get('/api/v1/internship/leaves/' + leave['id'], params=params))['enterpriseName'] == expected
        assert ok(student.get(PC + '/' + leave['id'], params=context(flow)))['enterpriseName'] == expected
        assert ok(student.get(PC, params=context(flow)))[0]['enterpriseName'] == expected
        exported = ok(school.post('/api/v1/internship/leaves/export', params=params))
        sheet = load_workbook(BytesIO(b64decode(exported['contentBase64']))).active
        assert (sheet.cell(3, 3).value or '') == expected
        rejected = ok(review(teacher, flow, leave, action='REJECT', comment='请补充原请假说明'))
        resubmitted = ok(student.post(PC + '/' + leave['id'] + '/resubmit', json=values(
            flow, leaveType='PERSONAL', endDate='2026-10-03', expectedVersion=rejected['version'])))
        assert resubmitted['id'] == leave['id'] and resubmitted['enterpriseName'] == expected


def test_return_query_pagination_and_export_exclude_acknowledged_and_invalid_relations(flow):
    with get_sessionmaker()() as db:
        rows = [InternshipLeave(tenant_id=int(flow['tenantId']), internship_id=int(flow['recordId']),
            student_id=int(flow['studentId']), leave_type='PERSONAL', start_date='2026-10-01', end_date='2026-10-01',
            days=1, reason='合成返岗查询样本', status='RETURNED' if i % 2 else 'OVERDUE', version=1) for i in range(23)]
        db.add_all(rows); db.flush()
        expected = {str(row.id) for row in rows[1:]}
        db.add(InternshipAuditTrail(tenant_id=int(flow['tenantId']), target_type='LEAVE', target_id=rows[0].id,
            action='ACK_OVERDUE_RETURN_VERSIONED', operator_name='合成准备', detail_json={}, occurred_at=datetime.utcnow()))
        other = db.scalar(select(InternshipRecord).where(InternshipRecord.tenant_id == int(flow['tenantId']), InternshipRecord.id != int(flow['recordId'])))
        db.add(InternshipLeave(tenant_id=int(flow['tenantId']), internship_id=int(flow['recordId']), student_id=other.student_id,
            leave_type='PERSONAL', start_date='2026-10-01', end_date='2026-10-01', days=1, reason='错误关联样本', status='RETURNED', version=1))
        db.commit()
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        params = {'batchId': flow['batchId'], 'returnPending': True}
        first = ok(teacher.get('/api/v1/internship/leaves', params=params))
        second = ok(teacher.get('/api/v1/internship/leaves', params={**params, 'page': 2}))
        assert first['total'] == second['total'] == 22
        assert {item['id'] for item in first['items'] + second['items']} == expected
        mobile_first = ok(teacher.get(TEACHER + '/overdue', params={'batchId': flow['batchId']}))
        mobile_second = ok(teacher.get(TEACHER + '/overdue', params={'batchId': flow['batchId'], 'page': 2}))
        assert mobile_first['total'] == 22 and len(mobile_first['list']) == 20 and mobile_first['hasMore']
        assert len(mobile_second['list']) == 2 and not mobile_second['hasMore']
        assert {item['id'] for item in mobile_first['list'] + mobile_second['list']} == expected
        exported = ok(school.post('/api/v1/internship/leaves/export', params=params))
        assert exported['rowCount'] == 22


@pytest.mark.parametrize('legacy_purpose', ['ATTACHMENT', 'INTERNSHIP_CHECKIN', 'INTERNSHIP'])
def test_dashboard_and_proof_respect_leave_permission_advisor_batch_and_school(flow, legacy_purpose):
    foreign = seed_flow()
    with get_sessionmaker()() as db:
        tenant = int(flow['tenantId'])
        leader_role = db.scalar(select(Role).where(Role.tenant_id == tenant, Role.role_code == 'PROFESSIONAL_LEADER'))
        mentor_role = db.scalar(select(Role).where(Role.tenant_id == tenant, Role.role_code == 'INTERN_MENTOR'))
        leader = User(tenant_id=tenant, login_name='professional_leader', real_name='合成专业负责人', password_hash=hash_password(PASSWORD),
                      user_type='PROFESSIONAL_LEADER', status='ACTIVE', credential_version=0, must_change_password=False)
        other = User(tenant_id=tenant, login_name='other_mentor', real_name='合成非指导教师', password_hash=hash_password(PASSWORD),
                     user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
        db.add_all([leader, other]); db.flush()
        db.add_all([UserRole(tenant_id=tenant, user_id=leader.id, role_id=leader_role.id, status='ACTIVE'),
                    UserRole(tenant_id=tenant, user_id=other.id, role_id=mentor_role.id, status='ACTIVE'),
                    TeacherStudentScope(tenant_id=tenant, user_id=leader.id, teacher_key='db-' + str(leader.id), role_code='PROFESSIONAL_LEADER',
                                        scope_type='STUDENT', ref_id=int(flow['studentId']), ref_value='S0', status='ACTIVE')]); db.commit()
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher, login(flow, 'PROFESSIONAL_LEADER') as leader, login(flow, 'OTHER_MENTOR') as other, login(foreign, 'INTERN_MENTOR') as outsider:
        proof = upload(student, biz='INTERNSHIP_LEAVE')
        row = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', endDate='2026-10-03', fileId=proof['fileId'])))
        own = dashboard_leaves(teacher, flow, row['id'], 'LEAVE_REVIEW')
        assert own['workItemTotal'] == 1
        limited = dashboard_leaves(leader, flow, row['id'])
        assert next(x['value'] for x in limited['stats'] if x['label'] == '本批学生') == '1'
        assert limited['workItemTotal'] == 0 and not limited['todos']
        assert leader.get('/api/v1/internship/leaves', params={'batchId': flow['batchId']}).status_code == 403
        for reader in (leader, other, outsider):
            for path in ('/api/v1/files/', '/api/v1/files/preview/', '/api/v1/files/download/'):
                response = reader.get(path + proof['fileId'])
                assert response.status_code == 404 and not response.json().get('data'), response.text
            assert reader.get('/api/v1/files/' + proof['fileId'] + '/url').status_code == 404
        assert teacher.get('/api/v1/files/download/' + proof['fileId']).content == png_bytes()
        from app.models.file import FileObject
        with get_sessionmaker()() as db:
            db.get(FileObject, int(proof['fileId'])).biz_type = legacy_purpose
            db.commit()
        assert teacher.get('/api/v1/files/preview/' + proof['fileId']).content == png_bytes()
        assert student.get('/api/v1/files/download/' + proof['fileId']).content == png_bytes()
        assert leader.get('/api/v1/files/' + proof['fileId']).status_code == 404
        with get_sessionmaker()() as db:
            binding = db.scalar(select(FileBinding).where(FileBinding.file_id == int(proof['fileId'])))
            binding.status = 'INACTIVE'; db.commit()
        assert student.get('/api/v1/files/' + proof['fileId']).status_code == 404
        assert teacher.get('/api/v1/files/download/' + proof['fileId']).status_code == 404
        assert dashboard_leaves(other, flow, row['id'])['workItemTotal'] == 0
        assert not ok(teacher.get('/api/v1/internship/dashboard', params={'batchId': flow['batches'][1]['id']}))['workItems']
        assert not dashboard_leaves(outsider, foreign, row['id'])['workItems']
        denied = outsider.get('/api/v1/internship/dashboard', params={'batchId': flow['batchId']})
        assert denied.json()['code'] != 0 and not denied.json().get('data')


def test_private_proof_cannot_be_borrowed_and_denial_creates_no_leave(flow):
    foreign = seed_flow()
    with login(flow, 'STUDENT') as student, login(flow) as school, login(foreign, 'STUDENT') as outsider:
        for owner in (school, outsider):
            proof = upload(owner, biz='INTERNSHIP_LEAVE')
            denied = student.post(PC, json=values(flow, fileId=proof['fileId']))
            assert denied.status_code in (403, 404) and denied.json()['code'] != 0, denied.text
        assert ok(student.get(PC, params=context(flow))) == []


def test_wrong_batch_and_foreign_school_cannot_read_or_change_original(flow):
    foreign = seed_flow()
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher, login(foreign, 'STUDENT') as outsider:
        proof = upload(student, biz='INTERNSHIP_LEAVE')
        row = ok(student.post(PC, json=values(flow, fileId=proof['fileId'])))
        before = persisted(flow, row['id'])
        for surface in (PC, MOBILE):
            assert student.get(surface + '/' + row['id']).status_code == 422
            denied = student.get(surface + '/' + row['id'], params={**context(flow), 'batchId': flow['batches'][1]['id']})
            assert denied.status_code == 404 and not denied.json().get('data'), denied.text
            assert outsider.get(surface + '/' + row['id'], params=context(foreign)).status_code == 404
        wrong = {**frozen(flow, row), 'batchId': flow['batches'][1]['id']}
        for suffix, body in [('/evidence-viewed', {**wrong, 'evidenceFileId': proof['fileId']}),
                             ('/review', {**wrong, 'action': 'REJECT', 'comment': '错误批次不得审批'}),
                             ('/ack-return', {**wrong, 'note': '错误批次不得办结'})]:
            conflict(teacher.post(TEACHER + '/' + row['id'] + suffix,
                params={'batchId': wrong['batchId']}, json=body))
        assert student.get(TEACHER + '/overdue', params={'batchId': flow['batchId']}).status_code == 403
        conflict(teacher.get('/api/v1/internship/leaves/' + row['id'], params={'batchId': wrong['batchId']}))
        conflict(teacher.post('/api/v1/internship/leaves/' + row['id'] + '/review', json={
            **wrong, 'action': 'REJECT', 'comment': '错误批次不得审批'}))
        assert persisted(flow, row['id']) == before


@pytest.mark.parametrize('comment', [123, ['不是文本'], {'text': '不是文本'}, '合' * 501])
def test_review_opinion_validation_has_no_business_side_effect(flow, comment):
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        row = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', endDate='2026-10-03')))
        before = persisted(flow, row['id'])
        denied = review(teacher, flow, row, action='REJECT', comment=comment)
        assert denied.status_code == 400 and denied.json()['bizCode'] == 'VALIDATION_ERROR', denied.text
        assert persisted(flow, row['id']) == before


@pytest.mark.parametrize('version', [None, True, 0.8, -1, '0.0', 'no-version'])
def test_frozen_versions_reject_coercion_without_side_effects(flow, version):
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        row = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', endDate='2026-10-03')))
        before = persisted(flow, row['id'])
        withdrawn = student.post(PC + '/' + row['id'] + '/withdraw', json={**context(flow), 'expectedVersion': version})
        rejected = review(teacher, flow, row, expectedVersion=version, action='REJECT', comment='无效版本不得审批')
        assert withdrawn.status_code in (400, 409) and rejected.status_code in (400, 409)
        assert persisted(flow, row['id']) == before


@pytest.mark.parametrize('extra', [{'reason': '合' * 501}, {'startDate': '2026-10-03T12:00'},
                                  {'endDate': '2026-02-30'}, {'startDate': '1000-01-01', 'endDate': '9999-12-31'}])
def test_invalid_form_is_rejected_before_database_write(flow, extra):
    with login(flow, 'STUDENT') as student:
        proof = upload(student, biz='INTERNSHIP_LEAVE')
        denied = student.post(PC, json=values(flow, **{'leaveType': 'PERSONAL', 'endDate': '2026-10-03', 'fileId': proof['fileId'], **extra}))
        assert denied.status_code == 400 and denied.json()['bizCode'] == 'VALIDATION_ERROR', denied.text
        assert ok(student.get(PC, params=context(flow))) == []


def test_withdraw_is_terminal_and_pending_uniqueness_remains(flow):
    with login(flow, 'STUDENT') as student:
        payload = values(flow, leaveType='PERSONAL', endDate='2026-10-03')
        row = ok(student.post(PC, json=payload))
        conflict(student.post(MOBILE, json=payload))
        withdrawn = ok(student.post(PC + '/' + row['id'] + '/withdraw', json={**context(flow), 'expectedVersion': row['version']}))
        conflict(student.post(MOBILE + '/' + row['id'] + '/resubmit', json={**payload, 'expectedVersion': withdrawn['version']}))
        conflict(student.post(PC + '/' + row['id'] + '/withdraw', json={**context(flow), 'expectedVersion': row['version']}))
        assert persisted(flow, row['id'])[0] == 'WITHDRAWN'
        assert all(status == 'DONE' for _, status in persisted(flow, row['id'])[-1])


def test_overdue_scan_cannot_overwrite_student_return(flow, monkeypatch):
    from app.modules.internship.services import internship_leave_service as service
    from app.services.db_service import session
    from contextlib import contextmanager
    read, release = Event(), Event()
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        row = ok(student.post(PC, json=values(flow, leaveType='PERSONAL', endDate='2026-10-03')))
        approved = ok(review(teacher, flow, row, action='APPROVE'))

        @contextmanager
        def paused_session():
            with session() as db:
                original = db.scalars
                def scalars(statement, *args, **kwargs):
                    result = original(statement, *args, **kwargs)
                    if statement.column_descriptions and statement.column_descriptions[0].get('entity') is InternshipLeave:
                        rows = result.all()
                        read.set()
                        assert release.wait(15)
                        class Rows:
                            def all(self): return rows
                        return Rows()
                    return result
                db.scalars = scalars
                yield db

        monkeypatch.setattr(service, 'session', paused_session)
        def scan():
            set_tenant(flow['tenantId'])
            try:
                return service.refresh_overdue(reference_date=date(2026, 10, 6), system=True)
            finally:
                set_tenant(None)
        with ThreadPoolExecutor(max_workers=1) as pool:
            future = pool.submit(scan)
            try:
                assert read.wait(15)
                ok(student.post(PC + '/' + row['id'] + '/return', json={**context(flow),
                    'expectedVersion': approved['version'], 'note': '并发测试：已正常返岗'}))
            finally:
                release.set()
            result = future.result(timeout=20)
        assert result['markedOverdue'] == result['risksCreated'] == 0
        assert persisted(flow, row['id'])[0] == 'RETURNED'
        with get_sessionmaker()() as db:
            assert db.scalar(select(RiskRecord).where(RiskRecord.source_id == int(row['id']),
                RiskRecord.tenant_id == int(flow['tenantId']), RiskRecord.source_type == 'LEAVE')) is None
