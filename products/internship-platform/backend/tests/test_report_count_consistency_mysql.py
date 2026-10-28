"""Cross-endpoint report-count regression on an isolated synthetic MySQL school."""
from base64 import b64decode
from datetime import datetime
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import get_sessionmaker
from app.models import InternshipBatch, InternshipRecord, StudentProfile, WeeklyReport
from test_yiyang_write_flows_mysql import flow, login

ROOT = '/api/v1/internship'


def read(client, path, data, **params):
    response = client.get(ROOT + path, params={'batchId': data['batchId'], **params})
    assert response.status_code == 200 and response.json()['code'] == 0, response.text
    return response.json()['data']


def prepare(data, required, *, dates=True):
    with get_sessionmaker()() as db:
        batch = db.get(InternshipBatch, int(data['batchId']))
        weekly = {'frequency': 'WEEKLY', 'deadlineWeekday': 7}
        if required is not None:
            weekly['requiredCount'] = required
        batch.rules_config = {'weeklyReport': weekly, 'processReport': {
            'dailyRequiredCount': 0, 'monthlyRequiredCount': 0, 'summaryRequiredCount': 0}}
        second = db.scalar(select(InternshipRecord).where(
            InternshipRecord.batch_id == int(data['batches'][1]['id'])))
        second.batch_id = batch.id
        second.advisor_user_id = data['users']['INTERN_MENTOR']
        if not dates:
            batch.start_date = batch.end_date = None
            first = db.get(InternshipRecord, int(data['recordId']))
            first.intern_start_date = first.intern_end_date = None
        for week, state, deleted in [(1, 'RETURNED', False), (2, 'DRAFT', False),
                                     (3, 'APPROVED', True), (4, 'APPROVED', False)]:
            db.add(WeeklyReport(tenant_id=int(data['tenantId']), internship_id=int(data['recordId']),
                week_number=week, status=state, is_deleted=deleted, report_version=3,
                submitted_at=datetime(2026, 9, 6), work_content='合成统计原单',
                harvest_content='合成收获', plan_content='合成计划'))
        db.commit()


@pytest.mark.parametrize('required', [None, 0, 1])
def test_count_config_is_shared_across_ledger_procurement_process_and_export(flow, required):
    prepare(flow, required)
    expected = (required or 0) * 2
    with login(flow) as school:
        ledger = read(school, '/report-obligations', flow)
        assert len(ledger['items']) == 2
        assert sum(row['weekly']['required'] for row in ledger['items']) == expected
        assert sum(row['weekly']['submitted'] for row in ledger['items']) == 2
        overview = read(school, '/stats/overview', flow)
        metric = next(x for x in overview['metrics'] if x['key'] == 'reportTaskCompletionRate')
        assert metric['denominator'] == expected
        assert metric['numerator'] == (1 if required else 0)
        assert metric['rate'] == (50.0 if required else None)
        procurement = read(school, '/stats/procurement-overview', flow)
        activity = next(x for x in procurement['groups'] if x['key'] == 'activity')
        assert activity['weeklyReportTasks']['expected'] == expected
        assert activity['weeklyReportTasks']['submitted'] == metric['numerator']
        drilldown = read(school, '/stats/metrics/reportTaskCompletionRate/drilldown', flow,
                         subset='denominator')
        assert drilldown['total'] == (2 if required else 0)
        assert sum(row['weeklyExpected'] for row in drilldown['items']) == expected
        process = read(school, '/stats/process-analytics', flow)
        assert sum(x['weeklyExpected'] for x in process['items']) == expected
        assert sum(x['weeklySubmitted'] for x in process['items']) == 2
        aggregate = read(school, '/stats/process-analytics', flow, groupBy='ADVISOR')
        assert aggregate['items'][0]['weeklyMissing'] == (1 if required else 0)
        assert aggregate['items'][0]['monthlyCompletionRate'] is None
        exported = school.post(ROOT + '/stats/process-analytics/export', params={
            'batchId': flow['batchId'], 'columns': 'studentNo,weeklyExpected,weeklySubmitted,weeklyCompletionRate'})
        assert exported.status_code == 200 and exported.json()['code'] == 0, exported.text
        values = list(load_workbook(BytesIO(b64decode(exported.json()['data']['contentBase64']))).active.values)
        assert all(any(tuple(row[:4]) == (item['studentNo'], item['weeklyExpected'],
            item['weeklySubmitted'], item['weeklyCompletionRate']) for row in values)
            for item in process['items'])
        exported_stats = school.post(ROOT + '/stats/export', params={'batchId': flow['batchId']})
        assert exported_stats.status_code == 200 and exported_stats.json()['code'] == 0, exported_stats.text
        metric_rows = list(load_workbook(BytesIO(b64decode(exported_stats.json()['data']['contentBase64']))).active.values)
        assert any(tuple(row[:4]) == ('周报任务完成率', metric['numerator'], expected,
            metric['rate'] if metric['rate'] is not None else '—') for row in metric_rows)


def test_missing_dates_keep_full_counts_without_inventing_current_period_obligations(flow):
    prepare(flow, 1, dates=False)
    with login(flow) as school:
        all_rows = read(school, '/stats/process-analytics', flow)['items']
        assert sum(row['weeklyExpected'] for row in all_rows) == 2
        month = read(school, '/stats/process-analytics', flow, period='MONTH', anchor='2026-10-01')
        assert sum(row['weeklyExpected'] for row in month['items']) == 0
        assert all(row['weeklyCompletionRate'] is None for row in month['items'])


def test_week_month_filters_and_deleted_student_share_ledger_population(flow):
    prepare(flow, 1)
    with login(flow) as school:
        week = read(school, '/stats/process-analytics', flow, period='WEEK', anchor='2026-09-01')
        assert sum(row['weeklyExpected'] for row in week['items']) == 2
        month = read(school, '/stats/process-analytics', flow, period='MONTH', anchor='2026-10-01')
        assert sum(row['weeklyExpected'] for row in month['items']) == 0
        with get_sessionmaker()() as db:
            db.get(StudentProfile, int(flow['studentId'])).is_deleted = True
            db.commit()
        ledger = read(school, '/report-obligations', flow)
        assert len(ledger['items']) == 1 and ledger['items'][0]['weekly']['submitted'] == 0
        overview = read(school, '/stats/overview', flow)
        metric = next(x for x in overview['metrics'] if x['key'] == 'reportTaskCompletionRate')
        assert (metric['numerator'], metric['denominator'], metric['rate']) == (0, 1, 0.0)


def test_teacher_scope_revocation_and_deleted_batch_do_not_leave_report_facts(flow):
    prepare(flow, 1)
    # The new preset permits own-scope analytics. Keep this negative case by
    # explicitly revoking the freshly created school's statistics permissions.
    from app.models import Role,RolePermission,Permission
    assert flow['tenantCode'].startswith('BULK-')
    with get_sessionmaker()() as db:
        role=db.scalar(select(Role).where(Role.tenant_id==int(flow['tenantId']),Role.role_code=='INTERN_MENTOR'))
        rows=db.scalars(select(RolePermission).join(Permission,Permission.id==RolePermission.permission_id).where(
            RolePermission.tenant_id==int(flow['tenantId']),RolePermission.role_id==role.id,
            Permission.permission_code.in_(['internship.stats.view','internship.stats.export']))).all()
        assert rows
        for row in rows:row.is_deleted=True
        db.commit()
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        assert len(read(teacher, '/report-obligations', flow)['items']) == 2
        statistics_paths = ['/stats/process-analytics', '/stats/overview', '/stats/procurement-overview']
        for path in statistics_paths:
            denied = teacher.get(ROOT + path, params={'batchId': flow['batchId']})
            assert denied.status_code == 403 and denied.json()['data'] is None, denied.text
        with get_sessionmaker()() as db:
            for record in db.scalars(select(InternshipRecord).where(
                InternshipRecord.batch_id == int(flow['batchId']))).all():
                record.advisor_user_id = flow['users']['SCHOOL_ADMIN']
            db.commit()
        ledger = read(teacher, '/report-obligations', flow)
        assert ledger['total'] == 0 and ledger['items'] == []
        # Explicitly revoked statistics permissions still deny every statistics route.
        for path in statistics_paths:
            denied = teacher.get(ROOT + path, params={'batchId': flow['batchId']})
            assert denied.status_code == 403 and denied.json()['data'] is None, denied.text
        with get_sessionmaker()() as db:
            db.get(InternshipBatch, int(flow['batchId'])).is_deleted = True
            db.commit()
        for path in ['/report-obligations', '/stats/overview', '/stats/procurement-overview', '/stats/process-analytics']:
            denied = school.get(ROOT + path, params={'batchId': flow['batchId']})
            assert denied.status_code == 404 and denied.json()['data'] is None, denied.text
        denied = teacher.get(ROOT + '/report-obligations', params={'batchId': flow['batchId']})
        assert denied.status_code == 404 and denied.json()['data'] is None, denied.text


def test_mid_month_finish_monthly_obligation_and_submission_share_period(flow):
    from app.models import InternshipProcessReport
    with get_sessionmaker()() as db:
        batch = db.get(InternshipBatch, int(flow['batchId']))
        record = db.get(InternshipRecord, int(flow['recordId']))
        batch.start_date = record.intern_start_date = datetime(2026, 9, 1)
        batch.end_date = record.intern_end_date = datetime(2026, 9, 15)
        batch.rules_config = {'processReport': {'monthlyRequiredCount': 1}}
        db.add(InternshipProcessReport(tenant_id=int(flow['tenantId']), internship_id=record.id,
            report_type='MONTHLY', period_key='2026-09', content='合成月中结束报告',
            status='PENDING_REVIEW', submitted_at=datetime(2026, 9, 15)))
        db.commit()
    with login(flow) as school:
        for anchor, expected in [('2026-09-14', 1), ('2026-09-28', 0)]:
            rows = read(school, '/stats/process-analytics', flow, period='WEEK', anchor=anchor)['items']
            assert len(rows) == 1
            assert (rows[0]['monthlyExpected'], rows[0]['monthlySubmitted']) == (expected, expected)


@pytest.mark.parametrize('scope_type', ['COLLEGE', 'MAJOR', 'CLASS', 'STUDENT'])
def test_stable_organization_scope_matches_all_report_stats_without_same_name_leak(flow, scope_type):
    from app.models import College, Major, SchoolClass, Role, User, UserRole, TeacherStudentScope
    prepare(flow, 1)
    with get_sessionmaker()() as db:
        tid = int(flow['tenantId'])
        records = db.scalars(select(InternshipRecord).where(
            InternshipRecord.batch_id == int(flow['batchId'])).order_by(InternshipRecord.id)).all()
        scope_ids = None
        for record in records:
            record.created_at = datetime.now()
            college = College(tenant_id=tid, college_name='同名院系', status='ACTIVE')
            db.add(college); db.flush()
            major = Major(tenant_id=tid, college_id=college.id, major_name='同名专业', status='ACTIVE')
            db.add(major); db.flush()
            school_class = SchoolClass(tenant_id=tid, major_id=major.id, class_name='同名班级',
                                       grade='2024', status='ACTIVE', class_status='NORMAL')
            db.add(school_class); db.flush()
            student = db.get(StudentProfile, record.student_id)
            student.class_id = school_class.id
            student.major_id = student.college_id = None
            if str(record.id) == flow['recordId']:
                scope_ids = {'COLLEGE': college.id, 'MAJOR': major.id,
                             'CLASS': school_class.id, 'STUDENT': student.id}
        role = db.scalar(select(Role).where(Role.tenant_id == tid,
            Role.role_code == 'COLLEGE_ADMIN', Role.is_deleted.is_(False)))
        assert role is not None and scope_ids is not None
        user = User(tenant_id=tid, login_name='college_admin', real_name='范围统计验收教师',
            password_hash=db.get(User, flow['users']['INTERN_MENTOR']).password_hash,
            user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
        db.add(user); db.flush()
        db.add(UserRole(tenant_id=tid, user_id=user.id, role_id=role.id, status='ACTIVE'))
        db.add(TeacherStudentScope(tenant_id=tid, teacher_key='college_admin', user_id=user.id,
            role_code='COLLEGE_ADMIN', scope_type=scope_type, ref_id=scope_ids[scope_type], status='ACTIVE'))
        db.commit()
    with login(flow, 'COLLEGE_ADMIN') as college:
        ledger = read(college, '/report-obligations', flow)
        assert [item['recordId'] for item in ledger['items']] == [flow['recordId']]
        overview = read(college, '/stats/overview', flow)
        metric = next(x for x in overview['metrics'] if x['key'] == 'reportTaskCompletionRate')
        assert (metric['numerator'], metric['denominator'], metric['rate']) == (1, 1, 100.0)
        assert next(x['value'] for x in overview['counters'] if x['key'] == 'totalStudents') == 1
        assert next(x['value'] for x in overview['counters'] if x['key'] == 'onboardStudents') == 1
        for subset in ['numerator', 'denominator']:
            detail = read(college, '/stats/metrics/reportTaskCompletionRate/drilldown', flow, subset=subset)
            assert [x['internshipId'] for x in detail['items']] == [flow['recordId']]
        procurement = read(college, '/stats/procurement-overview', flow)
        activity = next(x for x in procurement['groups'] if x['key'] == 'activity')
        assert activity['weeklyReportTasks']['expected'] == 1
        process = read(college, '/stats/process-analytics', flow)
        assert process['total'] == 1 and process['items'][0]['weeklyExpected'] == 1
        dimensions = read(college, '/stats/dimensions', flow)
        assert dimensions == {'colleges': ['同名院系'], 'majors': ['同名专业'], 'classes': ['同名班级']}
        trends = read(college, '/stats/trends', flow, college='同名院系')
        records = next(item for item in trends['series'] if item['key'] == 'records')
        assert sum(point['value'] for point in records['points']) == 1
        assert all(point['value'] == 0 for series in read(college, '/stats/trends', flow,
            college='无权院系')['series'] for point in series['points'])
        with get_sessionmaker()() as db:
            for scope in db.scalars(select(TeacherStudentScope).where(
                TeacherStudentScope.tenant_id == tid, TeacherStudentScope.user_id == user.id)).all():
                scope.is_deleted = True
            db.commit()
        assert read(college, '/stats/dimensions', flow) == {'colleges': [], 'majors': [], 'classes': []}
        assert all(point['value'] == 0 for series in read(college, '/stats/trends', flow)['series']
                   for point in series['points'])
