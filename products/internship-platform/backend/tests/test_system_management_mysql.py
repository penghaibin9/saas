"""真实 MySQL、正常账号的系统管理故事；只新增唯一合成学校，不清理共享办理历史。"""
import io
import os
from uuid import uuid4
from concurrent.futures import ThreadPoolExecutor
import pytest
from fastapi.testclient import TestClient
from openpyxl import load_workbook
from sqlalchemy import select, func
from sqlalchemy.engine import make_url
from app.config import settings
from app.main import app
from app.core.security import hash_password
from app.db.session import get_sessionmaker
from app.models import Tenant, User, Role, UserRole, StudentProfile, StaffProfile, StudentAccountLink, SecurityAuditLog
from app.services.system_authorization import bootstrap_school
from app.services.xlsx_util import build_template_xlsx
from app.services.system_identity_import import HEADERS

PASSWORD='System-Only-Test-2026!'
ROOT='/api/v1/system'


def test_preset_identity_matrix_never_combines_other_roles_or_exposes_school_management(school):
    """Same normal account, all staff identities: permission follows only the active role."""
    pages=['/overview','/organizations/colleges','/people/staff','/people/student','/accounts',
           '/roles','/scopes','/config/brand','/config/dictionary','/config/notification',
           '/identity-import/students/template','/tasks','/records/MIGRATION','/integrations',
           '/audit','/runtime','/records/BACKUP','/records/DOCUMENT','/leave-policies/options']
    college_pages={'/overview','/organizations/colleges','/people/staff','/people/student',
                   '/accounts','/identity-import/students/template','/tasks'}
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['MULTI','多身份边界合成人员','C01','原院','','','','','']])
        _,students=import_rows(owner,'students',[['S001','身份边界合成学生','C01','原院','M01','原专业','B01','原班','2024']])
        target=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='MULTI')
        # Test preconditions only: these memberships are not claimed as browser grant acceptance.
        for role in ok(owner.get(ROOT+'/roles'))['items']:
            if role['code']!='STUDENT':
                ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':target['id'],
                    'enabled':True,'expectedVersion':role['version'],'reason':'合成角色矩阵前置'}))
        with reset_for_test(owner,school,'MULTI',credentials['MULTI']) as actor:
            for code in ['SYSTEM_ADMIN','SCHOOL_ADMIN','COLLEGE_ADMIN','PROFESSIONAL_LEADER',
                         'MAJOR_LEADER','COUNSELOR','INTERN_MENTOR','LEADER','SECURITY_AUDITOR']:
                activate(actor,code)
                allowed=set(pages) if code=='SYSTEM_ADMIN' else college_pages if code=='COLLEGE_ADMIN' else {'/audit'} if code=='SECURITY_AUDITOR' else set()
                for path in pages:
                    response=actor.get(ROOT+path)
                    assert response.status_code==(200 if path in allowed else 403),(code,path,response.status_code)
                assert actor.get(ROOT+'/audit/export').status_code==(200 if code in {'SYSTEM_ADMIN','SECURITY_AUDITOR'} else 403)
                if code not in {'SYSTEM_ADMIN','COLLEGE_ADMIN'}:
                    assert actor.post(ROOT+'/organizations/colleges',json={'name':'拒绝新增','code':'DENIED','reason':'无权写入验收'}).status_code==403
                    assert actor.post(ROOT+'/exports/student',json={}).status_code==403
        with reset_for_test(owner,school,'S001',students['S001'],True) as student:
            for path in pages:
                assert student.get(ROOT+path).status_code==403,path
            # Staff-only switching rejects a student browser channel before selecting any role.
            assert student.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-system_admin','clientType':'PC'}).status_code==401
        assert all(r['code']!='DENIED' for r in ok(owner.get(ROOT+'/organizations/colleges'))['items'])


@pytest.mark.parametrize('member_role',['COMPANY_ADMIN','HR','MENTOR'])
@pytest.mark.parametrize('transport',['login','browser-login'])
def test_enterprise_member_cannot_enter_school_management_or_impersonate_staff(school,member_role,transport):
    from datetime import datetime, timedelta
    from app.models import EmpCompany
    from app.models.internship_enterprise_portal import InternshipEnterpriseMember
    with get_sessionmaker()() as db:
        company=EmpCompany(tenant_id=school['id'],name='角色隔离合成企业',status='ACTIVE',
            coop_status='ACTIVE',qualification_status='PASSED',access_valid_until=datetime.utcnow()+timedelta(days=30))
        account=User(tenant_id=school['id'],login_name='COMPANY',real_name='合成企业成员',
            user_type='ENTERPRISE_MENTOR',password_hash=hash_password(PASSWORD),status='ACTIVE',must_change_password=False)
        db.add_all([company,account]);db.flush()
        db.add(InternshipEnterpriseMember(tenant_id=school['id'],company_id=company.id,user_id=account.id,
            member_role=member_role,status='ACTIVE'));db.commit()
    prefix='/api/v1/internship/enterprise-portal'
    with TestClient(app) as actor:
        actor.headers['X-Browser-Session-Id']='enterprise-'+uuid4().hex
        result=ok(actor.post(prefix+'/auth/'+transport,json={'tenantCode':school['code'],'loginName':'COMPANY','password':PASSWORD}))
        actor.headers['Authorization']='Bearer '+result['accessToken']
        assert actor.get(prefix+'/campaigns').status_code==200
        with get_sessionmaker()() as db:
            event=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],SecurityAuditLog.action=='LOGIN'))
            assert event and event.operator_name=='合成企业成员'
            assert event.detail_json['identityName']=={'COMPANY_ADMIN':'企业管理员','HR':'企业人事成员','MENTOR':'企业导师'}[member_role]
            assert PASSWORD not in str(event.detail_json) and result['accessToken'] not in str(event.detail_json)
        for path in ['/context','/overview','/people/student','/people/staff','/accounts','/roles',
                     '/scopes','/audit','/exports/student','/tasks','/records/DOCUMENT','/config/brand']:
            assert actor.get(ROOT+path).status_code in (401,403),path
        assert actor.post(ROOT+'/organizations/colleges',json={'name':'拒绝新增','code':'DENIED','reason':'企业身份隔离验收'}).status_code in (401,403)
        assert actor.post('/api/v1/auth/browser-login',json={'tenantCode':school['code'],
            'loginName':'COMPANY','password':PASSWORD,'clientType':'PC'},headers={'X-Browser-Session-Id':'system-'+uuid4().hex}).status_code in (401,403)


def ok(response):
    # Never include credential-bearing response bodies in failed-test output.
    assert response.status_code==200,(response.status_code,response.json().get('message'))
    assert response.json()['code']==0,response.json().get('message')
    return response.json()['data']


def login(school,name='owner',password=PASSWORD,student=False):
    client=TestClient(app)
    client.headers['X-Browser-Session-Id']='system-'+uuid4().hex
    result=ok(client.post('/api/v1/auth/browser-login',json={'tenantCode':school['code'],
        'loginName':name,'password':password,'clientType':'STUDENT_PC' if student else 'PC'}))
    client.headers['Authorization']='Bearer '+result['accessToken']
    return client,result


@pytest.fixture
def school():
    if os.environ.get('SYSTEM_MYSQL_ACCEPTANCE')!='1': pytest.skip('requires explicit local MySQL opt-in')
    url=make_url(settings.DATABASE_URL)
    local = url.host in {'127.0.0.1','localhost'} and url.port==3318 and url.database=='internship_pr275_test_20261003'
    ci = os.environ.get('GITHUB_ACTIONS')=='true' and settings.APP_ENV=='ci' and url.host=='127.0.0.1' and url.port==3306 and url.database=='internship_auth_gate'
    assert local or ci, '系统管理验收只能连接已指定的本地库或流水线临时库'
    with get_sessionmaker()() as db:
        tenant=Tenant(tenant_code='SYS-'+uuid4().hex[:12],school_name='系统管理合成验收学校',status='ACTIVE')
        db.add(tenant);db.flush()
        owner=User(tenant_id=tenant.id,login_name='owner',real_name='学校初始化负责人',user_type='SCHOOL_ADMIN',
                   password_hash=hash_password(PASSWORD),status='ACTIVE',must_change_password=False)
        db.add(owner);db.flush();bootstrap_school(db,tenant.id);db.commit()
        return {'id':tenant.id,'code':tenant.tenant_code,'owner':owner.id}


def org(client,kind,name,code,parent=''):
    return ok(client.post(ROOT+'/organizations/'+kind,json={'name':name,'code':code,'parentId':parent,'reason':'合成学校初始化'}))


def preview(client,kind,rows):
    content=build_template_xlsx(list(HEADERS),samples=rows)
    return ok(client.post(ROOT+f'/identity-import/{kind}/preview',files={'file':('人员.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}))


def import_rows(client,kind,rows):
    plan=preview(client,kind,rows)
    assert plan['invalidRows']==0,plan['errors']
    result=ok(client.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':plan['version']}))
    response=client.get(ROOT+f'/identity-import/jobs/{plan["id"]}/receipt');assert response.status_code==200
    wb=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
    credentials={r[2]:r[3] for r in list(wb.active.iter_rows(values_only=True))[2:]};wb.close()
    return result,credentials


def reset_for_test(client,school,username,temporary,student=False):
    person,logged=login(school,username,temporary,student)
    assert logged['mustChangePassword']
    assert person.get(ROOT+'/overview').status_code==403
    ok(person.post('/api/v1/auth/change-password',json={'currentPassword':temporary,'newPassword':PASSWORD}))
    person.close()
    return login(school,username,PASSWORD,student)[0]


def test_product_centers_discover_assigned_roles_without_merging_permissions(school):
    with login(school)[0] as owner:
        initial=ok(owner.get('/api/v1/auth/me'))['contexts']
        assert len(initial)==1 and initial[0]['roleCode']=='SYSTEM_ADMIN'
        assert 'system.role.manage' in initial[0]['permissionPatterns']
        assert not any(p.startswith('internship.') for p in initial[0]['permissionPatterns'])
        roles=ok(owner.get(ROOT+'/roles'))['items']
        business=next(r for r in roles if r['code']=='SCHOOL_ADMIN')
        assigned=ok(owner.post(ROOT+f'/roles/{business["id"]}/members',json={
            'userId':str(school['owner']),'enabled':True,'expectedVersion':business['version'],
            'reason':'合成学校负责人兼任实习管理的授权验收'}))
        contexts=ok(owner.get('/api/v1/auth/me'))['contexts']
        target=next(c for c in contexts if c['roleCode']=='SCHOOL_ADMIN')
        assert 'internship.student.view' in target['permissionPatterns']
        current=ok(owner.get('/api/v1/rbac/current-context'))
        assert not any(p.startswith('internship.') for p in current['permissionPatterns'])
        ok(owner.post(ROOT+f'/roles/{business["id"]}/members',json={
            'userId':str(school['owner']),'enabled':False,'expectedVersion':assigned['version'],
            'reason':'合成兼任结束，入口应立即移除身份'}))
        assert all(c['roleCode']!='SCHOOL_ADMIN' for c in ok(owner.get('/api/v1/auth/me'))['contexts'])
        assert owner.post('/api/v1/auth/browser-switch-role',json={
            'contextId':target['contextId'],'clientType':'PC'}).status_code==403


def test_empty_school_identity_preview_commit_idempotence_first_login_and_binding(school):
    with login(school)[0] as client:
        overview=ok(client.get(ROOT+'/overview'))
        assert overview['studentCount']==0
        assert overview['backup']['path']=='/admin/system/runtime'
        assert overview['backup']['statusLabel']=='尚未取得备份与恢复演练回执'
        bad=preview(client,'students',[['000001','合成学生','C01','工程学院','','','','','2024']])
        assert bad['invalidRows']==1
        assert client.post(ROOT+f'/identity-import/jobs/{bad["id"]}/confirm',json={'expectedVersion':bad['version']}).status_code==409
        result,credentials=import_rows(client,'students',[['000001','合成学生','C01','工程学院','M01','软件技术','B01','软件一班','2024']])
        again=ok(client.post(ROOT+f'/identity-import/jobs/{result["id"]}/confirm',json={'expectedVersion':0}))
        assert again['status']=='IMPORTED'
        students=ok(client.get(ROOT+'/people/student'))['items'];assert len(students)==1 and students[0]['number']=='000001'
        moved=preview(client,'students',[['000001','合成学生','C02','另一学院','M02','另一专业','B02','另一班级','2024']])
        assert moved['invalidRows']==1 and any('组织归属发生变化' in r['message'] for r in moved['errors'])
        assert client.post(ROOT+f'/identity-import/jobs/{moved["id"]}/confirm',json={'expectedVersion':moved['version']}).status_code==409
        assert not ok(client.get(ROOT+'/organizations/colleges',params={'keyword':'C02'}))['items']
        with reset_for_test(client,school,'000001',credentials['000001'],True) as student:
            assert student.get('/api/v1/mobile/me/portal-config').status_code==200
            assert student.get(ROOT+'/people/student').status_code==403
        with get_sessionmaker()() as db:
            assert db.scalar(select(func.count()).select_from(StudentAccountLink).where(StudentAccountLink.tenant_id==school['id']))==1
            logs=db.scalars(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'])).all()
            assert not any(credentials['000001'] in str(r.detail_json) for r in logs)


@pytest.mark.parametrize('collision',['renamed_login','other_staff','changed_after_preview'])
def test_staff_import_binding_conflict_never_replaces_existing_identity(school,collision):
    values=['T001','合成原教师','C01','工程学院','','','','','']
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[values])
        with reset_for_test(owner,school,'T001',credentials['T001']) as teacher:
            person=ok(owner.get(ROOT+'/people/staff'))['items'][0]
            plan=preview(owner,'staff',[values])
            assert plan['invalidRows']==0
            with get_sessionmaker()() as db:
                account=db.get(User,int(person['userId']))
                original=(account.id,account.password_hash,account.credential_version)
                roles=list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==account.id)))
                if collision in ('renamed_login','changed_after_preview'):
                    account.login_name='teacher-existing-login'
                    account.version+=1
                db.commit()
            if collision=='other_staff':
                values=[*values];values[0]='T002'
                # A legacy same-name login is already associated with another staff file.
                with get_sessionmaker()() as db:
                    account=db.get(User,original[0]);account.login_name='T002';account.version+=1;db.commit()
            if collision!='changed_after_preview':
                plan=preview(owner,'staff',[values])
                assert plan['invalidRows']==1
                assert any('绑定' in error['message'] for error in plan['errors'])
                assert plan['entries'][0]['accountAction']=='需核对，暂不处理'
            result=owner.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':plan['version']})
            assert result.status_code==409
            with get_sessionmaker()() as db:
                current=db.get(StaffProfile,int(person['id']))
                account=db.get(User,original[0])
                assert current.user_id==original[0]
                assert (account.id,account.password_hash,account.credential_version)==original
                assert list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==account.id)))==roles
                assert db.scalar(select(func.count()).select_from(StaffProfile).where(StaffProfile.tenant_id==school['id']))==1
                assert db.scalar(select(func.count()).select_from(User).where(User.tenant_id==school['id']))==2
            assert ok(teacher.get('/api/v1/auth/me'))['userId']=='db-'+str(original[0])


def test_existing_staff_update_does_not_require_or_recreate_disabled_opening_role(school):
    values=['T001','合成原教师','C01','工程学院','','','','','']
    with login(school)[0] as owner:
        import_rows(owner,'staff',[values])
        with get_sessionmaker()() as db:
            account=db.scalar(select(User).where(User.tenant_id==school['id'],User.login_name=='T001'))
            original=(account.id,account.password_hash,account.credential_version)
            roles=list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==account.id)))
        preset=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='INTERN_MENTOR')
        ok(owner.put(ROOT+'/roles/'+preset['id'],json={'name':preset['name'],'baseRole':'INTERN_MENTOR',
            'permissions':preset['permissions'],'status':'DISABLED','expectedVersion':preset['version'],'sensitiveConfirmed':True,
            'reason':'停止用默认指导身份开通新账号，已有资料仍须可维护'}))
        plan=preview(owner,'staff',[values])
        assert plan['invalidRows']==0
        result=ok(owner.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':plan['version']}))
        assert result['status']=='IMPORTED'
        with get_sessionmaker()() as db:
            account=db.get(User,original[0])
            assert (account.id,account.password_hash,account.credential_version)==original
            assert list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==account.id)))==roles
            assert db.get(Role,int(preset['id'])).status=='DISABLED'
        bad=preview(owner,'staff',[['T002',*values[1:]]])
        assert bad['invalidRows']==1
        assert owner.post(ROOT+f'/identity-import/jobs/{bad["id"]}/confirm',json={'expectedVersion':bad['version']}).status_code==409


def test_login_audit_persists_denials_without_attributing_or_storing_credentials(school):
    with login(school)[0] as owner:
        with TestClient(app) as client:
            client.headers['X-Browser-Session-Id']='audit-'+uuid4().hex
            for route,kind in [('browser-login','PC'),('login','TEACHER_MINI')]:
                body={'tenantCode':school['code'],'loginName':'owner','password':'deliberately-invalid','clientType':kind}
                assert client.post('/api/v1/auth/'+route,json=body).status_code==401
                assert client.post('/api/v1/auth/'+route,json={**body,'loginName':'unknown-user'}).status_code==401
            ok(client.post('/api/v1/auth/login',json={**body,'password':PASSWORD}))
        rows=ok(owner.get(ROOT+'/audit',params={'keyword':'登录'}))['items']
        assert len(rows)==6
        assert sum(r['action']=='LOGIN_DENIED' and r['operator']=='未认证访问' for r in rows)==4
        with get_sessionmaker()() as db:
            logs=db.scalars(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'])).all()
            assert all(r.operator_id is None for r in logs if r.action=='LOGIN_DENIED')
            assert all(PASSWORD not in str(r.detail_json) and 'deliberately-invalid' not in str(r.detail_json) for r in logs)
        with login(school)[0] as other_context:
            assert ok(other_context.get(ROOT+'/audit',params={'keyword':'登录被拒绝'}))['total']==4


def test_overview_unbound_people_links_filter_pagination_and_export_consistently(school):
    from app.models import ExcelImportJob
    with login(school)[0] as owner:
        import_rows(owner,'staff',[['T001','已开户教师','C01','核对学院','','','','','']])
        import_rows(owner,'students',[['S001','已绑定学生','C01','核对学院','M01','核对专业','B01','核对班级','2026']])
        student=ok(owner.get(ROOT+'/people/student'))['items'][0]
        with get_sessionmaker()() as db:
            people=[StudentProfile(tenant_id=school['id'],student_no=number,real_name='未绑定'+number,
                    college_id=int(student['collegeId']),major_id=int(student['majorId']),class_id=int(student['classId']),status='ACTIVE')
                for number in ['000002','000003']]
            db.add_all(people);db.flush()
            for person,state,deleted in [(people[0],'REVOKED',False),(people[1],'ACTIVE',True)]:
                db.add(StudentAccountLink(tenant_id=school['id'],student_id=person.id,user_id=int(student['userId']),link_status=state,is_deleted=deleted))
            db.add_all([StaffProfile(tenant_id=school['id'],staff_no=number,real_name='未开户'+number,
                college_id=int(student['collegeId']),status='ACTIVE') for number in ['000002','000003']])
            db.commit()
        overview=ok(owner.get(ROOT+'/overview'))
        issues={row['name']:row for row in overview['issues']}
        for name,kind,path in [('学生尚未绑定账号','student','students'),('教职工尚未开通账号','staff','staff')]:
            assert issues[name]['count']==2
            assert issues[name]['path']==f'/admin/system/{path}?binding=UNBOUND'
            params={'accountStatus':'UNBOUND','pageSize':1}
            first=ok(owner.get(ROOT+'/people/'+kind,params=params))
            second=ok(owner.get(ROOT+'/people/'+kind,params={**params,'page':2}))
            assert first['total']==second['total']==2
            assert [first['items'][0]['number'],second['items'][0]['number']]==['000002','000003']
            bound=ok(owner.get(ROOT+'/people/'+kind,params={'accountStatus':'BOUND'}))
            assert bound['total']==1 and bound['items'][0]['userId']
            params={'accountStatus':'UNBOUND','keyword':'000003'}
            filtered=ok(owner.get(ROOT+'/people/'+kind,params=params))
            assert filtered['total']==1 and filtered['items'][0]['number']=='000003'
            exported=owner.get(ROOT+'/exports/'+kind,params=params);assert exported.status_code==200
            workbook=load_workbook(io.BytesIO(exported.content),read_only=True,data_only=True)
            rows=list(workbook.active.iter_rows(values_only=True));workbook.close()
            assert len(rows)==2 and rows[1][0]=='000003' and rows[1][3]=='尚未开通'
            assert owner.get(ROOT+'/people/'+kind,params={'accountStatus':'invalid'}).status_code==422
            assert owner.get(ROOT+'/exports/'+kind,params={'accountStatus':'invalid'}).status_code==422
        assert issues['本人失败的导入任务']['path']=='/admin/system/tasks?status=FAILED'
        with get_sessionmaker()() as db:
            exports=db.scalars(select(ExcelImportJob).where(ExcelImportJob.tenant_id==school['id'],
                ExcelImportJob.biz_type=='DIRECTORY_EXPORT')).all()
            assert len(exports)==2 and all(job.preview_json['accountStatus']=='UNBOUND' and job.total_rows==1 for job in exports)


def test_audit_time_filter_and_export_use_same_instant_and_beijing_display(school):
    from datetime import timezone,timedelta
    with login(school)[0] as owner:
        org(owner,'colleges','时间核对一院','TIME01')
        org(owner,'colleges','时间核对二院','TIME02')
        with get_sessionmaker()() as db:
            events=db.scalars(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],SecurityAuditLog.action=='ORG_SAVE').order_by(SecurityAuditLog.id)).all()
            expected={str(e.id) for e in events}
            begin=min(e.created_at for e in events).replace(tzinfo=timezone.utc)
            end=max(e.created_at for e in events).replace(tzinfo=timezone.utc)
            displayed={e.resource_id:e.created_at.replace(tzinfo=timezone.utc).astimezone(timezone(timedelta(hours=8))).replace(tzinfo=None) for e in events}
        assert len(expected)==2
        utc={'keyword':'ORG_SAVE','startAt':begin.isoformat(),'endAt':end.isoformat()}
        beijing={**utc,'startAt':begin.astimezone(timezone(timedelta(hours=8))).isoformat(),'endAt':end.astimezone(timezone(timedelta(hours=8))).isoformat()}
        for filters in [utc,beijing]:
            listed=ok(owner.get(ROOT+'/audit',params=filters))
            assert listed['total']==2 and {r['id'] for r in listed['items']}==expected
            response=owner.get(ROOT+'/audit/export',params=filters)
            assert response.status_code==200
            workbook=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
            rows=list(workbook.active.iter_rows(values_only=True));workbook.close()
            assert rows[0][5]=='时间（北京时间）' and len(rows)==3
            assert {str(r[3]):r[5] for r in rows[1:]}==displayed
        outside={**utc,'startAt':(end+timedelta(seconds=1)).isoformat(),'endAt':(end+timedelta(seconds=2)).isoformat()}
        assert ok(owner.get(ROOT+'/audit',params=outside))['total']==0
        reversed_range={**utc,'startAt':(end+timedelta(seconds=1)).isoformat()}
        assert owner.get(ROOT+'/audit',params=reversed_range).status_code==400
        assert owner.get(ROOT+'/audit/export',params=reversed_range).status_code==400


def test_preflight_persists_and_detects_changed_organization_without_partial_accounts(school):
    with login(school)[0] as client:
        college=org(client,'colleges','工程学院','C01')
        plan=preview(client,'staff',[['T001','教师甲','C01','工程学院','','','','','']])
        reread=ok(client.get(ROOT+f'/identity-import/jobs/{plan["id"]}'));assert reread['entries'][0]['number']=='T001'
        ok(client.put(ROOT+f'/organizations/colleges/{college["id"]}',json={'name':'新工程学院','code':'C01','expectedVersion':college['version'],'reason':'模拟预检后的名称变更'}))
        conflict=client.put(ROOT+f'/organizations/colleges/{college["id"]}',json={'name':'未保存的输入','code':'C01','expectedVersion':college['version'],'reason':'模拟并发维护'})
        assert conflict.status_code==409 and conflict.json()['details']['current']['name']=='新工程学院'
        assert client.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':plan['version']}).status_code==409
        assert not ok(client.get(ROOT+'/accounts',params={'keyword':'T001'}))['items']


def test_durable_queued_import_rolls_back_every_row_then_explicit_retry(school,monkeypatch):
    from app.services import system_identity_import as identity
    from app.services.system_import_runner import process_next
    client=login(school)[0]  # Execute the persisted job explicitly, without starting an in-process runner.
    try:
        plan=preview(client,'staff',[[f'QUEUE{i:03d}',f'合成教师{i}','Q01','队列学院','','','','',''] for i in range(26)])
        queued=ok(client.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':plan['version']}))
        assert queued['status']=='QUEUED'
        original=identity.hash_password
        calls=0
        def interrupted(value):
            nonlocal calls
            calls+=1
            if calls==3: raise RuntimeError('simulated worker interruption')
            return original(value)
        monkeypatch.setattr(identity,'hash_password',interrupted)
        assert process_next(int(plan['id']))
        failed=ok(client.get(ROOT+f'/identity-import/jobs/{plan["id"]}'))
        assert failed['status']=='FAILED'
        assert not ok(client.get(ROOT+'/accounts',params={'keyword':'QUEUE'}))['items']
        assert not ok(client.get(ROOT+'/organizations/colleges',params={'keyword':'Q01'}))['items']
        monkeypatch.setattr(identity,'hash_password',original)
        ok(client.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':failed['version']}))
        assert process_next(int(plan['id']))
        assert ok(client.get(ROOT+f'/identity-import/jobs/{plan["id"]}'))['status']=='IMPORTED'
        assert ok(client.get(ROOT+'/accounts',params={'keyword':'QUEUE'}))['total']==26
        assert not process_next(int(plan['id']))
    finally: client.close()


def test_binding_repair_retains_multiple_history_rows_and_revokes_old_sessions(school):
    from sqlalchemy.exc import IntegrityError
    with login(school)[0] as client:
        _,credentials=import_rows(client,'students',[[n,'同名合成学生','C01','工程学院','M01','软件技术','B01','软件一班','2026'] for n in ['S001','S002','S003']])
        student=ok(client.get(ROOT+'/people/student',params={'keyword':'S001'}))['items'][0]
        accounts={r['loginName']:r for r in ok(client.get(ROOT+'/accounts'))['items']}
        old=reset_for_test(client,school,'S001',credentials['S001'],True)
        same=ok(client.get(ROOT+f'/people/student/{student["id"]}/binding-preview',params={'userId':accounts['S001']['id']}))
        assert same['unchanged'] is True and '无需调整' in same['note']
        for name in ['S002','S003','S001']:
            plan=ok(client.get(ROOT+f'/people/student/{student["id"]}/binding-preview',params={'userId':accounts[name]['id']}))
            assert plan['unchanged'] is False
            body={'userId':accounts[name]['id'],'fingerprint':plan['fingerprint'],'expectedVersion':plan['expectedVersion'],'reason':'根据合成数据核对误绑并保留历次历史'}
            if name=='S002':
                assert client.post(ROOT+f'/people/student/{student["id"]}/binding',json={**body,'fingerprint':'0'*64}).status_code==409
            ok(client.post(ROOT+f'/people/student/{student["id"]}/binding',json=body))
        assert old.get('/api/v1/auth/me').status_code==401;old.close()
        with get_sessionmaker()() as db:
            links=db.scalars(select(StudentAccountLink).where(StudentAccountLink.tenant_id==school['id'],StudentAccountLink.student_id==int(student['id']))).all()
            assert len([r for r in links if r.link_status=='REVOKED'])==3
            assert len([r for r in links if r.link_status=='ACTIVE'])==1
            db.add(StudentAccountLink(tenant_id=school['id'],student_id=int(student['id']),user_id=int(accounts['S003']['id']),link_status='ACTIVE'))
            with pytest.raises(IntegrityError): db.flush()
            db.rollback()


def test_department_only_import_export_and_error_file_keep_text_ids(school):
    with login(school)[0] as client:
        bad=preview(client,'staff',[['000007','部门教师','','','','','','','','D01','']])
        assert bad['invalidRows']==1
        errors=client.get(ROOT+f'/identity-import/jobs/{bad["id"]}/errors')
        assert errors.status_code==200
        wb=load_workbook(io.BytesIO(errors.content),read_only=True)
        assert any(r[0].value=='000007' for r in wb.active.iter_rows());wb.close()
        _,credentials=import_rows(client,'staff',[['000007','部门教师','','','','','','','','D01','实习管理处']])
        person=ok(client.get(ROOT+'/people/staff'))['items'][0]
        assert person['collegeId']=='' and person['departmentId']
        assert person['organizationLabel']=='实习管理处'
        export=client.get(ROOT+'/exports/staff',params={'keyword':'000007'});assert export.status_code==200
        wb=load_workbook(io.BytesIO(export.content),read_only=True)
        rows=list(wb.active.iter_rows(values_only=True));wb.close()
        assert any(row[0]=='000007' for row in rows)
        exported=dict(zip(rows[0],rows[1]))
        assert exported['组织归属']=='实习管理处' and exported['部门代码']=='D01'
        assert credentials['000007'] not in str(rows)
        selected=ok(client.get(ROOT+'/selection',params={'kind':'departments','identifier':person['departmentId']}))
        assert selected['name']=='实习管理处'
        changed=ok(client.put(ROOT+'/people/staff/'+person['id'],json={'expectedVersion':person['version'],
            'name':'部门教师改名','departmentId':person['departmentId'],'reason':'资料维护验收'}))
        assert changed['version']==person['version']+1


def test_directory_export_preserves_organization_codes_hierarchy_and_major_fields(school):
    with login(school)[0] as owner:
        import_rows(owner,'students',[
            ['000001','同院名称学生甲','0001','同名学院','0011','同名专业','0111','同名班级','2026'],
            ['000002','同院名称学生乙','0002','同名学院','0012','同名专业','0112','同名班级','2026']])
        major=next(r for r in ok(owner.get(ROOT+'/organizations/majors'))['items'] if r['code']=='0011')
        ok(owner.put(ROOT+'/organizations/majors/'+major['id'],json={'name':major['name'],'code':major['code'],
            'parentId':major['parentId'],'nationalCode':'051234','trainingLevel':'HIGHER','educationYears':3,
            'expectedVersion':major['version'],'reason':'维护需导出的正式专业资料'}))
        def exported(kind):
            response=owner.get(ROOT+'/exports/'+kind);assert response.status_code==200
            book=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
            raw=list(book.active.iter_rows(values_only=True))
            assert all(c.data_type=='s' for row in book.active.iter_rows(min_row=2,max_col=1) for c in row)
            book.close();return [dict(zip(raw[0],r)) for r in raw[1:]],raw
        students,raw=exported('student')
        assert [(r['学号'],r['院系代码'],r['专业代码'],r['班级代码']) for r in students]==[
            ('000001','0001','0011','0111'),('000002','0002','0012','0112')]
        assert all(r['组织归属']=='同名学院 / 同名专业 / 同名班级' for r in students)
        job=next(r for r in ok(owner.get(ROOT+'/tasks'))['items'] if r['canDownload'])
        assert job['name']=='学生资料导出.xlsx' and job['createdAt']
        receipt=owner.get(ROOT+'/tasks/'+job['id']+'/result');assert receipt.status_code==200
        from urllib.parse import unquote
        assert f'任务{job["id"]}导出结果.xlsx' in unquote(receipt.headers['content-disposition'])
        book=load_workbook(io.BytesIO(receipt.content),read_only=True,data_only=True)
        assert list(book.active.iter_rows(values_only=True))==raw;book.close()
        majors,_=exported('majors');first=next(r for r in majors if r['组织代码']=='0011')
        assert first['上级组织代码']=='0001' and first['上级组织名称']=='同名学院'
        assert first['国家专业目录代码']=='051234' and first['学制（年）']==3 and first['培养层次']=='高职专科'
        classes,_=exported('classes');first=next(r for r in classes if r['组织代码']=='0111')
        assert first['上级组织代码']=='0011' and first['上级组织名称']=='同名专业' and first['年级']=='2026'
        assert not any('密码' in h or '电话' in h or '令牌' in h for h in raw[0])


def test_person_transfer_previews_relations_preserves_records_and_audits_class_change(school):
    from datetime import datetime
    from app.models import StaffProfile, InternshipBatch, InternshipRecord, UnifiedTodo
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','调岗教师','C01','工程学院','','','','','']])
        import_rows(owner,'students',[
            ['S001','调班学生','C01','工程学院','M01','软件技术','B01','原班','2024'],
            ['S002','跨院对照学生','C02','另一学院','M02','另一专业','B02','新班','2024'],
            ['S003','未分配学生','C01','工程学院','M01','软件技术','B01','原班','2024']])
        students={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        staff=ok(owner.get(ROOT+'/people/staff'))['items'][0]
        ok(owner.post(ROOT+'/scopes',json={'userId':staff['userId'],'roleCode':'INTERN_MENTOR','scopeType':'CLASS','refId':students['S001']['classId'],'reason':'保留原班负责关系'}))
        with get_sessionmaker()() as db:
            unbound=StaffProfile(tenant_id=school['id'],staff_no='UNBOUND',real_name='尚未开户教师',college_id=int(staff['collegeId']),status='ACTIVE')
            batch=InternshipBatch(tenant_id=school['id'],batch_name='人员调整合成批次',batch_no='TRANSFER',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add_all([unbound,batch]);db.flush()
            records=[InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=int(students[number]['id']),status='ONBOARD',
                advisor_user_id=int(staff['userId']) if number!='S003' else None,advisor_name='调岗教师' if number!='S003' else None) for number in students]
            db.add_all(records);db.flush()
            db.add(UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type='INTERN_WEEKLY',source_biz_id=records[0].id,
                todo_type='INTERN_WEEKLY_REVIEW',assignee_id=int(staff['userId']),student_id=int(students['S001']['id']),title='合成待批阅事项',status='PENDING'))
            db.commit();unbound_id=str(unbound.id);record_id=records[0].id
        empty=ok(owner.get(ROOT+'/people/staff/'+unbound_id+'/impact'))
        assert empty['internshipCount']==0 and empty['pendingCount']==0 and empty['scopeCount']==0
        assert ok(owner.get(ROOT+'/people/staff/'+staff['id']+'/impact'))['internshipCount']==2
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':staff['userId'],'enabled':True,'expectedVersion':role['version'],'reason':'合成院系资料管理员'}))
        ok(owner.post(ROOT+'/scopes',json={'userId':staff['userId'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':staff['collegeId'],'reason':'只负责本院资料维护'}))
        with reset_for_test(owner,school,'T001',credentials['T001']) as actor:
            activate(actor,'COLLEGE_ADMIN')
            assert ok(actor.get(ROOT+'/people/staff/'+staff['id']+'/impact'))['internshipCount']==1
        person=students['S001'];path=ROOT+'/people/student/'+person['id']
        impact=ok(owner.get(path+'/impact'))
        assert impact['internshipCount']==impact['pendingCount']==impact['batchCount']==1 and impact['scopeCount']==2
        assert impact['batches'][0]['name']=='人员调整合成批次'
        assert impact['responsibilities'][0]['refName']=='原班'
        body={'name':person['name'],'collegeId':students['S002']['collegeId'],'classId':students['S002']['classId'],'grade':'2024',
            'expectedVersion':person['version'],'reason':'已核对在办事项后换班'}
        assert owner.put(path,json=body).status_code==409
        ok(owner.put(path,json={**body,'confirmImpact':True}))
        current=ok(owner.get(ROOT+'/people/student',params={'keyword':'S001'}))['items'][0]
        assert current['classId']==students['S002']['classId'] and current['userId']==person['userId']
        event_view=ok(owner.get(ROOT+'/audit',params={'keyword':'维护人员'}))['items'][0]
        assert {'field':'classId','label':'班级编号','before':person['classId'],'after':students['S002']['classId']} in event_view['changes']
        exported=owner.get(ROOT+'/audit/export',params={'keyword':'维护人员'})
        assert exported.status_code==200
        workbook=load_workbook(io.BytesIO(exported.content),read_only=True,data_only=True)
        rows=list(workbook.active.iter_rows(values_only=True));workbook.close()
        assert rows[0][-1]=='变更明细' and f'班级编号：{person["classId"]} → {students["S002"]["classId"]}' in rows[1][-1]
        with get_sessionmaker()() as db:
            assert db.get(InternshipRecord,record_id).advisor_user_id==int(staff['userId'])
            event=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],SecurityAuditLog.action=='PERSON_SAVE'))
            assert event.detail_json['before']['classId']==person['classId']
            assert event.detail_json['after']['classId']==students['S002']['classId']


@pytest.mark.parametrize('kind',['student','staff'])
def test_person_name_change_updates_bound_account_without_changing_identity_or_history(school,kind):
    plural='students' if kind=='student' else 'staff'
    number='S001' if kind=='student' else 'T001'
    old_name='待核对姓名';new_name='已核对姓名'
    values=[number,old_name,'C01','工程学院','M01','软件技术','B01','一班','2024'] if kind=='student' else [number,old_name,'C01','工程学院','','','','','']
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,plural,[values])
        with reset_for_test(owner,school,number,credentials[number],student=kind=='student') as person_client:
            person=ok(owner.get(ROOT+'/people/'+kind))['items'][0]
            with get_sessionmaker()() as db:
                account=db.get(User,int(person['userId']))
                original=(account.login_name,account.password_hash,account.credential_version,account.status)
                roles=list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==account.id)).all())
                historical=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],SecurityAuditLog.operator_id==account.id).order_by(SecurityAuditLog.id))
                historical_id,historical_name=historical.id,historical.operator_name
            body={'name':new_name,'collegeId':person['collegeId'],'classId':person['classId'],'departmentId':person['departmentId'],
                'grade':person['grade'],'expectedVersion':person['version'],'reason':'核实后更正人员姓名'}
            ok(owner.put(ROOT+'/people/'+kind+'/'+person['id'],json=body))
            account_view=ok(owner.get(ROOT+'/accounts',params={'keyword':number}))['items'][0]
            assert account_view['name']==new_name
            assert ok(person_client.get('/api/v1/auth/me'))['realName']==new_name
            assert owner.put(ROOT+'/people/'+kind+'/'+person['id'],json={**body,'name':'过期表单姓名'}).status_code==409
            updated=list(values);updated[1]=new_name
            checked=preview(owner,plural,[updated]);assert checked['invalidRows']==0,checked['errors']
            ok(owner.post(ROOT+'/identity-import/jobs/'+checked['id']+'/confirm',json={'expectedVersion':checked['version']}))
            with get_sessionmaker()() as db:
                account=db.get(User,int(person['userId']))
                assert account.real_name==new_name
                assert (account.login_name,account.password_hash,account.credential_version,account.status)==original
                assert list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==account.id)).all())==roles
                assert db.get(SecurityAuditLog,historical_id).operator_name==historical_name==old_name
                model=StudentProfile if kind=='student' else StaffProfile
                assert db.get(model,int(person['id'])).real_name==new_name


def test_organization_import_reuses_preview_without_creating_accounts(school):
    from app.services.system_identity_import import ORG_HEADERS
    with login(school)[0] as client:
        before=ok(client.get(ROOT+'/accounts'))['total']
        content=build_template_xlsx(list(ORG_HEADERS),samples=[['001','工程学院','0001','软件技术','00001','软件一班','2026','','']])
        plan=ok(client.post(ROOT+'/identity-import/organizations/preview',files={'file':('组织.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}))
        assert plan['invalidRows']==0
        assert ok(client.get(ROOT+'/organizations/colleges'))['total']==0
        body={'expectedVersion':plan['version']}
        result=ok(client.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json=body))
        assert result['status']=='IMPORTED' and not result['receiptAvailable']
        ok(client.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json=body))
        assert ok(client.get(ROOT+'/organizations/classes'))['items'][0]['code']=='00001'
        assert ok(client.get(ROOT+'/accounts'))['total']==before


def test_organization_disable_rechecks_references_and_preserves_children_and_people(school):
    with login(school)[0] as owner:
        import_rows(owner,'staff',[['T001','合成教师','C01','待停用学院','','','','','']])
        import_rows(owner,'students',[['S001','合成学生','C01','待停用学院','M01','原专业','B01','原班','2026']])
        college=ok(owner.get(ROOT+'/organizations/colleges'))['items'][0]
        major=ok(owner.get(ROOT+'/organizations/majors'))['items'][0]
        path=ROOT+'/organizations/colleges/'+college['id']
        impact=ok(owner.get(path+'/impact'));counts={r['name']:r['count'] for r in impact['items']}
        assert impact['hasReferences'] and counts['下级组织']==2 and counts['学生主档']==counts['教职工资料']==1
        body={'name':college['name'],'code':college['code'],'status':'DISABLED','expectedVersion':college['version'],'reason':'合成组织引用停用验收'}
        assert owner.put(path,json=body).status_code==409
        import_rows(owner,'students',[['S002','新增合成学生','C01','待停用学院','M01','原专业','B01','原班','2026']])
        stale={**body,'confirmImpact':True,'impactFingerprint':impact['fingerprint']}
        assert owner.put(path,json=stale).status_code==409
        current=ok(owner.get(path+'/impact'))
        ok(owner.put(path,json={**stale,'impactFingerprint':current['fingerprint']}))
        assert ok(owner.get(ROOT+'/organizations/colleges'))['items'][0]['status']=='DISABLED'
        assert ok(owner.get(ROOT+'/people/student'))['total']==2
        assert ok(owner.get(ROOT+'/organizations/classes'))['total']==1
        # A stopped parent must not prevent orderly stopping of an already existing child.
        major_path=ROOT+'/organizations/majors/'+major['id']
        reviewed=ok(owner.get(major_path+'/impact'))
        ok(owner.put(major_path,json={'name':major['name'],'code':major['code'],'parentId':college['id'],'status':'DISABLED',
            'expectedVersion':major['version'],'reason':'继续停用原专业，保留资料','confirmImpact':True,'impactFingerprint':reviewed['fingerprint']}))
        assert ok(owner.get(ROOT+'/people/student'))['total']==2


@pytest.mark.parametrize('kind',['majors','classes'])
def test_organization_parent_move_requires_impact_and_updates_current_scope_without_rewriting_history(school,kind):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord,InternshipBatchParticipant,UnifiedTodo,Major,SchoolClass
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[['T001','原院管理员','C01','原学院','','','','','']])
        import_rows(owner,'students',[[n,n,c,cn,m,mn,b,bn,'2024'] for n,c,cn,m,mn,b,bn in [
            ('S001','C01','原学院','M01','原专业','B01','原班'),('S002','C02','目标学院','M02','目标专业','B02','目标班')]])
        people={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        original,target=people['S001'],people['S002']
        staff=ok(owner.get(ROOT+'/people/staff'))['items'][0]
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':staff['userId'],'enabled':True,'expectedVersion':role['version'],'reason':'核对原院当前范围'}))
        ok(owner.post(ROOT+'/scopes',json={'userId':staff['userId'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':original['collegeId'],'reason':'原院负责范围'}))
        with get_sessionmaker()() as db:
            batch=InternshipBatch(tenant_id=school['id'],batch_name='组织调整历史保留',batch_no='ORG-MOVE',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.flush()
            record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=int(original['id']),status='PREPARING',advisor_user_id=int(staff['userId']),advisor_name='原院管理员')
            db.add(record);db.flush()
            frozen=InternshipBatchParticipant(tenant_id=school['id'],batch_id=batch.id,student_id=int(original['id']),internship_id=record.id,snapshot_student_no='S001',snapshot_name='S001',snapshot_class_name='原班',snapshot_college_name='原学院')
            todo=UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type='INTERN_WEEKLY',source_biz_id=record.id,todo_type='INTERN_WEEKLY_REVIEW',assignee_id=int(staff['userId']),student_id=int(original['id']),title='合成历史责任',status='PENDING')
            db.add_all([frozen,todo]);db.commit();rid=record.id;fid=frozen.id;tid=todo.id
        item=next(r for r in ok(owner.get(ROOT+'/organizations/'+kind))['items'] if r['id']==original['majorId' if kind=='majors' else 'classId'])
        parent=target['collegeId' if kind=='majors' else 'majorId'];path=ROOT+'/organizations/'+kind+'/'+item['id']
        body={'name':item['name'],'code':item['code'],'parentId':parent,'grade':item['grade'] or '', 'status':'ACTIVE','expectedVersion':item['version'],'reason':'核对后调整组织归属'}
        reviewed=ok(owner.get(path+'/impact'))
        assert owner.put(path,json=body).status_code==409
        assert owner.put(path,json={**body,'confirmImpact':True,'impactFingerprint':'0'*64}).status_code==409
        with reset_for_test(owner,school,'T001',creds['T001']) as actor:
            activate(actor,'COLLEGE_ADMIN')
            assert ok(actor.get(ROOT+'/people/student'))['total']==1
            assert actor.put(path,json={**body,'confirmImpact':True,'impactFingerprint':reviewed['fingerprint']}).status_code==403
            moved=ok(owner.put(path,json={**body,'confirmImpact':True,'impactFingerprint':reviewed['fingerprint']}))
            assert moved['id']==item['id'] and moved['parentId']==parent and '1名学生' in moved['message']
            assert ok(actor.get(ROOT+'/people/student'))['total']==0
            current=ok(owner.get(ROOT+'/people/student',params={'keyword':'S001'}))['items'][0]
            assert current['id']==original['id'] and current['classId']==original['classId'] and current['userId']==original['userId']
            assert current['collegeId']==target['collegeId'] and current['majorId']==(original['majorId'] if kind=='majors' else target['majorId'])
            assert current['version']==original['version']+1
            with get_sessionmaker()() as db:
                record=db.get(InternshipRecord,rid);frozen=db.get(InternshipBatchParticipant,fid)
                assert record.student_id==int(original['id']) and record.status=='PREPARING' and record.advisor_user_id==int(staff['userId'])
                assert frozen.snapshot_college_name=='原学院' and frozen.snapshot_class_name=='原班'
                assert db.get(UnifiedTodo,tid).assignee_id==int(staff['userId']) and db.get(UnifiedTodo,tid).status=='PENDING'
                event=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],SecurityAuditLog.action=='ORG_SAVE',SecurityAuditLog.resource_id==item['id']).order_by(SecurityAuditLog.id.desc()))
                assert event.detail_json['before']['parentId']==item['parentId'] and event.detail_json['after']['parentId']==parent
                assert event.detail_json['updatedStudentCount']==1


@pytest.mark.parametrize('kind,stopped_kind',[
    ('student','colleges'),('student','majors'),('student','classes'),
    ('staff','colleges'),('staff','departments')])
def test_stopped_organization_allows_existing_person_correction_but_rejects_transfer(school,kind,stopped_kind):
    with login(school)[0] as owner:
        if stopped_kind=='departments':
            rows=[['P001','原部门人员','','','','','','','','D01','原部门'],
                  ['P002','对照部门人员','','','','','','','','D02','对照部门']]
        else:
            rows=[['P001','原组织人员','C01','原学院','M01','原专业','B01','原班','2026'],
                  ['P002','对照组织人员','C02','对照学院','M02','对照专业','B02','对照班','2026']]
        import_rows(owner,'students' if kind=='student' else 'staff',rows)
        people={r['number']:r for r in ok(owner.get(ROOT+'/people/'+kind))['items']}
        original,other=people['P001'],people['P002']
        target=next(r for r in ok(owner.get(ROOT+'/organizations/'+stopped_kind))['items'] if r['code'].endswith('01'))
        path=ROOT+'/organizations/'+stopped_kind+'/'+target['id']
        impact=ok(owner.get(path+'/impact'))
        ok(owner.put(path,json={'name':target['name'],'code':target['code'],'parentId':target['parentId'],
            'grade':target['grade'] or '', 'expectedVersion':target['version'],'status':'DISABLED',
            'confirmImpact':True,'impactFingerprint':impact['fingerprint'],'reason':'合成停用组织保留资料验收'}))
        body={key:original[key] for key in ('name','collegeId','classId','departmentId','grade','status')}
        correction={**body,'name':'原组织人员姓名更正','expectedVersion':original['version'],'reason':'原归属不变更正资料'}
        ok(owner.put(ROOT+'/people/'+kind+'/'+original['id'],json=correction))
        current=ok(owner.get(ROOT+'/people/'+kind,params={'keyword':'P001'}))['items'][0]
        assert current['name']==correction['name'] and current['userId']==original['userId']
        assert all(current[key]==original[key] for key in ('collegeId','classId','departmentId'))
        # The same stopped destination is forbidden for a different person's transfer.
        rejected=owner.put(ROOT+'/people/'+kind+'/'+other['id'],json={**body,'name':other['name'],
            'expectedVersion':other['version'],'confirmImpact':True,'reason':'不得调入停用组织'})
        assert rejected.status_code in {400,403}
        unchanged=ok(owner.get(ROOT+'/people/'+kind,params={'keyword':'P002'}))['items'][0]
        assert unchanged['version']==other['version'] and unchanged['collegeId']==other['collegeId']
        assert unchanged['classId']==other['classId'] and unchanged['departmentId']==other['departmentId']


def test_mobile_role_switch_after_refresh_uses_current_rotation_and_role_permissions(school):
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','移动身份核对教师','C01','工程学院','','','','','']])
        with reset_for_test(owner,school,'T001',credentials['T001']):
            pass
        target=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        teacher=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        ok(owner.post(ROOT+'/roles/'+target['id']+'/members',json={'userId':teacher['id'],
            'enabled':True,'expectedVersion':target['version'],'reason':'合成移动身份切换前置'}))
    with TestClient(app) as mobile:
        logged=ok(mobile.post('/api/v1/auth/login',json={'tenantCode':school['code'],'loginName':'T001',
            'password':PASSWORD,'clientType':'TEACHER_MINI'}))
        renewed=ok(mobile.post('/api/v1/auth/refresh',json={'refreshToken':logged['refreshToken']}))
        headers={'Authorization':'Bearer '+renewed['accessToken']}
        path='/api/v1/auth/switch-role'
        assert mobile.post(path,headers=headers,json={'roleCode':'COLLEGE_ADMIN','refreshToken':logged['refreshToken']}).status_code==401
        assert mobile.post(path,headers=headers,json={'roleCode':'SYSTEM_ADMIN','refreshToken':renewed['refreshToken']}).status_code==403
        switched=ok(mobile.post(path,headers=headers,json={'roleCode':'COLLEGE_ADMIN','refreshToken':renewed['refreshToken']}))
        context=ok(mobile.get('/api/v1/rbac/current-context',headers={'Authorization':'Bearer '+switched['accessToken']}))
        assert context['currentRole']['roleCode']=='COLLEGE_ADMIN'
        assert 'system.student.view' in context['permissionPatterns']
        assert 'system.role.manage' not in context['permissionPatterns']
        assert mobile.post(path,headers=headers,json={'roleCode':'INTERN_MENTOR','refreshToken':renewed['refreshToken']}).status_code==401


def test_import_rechecks_alias_uniqueness_after_organization_changes(school):
    with login(school)[0] as owner:
        import_rows(owner,'staff',[['T001','别名核对教师','C01','正式学院','','','','','']])
        config={'announcementCategories':['实习通知'],'aliases':[{'alias':'原学院简称','code':'C01'}],
            'reason':'核对导入别名实际消费','expectedVersion':0}
        ok(owner.put(ROOT+'/config/dictionary',json=config))
        rows=[['T001','别名核对教师','C01','原学院简称','','','','','']]
        accepted=preview(owner,'staff',rows)
        assert accepted['invalidRows']==0 and accepted['newOrganizations']==[]
        ok(owner.post(ROOT+f'/identity-import/jobs/{accepted["id"]}/confirm',json={'expectedVersion':accepted['version']}))
        pending=preview(owner,'staff',rows)
        org(owner,'departments','另一个同代码部门','C01')
        ambiguous=preview(owner,'staff',rows)
        assert ambiguous['invalidRows']==1
        assert any('别名' in item['message'] and '核对' in item['message'] for item in ambiguous['errors'])
        assert owner.post(ROOT+f'/identity-import/jobs/{pending["id"]}/confirm',json={'expectedVersion':pending['version']}).status_code==409
        canonical=preview(owner,'staff',[['T001','别名核对教师','C01','正式学院','','','','','']])
        assert canonical['invalidRows']==0,canonical['errors']
        with get_sessionmaker()() as db:
            person=db.scalar(select(StaffProfile).where(StaffProfile.tenant_id==school['id'],StaffProfile.staff_no=='T001'))
            assert person.department_id is None and person.college_id is not None
            assert db.scalar(select(func.count()).select_from(User).where(User.tenant_id==school['id']))==2


def test_account_unlock_and_enable_require_distinct_current_states(school):
    with login(school)[0] as owner:
        _, credentials=import_rows(owner,'staff',[['T001','解锁核对教师','C01','工程学院','','','','','']])
        with reset_for_test(owner,school,'T001',credentials['T001']):
            pass
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        path=ROOT+'/accounts/'+account['id']+'/actions'
        disabled=ok(owner.post(path,json={'action':'DISABLE','expectedVersion':account['version'],'reason':'合成账号停用验收'}))
        with get_sessionmaker()() as db:
            user=db.get(User,int(account['id']))
            before=(user.status,user.version,user.credential_version,user.password_hash,user.expires_at,user.must_change_password)
        rejected=owner.post(path,json={'action':'UNLOCK','expectedVersion':disabled['version'],'reason':'解锁不能代替启用停用账号'})
        assert rejected.status_code==409
        with get_sessionmaker()() as db:
            user=db.get(User,int(account['id']))
            assert (user.status,user.version,user.credential_version,user.password_hash,user.expires_at,user.must_change_password)==before
            assert db.scalar(select(func.count()).select_from(SecurityAuditLog).where(
                SecurityAuditLog.tenant_id==school['id'],SecurityAuditLog.action=='ACCOUNT_UNLOCK'))==0
        enabled=ok(owner.post(path,json={'action':'ENABLE','expectedVersion':disabled['version'],'reason':'明确启用已停用合成账号'}))
        assert enabled['version']==disabled['version']+1
        with login(school,'T001')[0] as teacher:
            assert teacher.get('/api/v1/auth/me').status_code==200


def test_account_unlock_preserves_credentials_and_expiry_but_revokes_old_sessions(school):
    from datetime import datetime
    with login(school)[0] as owner:
        _, credentials=import_rows(owner,'staff',[['T001','解锁会话核对教师','C01','工程学院','','','','','']])
        with reset_for_test(owner,school,'T001',credentials['T001']):
            pass
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        path=ROOT+'/accounts/'+account['id']+'/actions'
        changed=ok(owner.post(path,json={'action':'EXPIRY','expiresAt':'2030-10-08T10:30:00Z',
            'expectedVersion':account['version'],'reason':'合成解锁保留有效期前置'}))
        with login(school,'T001')[0] as teacher, TestClient(app) as mobile:
            mobile_login=ok(mobile.post('/api/v1/auth/login',json={'tenantCode':school['code'],
                'loginName':'T001','password':PASSWORD,'clientType':'TEACHER_MINI'}))
            mobile_headers={'Authorization':'Bearer '+mobile_login['accessToken']}
            with get_sessionmaker()() as db:
                user=db.get(User,int(account['id']))
                unchanged=(user.password_hash,user.expires_at,user.must_change_password)
                credential_version=user.credential_version
                roles=list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==user.id)).all())
                # Only establish the synthetic authentication precondition; no business record is changed.
                user.status='LOCKED';db.commit()
            assert teacher.get('/api/v1/auth/me').status_code==401
            assert mobile.get('/api/v1/auth/me',headers=mobile_headers).status_code==401
            wrong=owner.post(path,json={'action':'ENABLE','expectedVersion':changed['version'],'reason':'启用不能代替解锁'})
            assert wrong.status_code==409
            with get_sessionmaker()() as db:
                user=db.get(User,int(account['id']))
                assert user.status=='LOCKED' and user.version==changed['version'] and user.credential_version==credential_version
            unlocked=ok(owner.post(path,json={'action':'UNLOCK','expectedVersion':changed['version'],'reason':'合成锁定账号解锁验收'}))
            assert unlocked['version']==changed['version']+1 and '原登录会话失效' in unlocked['message']
            with get_sessionmaker()() as db:
                user=db.get(User,int(account['id']))
                assert user.status=='ACTIVE' and user.credential_version==credential_version+1
                assert (user.password_hash,user.expires_at,user.must_change_password)==unchanged
                assert list(db.execute(select(UserRole.role_id,UserRole.status).where(UserRole.user_id==user.id)).all())==roles
                event=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],
                    SecurityAuditLog.action=='ACCOUNT_UNLOCK',SecurityAuditLog.resource_id==account['id']))
                assert event.detail_json['before']['status']=='LOCKED' and event.detail_json['after']['status']=='ACTIVE'
                assert PASSWORD not in str(event.detail_json)
            assert teacher.get('/api/v1/auth/me').status_code==401
            assert teacher.post('/api/v1/auth/browser-refresh',headers={'X-Browser-Session':'staff'}).status_code==401
            assert mobile.get('/api/v1/auth/me',headers=mobile_headers).status_code==401
            assert mobile.post('/api/v1/auth/refresh',json={'refreshToken':mobile_login['refreshToken']}).status_code==401
            fresh=ok(mobile.post('/api/v1/auth/login',json={'tenantCode':school['code'],
                'loginName':'T001','password':PASSWORD,'clientType':'TEACHER_MINI'}))
            assert not fresh['mustChangePassword']
            assert mobile.get('/api/v1/auth/me',headers={'Authorization':'Bearer '+fresh['accessToken']}).status_code==200
        with login(school,'T001')[0] as fresh:
            assert fresh.get('/api/v1/auth/me').status_code==200
        # Unlocking changes the lock only, and cannot renew an already expired account.
        expired=ok(owner.post(path,json={'action':'EXPIRY','expiresAt':'2020-01-01T00:00:00Z',
            'expectedVersion':unlocked['version'],'reason':'合成解锁不得延长账号有效期'}))
        with get_sessionmaker()() as db:
            user=db.get(User,int(account['id']));user.status='LOCKED';db.commit()
        ok(owner.post(path,json={'action':'UNLOCK','expectedVersion':expired['version'],'reason':'只解除合成账号锁定'}))
        with TestClient(app) as guest:
            guest.headers['X-Browser-Session-Id']='unlock-expiry-'+uuid4().hex
            for endpoint,kind in [('browser-login','PC'),('login','TEACHER_MINI')]:
                denied=guest.post('/api/v1/auth/'+endpoint,json={'tenantCode':school['code'],
                    'loginName':'T001','password':PASSWORD,'clientType':kind})
                assert denied.status_code==401 and '有效期已结束' in denied.json()['message']
        with get_sessionmaker()() as db:
            assert db.get(User,int(account['id'])).expires_at==datetime(2020,1,1)


def test_account_expiry_round_trip_has_explicit_timezone_and_preserves_instant(school):
    from datetime import datetime
    with login(school)[0] as owner:
        _, credentials=import_rows(owner,'staff',[['T001','有效期核对教师','C01','工程学院','','','','','']])
        with reset_for_test(owner,school,'T001',credentials['T001']):
            pass
        def account():
            return next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        original=account()
        assert original['expiresAt'] is None
        path=ROOT+'/accounts/'+original['id']+'/actions'
        ok(owner.post(path,json={'action':'EXPIRY','expiresAt':'2030-10-08T18:30:00+08:00',
            'expectedVersion':original['version'],'reason':'合成账号有效期时区核对'}))
        saved=account()
        assert saved['expiresAt']=='2030-10-08T10:30:00Z'
        ok(owner.post(path,json={'action':'EXPIRY','expiresAt':saved['expiresAt'],
            'expectedVersion':saved['version'],'reason':'原样重新保存账号有效期'}))
        assert account()['expiresAt']==saved['expiresAt']
        with get_sessionmaker()() as db:
            assert db.get(User,int(saved['id'])).expires_at==datetime(2030,10,8,10,30)
        stale=owner.post(path,json={'action':'EXPIRY','expiresAt':None,
            'expectedVersion':saved['version'],'reason':'过期版本不能清空有效期'})
        assert stale.status_code==409
        assert stale.json()['details']['current']['expiresAt']==saved['expiresAt']
        current=account()
        ok(owner.post(path,json={'action':'EXPIRY','expiresAt':'2020-01-01T00:00:00Z',
            'expectedVersion':current['version'],'reason':'合成账号到期状态核对'}))
        expired=account()
        assert expired['status']=='ACTIVE' and expired['statusLabel']=='有效期已结束'
        with TestClient(app) as guest:
            guest.headers['X-Browser-Session-Id']='expiry-'+uuid4().hex
            for endpoint,kind in [('browser-login','PC'),('login','TEACHER_MINI')]:
                denied=guest.post('/api/v1/auth/'+endpoint,json={'tenantCode':school['code'],
                    'loginName':'T001','password':PASSWORD,'clientType':kind})
                assert denied.status_code==401
                assert '有效期已结束' in denied.json()['message']
        ok(owner.post(path,json={'action':'DISABLE','expectedVersion':expired['version'],'reason':'到期后停用合成账号'}))
        assert account()['statusLabel']=='停用'


def test_service_owner_rejects_external_or_expired_accounts_and_stops_display_after_expiry(school):
    from datetime import datetime, timedelta
    with get_sessionmaker()() as db:
        accounts=[]
        for kind in ['ENTERPRISE_MENTOR','GUARDIAN','PLATFORM_OP','STUDENT','TEACHER']:
            account=User(tenant_id=school['id'],login_name='service-'+kind,real_name='合成服务承办人',
                user_type=kind,password_hash=hash_password(PASSWORD),status='ACTIVE',must_change_password=True)
            db.add(account);db.flush();accounts.append(str(account.id))
        db.commit()
    with login(school)[0] as client:
        payload={'contactPhone':'','faq':[],'reminderText':'学校办理提醒','expectedVersion':0,'reason':'服务承办资格验收'}
        role=next(r for r in ok(client.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        for identifier in accounts[:-1]:
            assert client.put(ROOT+'/config/notification',json={**payload,'responsibleUserId':identifier}).status_code==400
            assert client.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':identifier,
                'enabled':True,'expectedVersion':role['version'],'reason':'校内管理身份边界验收'}).status_code==403
        with get_sessionmaker()() as db:
            assert db.scalar(select(func.count()).select_from(UserRole).where(UserRole.tenant_id==school['id'],
                UserRole.user_id.in_([int(identifier) for identifier in accounts[:-1]])))==0
        ok(client.put(ROOT+'/config/notification',json={**payload,'responsibleUserId':accounts[-1]}))
        assert ok(client.get(ROOT+'/school-service'))['service']['responsibleName']=='合成服务承办人'
        account=next(r for r in ok(client.get(ROOT+'/accounts'))['items'] if r['id']==accounts[-1])
        ok(client.post(ROOT+'/accounts/'+account['id']+'/actions',json={'action':'EXPIRY',
            'expiresAt':(datetime.utcnow()-timedelta(days=1)).isoformat()+'Z',
            'expectedVersion':account['version'],'reason':'合成承办人有效期验收'}))
        assert ok(client.get(ROOT+'/school-service'))['service']['responsibleName']==''
        assert client.put(ROOT+'/config/notification',json={**payload,'expectedVersion':1,'responsibleUserId':accounts[-1]}).status_code==400


def test_service_owner_staff_disable_and_read_only_configuration_without_account_access(school):
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[[n,name,'C01','合成院系','','','','',''] for n,name in [('T001','合成服务承办人'),('T002','合成配置查看人')]])
        accounts={row['loginName']:row for row in ok(owner.get(ROOT+'/accounts'))['items']}
        setting={'responsibleUserId':accounts['T001']['id'],'reminderText':'合成学校服务提醒','expectedVersion':0,'reason':'核对在职承办与只读配置'}
        ok(owner.put(ROOT+'/config/notification',json=setting))
        assert ok(owner.get(ROOT+'/config/notification'))['responsibleName']=='合成服务承办人'
        role=ok(owner.post(ROOT+'/roles',json={'name':'学校服务配置查看岗','baseRole':'LEADER','permissions':['system.notification.view'],'reason':'合成只读配置身份'}))
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':accounts['T002']['id'],'enabled':True,'expectedVersion':role['version'],'reason':'合成配置查询前置'}))
        role=next(row for row in ok(owner.get(ROOT+'/roles'))['items'] if row['id']==role['id'])
        with reset_for_test(owner,school,'T002',credentials['T002']) as reader:
            switched=ok(reader.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+role['code'].lower(),'clientType':'PC'}))
            reader.headers['Authorization']='Bearer '+switched['accessToken']
            readonly=ok(reader.get(ROOT+'/config/notification'))
            assert readonly['responsibleName']=='合成服务承办人' and readonly['version']==1
            assert readonly['responsibleValid'] is True
            assert reader.get(ROOT+'/accounts').status_code==403
            assert reader.get(ROOT+'/selection',params={'kind':'accounts','identifier':accounts['T001']['id']}).status_code==403
            assert reader.put(ROOT+'/config/notification',json={**setting,'expectedVersion':1}).status_code==403
            staff=next(row for row in ok(owner.get(ROOT+'/people/staff'))['items'] if row['number']=='T001')
            assert ok(owner.get(ROOT+'/people/staff/'+staff['id']+'/impact'))['serviceResponsibilityCount']==1
            body={key:staff[key] for key in ('name','collegeId','departmentId','classId','grade')}
            ok(owner.put(ROOT+'/people/staff/'+staff['id'],json={**body,'status':'DISABLED','expectedVersion':staff['version'],'reason':'合成承办人离职核对'}))
            assert next(row for row in ok(owner.get(ROOT+'/accounts'))['items'] if row['loginName']=='T001')['status']=='ACTIVE'
            assert ok(reader.get(ROOT+'/config/notification'))['responsibleName']==''
            assert ok(reader.get(ROOT+'/school-service'))['service']['responsibleName']==''
            assert all(row['responsibleName']=='学校尚未指定服务承办人' for row in ok(owner.get(ROOT+'/integrations'))['items'])
            rejected=owner.put(ROOT+'/config/notification',json={**setting,'expectedVersion':1})
            assert rejected.status_code==400 and '资料在职' in rejected.json()['message']
            current=ok(owner.get(ROOT+'/config/notification'))
            assert current['version']==1 and current['responsibleUserId']==accounts['T001']['id']
            assert current['responsibleValid'] is False
            assert ok(owner.get(ROOT+'/people/staff/'+staff['id']+'/impact'))['serviceResponsibilityCount']==1


def test_school_logo_current_binding_replacement_and_tenant_for_plain_student(school):
    from PIL import Image
    from app.models.file import FileObject
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'students',[['LOGO001','校徽查看合成学生','C01','工程学院','M01','软件技术','B01','一班','2024']])
        with reset_for_test(owner,school,'LOGO001',credentials['LOGO001'],True) as student:
            files=[]
            for color in ['white','blue']:
                picture=io.BytesIO();Image.new('RGB',(8,8),color).save(picture,format='PNG')
                image=ok(owner.post('/api/v1/files?bizType=SYSTEM_BRAND',files={'file':('合成校徽.png',picture.getvalue(),'image/png')}))
                assert image['readyForBusiness'];files.append(image['fileId'])
                assert student.get('/api/v1/files/download/'+image['fileId']).status_code in (403,404)
                brand=ok(owner.get(ROOT+'/config/brand'))
                ok(owner.put(ROOT+'/config/brand',json={'schoolName':'合成校徽学校','platformName':'合成校徽平台',
                    'logoFileId':image['fileId'],'expectedVersion':brand['version'],'reason':'校徽绑定更换验收'}))
                assert ok(student.get(ROOT+'/school-service'))['brand']['logoFileId']==image['fileId']
                assert ok(student.get('/api/v1/files/'+image['fileId']))['readyForBusiness']
                assert student.get('/api/v1/files/preview/'+image['fileId']).content==picture.getvalue()
                assert student.get('/api/v1/files/download/'+image['fileId']).content==picture.getvalue()
            for suffix in [files[0],'preview/'+files[0],'download/'+files[0]]:
                assert student.get('/api/v1/files/'+suffix).status_code in (403,404)
            assert student.get(ROOT+'/config/brand').status_code==403
            assert student.post('/api/v1/files?bizType=SYSTEM_BRAND',files={'file':('合成校徽.png',picture.getvalue(),'image/png')}).status_code==403
            with get_sessionmaker()() as db:
                foreign=Tenant(tenant_code='SYS-'+uuid4().hex[:12],school_name='校徽隔离合成学校',status='ACTIVE')
                db.add(foreign);db.flush()
                account=User(tenant_id=foreign.id,login_name='owner',real_name='校徽隔离负责人',user_type='SCHOOL_ADMIN',
                    password_hash=hash_password(PASSWORD),status='ACTIVE',must_change_password=False)
                db.add(account);db.flush();bootstrap_school(db,foreign.id);db.commit()
                other={'id':foreign.id,'code':foreign.tenant_code,'owner':account.id}
                assert all(not db.get(FileObject,int(id)).is_deleted for id in files)
            with login(other)[0] as outsider:
                for suffix in [files[1],'preview/'+files[1],'download/'+files[1]]:
                    assert outsider.get('/api/v1/files/'+suffix).status_code in (403,404)
            brand=ok(owner.get(ROOT+'/config/brand'))
            ok(owner.put(ROOT+'/config/brand',json={'schoolName':'合成校徽学校','platformName':'合成校徽平台',
                'logoFileId':'','expectedVersion':brand['version'],'reason':'清除当前校徽验收'}))
            assert ok(student.get(ROOT+'/school-service'))['brand']['logoFileId']==''
            assert student.get('/api/v1/files/download/'+files[1]).status_code in (403,404)


def test_school_config_is_consumed_and_document_download_rechecks_role(school):
    from PIL import Image
    with login(school)[0] as client:
        picture=io.BytesIO();Image.new('RGB',(8,8),'white').save(picture,format='PNG')
        uploaded=ok(client.post('/api/v1/files?bizType=SYSTEM_BRAND',files={'file':('合成校徽.png',picture.getvalue(),'image/png')}))
        assert uploaded['readyForBusiness']
        brand=ok(client.get(ROOT+'/config/brand'))
        ok(client.put(ROOT+'/config/brand',json={'schoolName':'合成配置学校','platformName':'合成学校实习平台',
            'contactPhone':'000-0000001','supervisionPhone':'000-0000002',
            'logoFileId':uploaded['fileId'],'expectedVersion':brand['version'],'reason':'品牌消费验收'}))
        service=ok(client.get(ROOT+'/school-service'))
        assert service['service']['contactPhone']=='000-0000001'
        assert service['service']['supervisionPhone']=='000-0000002'
        ok(client.put(ROOT+'/config/notification',json={'contactPhone':'000-0000000','faq':[{'question':'如何开始','answer':'从实习工作台查看学校安排'}],
            'reminderText':'请关注学校办理提醒','expectedVersion':0,'reason':'咨询配置消费验收'}))
        service=ok(client.get(ROOT+'/school-service'))
        assert service['brand']['platformName']=='合成学校实习平台'
        assert service['service']['contactPhone']=='000-0000000'
        assert service['service']['supervisionPhone']=='000-0000002'
        assert service['service']['faq'][0]['answer']=='从实习工作台查看学校安排'
        changes=ok(client.get(ROOT+'/audit',params={'keyword':'修改学校配置'}))['items'][0]['changes']
        reminder_change=next(r for r in changes if r['field']=='reminderText')
        assert reminder_change=={'field':'reminderText','label':'提醒附言',
            'before':'请及时处理岗位实习待办。','after':'请关注学校办理提醒'}
        assert '000-0000000' not in str(changes)
        exported=client.get(ROOT+'/audit/export',params={'keyword':'修改学校配置'})
        assert exported.status_code==200
        book=load_workbook(io.BytesIO(exported.content),read_only=True,data_only=True)
        export_rows=list(book.active.iter_rows(values_only=True));book.close()
        assert any('提醒附言：请及时处理岗位实习待办。 → 请关注学校办理提醒' in str(r) for r in export_rows)
        assert export_rows[0][5]=='时间（北京时间）'
        from datetime import timedelta
        with get_sessionmaker()() as db:
            recorded=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],
                SecurityAuditLog.action=='CONFIG_SAVE').order_by(SecurityAuditLog.id.desc()))
            assert export_rows[1][5]==recorded.created_at+timedelta(hours=8)
        ok(client.put(ROOT+'/config/notification',json={'contactPhone':'','faq':[{'question':'如何开始','answer':'从实习工作台查看学校安排'}],
            'reminderText':'请关注学校办理提醒','expectedVersion':1,'reason':'服务咨询电话恢复学校默认'}))
        assert ok(client.get(ROOT+'/school-service'))['service']['contactPhone']=='000-0000001'
        invalid_alias={'announcementCategories':['实习通知'],'aliases':[{'alias':'旧称','code':'MISSING'}],'reason':'校验组织匹配别名'}
        assert client.put(ROOT+'/config/dictionary',json=invalid_alias).status_code==400
        org(client,'colleges','别名核对学院','ALIAS01')
        valid_alias={**invalid_alias,'aliases':[{'alias':'旧称','code':'ALIAS01'}]}
        ok(client.put(ROOT+'/config/dictionary',json=valid_alias))
        org(client,'departments','另一个同代码部门','ALIAS01')
        assert client.put(ROOT+'/config/dictionary',json={**valid_alias,'expectedVersion':1}).status_code==400
        from app.services.message_event_outbox_service import emit_message_event
        with get_sessionmaker()() as db:
            reminder=emit_message_event(db,event_code='INTERNSHIP.WEEKLY_REMIND',source_module='internship',source_biz_type='WEEKLY',
                source_biz_id=123,recipient_refs=[{'userId':school['owner']}],content='原办理事项',tenant_id=school['id'])
            assert reminder.payload_json['content']=='原办理事项\n\n请关注学校办理提醒'
            db.rollback()
        for route in ['/integrations','/runtime','/records/MIGRATION','/audit']:
            assert isinstance(ok(client.get(ROOT+route))['items'],list)
        assert ok(client.get('/api/v1/rbac/current-context'))['brand']['logoFileId']==uploaded['fileId']
        document=ok(client.post('/api/v1/files?bizType=SYSTEM_DOCUMENT',files={'file':('合成培训图.png',picture.getvalue(),'image/png')}))
        record=ok(client.post(ROOT+'/records/DOCUMENT',json={'title':'合成培训资料','fileIds':[document['fileId']],
            'reason':'受控附件下载验收'}))
        _,credentials=import_rows(client,'staff',[['VIEWER','资料查看人','C01','工程学院','','','','','']])
        reader=reset_for_test(client,school,'VIEWER',credentials['VIEWER'])
        account=ok(client.get(ROOT+'/accounts',params={'keyword':'VIEWER'}))['items'][0]
        role=ok(client.post(ROOT+'/roles',json={'name':'资料查看','baseRole':'INTERN_MENTOR',
            'permissions':['system.delivery.view'],'reason':'最小权限文件验收'}))
        role=ok(client.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],
            'enabled':True,'expectedVersion':role['version'],'reason':'配置资料查看职责'}))
        current=next(r for r in ok(client.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
        assert ok(client.get(ROOT+f'/roles/{role["id"]}/members'))['items'][0]['name']=='资料查看人'
        switched=ok(reader.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+current['code'].lower(),'clientType':'PC'}))
        reader.headers['Authorization']='Bearer '+switched['accessToken']
        assert reader.get('/api/v1/files/download/'+document['fileId']).status_code==200
        ok(client.put(ROOT+f'/roles/{role["id"]}',json={'name':'资料查看','baseRole':'INTERN_MENTOR',
            'permissions':[],'expectedVersion':role['version'],'reason':'回收资料查看权限'}))
        assert reader.get('/api/v1/files/download/'+document['fileId']).status_code in (403,404)
        assert reader.get(ROOT+'/records/DOCUMENT').status_code==403
        revoked=ok(client.get(ROOT+'/audit',params={'keyword':'调整角色权限'}))['items'][0]
        assert {'field':'removed','label':'撤销功能权限','before':'查看交付资料','after':'无'} in revoked['changes']
        membership=ok(client.get(ROOT+'/audit',params={'keyword':'调整角色成员'}))['items'][0]
        assert {'field':'enabled','label':'成员授权','before':'未授权','after':'已授权'} in membership['changes']
        reader.close()


def test_task_and_delivery_search_filter_in_database_and_keep_operator_school_boundaries(school):
    from app.models import ExcelImportJob
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','任务查看合成教师','C01','合成院系','','','','','']])
        reset_for_test(owner,school,'T001',credentials['T001']).close()
        failed=preview(owner,'students',[['S001','缺少专业的合成学生','C01','合成院系','','','','','2024']])
        assert failed['status']=='FAILED'
        assert failed['newOrganizations']==[] and failed['newOrganizationCount']==0
        assert owner.get(ROOT+'/exports/departments').status_code==200
        filtered=ok(owner.get(ROOT+'/tasks',params={'keyword':'人员.xlsx','status':'FAILED'}))
        assert filtered['total']==1 and filtered['items'][0]['id']==failed['id']
        assert filtered['items'][0]['typeName']=='学生开户' and filtered['items'][0]['canOpen']
        assert failed['id'] in filtered['items'][0]['path']
        assert ok(owner.get(ROOT+'/tasks',params={'keyword':failed['id'],'status':'SUCCESS'}))['total']==0
        assert ok(owner.get(ROOT+'/tasks',params={'keyword':failed['id']}))['items'][0]['id']==failed['id']
        assert ok(owner.get(ROOT+'/tasks',params={'keyword':'学校初始化负责人'}))['total']==3
        for page in [1,2]:
            result=ok(owner.get(ROOT+'/tasks',params={'keyword':'人员.xlsx','page':page,'pageSize':1}))
            assert result['total']==2 and len(result['items'])==1
        assert ok(owner.get(ROOT+'/tasks',params={'keyword':'%'}))['total']==0
        assert owner.get(ROOT+'/tasks',params={'status':'not-a-status'}).status_code==422
        exported=ok(owner.get(ROOT+'/tasks',params={'status':'SUCCESS'}))['items'][0]
        assert exported['typeName']=='基础资料导出' and exported['canDownload']
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        role=ok(owner.post(ROOT+'/roles',json={'name':'任务记录只读岗','baseRole':'INTERN_MENTOR',
            'permissions':['system.task.view'],'reason':'合成任务范围查询'}))
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'合成任务范围查询'}))
        current=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
        with get_sessionmaker()() as db:
            own_job=ExcelImportJob(tenant_id=school['id'],module_key='internship',biz_type='enterprise',
                file_name='本人合成上传.xlsx',status='UPLOADED',operator_id=int(account['id']),operator_name='任务查看合成教师')
            other_school=Tenant(tenant_code='SYS-'+uuid4().hex[:12],school_name='其他合成学校',status='ACTIVE')
            db.add_all([own_job,other_school]);db.flush()
            db.add(ExcelImportJob(tenant_id=other_school.id,module_key='internship',biz_type='enterprise',
                file_name='另一学校独有文件.xlsx',status='UPLOADED',operator_name='其他学校合成人员'))
            own_id=str(own_job.id);db.commit()
        with login(school,'T001')[0] as reader:
            switched=ok(reader.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+current['code'].lower(),'clientType':'PC'}))
            reader.headers['Authorization']='Bearer '+switched['accessToken']
            own=ok(reader.get(ROOT+'/tasks',params={'keyword':'合成上传','status':'UPLOADED'}))
            assert own['total']==1 and own['items'][0]['id']==own_id
            assert own['items'][0]['typeName']=='企业资料导入'
            assert own['items'][0]['statusLabel']=='已上传'
            assert ok(reader.get(ROOT+'/tasks',params={'keyword':failed['id']}))['total']==0
            assert ok(reader.get(ROOT+'/tasks',params={'keyword':'学校初始化负责人'}))['total']==0
        assert ok(owner.get(ROOT+'/tasks',params={'keyword':'另一学校独有文件'}))['total']==0
        for kind in ['MIGRATION','DOCUMENT']:
            endpoint=ROOT+'/records/'+kind
            created=ok(owner.post(endpoint,json={'title':'合成资料100%核对','source':'来源__A','reason':'合成资料查询'}))
            ok(owner.post(endpoint,json={'title':'其他合成资料','reason':'合成资料查询'}))
            for keyword in ['100%','来源__A','%']:
                result=ok(owner.get(endpoint,params={'keyword':keyword,'pageSize':1}))
                assert result['total']==1 and result['items'][0]['id']==created['id']
            assert ok(owner.get(endpoint,params={'keyword':'不存在的来源'}))['total']==0


def test_migration_registration_requires_mapping_reconciliation_and_allows_read_only_review(school):
    with login(school)[0] as owner:
        endpoint=ROOT+'/records/MIGRATION'
        draft={'title':'合成迁移登记校验','source':'虚构资料，仅验证登记约束','notes':'不代表真实迁移验收',
            'expectedPeople':1,'actualPeople':1,'expectedRecords':2,'actualRecords':2,
            'expectedFiles':3,'actualFiles':3,'continuationChecked':True,'status':'VERIFIED','reason':'合成登记校验'}
        assert owner.post(endpoint,json=draft).status_code==400
        assert not ok(owner.get(endpoint))['items']
        mapping={'mappingRules':'按旧组织代码和工学号对应，业务状态沿原迁移规则',
            'identityMapping':'旧院系C01对应本校C01，旧学生S001对应本校S001',
            'exceptionResolution':'无未处理异常；本记录仅为虚构登记校验'}
        created=ok(owner.post(endpoint,json={**draft,**mapping,'status':'REVIEW'}))
        target=endpoint+'/'+created['id']
        mismatch={**draft,**mapping,'expectedVersion':created['version'],'actualFiles':2}
        assert owner.put(target,json=mismatch).status_code==400
        assert owner.put(target,json={**draft,**mapping,'expectedVersion':created['version'],'continuationChecked':False}).status_code==400
        for field in ['source',*mapping]:
            assert owner.put(target,json={**draft,**mapping,'expectedVersion':created['version'],field:'  '}).status_code==400
        unchanged=ok(owner.get(endpoint))['items'][0]
        assert unchanged['status']=='REVIEW' and unchanged['version']==created['version'] and unchanged['checkedAt'] is None
        assert all(unchanged['details'][key]==value for key,value in mapping.items())
        verified=ok(owner.put(target,json={**draft,**mapping,'expectedVersion':created['version']}))
        stale=owner.put(target,json={**draft,**mapping,'expectedVersion':created['version'],'mappingRules':'不得覆盖的新输入'})
        assert stale.status_code==409
        assert stale.json()['details']['current']['mappingRules']==mapping['mappingRules']
        _,credentials=import_rows(owner,'staff',[['T001','迁移核对合成教师','C01','合成院系','','','','','']])
        reset_for_test(owner,school,'T001',credentials['T001']).close()
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        role=ok(owner.post(ROOT+'/roles',json={'name':'迁移登记只读岗','baseRole':'INTERN_MENTOR',
            'permissions':['system.migration.view'],'reason':'合成只读核对'}))
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'合成只读核对'}))
        current=next(row for row in ok(owner.get(ROOT+'/roles'))['items'] if row['id']==role['id'])
        with login(school,'T001')[0] as reader:
            switched=ok(reader.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+current['code'].lower(),'clientType':'PC'}))
            reader.headers['Authorization']='Bearer '+switched['accessToken']
            reviewed=ok(reader.get(endpoint))['items'][0]
            assert reviewed['version']==verified['version'] and reviewed['checkedAt']
            assert reviewed['evidenceLabel']=='学校登记资料，须另行核验'
            assert reviewed['statusLabel']=='学校登记已核验'
            assert all(reviewed['details'][key]==value for key,value in mapping.items())
            assert reader.put(target,json={**draft,**mapping,'expectedVersion':verified['version']}).status_code==403
        assert not ok(owner.get(endpoint))['receipts'],'人工登记不得生成迁移脚本通过回执'
        # The same shared document schema stays usable without migration-only fields.
        ok(owner.post(ROOT+'/records/DOCUMENT',json={'title':'合成交付资料','reason':'通用资料登记回归'}))


def test_notification_config_delivers_frozen_messages_to_only_the_recipient(school):
    from app.core.context import get_tenant,set_tenant
    from app.services.message_event_outbox_service import emit_message_event,process_pending_outbox
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'students',[
            ['NOTICE1','通知接收学生','C01','工程学院','M01','专业','B01','班级','2026'],
            ['NOTICE2','其他学生','C01','工程学院','M01','专业','B01','班级','2026']])
        student_id=next(row['id'] for row in ok(owner.get(ROOT+'/people/student'))['items'] if row['number']=='NOTICE1')
        config={'reminderText':'学校补充办理说明','expectedVersion':0,'reason':'验证站内提醒实际投递'}
        ok(owner.put(ROOT+'/config/notification',json=config))
        with get_sessionmaker()() as db:
            events=[emit_message_event(db,event_code=code,source_module='internship',source_biz_type='SYSTEM_NOTIFICATION_TEST',
                source_biz_id=123,recipient_refs=[{'studentId':int(student_id)}],content=content,tenant_id=school['id'])
                for code,content in [('INTERNSHIP.WEEKLY_REMIND','原提醒事实'),('INTERNSHIP.RISK_CREATED','原风险事实')]]
            event_ids=[row.id for row in events]
            db.commit()
        ok(owner.put(ROOT+'/config/notification',json={**config,'reminderText':'之后才修改的说明','expectedVersion':1}))
        previous=get_tenant()
        try:
            set_tenant(school['id'])
            assert process_pending_outbox(outbox_ids=event_ids)==2
            assert process_pending_outbox(outbox_ids=event_ids)==0
        finally: set_tenant(previous)
        with reset_for_test(owner,school,'NOTICE1',credentials['NOTICE1'],True) as student:
            messages=ok(student.get('/api/v1/mobile/me/messages'))
            assert messages['total']==2
            contents={row['title']:ok(student.get('/api/v1/mobile/me/messages/'+row['messageId']))['content'] for row in messages['items']}
            assert contents['实习周报提醒']=='原提醒事实\n\n学校补充办理说明'
            assert contents['实习风险提醒']=='原风险事实'
            message_id=messages['items'][0]['messageId']
        with reset_for_test(owner,school,'NOTICE2',credentials['NOTICE2'],True) as other:
            assert ok(other.get('/api/v1/mobile/me/messages'))['total']==0
            assert other.get('/api/v1/mobile/me/messages/'+message_id).status_code==404


def test_concurrent_admin_disable_preserves_one_and_preset_auditor_is_audit_only(school):
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['ADMIN2','第二管理员','C01','工程学院','','','','','']])
        second=reset_for_test(owner,school,'ADMIN2',credentials['ADMIN2'])
        accounts=ok(owner.get(ROOT+'/accounts'))['items']
        target=next(a for a in accounts if a['loginName']=='ADMIN2')
        roles=ok(owner.get(ROOT+'/roles'))['items']
        auditor=next(r for r in roles if r['code']=='SECURITY_AUDITOR')
        assert set(auditor['permissions'])=={'system.audit.view','system.audit.export'}
        admin=next(r for r in roles if r['code']=='SYSTEM_ADMIN')
        ok(owner.post(ROOT+f'/roles/{admin["id"]}/members',json={'userId':target['id'],
            'enabled':True,'expectedVersion':admin['version'],'reason':'并发保护验收'}))
        switch=ok(second.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-system_admin','clientType':'PC'}))
        second.headers['Authorization']='Bearer '+switch['accessToken']
        accounts=ok(owner.get(ROOT+'/accounts'))['items']
        primary=next(a for a in accounts if a['loginName']=='owner');target=next(a for a in accounts if a['loginName']=='ADMIN2')
        def disable(client,account):
            return client.post(ROOT+f'/accounts/{account["id"]}/actions',json={'action':'DISABLE',
                'expectedVersion':account['version'],'reason':'并发停用对方验收'}).status_code
        with ThreadPoolExecutor(max_workers=2) as pool:
            a=pool.submit(disable,owner,target);b=pool.submit(disable,second,primary)
            statuses=[a.result(),b.result()]
        assert statuses.count(200)==1 and all(s in {200,401,403,409} for s in statuses)
        with get_sessionmaker()() as db:
            assert db.scalar(select(func.count()).select_from(User).where(User.tenant_id==school['id'],
                User.id.in_([int(primary['id']),int(target['id'])]),User.status=='ACTIVE'))==1
        second.close()


def test_preset_comparison_preserves_school_changes_and_applies_only_reviewed_permissions(school):
    with login(school)[0] as owner:
        data=ok(owner.get(ROOT+'/roles'))
        role=next(r for r in data['items'] if r['code']=='INTERN_MENTOR')
        template=next(r for r in data['templates'] if r['code']=='INTERN_MENTOR')
        assert sorted(template['permissions'])==sorted(role['permissions'])
        assert all('*' not in p for p in template['permissions'])
        chosen=['internship.leave.view','internship.leave.review']
        changed=ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':'本校请假指导岗','baseRole':role['baseRole'],
            'permissions':chosen,'expectedVersion':role['version'],'reason':'学校裁剪指导岗位'}))
        with get_sessionmaker()() as db:
            bootstrap_school(db,school['id']);db.commit()
        after=ok(owner.get(ROOT+'/roles'))
        retained=next(r for r in after['items'] if r['id']==role['id'])
        current_template=next(r for r in after['templates'] if r['code']=='INTERN_MENTOR')
        assert retained['name']=='本校请假指导岗' and set(retained['permissions'])==set(chosen)
        assert retained['version']==changed['version']
        assert current_template==template
        assert 'internship.guidance.view' in current_template['permissions']
        applied=ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':retained['name'],'baseRole':role['baseRole'],
            'permissions':chosen+['internship.guidance.view'],'expectedVersion':changed['version'],
            'reason':'核对预配对照后仅恢复指导查看'}))
        reread=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
        assert set(reread['permissions'])==set(chosen+['internship.guidance.view'])
        assert 'internship.student.password.reset' not in reread['permissions']
        assert owner.put(ROOT+'/roles/'+role['id'],json={'name':retained['name'],'baseRole':role['baseRole'],
            'permissions':template['permissions'],'sensitiveConfirmed':True,'expectedVersion':changed['version'],
            'reason':'旧页面不能用模板覆盖已保存权限'}).status_code==409
        assert next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])['version']==applied['version']
        with get_sessionmaker()() as db:
            event=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],
                SecurityAuditLog.action=='ROLE_SAVE').order_by(SecurityAuditLog.id.desc()))
            assert event.detail_json['added']==['internship.guidance.view'] and event.detail_json['removed']==[]


def test_browser_refresh_keeps_token_compact_and_role_switch_revokes_old_session(school):
    from app.core.security import decode_token
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','刷新后切换身份教师','C01','工程学院','','','','','']])
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        role=ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':account['id'],'enabled':True,'expectedVersion':role['version'],'reason':'合成已有院系管理员身份'}))
        with reset_for_test(owner,school,'T001',credentials['T001']) as teacher:
            for current,target in [('INTERN_MENTOR','COLLEGE_ADMIN'),('COLLEGE_ADMIN','INTERN_MENTOR'),('INTERN_MENTOR','COLLEGE_ADMIN')]:
                renewed=ok(teacher.post('/api/v1/auth/browser-refresh',headers={'X-Browser-Session':'staff'}))
                assert len(renewed['accessToken'])<2048
                claims=decode_token(renewed['accessToken'])
                assert claims['currentRoleCode']==current and 'permissions' not in claims
                old='Bearer '+renewed['accessToken'];teacher.headers['Authorization']=old
                assert ok(teacher.get('/api/v1/rbac/current-context'))['permissionPatterns']
                activate(teacher,target)
                assert teacher.get('/api/v1/auth/me',headers={'Authorization':old}).status_code==401
            ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':'院系实习管理员','baseRole':'COLLEGE_ADMIN','permissions':[],
                'expectedVersion':role['version'],'reason':'刷新后旧会话仍按当前数据库撤权'}))
            assert ok(teacher.get('/api/v1/rbac/current-context'))['permissionPatterns']==[]
            assert teacher.get(ROOT+'/people/student').status_code==403


def test_custom_active_role_empty_permission_revocation_and_last_admin(school):
    with login(school)[0] as client:
        catalog=ok(client.get(ROOT+'/roles'))['catalog']
        assert catalog and all('*' not in r['code'] and not r['code'].endswith('.') for r in catalog)
        assert all('实习业务' not in r['name'] and '未登记' not in r['name'] for r in catalog if r['code'].startswith('internship.'))
        _,credentials=import_rows(client,'staff',[['T001','教师甲','C01','工程学院','','','','','']])
        teacher=reset_for_test(client,school,'T001',credentials['T001'])
        account=ok(client.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        empty=ok(client.post(ROOT+'/roles',json={'name':'空权限教师','baseRole':'INTERN_MENTOR','permissions':[],'reason':'验证空权限不回退'}))
        empty=ok(client.post(ROOT+f'/roles/{empty["id"]}/members',json={'userId':account['id'],'enabled':True,'expectedVersion':empty['version'],'reason':'合成身份分配'}))
        roles=ok(client.get(ROOT+'/roles'))['items'];target=next(r for r in roles if r['id']==empty['id'])
        old=teacher.headers['Authorization']
        switched=ok(teacher.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+target['code'].lower(),'clientType':'PC'}))
        teacher.headers['Authorization']='Bearer '+switched['accessToken']
        assert ok(teacher.get('/api/v1/rbac/current-context'))['permissionPatterns']==[]
        assert teacher.get(ROOT+'/overview').status_code==403
        assert teacher.get('/api/v1/auth/me',headers={'Authorization':old}).status_code==401
        ok(client.post(ROOT+f'/roles/{target["id"]}/members',json={'userId':account['id'],'enabled':False,'expectedVersion':target['version'],'reason':'撤销该身份'}))
        assert teacher.get('/api/v1/auth/me').status_code in (401,403)
        owner=next(r for r in ok(client.get(ROOT+'/accounts'))['items'] if r['id']==str(school['owner']))
        assert client.post(ROOT+f'/accounts/{owner["id"]}/actions',json={'action':'DISABLE','expectedVersion':owner['version'],'reason':'测试最后管理员保护'}).status_code==409
        teacher.close()


def test_custom_mentor_identity_consumes_leave_policy_and_live_permission_revocation(school):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord,InternshipLeave,InternshipCheckin,UnifiedTodo
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[[n,n,'C01','工程学院','','','','',''] for n in ['T001','T002']])
        _,students=import_rows(owner,'students',[[n,n,'C01','工程学院','M01','软件技术','B01','软件一班','2024'] for n in ['S001','S002']])
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        permissions=['internship.batch.view','internship.student.view','internship.leave.view','internship.leave.review']
        custom=ok(owner.post(ROOT+'/roles',json={'name':'本校请假指导岗','baseRole':'INTERN_MENTOR','permissions':permissions,'reason':'配置本校裁剪指导身份'}))
        for name in ['T001','T002']:
            role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==custom['id'])
            ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':accounts[name]['id'],'enabled':True,'expectedVersion':role['version'],'reason':'配置合成指导人员'}))
        preset=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='INTERN_MENTOR')
        ok(owner.post(ROOT+'/roles/'+preset['id']+'/members',json={'userId':accounts['T001']['id'],'enabled':False,'expectedVersion':preset['version'],'reason':'仅使用自定义身份验收'}))
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==custom['id'])
        with get_sessionmaker()() as db:
            profiles={s.student_no:s for s in db.scalars(select(StudentProfile).where(StudentProfile.tenant_id==school['id']))}
            batch=InternshipBatch(tenant_id=school['id'],batch_name='自定义指导身份',batch_no='CUSTOM-MENTOR',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.flush();records={}
            for student_name,teacher_name in [('S001','T001'),('S002','T002')]:
                record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=profiles[student_name].id,status='ONBOARD',eligibility_status='QUALIFIED',advisor_user_id=int(accounts[teacher_name]['id']),advisor_name=teacher_name,intern_start_date=datetime(2026,9,1),intern_end_date=datetime(2026,12,31))
                db.add(record);db.flush();records[student_name]=str(record.id)
            db.commit();bid=str(batch.id)
        policy=ok(owner.get(ROOT+'/leave-policies/'+bid))
        nodes=[{'type':'GUIDE','roleCode':role['code'],'userId':''}]
        assert ok(owner.post(ROOT+'/leave-policies/'+bid+'/preview',json={'nodes':nodes}))['canPublish']
        ok(owner.post(ROOT+'/leave-policies/'+bid+'/publish',json={'nodes':nodes,'expectedVersion':policy['expectedVersion'],'reason':'使用学校自定义请假指导岗'}))
        with reset_for_test(owner,school,'T001',creds['T001']) as teacher, reset_for_test(owner,school,'S001',students['S001'],True) as student:
            context=ok(teacher.get('/api/v1/rbac/current-context'))
            assert context['currentRole']['roleCode']==role['code'] and set(context['permissionPatterns'])==set(permissions)
            roster=ok(teacher.get('/api/v1/internship/intern-students',params={'batchId':bid}))
            assert [r['id'] for r in roster['items']]==[records['S001']]
            assert teacher.get('/api/v1/internship/intern-students/'+records['S002']).status_code==403
            route='/api/v1/portal/internship/context/leaves'
            body={'batchId':bid,'internshipId':records['S001'],'leaveType':'PERSONAL','startDate':'2026-10-11','endDate':'2026-10-11','reason':'自定义指导身份办理验收'}
            leave=ok(student.post(route,json=body));lid=leave['id']
            current=ok(teacher.get('/api/v1/internship/leaves/'+lid,params={'batchId':bid}))
            assert current['approvalProgress']['canReview']
            result=ok(teacher.post('/api/v1/internship/leaves/'+lid+'/review',params={'batchId':bid},json={'action':'APPROVE','comment':'本人指导范围审核通过','batchId':bid,'expectedVersion':current['version']}))
            assert result['status']=='APPROVED'
            body.update(startDate='2026-10-12',endDate='2026-10-12')
            waiting=ok(student.post(route,json=body));wid=waiting['id']
            current=ok(teacher.get('/api/v1/internship/leaves/'+wid,params={'batchId':bid}))
            ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':role['name'],'baseRole':'INTERN_MENTOR','permissions':[p for p in permissions if p!='internship.leave.review'],'expectedVersion':role['version'],'reason':'收回自定义身份审批权限'}))
            denied=teacher.post('/api/v1/internship/leaves/'+wid+'/review',params={'batchId':bid},json={'action':'APPROVE','comment':'旧会话不应继续审核','batchId':bid,'expectedVersion':current['version']})
            assert denied.status_code==403
            final=ok(student.get(route+'/'+wid,params={'batchId':bid,'internshipId':records['S001']}))
            assert final['approvalProgress']['steps'][0]['statusLabel']=='待重新指定责任人'
            with get_sessionmaker()() as db:
                assert db.get(InternshipLeave,int(wid)).status=='PENDING'
                assert db.scalar(select(func.count()).select_from(InternshipCheckin).where(InternshipCheckin.internship_id==int(records['S001']),InternshipCheckin.result=='LEAVE'))==1
                assert db.scalar(select(func.count()).select_from(UnifiedTodo).where(UnifiedTodo.tenant_id==school['id'],UnifiedTodo.source_biz_id==int(wid),UnifiedTodo.status=='PENDING'))==1


@pytest.mark.parametrize('base,scope_kind',[('PROFESSIONAL_LEADER','MAJOR'),('COUNSELOR','CLASS')])
def test_narrow_scope_can_maintain_own_student_without_gaining_parent_or_sibling_authority(school,base,scope_kind):
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[['T001','专业班级资料负责人','C01','工程学院','','','','','']])
        import_rows(owner,'students',[
            ['S001','负责学生','C01','工程学院','M01','负责专业','B01','负责班级','2024'],
            ['S002','同院其他专业学生','C01','工程学院','M02','其他专业','B02','其他班级','2024'],
            ['S003','其他学院学生','C02','其他学院','M03','外院专业','B03','外院班级','2024']])
        teacher=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        role=ok(owner.post(ROOT+'/roles',json={'name':'负责范围资料维护','baseRole':base,
            'permissions':['system.org.view','system.org.manage','system.student.view','system.student.manage'],
            'reason':'专业和班级稳定范围验收'}))
        assigned=ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':teacher['id'],
            'enabled':True,'expectedVersion':role['version'],'reason':'分配合成资料职责'}))
        definition=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==assigned['id'])
        students={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        own,other,foreign=(students[n] for n in ['S001','S002','S003'])
        ok(owner.post(ROOT+'/scopes',json={'userId':teacher['id'],'roleCode':definition['code'],
            'scopeType':scope_kind,'refId':own['majorId'] if scope_kind=='MAJOR' else own['classId'],'reason':'仅负责一个专业或班级'}))
        with reset_for_test(owner,school,'T001',creds['T001']) as actor:
            activate(actor,definition['code'])
            assert [r['number'] for r in ok(actor.get(ROOT+'/people/student'))['items']]==['S001']
            body={'name':'负责学生（资料核对）','collegeId':own['collegeId'],'classId':own['classId'],
                'grade':'2024','expectedVersion':own['version'],'confirmImpact':True,'reason':'仅维护负责学生资料'}
            saved=ok(actor.put(ROOT+'/people/student/'+own['id'],json=body))
            assert ok(actor.get(ROOT+'/selection',params={'kind':'colleges','identifier':own['collegeId']}))['name']=='工程学院'
            assert ok(actor.get(ROOT+'/selection',params={'kind':'majors','identifier':own['majorId']}))['name']=='负责专业'
            assert actor.get(ROOT+'/selection',params={'kind':'colleges','identifier':foreign['collegeId']}).status_code==404
            for target in (other,foreign):
                assert actor.put(ROOT+'/people/student/'+own['id'],json={**body,'expectedVersion':saved['version'],
                    'collegeId':target['collegeId'],'classId':target['classId']}).status_code in (403,404)
                assert actor.get(ROOT+'/people/student/'+target['id']+'/impact').status_code==404
            assert actor.get(ROOT+'/organizations/colleges/'+own['collegeId']+'/impact').status_code==404
            kind='majors' if scope_kind=='MAJOR' else 'classes'
            organization=ok(actor.get(ROOT+'/organizations/'+kind))['items'][0]
            parent_kind='colleges' if scope_kind=='MAJOR' else 'majors'
            parent=next(r for r in ok(owner.get(ROOT+'/organizations/'+parent_kind))['items'] if r['id']==organization['parentId'])
            assert actor.put(ROOT+'/organizations/'+parent_kind+'/'+parent['id'],json={'name':'禁止越级修改',
                'code':parent['code'],'parentId':parent['parentId'],'expectedVersion':parent['version'],'reason':'只读祖先不能获得维护权'}).status_code==404
            assert actor.post(ROOT+'/organizations/'+kind,json={'name':'禁止新增同级组织','code':'OUTSIDE',
                'parentId':parent['id'],'reason':'禁止从祖先展示扩大范围'}).status_code==403
            ok(actor.put(ROOT+'/organizations/'+kind+'/'+organization['id'],json={'code':organization['code'],'parentId':organization['parentId'],
                'expectedVersion':organization['version'],'reason':'维护本人负责组织的名称','name':organization['name']+'（核对）'}))
            with get_sessionmaker()() as db:
                current=db.get(StudentProfile,int(own['id']))
                assert str(current.class_id)==own['classId'] and str(current.college_id)==own['collegeId']


@pytest.mark.parametrize('scope_kind',['COLLEGE','MAJOR','CLASS'])
def test_batch_participant_preview_consumes_stable_scope_after_rename_and_revocation(school,scope_kind):
    from datetime import datetime
    from app.models import InternshipBatch, InternshipBatchParticipant, InternshipRecord
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','批次负责教师','C01','同名学院','','','','','']])
        import_rows(owner,'students',[
            ['S001','本范围学生','C01','同名学院','M01','同名专业','B01','同名班级','2024'],
            ['S002','范围外学生','C02','同名学院','M02','同名专业','B02','同名班级','2024']])
        teacher=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':teacher['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'合成批次职责验收'}))
        people={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        own=people['S001']
        reference=own[{'COLLEGE':'collegeId','MAJOR':'majorId','CLASS':'classId'}[scope_kind]]
        grant={'userId':teacher['id'],'roleCode':'COLLEGE_ADMIN','scopeType':scope_kind,'refId':reference,'reason':'只负责稳定编号指定对象'}
        saved=ok(owner.post(ROOT+'/scopes',json=grant))
        with get_sessionmaker()() as db:
            batch=InternshipBatch(tenant_id=school['id'],batch_name='稳定范围选人验收',batch_no='STABLE-SCOPE',status='DRAFT',
                start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.commit();batch_id=str(batch.id)
        path='/api/v1/internship/batches/'+batch_id+'/participants/preview'
        with reset_for_test(owner,school,'T001',credentials['T001']) as actor:
            activate(actor,'COLLEGE_ADMIN')
            result=ok(actor.post(path,json={'grades':['2024']}))
            assert [r['studentNo'] for r in result['rows']]==['S001']
            assert result['matchedCount']==result['outOfScopeCount']==1
            frozen=ok(actor.post(path.removesuffix('/preview')+'/freeze',json={'rule':result['rule']}))
            assert frozen['total']==frozen['createdRecords']==1
            listed=ok(actor.get(path.removesuffix('/preview')))
            assert [r['studentNo'] for r in listed['items']]==['S001']
            rejected=ok(actor.post(path.removesuffix('/preview')+'/add',json={
                'studentIds':[people['S002']['id']],'reason':'范围外学生不得补录'}))
            assert rejected['added']==0 and rejected['rejectedOutOfScope']==[int(people['S002']['id'])]
            for kind,key,name in [('colleges','collegeId','更名后的学院'),('majors','majorId','更名后的专业'),('classes','classId','更名后的班级')]:
                item=next(r for r in ok(owner.get(ROOT+'/organizations/'+kind))['items'] if r['id']==own[key])
                ok(owner.put(ROOT+'/organizations/'+kind+'/'+item['id'],json={'name':name,
                    'code':item['code'],'parentId':item['parentId'],'grade':item['grade'] or '',
                    'expectedVersion':item['version'],'reason':'组织改名保持批次授权'}))
            renamed=ok(actor.post(path,json={'grades':['2024']}))
            assert [r['studentId'] for r in renamed['rows']]==[own['id']]
            assert renamed['rows'][0]['collegeName']=='更名后的学院'
            with get_sessionmaker()() as db:
                participant=db.scalar(select(InternshipBatchParticipant).where(InternshipBatchParticipant.batch_id==int(batch_id)))
                assert participant.snapshot_college_name=='同名学院' and participant.snapshot_class_name=='同名班级'
                assert db.get(InternshipRecord,participant.internship_id).student_id==int(own['id'])
            ok(owner.put(ROOT+'/scopes/'+saved['id'],json={**grant,'status':'DISABLED',
                'expectedVersion':saved['version'],'reason':'核对旧会话选人范围立即收回'}))
            revoked=ok(actor.post(path,json={'grades':['2024']}))
            assert not revoked['rows'] and revoked['outOfScopeCount']==2
            assert not ok(actor.get(path.removesuffix('/preview')))['items']


def test_disabled_student_cannot_enter_new_rosters_but_retains_historical_records(school):
    from datetime import datetime
    from app.models import InternshipBatch, InternshipRecord
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','校级实习负责人','C01','工程学院','','','','','']])
        import_rows(owner,'students',[
            ['S001','已有实习学生','C01','工程学院','M01','软件专业','B01','软件班','2024'],
            ['S002','未建档学生','C01','工程学院','M01','软件专业','B01','软件班','2024']])
        teacher=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SCHOOL_ADMIN')
        ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':teacher['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'合成校级实习职责验收'}))
        people=ok(owner.get(ROOT+'/people/student'))['items']
        with get_sessionmaker()() as db:
            batches=[InternshipBatch(tenant_id=school['id'],batch_name=name,batch_no=code,status='DRAFT',
                start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
                for name,code in [('保留历史名单','HISTORY'),('停用学生新建档核对','DISABLED')]]
            db.add_all(batches);db.commit();old_batch,new_batch=[str(b.id) for b in batches]
        with reset_for_test(owner,school,'T001',credentials['T001']) as actor:
            activate(actor,'SCHOOL_ADMIN')
            old_path='/api/v1/internship/batches/'+old_batch+'/participants'
            new_path='/api/v1/internship/batches/'+new_batch+'/participants'
            ok(actor.post(old_path+'/freeze',json={'rule':{'studentIds':[people[0]['id']]}}))
            for person in people:
                ok(owner.put(ROOT+'/people/student/'+person['id'],json={'name':person['name'],
                    'collegeId':person['collegeId'],'classId':person['classId'],'grade':person['grade'],
                    'expectedVersion':person['version'],'status':'DISABLED','reason':'资料停用不删除历史'}))
            assert not ok(actor.get('/api/v1/directory/student-options'))['items']
            assert [r['studentId'] for r in ok(actor.get(old_path))['items']]==[people[0]['id']]
            pre=ok(actor.post(new_path+'/preview',json={'grades':['2024']}))
            assert pre['matchedCount']==0 and not pre['rows']
            assert actor.post(new_path+'/freeze',json={'rule':{'grades':['2024']}}).status_code==400
            added=ok(actor.post(old_path+'/add',json={'studentIds':[people[1]['id']],'reason':'停用资料不可新补录'}))
            assert added['added']==0 and added['rejectedOutOfScope']==[int(people[1]['id'])]
            assert actor.post('/api/v1/internship/intern-students',json={'batchId':new_batch,'studentId':people[0]['id']}).status_code==400
            body={'batchId':new_batch,'rows':[{'studentNo':people[0]['number']}]}
            checked=ok(actor.post('/api/v1/internship/intern-students/import/dry-run',json=body))
            assert checked['invalidRows']==1 and '停用' in checked['errors'][0]['message']
            assert actor.post('/api/v1/internship/intern-students/import/confirm',json=body).status_code==400
            with get_sessionmaker()() as db:
                records=db.scalars(select(InternshipRecord).where(InternshipRecord.tenant_id==school['id'])).all()
                assert len(records)==1 and not records[0].is_deleted and records[0].status=='PREPARING'
                assert db.get(InternshipBatch,int(new_batch)).status=='DRAFT'
                assert all(db.get(User,int(person['userId'])).status=='ACTIVE' for person in people)


def test_roster_import_preflight_uses_same_school_identity_guard_as_confirm(school):
    from datetime import datetime
    from app.models import InternshipBatch, InternshipRecord
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','兼任实习负责人','C01','一院','','','','','']])
        import_rows(owner,'students',[['S002','二院学生','C02','二院','M02','二院专业','B02','二院班级','2024']])
        account=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        for code in ['COLLEGE_ADMIN','SCHOOL_ADMIN']:
            role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']==code)
            ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':account['id'],'enabled':True,
                'expectedVersion':role['version'],'reason':'合成多身份导入边界验收'}))
        college=next(r for r in ok(owner.get(ROOT+'/organizations/colleges'))['items'] if r['code']=='C01')
        ok(owner.post(ROOT+'/scopes',json={'userId':account['id'],'roleCode':'COLLEGE_ADMIN',
            'scopeType':'COLLEGE','refId':college['id'],'reason':'院系身份只负责一院'}))
        with get_sessionmaker()() as db:
            batch=InternshipBatch(tenant_id=school['id'],batch_name='业务导入身份边界',batch_no='IMPORT-ROLE',status='DRAFT',
                start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.commit();batch_id=str(batch.id)
        path='/api/v1/internship/intern-students/import'
        content=build_template_xlsx(['学号','指导教师','实习开始日期','实习结束日期','备注'],samples=[['S002','','','','合成业务导入']])
        with reset_for_test(owner,school,'T001',credentials['T001']) as actor:
            activate(actor,'COLLEGE_ADMIN')
            for number in ['S002','UNKNOWN']:
                response=actor.post(path+'/dry-run',json={'batchId':batch_id,'rows':[{'studentNo':number}]})
                assert response.status_code==403 and '仅校级管理员' in response.json()['message']
            denied=actor.post(path+'/xlsx',params={'batchId':batch_id},files={'file':('合成名单.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
            assert denied.status_code==403
            body={'batchId':batch_id,'rows':[{'studentNo':'S002'}]}
            assert actor.post(path+'/confirm',json=body).status_code==403
            activate(actor,'SCHOOL_ADMIN')
            assert ok(actor.post(path+'/dry-run',json=body))['validRows']==1
            uploaded=ok(actor.post(path+'/xlsx',params={'batchId':batch_id},files={'file':('合成名单.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')}))
            assert uploaded['validRows']==1 and uploaded['invalidRows']==0
            assert ok(actor.post(path+'/confirm',json=body))['created']==1
            with get_sessionmaker()() as db:
                assert db.scalar(select(func.count()).select_from(InternshipRecord).where(InternshipRecord.tenant_id==school['id']))==1
                assert db.scalar(select(func.count()).select_from(User).where(User.tenant_id==school['id']))==3


def test_responsibility_issues_page_to_stable_object_and_resolve_duplicate(school):
    from app.models import College
    with login(school)[0] as owner:
        for index in range(31):
            org(owner,'colleges',f'待分配学院{index+1:02d}',f'C{index+1:02d}')
        _,credentials=import_rows(owner,'staff',[
            ['T001','合成第一负责人','C99','重复责任学院','','','','',''],
            ['T002','合成第二负责人','C99','重复责任学院','','','','','']])
        colleges=ok(owner.get(ROOT+'/organizations/colleges',params={'pageSize':100}))['items']
        target=next(r for r in colleges if r['code']=='C99')
        second=org(owner,'colleges','待页面核对学院','C98')
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        for name in ['T001','T002']:
            role=ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':accounts[name]['id'],
                'enabled':True,'expectedVersion':role['version'],'reason':'合成责任重复核对'}))
            for college in [target,second]:
                ok(owner.post(ROOT+'/scopes',json={'userId':accounts[name]['id'],'roleCode':'COLLEGE_ADMIN',
                    'scopeType':'COLLEGE','refId':college['id'],'reason':'准备同对象两名负责人'}))
        first=ok(owner.get(ROOT+'/scopes'))['responsibilityCheck']
        later=ok(owner.get(ROOT+'/scopes',params={'issuePage':2}))['responsibilityCheck']
        assert first['missing']==31 and first['duplicates']==2 and first['total']==33
        assert len(first['items'])==30 and len(later['items'])==3
        assert len({(r['scopeType'],r['id']) for r in first['items']+later['items']})==33
        assert target['id'] in {r['id'] for r in later['items']}
        duplicates=ok(owner.get(ROOT+'/scopes',params={'issueType':'DUPLICATE'}))['responsibilityCheck']
        assert duplicates['total']==2 and all(r['count']==2 for r in duplicates['items'])
        assert ok(owner.get(ROOT+'/scopes',params={'issueType':'MISSING','issuePage':2}))['responsibilityCheck']['total']==31
        selected=ok(owner.get(ROOT+'/scopes',params={'scopeType':'COLLEGE','refId':target['id']}))
        assert selected['selectedObject']=={'scopeType':'COLLEGE','refId':target['id'],'name':target['name']}
        assert selected['total']==2 and all(r['refId']==target['id'] for r in selected['items'])
        retiring=next(r for r in selected['items'] if r['userId']==accounts['T002']['id'])
        impact=ok(owner.get(ROOT+'/scopes/'+retiring['id']+'/impact'))
        assert impact['original']['refId']==target['id'] and impact['roleImpact']['pendingApprovalCount']==0
        body={k:retiring[k] for k in ['userId','roleCode','scopeType','refId']}
        ok(owner.put(ROOT+'/scopes/'+retiring['id'],json={**body,'status':'DISABLED','expectedVersion':retiring['version'],
            'reason':'核对没有在办事项后撤销重复责任'}))
        refreshed=ok(owner.get(ROOT+'/scopes',params={'issueType':'DUPLICATE'}))['responsibilityCheck']
        assert refreshed['total']==1 and refreshed['items'][0]['id']==second['id']
        current=ok(owner.get(ROOT+'/scopes',params={'scopeType':'COLLEGE','refId':target['id']}))
        assert current['total']==2 and {r['status'] for r in current['items']}=={'ACTIVE','DISABLED'}
        overview=ok(owner.get(ROOT+'/overview'))
        assert next(r for r in overview['issues'] if r['name']=='多名责任人需要核对')['path']=='/admin/system/scopes?issue=DUPLICATE'
        assert owner.get(ROOT+'/scopes',params={'scopeType':'COLLEGE'}).status_code==400
        assert owner.get(ROOT+'/scopes',params={'issueType':'invalid'}).status_code==422
        with get_sessionmaker()() as db:
            foreign=Tenant(tenant_code='SCOPE-'+uuid4().hex[:12],school_name='另一合成学校',status='ACTIVE');db.add(foreign);db.flush()
            foreign_college=College(tenant_id=foreign.id,college_name='重复责任学院',code='C99',status='ACTIVE');db.add(foreign_college);db.commit()
            foreign_id=str(foreign_college.id)
        assert owner.get(ROOT+'/scopes',params={'scopeType':'COLLEGE','refId':foreign_id}).status_code==404
        with reset_for_test(owner,school,'T001',credentials['T001']) as actor:
            assert actor.get(ROOT+'/scopes',params={'scopeType':'COLLEGE','refId':target['id']}).status_code==403
        with get_sessionmaker()() as db:
            log=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],
                SecurityAuditLog.action=='SCOPE_SAVE',SecurityAuditLog.resource_id==retiring['id']).order_by(SecurityAuditLog.id.desc()))
            assert log.detail_json['after']['status']=='DISABLED'


def test_legacy_scope_review_can_resolve_or_retire_without_rebuilding_identity(school):
    from app.models import TeacherStudentScope,College
    with login(school)[0] as owner:
        college=org(owner,'colleges','已核对的真实对象','C01')
        other=org(owner,'colleges','已有正常责任的学院','C02')
        with get_sessionmaker()() as db:
            pending=[TeacherStudentScope(tenant_id=school['id'],teacher_key=f'legacy-{index}',
                teacher_name=f'历史待核对教师{index+1:02d}',role_code='COLLEGE_ADMIN',scope_type='COLLEGE',
                ref_value='历史同名学院',status='REVIEW' if index else 'ACTIVE') for index in range(33)]
            db.add_all(pending)
            stable=TeacherStudentScope(tenant_id=school['id'],teacher_key=f'db-{school["owner"]}',
                teacher_name='学校初始化负责人',role_code='SYSTEM_ADMIN',scope_type='COLLEGE',
                ref_value=other['name'],user_id=school['owner'],ref_id=int(other['id']),status='ACTIVE')
            db.add(stable)
            foreign=Tenant(tenant_code='LEGACY-'+uuid4().hex[:12],school_name='其他学校',status='ACTIVE');db.add(foreign);db.flush()
            outsider=TeacherStudentScope(tenant_id=foreign.id,teacher_key='legacy',teacher_name='其他学校旧关系',
                role_code='COLLEGE_ADMIN',scope_type='COLLEGE',ref_value='历史同名学院',status='REVIEW');db.add(outsider);db.commit()
            pending_ids=[str(r.id) for r in pending];stable_id=str(stable.id);foreign_id=str(outsider.id)
        overview=next(r for r in ok(owner.get(ROOT+'/overview'))['issues'] if r['name']=='历史责任范围待核对')
        assert overview['count']==33 and overview['path']=='/admin/system/scopes?review=1'
        first=ok(owner.get(ROOT+'/scopes',params={'reviewOnly':True}))
        second=ok(owner.get(ROOT+'/scopes',params={'reviewOnly':True,'page':2}))
        assert first['total']==33 and len(first['items'])==30 and len(second['items'])==3
        assert {r['id'] for r in first['items']+second['items']}==set(pending_ids)
        original=first['items'][0]
        impact=ok(owner.get(ROOT+'/scopes/'+original['id']+'/impact'))
        assert impact['roleImpact'] is None and '不能将此情况视为没有待办' in impact['note']
        body={'expectedVersion':original['version'],'reason':'核对原资料后此历史关系不再使用'}
        result=ok(owner.post(ROOT+'/scopes/'+original['id']+'/retire',json=body))
        assert owner.post(ROOT+'/scopes/'+original['id']+'/retire',json=body).status_code==409
        assert owner.post(ROOT+'/scopes/'+foreign_id+'/retire',json=body).status_code==404
        assert owner.post(ROOT+'/scopes/'+stable_id+'/retire',json=body).status_code==409
        assert owner.post(ROOT+'/scopes/'+pending_ids[1]+'/retire',json={**body,'userId':str(school['owner'])}).status_code==422
        repaired=next(r for r in first['items'] if r['id']==pending_ids[1])
        ok(owner.put(ROOT+'/scopes/'+repaired['id'],json={'userId':str(school['owner']),'roleCode':'SYSTEM_ADMIN',
            'scopeType':'COLLEGE','refId':college['id'],'status':'ACTIVE','expectedVersion':repaired['version'],
            'reason':'已人工核对原资料并确认到现有人员和对象'}))
        assert ok(owner.get(ROOT+'/scopes',params={'reviewOnly':True}))['total']==31
        assert next(r for r in ok(owner.get(ROOT+'/overview'))['issues'] if r['name']=='历史责任范围待核对')['count']==31
        with get_sessionmaker()() as db:
            retired=db.get(TeacherStudentScope,int(original['id']))
            assert retired.status=='DISABLED' and retired.user_id is None and retired.ref_id is None
            assert retired.teacher_key=='legacy-0' and retired.teacher_name==original['name'] and retired.ref_value=='历史同名学院'
            assert retired.version==result['version'] and not retired.is_deleted
            assert db.scalar(select(func.count()).select_from(User).where(User.tenant_id==school['id']))==1
            event=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],
                SecurityAuditLog.action=='SCOPE_SAVE',SecurityAuditLog.resource_id==original['id']).order_by(SecurityAuditLog.id.desc()))
            assert event.detail_json['before']['status']=='ACTIVE' and event.detail_json['after']['status']=='DISABLED'
            assert event.detail_json['before']['refName']==event.detail_json['after']['refName']=='历史同名学院'
        _,credentials=import_rows(owner,'staff',[['T001','无管理权限教师','C01','已核对的真实对象','','','','','']])
        with reset_for_test(owner,school,'T001',credentials['T001']) as teacher:
            assert teacher.post(ROOT+'/scopes/'+pending_ids[2]+'/retire',json=body).status_code==403


def test_scope_by_stable_id_refuses_other_college_and_stale_updates(school):
    with login(school)[0] as client:
        first,creds=import_rows(client,'staff',[['T001','院系管理员','C01','同名学院','','','','','']])
        import_rows(client,'students',[['S001','一院学生','C01','同名学院','M01','专业一','B01','班级一','2024'],
                                      ['S002','二院学生','C02','同名学院','M02','专业二','B02','班级二','2024']])
        account=ok(client.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        role=ok(client.post(ROOT+'/roles',json={'name':'本院资料管理员','baseRole':'COLLEGE_ADMIN',
            'permissions':['system.student.view','system.student.manage','system.org.view','system.person.export','system.task.view'],'reason':'按学院隔离维护资料'}))
        role=ok(client.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],'enabled':True,'expectedVersion':role['version'],'reason':'分配院系职责'}))
        definition=next(r for r in ok(client.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
        college=next(r for r in ok(client.get(ROOT+'/organizations/colleges'))['items'] if r['code']=='C01')
        scope=ok(client.post(ROOT+'/scopes',json={'userId':account['id'],'roleCode':definition['code'],'scopeType':'COLLEGE','refId':college['id'],'reason':'明确负责一院'}))
        teacher=reset_for_test(client,school,'T001',creds['T001'])
        switched=ok(teacher.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+definition['code'].lower(),'clientType':'PC'}));teacher.headers['Authorization']='Bearer '+switched['accessToken']
        visible=ok(teacher.get(ROOT+'/people/student'));assert [r['number'] for r in visible['items']]==['S001']
        ok(client.put(ROOT+f'/organizations/colleges/{college["id"]}',json={'name':'更名后的学院','code':'C01','expectedVersion':college['version'],'reason':'核对改名不改变授权'}))
        assert ok(teacher.get(ROOT+'/people/student'))['total']==1
        other=next(r for r in ok(client.get(ROOT+'/people/student'))['items'] if r['number']=='S002')
        assert teacher.get(ROOT+f'/people/student/{other["id"]}/impact').status_code==404
        assert teacher.get(ROOT+'/exports/student').status_code==200
        job=next(r for r in ok(teacher.get(ROOT+'/tasks'))['items'] if r['canDownload'])
        assert teacher.get(ROOT+'/tasks/'+job['id']+'/result').status_code==200
        checks=ok(client.get(ROOT+'/scopes'))['responsibilityCheck']
        assert checks['missing']>=1
        ok(client.put(ROOT+'/scopes/'+scope['id'],json={'userId':account['id'],'roleCode':definition['code'],'scopeType':'COLLEGE',
            'refId':college['id'],'expectedVersion':scope['version'],'status':'DISABLED','reason':'回收院系管理范围后旧导出失效'}))
        assert teacher.get(ROOT+'/tasks/'+job['id']+'/result').status_code==403
        assert ok(teacher.get(ROOT+'/people/student'))['total']==0
        teacher.close()


@pytest.mark.parametrize('revoked',['membership','role','account','expired'])
def test_revoked_responsibility_can_only_be_disabled_without_restoring_access(school,revoked):
    with login(school)[0] as owner:
        import_rows(owner,'staff',[['T001','离岗范围核对教师','C01','原负责学院','','','','','']])
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        college=ok(owner.get(ROOT+'/organizations/colleges'))['items'][0]
        other=org(owner,'colleges','另一学院','C02')
        role=ok(owner.post(ROOT+'/roles',json={'name':'离岗责任核对角色','baseRole':'COLLEGE_ADMIN',
            'permissions':['system.student.view'],'reason':'合成责任撤回验收'}))
        role=ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'合成原有任职'}))
        definition=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
        body={'userId':account['id'],'roleCode':definition['code'],'scopeType':'COLLEGE',
            'refId':college['id'],'reason':'合成离岗范围原关系'}
        scope=ok(owner.post(ROOT+'/scopes',json=body))
        preview=ok(owner.get(ROOT+'/scopes/'+scope['id']+'/impact'))
        assert preview['scopeId']==scope['id'] and preview['original']['refName']==college['name']
        assert preview['roleName']==definition['name'] and preview['roleImpact']['memberCount']==1
        if revoked=='membership':
            ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],'enabled':False,
                'expectedVersion':role['version'],'reason':'先撤销离岗人员身份'}))
        elif revoked=='role':
            ok(owner.put(ROOT+f'/roles/{role["id"]}',json={'name':definition['name'],'baseRole':definition['baseRole'],
                'permissions':definition['permissions'],'status':'DISABLED','expectedVersion':role['version'],'reason':'先停用原职责角色'}))
        else:
            current=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
            ok(owner.post(ROOT+f'/accounts/{account["id"]}/actions',json={'action':'DISABLE' if revoked=='account' else 'EXPIRY',
                'expiresAt':None if revoked=='account' else '2020-01-01T00:00:00Z',
                'expectedVersion':current['version'],'reason':'先结束人员登录资格'}))
        disabled={**body,'status':'DISABLED','expectedVersion':scope['version'],'reason':'停用离岗人员遗留范围'}
        result=ok(owner.put(ROOT+'/scopes/'+scope['id'],json=disabled))
        assert result['version']==scope['version']+1
        assert owner.put(ROOT+'/scopes/'+scope['id'],json=disabled).status_code==409
        assert owner.put(ROOT+'/scopes/'+scope['id'],json={**disabled,'status':'ACTIVE','expectedVersion':result['version']}).status_code==403
        assert owner.put(ROOT+'/scopes/'+scope['id'],json={**disabled,'refId':other['id'],'expectedVersion':result['version']}).status_code==403
        assert owner.post(ROOT+'/scopes',json={**body,'refId':other['id']}).status_code==403
        current=next(r for r in ok(owner.get(ROOT+'/scopes'))['items'] if r['id']==scope['id'])
        assert current['status']=='DISABLED' and current['refId']==college['id'] and current['version']==result['version']
        preview=ok(owner.get(ROOT+'/scopes/'+scope['id']+'/impact'))
        assert preview['original']['status']=='DISABLED' and preview['roleImpact']['scopeCount']==0
        with get_sessionmaker()() as db:
            audit=db.scalar(select(SecurityAuditLog).where(SecurityAuditLog.tenant_id==school['id'],
                SecurityAuditLog.action=='SCOPE_SAVE',SecurityAuditLog.resource_id==scope['id']).order_by(SecurityAuditLog.id.desc()))
            assert audit.detail_json['before']['status']=='ACTIVE' and audit.detail_json['after']['status']=='DISABLED'


def activate(client, code):
    result=ok(client.post('/api/v1/auth/browser-switch-role',json={'contextId':'role-'+code.lower(),'clientType':'PC'}))
    client.headers['Authorization']='Bearer '+result['accessToken']


def test_narrow_business_role_batch_context_keeps_scope_without_batch_management_access(school):
    from datetime import datetime
    from app.models import InternshipBatch, InternshipRecord
    with login(school)[0] as owner:
        _,credentials=import_rows(owner,'staff',[['T001','本院请假承办人','C01','本院','','','','','']])
        import_rows(owner,'students',[
            ['S001','本院合成学生','C01','本院','M01','本院专业','B01','本院班级','2024'],
            ['S002','外院合成学生','C02','外院','M02','外院专业','B02','外院班级','2024']])
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        students={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        definition=ok(owner.post(ROOT+'/roles',json={'name':'仅本院请假办理','baseRole':'COLLEGE_ADMIN',
            'permissions':['internship.leave.view','internship.leave.review'],'reason':'最小业务权限上下文验收'}))
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==definition['id'])
        membership=ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':account['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'配置本院请假职责'}))
        scope=ok(owner.post(ROOT+'/scopes',json={'userId':account['id'],'roleCode':role['code'],'scopeType':'COLLEGE',
            'refId':students['S001']['collegeId'],'reason':'仅负责本院学生'}))
        with get_sessionmaker()() as db:
            other=Tenant(tenant_code='SYS-'+uuid4().hex[:12],school_name='目录隔离合成学校',status='ACTIVE')
            db.add(other);db.flush()
            batches=[]
            for key,tenant_id,status in [('OWN_LITERAL',school['id'],'RUNNING'),('OWN_SECOND',school['id'],'CLOSED'),
                                         ('OTHER_COLLEGE',school['id'],'RUNNING'),('FOREIGN',other.id,'RUNNING'),
                                         ('VOIDED',school['id'],'VOIDED'),('EMPTY',school['id'],'DRAFT')]:
                row=InternshipBatch(tenant_id=tenant_id,batch_no=key,batch_name='审批%_批次' if key=='OWN_LITERAL' else key,
                    status=status,start_date=datetime(2026,9,1),end_date=datetime(2026,12,31),
                    rules_config={'privateRule':'not-context-data'},remark='不应出现在选择目录')
                db.add(row);db.flush();batches.append(row)
            ids={row.batch_no:str(row.id) for row in batches}
            for key,student_no in [('OWN_LITERAL','S001'),('OWN_SECOND','S001'),('OTHER_COLLEGE','S002'),('VOIDED','S001')]:
                db.add(InternshipRecord(tenant_id=school['id'],student_id=int(students[student_no]['id']),batch_id=int(ids[key]),status='IN_PROGRESS'))
            db.commit()
        path='/api/v1/internship/batch-options'
        with reset_for_test(owner,school,'T001',credentials['T001']) as actor:
            activate(actor,role['code'])
            page=ok(actor.get(path,params={'pageSize':1,'selectedId':ids['OWN_LITERAL']}))
            assert page['total']==2 and [r['id'] for r in page['items']]==[ids['OWN_SECOND']]
            assert page['selected']['id']==ids['OWN_LITERAL']
            assert set(page['selected'])=={'id','batchName','batchNo','status','startDate','endDate'}
            assert [r['id'] for r in ok(actor.get(path,params={'page':2,'pageSize':1}))['items']]==[ids['OWN_LITERAL']]
            assert [r['id'] for r in ok(actor.get(path,params={'keyword':'%_'}))['items']]==[ids['OWN_LITERAL']]
            for excluded in ['OTHER_COLLEGE','FOREIGN','VOIDED','EMPTY']:
                assert ok(actor.get(path,params={'selectedId':ids[excluded]}))['selected'] is None
            assert actor.get(path,params={'selectedId':'not-a-number'}).status_code==422
            assert actor.get('/api/v1/internship/batches').status_code==403
            assert actor.get('/api/v1/internship/batches/'+ids['OWN_LITERAL']).status_code==403
            assert actor.get('/api/v1/internship/intern-students',params={'batchId':ids['OWN_LITERAL']}).status_code==403
            assert actor.get('/api/v1/internship/leaves',params={'batchId':ids['OWN_LITERAL']}).status_code==200
            ok(owner.put(ROOT+'/scopes/'+scope['id'],json={'userId':account['id'],'roleCode':role['code'],
                'scopeType':'COLLEGE','refId':students['S001']['collegeId'],'status':'DISABLED',
                'expectedVersion':scope['version'],'reason':'撤回负责范围，旧会话目录立即收敛'}))
            withdrawn=ok(actor.get(path,params={'selectedId':ids['OWN_LITERAL']}))
            assert withdrawn['total']==0 and withdrawn['items']==[] and withdrawn['selected'] is None
            current=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
            ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':role['name'],'baseRole':role['baseRole'],
                'permissions':[],'expectedVersion':current['version'],'reason':'清空权限后不得借用另一身份的批次权限'}))
            assert actor.get(path).status_code==403
        assert owner.get(path).status_code==403, '系统配置身份不自动获取业务批次目录'


def test_mentor_system_directory_and_account_actions_follow_actual_assignments(school):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[
            ['T001','本人指导教师','C01','工程学院','','','','',''],
            ['T002','另一指导教师','C01','工程学院','','','','','']])
        import_rows(owner,'students',[
            ['S001','本人指导学生','C01','工程学院','M01','软件技术','B01','本人指导班级','2026'],
            ['S002','其他教师学生','C01','工程学院','M02','其他专业','B02','其他班级','2026']])
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        students={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        staff={r['number']:r for r in ok(owner.get(ROOT+'/people/staff'))['items']}
        with get_sessionmaker()() as db:
            batch=InternshipBatch(tenant_id=school['id'],batch_name='指导身份范围验收',batch_no='MENTOR-SCOPE',status='RUNNING',
                start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.flush()
            records=[InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=int(students[s]['id']),status='ONBOARD',
                advisor_user_id=int(accounts[t]['id']),advisor_name=staff[t]['name']) for s,t in [('S001','T001'),('S002','T002')]]
            db.add_all(records);db.commit();record_id=str(records[0].id);record_version=records[0].version
        role=ok(owner.post(ROOT+'/roles',json={'name':'受控指导资料角色','baseRole':'INTERN_MENTOR',
            'permissions':['system.student.view','system.staff.view','system.org.view','system.account.view','system.account.reset','system.person.export','system.task.view'],
            'sensitiveConfirmed':True,'reason':'验证指导身份只能管理正式分配对象'}))
        role=ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':accounts['T001']['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'合成指导资料职责'}))
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
        # An old/broad responsibility entry must not turn a mentor identity into a college administrator.
        ok(owner.post(ROOT+'/scopes',json={'userId':accounts['T001']['id'],'roleCode':role['code'],'scopeType':'COLLEGE',
            'refId':students['S001']['collegeId'],'reason':'核验宽组织关系不能扩大指导边界'}))
        with reset_for_test(owner,school,'T001',creds['T001']) as teacher:
            activate(teacher,role['code'])
            assert [r['number'] for r in ok(teacher.get(ROOT+'/people/student'))['items']]==['S001']
            assert [r['number'] for r in ok(teacher.get(ROOT+'/people/staff'))['items']]==['T001']
            assert [r['loginName'] for r in ok(teacher.get(ROOT+'/accounts'))['items']]==['S001']
            assert [r['code'] for r in ok(teacher.get(ROOT+'/organizations/classes'))['items']]==['B01']
            assert teacher.get(ROOT+'/people/student/'+students['S002']['id']+'/impact').status_code==404
            for name in ['S002','T001','T002']:
                assert teacher.post(ROOT+'/accounts/'+accounts[name]['id']+'/actions',json={'action':'RESET',
                    'expectedVersion':accounts[name]['version'],'reason':'拒绝对指导范围外账号执行重置'}).status_code==404
            own=ok(teacher.get(ROOT+'/accounts'))['items'][0]
            assert teacher.post(ROOT+'/accounts/'+own['id']+'/actions',json={'action':'RESET',
                'expectedVersion':own['version'],'reason':'受控重置本人正式指导学生'}).status_code==200
            response=teacher.get(ROOT+'/exports/student');assert response.status_code==200
            workbook=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
            rows=list(workbook.active.iter_rows(values_only=True));workbook.close()
            assert len(rows)==2 and rows[1][0]=='S001'
            job=next(r for r in ok(teacher.get(ROOT+'/tasks'))['items'] if r['canDownload'])
            school_role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SCHOOL_ADMIN')
            ok(owner.post(ROOT+'/roles/'+school_role['id']+'/members',json={'userId':str(school['owner']),'enabled':True,
                'expectedVersion':school_role['version'],'reason':'合成导师正式交接身份'}))
            activate(owner,'SCHOOL_ADMIN')
            ok(owner.post('/api/v1/internship/intern-students/'+record_id+'/advisor',json={'advisorUserId':accounts['T002']['id'],
                'expectedVersion':record_version,'reason':'正式交接后原导师资料范围同步收回'}))
            assert ok(teacher.get(ROOT+'/people/student'))['total']==0
            assert ok(teacher.get(ROOT+'/accounts'))['total']==0
            assert teacher.get(ROOT+'/tasks/'+job['id']+'/result').status_code==403


def test_two_college_admins_isolate_same_name_organizations_students_and_exports(school):
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[
            ['T001','一院管理老师','C01','同名学院','','','','',''],
            ['T002','二院管理老师','C02','同名学院','','','','','']])
        import_rows(owner,'students',[
            ['S001','一院核验学生','C01','同名学院','M01','一院专业','B01','一院班级','2026'],
            ['S002','二院核验学生','C02','同名学院','M02','二院专业','B02','二院班级','2026']])
        role=ok(owner.post(ROOT+'/roles',json={'name':'双院隔离资料管理员','baseRole':'COLLEGE_ADMIN',
            'permissions':['system.org.view','system.student.view','system.student.manage','system.person.export','system.task.view'],
            'reason':'双院合成场景按稳定组织范围核验'}))
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        colleges={r['code']:r for r in ok(owner.get(ROOT+'/organizations/colleges'))['items']}
        students={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        for username,code in [('T001','C01'),('T002','C02')]:
            current=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==role['id'])
            role=current
            ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':accounts[username]['id'],
                'enabled':True,'expectedVersion':current['version'],'reason':'合成院系管理员任职'}))
            ok(owner.post(ROOT+'/scopes',json={'userId':accounts[username]['id'],'roleCode':role['code'],
                'scopeType':'COLLEGE','refId':colleges[code]['id'],'reason':'按同名学院稳定编号负责'}))
        with reset_for_test(owner,school,'T001',creds['T001']) as first, reset_for_test(owner,school,'T002',creds['T002']) as second:
            jobs=[]
            for actor,own,other,college in [(first,'S001','S002','C01'),(second,'S002','S001','C02')]:
                activate(actor,role['code'])
                assert [r['number'] for r in ok(actor.get(ROOT+'/people/student'))['items']]==[own]
                assert [r['number'] for r in ok(actor.get(ROOT+'/people/student',params={'accountStatus':'BOUND'}))['items']]==[own]
                assert ok(actor.get(ROOT+'/people/student',params={'accountStatus':'UNBOUND'}))['total']==0
                assert ok(actor.get(ROOT+'/people/student',params={'accountStatus':'BOUND','keyword':other}))['total']==0
                assert [r['code'] for r in ok(actor.get(ROOT+'/organizations/colleges'))['items']]==[college]
                assert actor.get(ROOT+f'/people/student/{students[other]["id"]}/impact').status_code==404
                response=actor.get(ROOT+'/exports/student');assert response.status_code==200
                workbook=load_workbook(io.BytesIO(response.content),read_only=True,data_only=True)
                rows=list(workbook.active.iter_rows(values_only=True));workbook.close()
                assert len(rows)==2 and rows[1][0]==own
                jobs.append(next(r for r in ok(actor.get(ROOT+'/tasks'))['items'] if r['canDownload'])['id'])
            assert first.get(ROOT+'/tasks/'+jobs[1]+'/result').status_code==403
            assert second.get(ROOT+'/tasks/'+jobs[0]+'/result').status_code==403


def test_college_identity_import_rejects_cross_scope_collisions_and_revoked_receipts(school):
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[['T001','院系开户管理员','C01','一院','','','','','']])
        import_rows(owner,'students',[['S002','二院学生','C02','二院','M02','二院专业','B02','二院班级','2026']])
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':account['id'],'enabled':True,
            'expectedVersion':role['version'],'reason':'负责本院开户'}))
        college=next(r for r in ok(owner.get(ROOT+'/organizations/colleges'))['items'] if r['code']=='C01')
        scope=ok(owner.post(ROOT+'/scopes',json={'userId':account['id'],'roleCode':'COLLEGE_ADMIN',
            'scopeType':'COLLEGE','refId':college['id'],'reason':'限定在第一学院'}))
        with reset_for_test(owner,school,'T001',creds['T001']) as actor:
            activate(actor,'COLLEGE_ADMIN')
            result,_=import_rows(actor,'students',[['S001','一院学生','C01','一院','M01','一院专业','B01','一院班级','2026']])
            import_rows(actor,'staff',[['T002','一院指导教师','C01','一院','','','','','']])
            for row in [
                ['S003','跨院学生','C02','二院','M02','二院专业','B02','二院班级','2026'],
                ['S002','二院学生','C01','一院','M01','一院专业','B01','一院班级','2026'],
                ['S003','错挂专业学生','C01','一院','M02','二院专业','B02','二院班级','2026']]:
                bad=preview(actor,'students',[row]);assert bad['invalidRows']==1
                assert not bad['entries'][0].get('personId') and not bad['entries'][0].get('accountId') and not bad['entries'][0]['refs']
                assert actor.post(ROOT+f'/identity-import/jobs/{bad["id"]}/confirm',json={'expectedVersion':bad['version']}).status_code==409
            student=ok(owner.get(ROOT+'/people/student',params={'keyword':'S001'}))['items'][0]
            other=ok(owner.get(ROOT+'/people/student',params={'keyword':'S002'}))['items'][0]
            invalid=preview(actor,'students',[['S001','一院学生','C01','一院','M01','一院专业','B01','一院班级','']])
            assert invalid['invalidRows']==1
            assert actor.get(ROOT+f'/identity-import/jobs/{invalid["id"]}/errors').status_code==200
            assert actor.get(ROOT+f'/identity-import/jobs/{result["id"]}').status_code==200
            ok(owner.put(ROOT+'/people/student/'+student['id'],json={'name':student['name'],
                'collegeId':other['collegeId'],'classId':other['classId'],'grade':'2026','expectedVersion':student['version'],
                'confirmImpact':True,'reason':'学生换班到另一院系，原开户回执不得继续领取'}))
            assert actor.get(ROOT+f'/identity-import/jobs/{result["id"]}/receipt').status_code==403
            assert actor.get(ROOT+f'/identity-import/jobs/{result["id"]}').status_code==403
            assert actor.get(ROOT+f'/identity-import/jobs/{invalid["id"]}').status_code==403
            assert actor.get(ROOT+f'/identity-import/jobs/{invalid["id"]}/errors').status_code==403
            plan=preview(actor,'students',[['S004','待确认学生','C01','一院','M01','一院专业','B01','一院班级','2026']])
            assert plan['invalidRows']==0
            ok(owner.put(ROOT+'/scopes/'+scope['id'],json={'userId':account['id'],'roleCode':'COLLEGE_ADMIN',
                'scopeType':'COLLEGE','refId':college['id'],'expectedVersion':scope['version'],'status':'DISABLED','reason':'撤销该院系负责范围'}))
            assert actor.get(ROOT+f'/identity-import/jobs/{result["id"]}/receipt').status_code==403
            assert actor.get(ROOT+f'/identity-import/jobs/{plan["id"]}').status_code==403
            assert actor.post(ROOT+f'/identity-import/jobs/{plan["id"]}/confirm',json={'expectedVersion':plan['version']}).status_code==403
            assert not ok(owner.get(ROOT+'/accounts',params={'keyword':'S004'}))['items']


def test_system_batch_picker_searches_and_pages_past_200_without_cross_school_access(school):
    from app.models import InternshipBatch
    with get_sessionmaker()() as db:
        old=InternshipBatch(tenant_id=school['id'],batch_no='B-000',batch_name='历史核对批次',status='RUNNING')
        foreign_school=Tenant(tenant_code='SYS-'+uuid4().hex[:12],school_name='另一合成批次学校',status='ACTIVE')
        db.add_all([old,foreign_school]);db.flush();old_id=str(old.id)
        db.add_all([InternshipBatch(tenant_id=school['id'],batch_no=f'B-{i:03}',batch_name=f'后续合成批次{i}',status='RUNNING') for i in range(1,201)])
        foreign=InternshipBatch(tenant_id=foreign_school.id,batch_no='FOREIGN',batch_name='其他学校批次',status='RUNNING')
        db.add(foreign);db.commit();foreign_id=str(foreign.id)
    with login(school)[0] as owner:
        page=ok(owner.get(ROOT+'/batches',params={'pageSize':50}))
        assert page['total']==201 and len(page['items'])==50
        assert ok(owner.get(ROOT+'/batches',params={'page':5,'pageSize':50}))['items'][0]['id']==old_id
        for keyword in ['B-000','历史核对批次']:
            found=ok(owner.get(ROOT+'/batches',params={'keyword':keyword}))
            assert found['total']==1 and found['items'][0]['id']==old_id
        assert ok(owner.get(ROOT+'/selection',params={'kind':'batches','identifier':old_id}))['name']=='历史核对批次'
        assert owner.get(ROOT+'/selection',params={'kind':'batches','identifier':foreign_id}).status_code==404
        assert ok(owner.get(ROOT+'/batches',params={'keyword':'FOREIGN'}))['total']==0
        assert ok(owner.get(ROOT+'/leave-policies/'+old_id))['batchName']=='历史核对批次'
        assert ok(owner.get(ROOT+'/handoffs',params={'batchId':old_id,'userId':school['owner']}))['total']==0
        _,credentials=import_rows(owner,'staff',[['T001','普通指导教师','C01','工程学院','','','','','']])
        with reset_for_test(owner,school,'T001',credentials['T001']) as teacher:
            assert teacher.get(ROOT+'/batches').status_code==403
            assert teacher.get(ROOT+'/selection',params={'kind':'batches','identifier':old_id}).status_code==403


def test_existing_database_initialization_refuses_without_receipt_or_data_change(school,tmp_path):
    import importlib.util
    from pathlib import Path
    spec=importlib.util.spec_from_file_location('school_initializer',Path(__file__).resolve().parents[2]/'scripts/initialize_school.py')
    module=importlib.util.module_from_spec(spec);spec.loader.exec_module(module)
    with get_sessionmaker()() as db: before=db.scalar(select(func.count()).select_from(Tenant))
    receipt=tmp_path/'private.json'
    with pytest.raises(ValueError,match='禁止重复初始化'):
        module.initialize('SHOULD-NOT-CREATE','禁止覆盖测试学校','admin',receipt)
    assert not receipt.exists()
    with get_sessionmaker()() as db: assert db.scalar(select(func.count()).select_from(Tenant))==before


@pytest.mark.parametrize('unavailable',['staff_disabled','account_expired'])
def test_unavailable_staff_cannot_receive_new_students_but_history_can_be_handed_over(school,unavailable):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord
    with login(school)[0] as owner:
        import_rows(owner,'staff',[[n,n,'C01','工程学院','','','','',''] for n in ['T001','T002']])
        import_rows(owner,'students',[[n,n,'C01','工程学院','M01','软件技术','B01','软件一班','2024'] for n in ['S001','S002']])
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        staff=next(r for r in ok(owner.get(ROOT+'/people/staff'))['items'] if r['number']=='T002')
        with get_sessionmaker()() as db:
            students={s.student_no:s.id for s in db.scalars(select(StudentProfile).where(StudentProfile.tenant_id==school['id']))}
            batch=InternshipBatch(tenant_id=school['id'],batch_name='导师资格校验',batch_no='ADVISOR-STATUS',status='DRAFT',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.flush()
            record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=students['S001'],status='PREPARING',advisor_user_id=int(accounts['T002']['id']),advisor_name='T002')
            db.add(record);db.commit();rid=str(record.id);bid=str(batch.id);version=record.version
        if unavailable=='staff_disabled':
            body={key:staff[key] for key in ('name','collegeId','departmentId','classId','grade')}
            ok(owner.put(ROOT+'/people/staff/'+staff['id'],json={**body,'status':'DISABLED','expectedVersion':staff['version'],'reason':'合成停用资料验收'}))
        else:
            ok(owner.post(ROOT+'/accounts/'+accounts['T002']['id']+'/actions',json={'action':'EXPIRY','expiresAt':'2020-01-01T00:00:00Z','expectedVersion':accounts['T002']['version'],'reason':'合成账号过期验收'}))
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SCHOOL_ADMIN')
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':str(school['owner']),'enabled':True,'expectedVersion':role['version'],'reason':'测试校级实习办理身份'}))
        activate(owner,'SCHOOL_ADMIN')
        root='/api/v1/internship/intern-students'
        advisors=ok(owner.get(root+'/advisors'))
        assert {r['id'] for r in advisors}=={accounts['T001']['id']}
        failed=owner.post(root,json={'studentId':str(students['S002']),'batchId':bid,'advisorUserId':accounts['T002']['id']})
        assert failed.status_code==400
        assert '所选教师' in failed.json()['message']
        plan=ok(owner.post(root+'/import/dry-run',json={'batchId':bid,'rows':[{'studentNo':'S002','advisorName':'T002'}]}))
        assert plan['invalidRows']==1 and plan['errors'][0]['field']=='advisorName'
        assert owner.post(root+'/'+rid+'/advisor',json={'advisorUserId':accounts['T002']['id'],'expectedVersion':version,'reason':'不能再次分配失效导师'}).status_code==400
        with get_sessionmaker()() as db:
            preserved=db.get(InternshipRecord,int(rid))
            assert preserved.advisor_user_id==int(accounts['T002']['id']) and preserved.version==version
            assert db.scalar(select(func.count()).select_from(InternshipRecord).where(InternshipRecord.tenant_id==school['id']))==1
            assert db.get(User,int(accounts['T002']['id'])).status=='ACTIVE'
        ok(owner.post(root+'/'+rid+'/advisor',json={'advisorUserId':accounts['T001']['id'],'expectedVersion':version,'reason':'由有效导师接手原记录'}))
        with get_sessionmaker()() as db:
            transferred=db.get(InternshipRecord,int(rid))
            assert transferred.advisor_user_id==int(accounts['T001']['id']) and transferred.status=='PREPARING'
        created=ok(owner.post(root,json={'studentId':str(students['S002']),'batchId':bid,'advisorUserId':accounts['T001']['id']}))
        assert created['advisorUserId']==accounts['T001']['id']


@pytest.mark.parametrize('mode',['permission','membership','role','alternate'])
def test_revoked_advisor_permissions_surface_original_todos_and_handoff(school,mode):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord,WeeklyReport,UnifiedTodo
    from app.models.internship_employment_registration import InternshipEmploymentRegistration
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[[n,n,'C01','工程学院','','','','',''] for n in ['T001','T002']])
        import_rows(owner,'students',[['S001','撤权交接学生','C01','工程学院','M01','软件技术','B01','软件一班','2024']])
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        permissions=['internship.report.view','internship.report.review','internship.employment.view','internship.employment.review']
        original=ok(owner.post(ROOT+'/roles',json={'name':'原教师办理岗','baseRole':'INTERN_MENTOR','permissions':permissions,'reason':'合成撤权办理身份'}))
        original=ok(owner.post(ROOT+'/roles/'+original['id']+'/members',json={'userId':accounts['T001']['id'],'enabled':True,'expectedVersion':original['version'],'reason':'合成原教师身份'}))
        original=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['id']==original['id'])
        preset=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='INTERN_MENTOR')
        ok(owner.post(ROOT+'/roles/'+preset['id']+'/members',json={'userId':accounts['T001']['id'],'enabled':False,'expectedVersion':preset['version'],'reason':'仅用合成自定义办理身份'}))
        if mode=='alternate':
            extra=ok(owner.post(ROOT+'/roles',json={'name':'仍有效的另一教师岗','baseRole':'INTERN_MENTOR','permissions':permissions,'reason':'合成另一有效身份'}))
            ok(owner.post(ROOT+'/roles/'+extra['id']+'/members',json={'userId':accounts['T001']['id'],'enabled':True,'expectedVersion':extra['version'],'reason':'另一个有效办理身份'}))
        with get_sessionmaker()() as db:
            student=db.scalar(select(StudentProfile).where(StudentProfile.tenant_id==school['id']))
            batch=InternshipBatch(tenant_id=school['id'],batch_name='撤权待办交接',batch_no='REVOKED-HANDOFF',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31));db.add(batch);db.flush()
            record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=student.id,status='ONBOARD',advisor_user_id=int(accounts['T001']['id']),advisor_name='T001');db.add(record);db.flush()
            report=WeeklyReport(tenant_id=school['id'],internship_id=record.id,week_number=1,status='PENDING_REVIEW');db.add(report)
            employment=InternshipEmploymentRegistration(tenant_id=school['id'],internship_id=record.id,student_id=student.id,batch_id=batch.id,status='SUBMITTED');db.add(employment);db.flush()
            for item,kind,typ in [(report,'INTERN_WEEKLY','INTERN_WEEKLY_REVIEW'),(employment,'INTERNSHIP_EMPLOYMENT','INTERN_EMPLOYMENT_REPORT')]:
                db.add(UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type=kind,source_biz_id=item.id,todo_type=typ,assignee_id=int(accounts['T001']['id']),student_id=student.id,title='撤权前原待办',status='PENDING'))
            db.commit();rid=str(record.id);bid=str(batch.id);version=record.version;report_id=report.id
        assert ok(owner.get(ROOT+'/scopes'))['blockedTasks']['total']==0
        with reset_for_test(owner,school,'T001',creds['T001']) as teacher:
            activate(teacher,original['code'])
            if mode=='membership':
                ok(owner.post(ROOT+'/roles/'+original['id']+'/members',json={'userId':accounts['T001']['id'],'enabled':False,'expectedVersion':original['version'],'reason':'撤销原岗位成员'}))
            else:
                ok(owner.put(ROOT+'/roles/'+original['id'],json={'name':'原教师办理岗','baseRole':'INTERN_MENTOR','permissions':permissions if mode=='role' else [],
                    'status':'DISABLED' if mode=='role' else 'ACTIVE','expectedVersion':original['version'],'reason':'收回原岗位办理能力'}))
            assert teacher.post(f'/api/v1/internship/reports/{report_id}/review',json={'action':'APPROVE','comment':'撤权旧会话不得办理','expectedVersion':0,'ratingLevel':3}).status_code in (401,403)
        blocked=ok(owner.get(ROOT+'/scopes'))['blockedTasks']
        assert blocked['total']==(0 if mode=='alternate' else 2)
        assert all(t['path']==f'/admin/internship/students?batchId={bid}&advisorId={rid}' for t in blocked['items'])
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SCHOOL_ADMIN')
        for account in [str(school['owner']),accounts['T002']['id']]:
            role=ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':account,'enabled':True,'expectedVersion':role['version'],'reason':'合成有权交接与接收人员'}))
        activate(owner,'SCHOOL_ADMIN')
        ok(owner.post('/api/v1/internship/intern-students/'+rid+'/advisor',json={'advisorUserId':accounts['T002']['id'],'expectedVersion':version,'reason':'原岗位撤权后正常交接'}))
        activate(owner,'SYSTEM_ADMIN');assert ok(owner.get(ROOT+'/scopes'))['blockedTasks']['total']==0
        with get_sessionmaker()() as db:
            todos=list(db.scalars(select(UnifiedTodo).where(UnifiedTodo.tenant_id==school['id'])))
            assert len(todos)==4
            assert {(t.assignee_id,t.status) for t in todos}=={(int(accounts['T001']['id']),'CANCELLED'),(int(accounts['T002']['id']),'PENDING')}
            assert db.get(WeeklyReport,report_id).status=='PENDING_REVIEW'


def test_disabled_advisor_new_report_keeps_original_and_handoff_recovers_missing_reports(school):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord,WeeklyReport,InternshipProcessReport,UnifiedTodo
    with login(school)[0] as owner:
        import_rows(owner,'staff',[[n,'同名指导教师','C01','工程学院','','','','',''] for n in ['T001','T002']])
        _,passwords=import_rows(owner,'students',[['S001','交接学生','C01','工程学院','M01','软件技术','B01','软件一班','2024']])
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        with get_sessionmaker()() as db:
            student=db.scalar(select(StudentProfile).where(StudentProfile.tenant_id==school['id']))
            batch=InternshipBatch(tenant_id=school['id'],batch_name='停用后交接',batch_no='LATE-HANDOFF',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31));db.add(batch);db.flush()
            record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=student.id,status='ONBOARD',advisor_user_id=int(accounts['T001']['id']),advisor_name='同名指导教师');db.add(record);db.flush()
            missing=WeeklyReport(tenant_id=school['id'],internship_id=record.id,week_number=2,status='PENDING_REVIEW');db.add(missing)
            monthly=InternshipProcessReport(tenant_id=school['id'],internship_id=record.id,report_type='MONTHLY',period_key='2026-09',status='PENDING_REVIEW');db.add(monthly)
            history=WeeklyReport(tenant_id=school['id'],internship_id=record.id,week_number=3,status='APPROVED',reviewed_by_name='原教师');db.add(history);db.commit()
            bid=str(batch.id);rid=str(record.id);version=record.version;missing_id=missing.id;monthly_id=monthly.id;history_id=history.id
        old=accounts['T001']
        ok(owner.post(ROOT+'/accounts/'+old['id']+'/actions',json={'action':'DISABLE','expectedVersion':old['version'],'reason':'紧急停用原指导教师'}))
        with reset_for_test(owner,school,'S001',passwords['S001'],True) as student_client:
            submitted=ok(student_client.post('/api/v1/portal/internship/context/weekly-reports',json={'batchId':bid,'internshipId':rid,'weekNo':1,'expectedVersion':0,
                'workContent':'合成周报工作记录'*8,'harvestContent':'合成收获记录','planContent':'合成下周计划','attachmentFileIds':[]}))
        blocked=ok(owner.get(ROOT+'/scopes'))['blockedTasks']
        assert blocked['total']==1 and blocked['items'][0]['assigneeId']==old['id']
        assert blocked['items'][0]['path']==f'/admin/internship/students?batchId={bid}&advisorId={rid}'
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SCHOOL_ADMIN')
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':str(school['owner']),'enabled':True,'expectedVersion':role['version'],'reason':'合成学校办理人员'}));activate(owner,'SCHOOL_ADMIN')
        ok(owner.post('/api/v1/internship/intern-students/'+rid+'/advisor',json={'advisorUserId':accounts['T002']['id'],'reason':'同名不同账号正常交接','expectedVersion':version}))
        with get_sessionmaker()() as db:
            todo_rows=list(db.scalars(select(UnifiedTodo).where(UnifiedTodo.tenant_id==school['id'])))
            assert {(t.source_biz_id,t.todo_type) for t in todo_rows if t.status=='PENDING'}=={
                (int(submitted['id']),'INTERN_WEEKLY_REVIEW'),(missing_id,'INTERN_WEEKLY_REVIEW'),(monthly_id,'INTERN_MONTHLY_REVIEW')}
            assert all(t.assignee_id==int(accounts['T002']['id']) for t in todo_rows if t.status=='PENDING')
            assert any(t.assignee_id==int(old['id']) and t.status=='CANCELLED' for t in todo_rows)
            assert db.get(WeeklyReport,history_id).reviewed_by_name=='原教师'
            assert db.get(WeeklyReport,missing_id).status=='PENDING_REVIEW'


@pytest.mark.parametrize('many',[False,True])
def test_advisor_command_transfers_pending_todo_but_preserves_completed_history(school,many):
    from datetime import datetime
    from app.models import InternshipBatch,InternshipRecord,WeeklyReport,UnifiedTodo,InternshipApplication
    with login(school)[0] as owner:
        import_rows(owner,'staff',[[n,n,'C01','工程学院','','','','',''] for n in ['T001','T002']])
        import_rows(owner,'students',[['S001','交接学生','C01','工程学院','M01','软件技术','B01','软件一班','2024']])
        accounts={r['loginName']:int(r['id']) for r in ok(owner.get(ROOT+'/accounts'))['items']}
        with get_sessionmaker()() as db:
            student=db.scalar(select(StudentProfile).where(StudentProfile.tenant_id==school['id']))
            batch=InternshipBatch(tenant_id=school['id'],batch_name='交接测试',batch_no='HANDOFF',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31));db.add(batch);db.flush()
            record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=student.id,status='ONBOARD',advisor_user_id=accounts['T001'],advisor_name='T001');db.add(record);db.flush()
            report=WeeklyReport(tenant_id=school['id'],internship_id=record.id,week_number=1,status='PENDING_REVIEW');db.add(report);db.flush()
            original=UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type='INTERN_WEEKLY',source_biz_id=report.id,todo_type='INTERN_WEEKLY_REVIEW',assignee_id=accounts['T001'],student_id=student.id,title='周报待批',status='PENDING');db.add(original)
            application=InternshipApplication(tenant_id=school['id'],record_id=record.id,student_id=student.id,application_type='SELF_ARRANGED',status='PENDING');db.add(application);db.flush()
            db.add(UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type='INTERN_APPLICATION',source_biz_id=application.id,todo_type='INTERN_APPLICATION_REVIEW',assignee_id=accounts['T001'],student_id=student.id,title='自主实习申请待审',status='PENDING'))
            if many:
                for number in range(3,34):
                    extra=WeeklyReport(tenant_id=school['id'],internship_id=record.id,week_number=number,status='PENDING_REVIEW');db.add(extra);db.flush()
                    db.add(UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type='INTERN_WEEKLY',source_biz_id=extra.id,todo_type='INTERN_WEEKLY_REVIEW',assignee_id=accounts['T001'],student_id=student.id,title=f'分页周报{number}',status='PENDING'))
                db.add(UnifiedTodo(tenant_id=school['id'],source_module='internship',source_biz_type='UNKNOWN',source_biz_id=999999999,
                    todo_type='UNKNOWN',assignee_id=accounts['T001'],student_id=student.id,title='无法定位原单的历史待办',status='PENDING'))
            history=WeeklyReport(tenant_id=school['id'],internship_id=record.id,week_number=2,status='APPROVED',reviewed_by_name='T001');db.add(history);db.commit();rid=str(record.id);version=record.version;report_id=report.id;history_id=history.id
        old=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T001')
        ok(owner.post(ROOT+'/accounts/'+old['id']+'/actions',json={'action':'DISABLE','expectedVersion':old['version'],'reason':'原导师离任后交接待办'}))
        first=ok(owner.get(ROOT+'/scopes'))['blockedTasks']
        assert first['total']==(34 if many else 2)
        all_items=first['items']
        if many:
            second=ok(owner.get(ROOT+'/scopes',params={'blockedPage':2}))['blockedTasks']
            assert len(first['items'])==30 and len(second['items'])==4
            all_items+=second['items']
            assert len({r['id'] for r in all_items})==34
            missing=next(r for r in all_items if r['objectId']=='999999999')
            assert not missing['path'] and missing['actionNote']
        assert all(r['path']==f'/admin/internship/students?batchId={batch.id}&advisorId={rid}' for r in all_items if r['objectId']!='999999999')
        # School configuration identity is separate from the internship management identity.
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SCHOOL_ADMIN')
        ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':str(school['owner']),'enabled':True,
            'expectedVersion':role['version'],'reason':'由正式角色入口授权测试实习办理身份'}))
        activate(owner,'SCHOOL_ADMIN')
        payload={'advisorUserId':str(accounts['T002']),'reason':'原导师交接在办周报','expectedVersion':version}
        ok(owner.post('/api/v1/internship/intern-students/'+rid+'/advisor',json=payload))
        assert owner.post('/api/v1/internship/intern-students/'+rid+'/advisor',json=payload).status_code==409
        activate(owner,'SYSTEM_ADMIN')
        assert ok(owner.get(ROOT+'/scopes'))['blockedTasks']['total']==(1 if many else 0)
        with get_sessionmaker()() as db:
            items=db.scalars(select(UnifiedTodo).where(UnifiedTodo.tenant_id==school['id'],UnifiedTodo.source_biz_type=='INTERN_WEEKLY',UnifiedTodo.source_biz_id==report_id)).all()
            assert {(r.assignee_id,r.status) for r in items}=={(accounts['T001'],'CANCELLED'),(accounts['T002'],'PENDING')}
            applications=db.scalars(select(UnifiedTodo).where(UnifiedTodo.tenant_id==school['id'],UnifiedTodo.source_biz_type=='INTERN_APPLICATION')).all()
            assert {(r.assignee_id,r.status) for r in applications}=={(accounts['T001'],'CANCELLED'),(accounts['T002'],'PENDING')}
            assert db.get(WeeklyReport,report_id).status=='PENDING_REVIEW'
            assert db.get(WeeklyReport,history_id).reviewed_by_name=='T001'


@pytest.mark.parametrize('with_handoff',[False,'disabled','scope','membership','permission'])
def test_three_stage_leave_frozen_policy_return_resubmit_and_final_attendance(school,with_handoff):
    from datetime import datetime
    from app.models import InternshipBatch, InternshipRecord, InternshipCheckin, UnifiedTodo
    with login(school)[0] as owner:
        _,creds=import_rows(owner,'staff',[[n,label,'C01','工程学院','','','','',''] for n,label in [('T001','指导教师'),('T002','院系责任人'),('T003','校级责任人'),('T004','交接责任人')]])
        _,students=import_rows(owner,'students',[['S001','实习学生','C01','工程学院','M01','软件技术','B01','软件一班','2024']])
        accounts={r['loginName']:r for r in ok(owner.get(ROOT+'/accounts'))['items']}
        roles={r['code']:r for r in ok(owner.get(ROOT+'/roles'))['items']}
        for username,code in [('T002','COLLEGE_ADMIN'),('T003','SCHOOL_ADMIN')]:
            role=roles[code]
            ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':accounts[username]['id'],'enabled':True,'expectedVersion':role['version'],'reason':'合成审批任职'}))
        college=ok(owner.get(ROOT+'/organizations/colleges'))['items'][0]
        ok(owner.post(ROOT+'/scopes',json={'userId':accounts['T002']['id'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':college['id'],'reason':'合成院系责任关系'}))
        with get_sessionmaker()() as db:
            student=db.scalar(select(StudentProfile).where(StudentProfile.tenant_id==school['id'],StudentProfile.student_no=='S001'))
            batch=InternshipBatch(tenant_id=school['id'],batch_name='多级审批合成批次',batch_no='LEAVE-TEST',status='RUNNING',start_date=datetime(2026,9,1),end_date=datetime(2026,12,31))
            db.add(batch);db.flush()
            record=InternshipRecord(tenant_id=school['id'],batch_id=batch.id,student_id=student.id,status='ONBOARD',eligibility_status='QUALIFIED',
                advisor_user_id=int(accounts['T001']['id']),advisor_name='指导教师',intern_start_date=datetime(2026,9,1),intern_end_date=datetime(2026,12,31))
            db.add(record);db.commit();bid=str(batch.id);rid=str(record.id)
        nodes=[{'type':kind,'roleCode':code,'userId':''} for kind,code in [('GUIDE','INTERN_MENTOR'),('COLLEGE','COLLEGE_ADMIN'),('SCHOOL','SCHOOL_ADMIN')]]
        policy=ok(owner.get(ROOT+'/leave-policies/'+bid))
        assert ok(owner.post(ROOT+'/leave-policies/'+bid+'/preview',json={'nodes':nodes}))['canPublish']
        published=ok(owner.post(ROOT+'/leave-policies/'+bid+'/publish',json={'nodes':nodes,'expectedVersion':policy['expectedVersion'],'reason':'确认三级顺序审批'}))
        conflict=owner.post(ROOT+'/leave-policies/'+bid+'/publish',json={'nodes':nodes[:1],'expectedVersion':policy['expectedVersion'],'reason':'并发旧版本不得覆盖'})
        assert conflict.status_code==409
        assert conflict.json()['details']=={'version':published['expectedVersion'],'current':{'nodes':nodes}}
        with reset_for_test(owner,school,'S001',students['S001'],True) as student, reset_for_test(owner,school,'T001',creds['T001']) as mentor, reset_for_test(owner,school,'T002',creds['T002']) as college_actor, reset_for_test(owner,school,'T003',creds['T003']) as school_actor:
            activate(college_actor,'COLLEGE_ADMIN');activate(school_actor,'SCHOOL_ADMIN')
            body={'batchId':bid,'internshipId':rid,'leaveType':'PERSONAL','startDate':'2026-10-10','endDate':'2026-10-10','reason':'合成验收：办理个人事项'}
            route='/api/v1/portal/internship/context/leaves'
            leave=ok(student.post(route,json=body));lid=leave['id']
            college_role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
            impact_path=ROOT+'/roles/'+college_role['id']+'/impact'
            waiting=ok(owner.get(impact_path,params={'userId':accounts['T002']['id']}))
            assert waiting['memberCount']==waiting['scopeCount']==waiting['waitingApprovalCount']==1 and waiting['pendingApprovalCount']==0
            assert waiting['items'][0]['objectId']==lid and waiting['items'][0]['statusLabel']=='等待前级'
            assert not ok(owner.get(impact_path,params={'userId':accounts['T001']['id']}))['items']
            # Changing the published policy cannot shorten the already submitted instance.
            ok(owner.post(ROOT+'/leave-policies/'+bid+'/publish',json={'nodes':nodes[:1],'expectedVersion':published['expectedVersion'],'reason':'后续申请改为一级，原单保留三级'}))
            def review_by(client,action='APPROVE'):
                current=ok(client.get('/api/v1/internship/leaves/'+lid,params={'batchId':bid}))
                path='/api/v1/mobile/teacher/internship/context/leaves/'+lid+'/review' if action=='REQUEST_CHANGES' else '/api/v1/internship/leaves/'+lid+'/review'
                return ok(client.post(path,params={'batchId':bid},json={'action':action,'comment':'材料需补充具体事由' if action!='APPROVE' else '本级核对通过','batchId':bid,'expectedVersion':current['version']}))
            assert review_by(mentor)['status']=='PENDING'
            ready=ok(owner.get(impact_path))
            assert ready['pendingApprovalCount']==1 and ready['waitingApprovalCount']==0
            assert ready['memberTodoCount']==1 and ready['items'][0]['assigneeName']=='院系责任人'
            original_scope=next(r for r in ok(owner.get(ROOT+'/scopes'))['items'] if r['userId']==accounts['T002']['id'])
            scope_preview=ok(owner.get(ROOT+'/scopes/'+original_scope['id']+'/impact'))
            assert scope_preview['roleImpact']['pendingApprovalCount']==1
            assert scope_preview['roleImpact']['items'][0]['objectId']==lid
            assert scope_preview['original']['refName']==college['name']
            assert mentor.get(ROOT+'/scopes/'+original_scope['id']+'/impact').status_code==403
            with get_sessionmaker()() as db:
                assert db.scalar(select(func.count()).select_from(InternshipCheckin).where(InternshipCheckin.internship_id==int(rid)))==0
                assert db.scalar(select(UnifiedTodo.assignee_id).where(UnifiedTodo.tenant_id==school['id'],UnifiedTodo.source_biz_id==int(lid),UnifiedTodo.status=='PENDING'))==int(accounts['T002']['id'])
            replacement=None
            if with_handoff:
                old=next(r for r in ok(owner.get(ROOT+'/accounts'))['items'] if r['loginName']=='T002')
                role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
                if with_handoff=='disabled':
                    ok(owner.post(ROOT+f'/accounts/{old["id"]}/actions',json={'action':'DISABLE','expectedVersion':old['version'],'reason':'紧急停用并交接责任'}))
                    account_audit=ok(owner.get(ROOT+'/audit',params={'keyword':'停用账号'}))['items'][0]
                    assert {'field':'status','label':'状态','before':'正常','after':'停用'} in account_audit['changes']
                elif with_handoff=='scope':
                    scope=next(r for r in ok(owner.get(ROOT+'/scopes'))['items'] if r['userId']==old['id'])
                    ok(owner.put(ROOT+'/scopes/'+scope['id'],json={'userId':old['id'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':college['id'],'status':'DISABLED','expectedVersion':scope['version'],'reason':'收回原院系负责范围'}))
                    scope_audit=ok(owner.get(ROOT+'/audit',params={'keyword':'调整责任范围'}))['items'][0]
                    assert {'field':'status','label':'状态','before':'正常','after':'停用'} in scope_audit['changes']
                elif with_handoff=='membership':
                    ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':old['id'],'enabled':False,'expectedVersion':role['version'],'reason':'撤销原院系审批身份'}))
                else:
                    ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':role['name'],'baseRole':role['baseRole'],'permissions':[],'status':'ACTIVE','expectedVersion':role['version'],'reason':'收回院系审批功能权限','sensitiveConfirmed':True}))
                current=ok(student.get(route+'/'+lid,params={'batchId':bid,'internshipId':rid}))
                assert current['approvalProgress']['steps'][1]['statusLabel']=='待重新指定责任人'
                blocked=ok(owner.get(ROOT+'/scopes'))['blockedTasks']
                assert blocked['total']==1 and blocked['items'][0]['objectId']==lid
                from urllib.parse import urlsplit,parse_qs
                target=parse_qs(urlsplit(blocked['items'][0]['path']).query)
                assert target['handoffBatchId']==[bid] and target['handoffSourceId']==[old['id']]
                exact=ok(owner.get(ROOT+'/handoffs',params={'batchId':bid,'userId':old['id'],'taskId':target['handoffTaskId'][0]}))
                assert exact['total']==1 and exact['items'][0]['leaveId']==lid
                assert ok(owner.get(ROOT+'/handoffs',params={'batchId':bid,'userId':accounts['T001']['id'],'taskId':target['handoffTaskId'][0]}))['total']==0
                role_impact=ok(owner.get(impact_path,params={'userId':old['id']}))
                assert role_impact['blockedApprovalCount']==1 and role_impact['items'][0]['statusLabel']=='待重新指定责任人'
                assert college_actor.post('/api/v1/internship/leaves/'+lid+'/review',params={'batchId':bid},json={'action':'APPROVE','comment':'撤权后不得继续办理','batchId':bid,'expectedVersion':current['version']}).status_code in (401,403)
                if with_handoff=='permission':
                    # Restore the role template for the replacement; the old actor's scope stays revoked.
                    scope=next(r for r in ok(owner.get(ROOT+'/scopes'))['items'] if r['userId']==old['id'])
                    ok(owner.put(ROOT+'/scopes/'+scope['id'],json={'userId':old['id'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':college['id'],'status':'DISABLED','expectedVersion':scope['version'],'reason':'交接前收回原负责范围'}))
                    current_role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
                    ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':role['name'],'baseRole':role['baseRole'],'permissions':['internship.leave.view','internship.leave.review'],'status':'ACTIVE','expectedVersion':current_role['version'],'reason':'恢复接收人所需审批权限','sensitiveConfirmed':True}))
                role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
                ok(owner.post(ROOT+f'/roles/{role["id"]}/members',json={'userId':accounts['T004']['id'],'enabled':True,'expectedVersion':role['version'],'reason':'接收院系审批责任'}))
                ok(owner.post(ROOT+'/scopes',json={'userId':accounts['T004']['id'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':college['id'],'reason':'接收同院系管理责任'}))
                task=ok(owner.get(ROOT+'/handoffs',params={'batchId':bid,'userId':old['id']}))['items'][0]
                assert task['statusLabel']=='待重新指定责任人'
                transfer={'batchId':bid,'userId':accounts['T004']['id'],'roleCode':'COLLEGE_ADMIN','expectedVersion':task['version'],'reason':'原人员停用后的真实任务交接'}
                ok(owner.post(ROOT+'/handoffs/'+task['id'],json=transfer))
                assert owner.post(ROOT+'/handoffs/'+task['id'],json=transfer).status_code==409
                assert ok(owner.get(ROOT+'/scopes'))['blockedTasks']['total']==0
                if with_handoff in ('disabled','membership'): assert college_actor.get('/api/v1/auth/me').status_code==401
                replacement=reset_for_test(owner,school,'T004',creds['T004']);activate(replacement,'COLLEGE_ADMIN');college_actor=replacement
            action='REQUEST_CHANGES' if with_handoff else 'REJECT'
            assert review_by(college_actor,action)['status']==('NEEDS_REVISION' if with_handoff else 'REJECTED')
            original=ok(student.get(route+'/'+lid,params={'batchId':bid,'internshipId':rid}))
            assert original['approvalProgress']['steps'][1]['statusLabel']==('已退回补正' if with_handoff else '已驳回')
            ok(student.post(route+'/'+lid+'/resubmit',json={**body,'reason':'已补充个人事项具体原因','expectedVersion':original['version']}))
            restarted=ok(student.get(route+'/'+lid,params={'batchId':bid,'internshipId':rid}))
            assert [s['statusLabel'] for s in restarted['approvalProgress']['steps']]==['待审批','等待前级','等待前级']
            assert restarted['submissionHistory'][0]['reviewComment']=='材料需补充具体事由'
            assert review_by(mentor)['status']=='PENDING'
            assert review_by(college_actor)['status']=='PENDING'
            assert review_by(school_actor)['status']=='APPROVED'
            for surface in [route,'/api/v1/mobile/internship/context/leaves']:
                final=ok(student.get(surface+'/'+lid,params={'batchId':bid,'internshipId':rid}))
                assert len(final['approvalProgress']['steps'])==3
                assert all(s['statusLabel']=='已通过' for s in final['approvalProgress']['steps'])
            with get_sessionmaker()() as db:
                assert db.scalar(select(func.count()).select_from(InternshipCheckin).where(InternshipCheckin.internship_id==int(rid),InternshipCheckin.result=='LEAVE'))==1
            if replacement: replacement.close()


def test_dictionary_provenance_roundtrip_audit_conflict_and_older_client_preservation(school):
    with login(school)[0] as client:
        path=ROOT+'/config/dictionary'
        initial=ok(client.get(path));assert initial['standardReferences']==[]
        reference={'kind':'major','source':'合成学校实际使用目录','edition':'合成测试版一'}
        body={'announcementCategories':['实习通知'],'aliases':[],'standardReferences':[reference],
              'expectedVersion':initial['version'],'reason':'合成来源版本验收'}
        saved=ok(client.put(path,json=body))
        assert ok(client.get(path))['standardReferences']==[reference]
        assert client.put(path,json=body).status_code==409
        body.pop('standardReferences');body['expectedVersion']=saved['version']
        ok(client.put(path,json=body))
        assert ok(client.get(path))['standardReferences']==[reference]
        audit_rows=ok(client.get(ROOT+'/audit',params={'keyword':'修改学校配置'}))['items']
        assert any('合成测试版一' in str(row['changes']) for row in audit_rows)
        student_result,credentials=import_rows(client,'students',[['000001','合成学生','C01','学院','M01','专业','B01','班级','2024']])
        student=reset_for_test(client,school,'000001',credentials['000001'],True)
        assert student.get(path).status_code==403
        assert student.put(path,json=body).status_code==403
        student.close()


def test_directory_export_5000_boundary_keeps_text_ids_formula_safety_and_no_partial_job(school):
    from app.models import ExcelImportJob
    with login(school)[0] as owner:
        college=org(owner,'colleges','容量合成学院','CAPACITY')
        major=org(owner,'majors','容量合成专业','CAP-M',college['id'])
        classroom=org(owner,'classes','容量合成班级','CAP-B',major['id'])
        with get_sessionmaker()() as db:
            db.add_all([StudentProfile(tenant_id=school['id'],student_no=f'00000000000000{i:06d}',
                real_name='=1+1' if i==0 else f'容量合成学生{i}',college_id=int(college['id']),
                major_id=int(major['id']),class_id=int(classroom['id']),grade='2024') for i in range(5001)])
            db.commit()
        before=ok(owner.get(ROOT+'/tasks'))['total']
        rejected=owner.get(ROOT+'/exports/student')
        assert rejected.status_code==400 and '5000' in rejected.json()['message']
        assert ok(owner.get(ROOT+'/tasks'))['total']==before
        # Match exactly 5000 of the 5001 rows using the same bound account filter as the page.
        with get_sessionmaker()() as db:
            last=db.scalar(select(StudentProfile).where(StudentProfile.tenant_id==school['id'],StudentProfile.student_no=='00000000000000005000'))
            assert last is not None
            account=User(tenant_id=school['id'],login_name='CAP-BOUND',real_name='容量合成绑定学生',user_type='STUDENT',
                password_hash=hash_password(PASSWORD),status='ACTIVE',must_change_password=True)
            db.add(account);db.flush()
            db.add(StudentAccountLink(tenant_id=school['id'],student_id=last.id,user_id=account.id,link_status='ACTIVE'))
            db.commit()
        params={'accountStatus':'UNBOUND'}
        assert ok(owner.get(ROOT+'/people/student',params=params))['total']==5000
        response=owner.get(ROOT+'/exports/student',params=params);assert response.status_code==200
        book=load_workbook(io.BytesIO(response.content),read_only=True,data_only=False)
        rows=list(book.active.iter_rows());book.close()
        assert len(rows)==5001
        assert rows[1][0].value=='00000000000000000000' and rows[1][0].data_type=='s'
        assert rows[1][1].value=="'=1+1" and rows[1][1].data_type=='s'
        assert len({r[0].value for r in rows[1:]})==5000
        assert rows[-1][0].value=='00000000000000004999'
        task=ok(owner.get(ROOT+'/tasks'))['items'][0]
        assert task['total']==task['succeeded']==5000 and task['canDownload']
        result=owner.get(ROOT+'/tasks/'+task['id']+'/result');assert result.status_code==200
        book=load_workbook(io.BytesIO(result.content),read_only=True,data_only=False)
        assert len(list(book.active.iter_rows()))==5001;book.close()
        assert owner.get('/api/v1/system/exports/student',params={'keyword':'not-a-real-student'}).status_code==200


def test_large_directory_export_queue_retry_snapshot_and_revoked_file_access(school,monkeypatch):
    from app.models import ExcelImportJob,FileObject
    from app.services import system_management as management
    from app.services.system_import_runner import process_next
    owner=login(school)[0]
    try:
        with get_sessionmaker()() as db:
            db.add_all([StudentProfile(tenant_id=school['id'],student_no=f'00000000000000{i:06d}',
                real_name='=1+1' if i==0 else f'大任务合成学生{i}') for i in range(5001)])
            db.commit()
        task=ok(owner.post(ROOT+'/exports/student',json={'accountStatus':'UNBOUND'}));identifier=int(task['id'])
        assert task['total']==5001 and task['status']=='QUEUED'
        assert owner.get(ROOT+f'/tasks/{identifier}/result').status_code==409
        original=management.directory_rows;calls=0
        def failed_page(*args):
            nonlocal calls
            calls+=1
            if calls==2:raise RuntimeError('synthetic export interruption')
            return original(*args)
        monkeypatch.setattr(management,'directory_rows',failed_page)
        assert process_next(identifier)
        failed=next(r for r in ok(owner.get(ROOT+'/tasks'))['items'] if r['id']==str(identifier))
        assert failed['status']=='FAILED' and failed['canRetry'] and not failed['canDownload']
        with get_sessionmaker()() as db:
            assert not db.get(ExcelImportJob,identifier).preview_json.get('fileId')
            assert db.scalar(select(func.count()).select_from(FileObject).where(FileObject.tenant_id==school['id'],FileObject.biz_type=='SYSTEM_DIRECTORY_EXPORT'))==0
        monkeypatch.setattr(management,'directory_rows',original)
        ok(owner.post(ROOT+f'/tasks/{identifier}/retry-export',json={'expectedVersion':failed['version']}))
        assert owner.post(ROOT+f'/tasks/{identifier}/retry-export',json={'expectedVersion':failed['version']}).status_code==409
        with ThreadPoolExecutor(max_workers=2) as executor:
            assert any(executor.map(process_next,[identifier,identifier]))
        assert not process_next(identifier)
        completed=next(r for r in ok(owner.get(ROOT+'/tasks'))['items'] if r['id']==str(identifier))
        assert completed['total']==completed['succeeded']==5001 and completed['canDownload'] and not completed['canRetry']
        downloaded=owner.get(ROOT+f'/tasks/{identifier}/result');assert downloaded.status_code==200
        book=load_workbook(io.BytesIO(downloaded.content),read_only=True,data_only=False)
        rows=list(book.active.iter_rows());book.close()
        assert len(rows)==5002 and rows[1][0].value=='00000000000000000000'
        assert rows[-1][0].value=='00000000000000005000' and rows[1][1].value=="'=1+1"
        assert len({r[0].value for r in rows[1:]})==5001
        with get_sessionmaker()() as db:
            job=db.get(ExcelImportJob,identifier);file_id=job.preview_json['fileId']
            assert 'rows' not in job.preview_json
            assert db.scalar(select(func.count()).select_from(FileObject).where(FileObject.tenant_id==school['id'],
                FileObject.biz_type=='SYSTEM_DIRECTORY_EXPORT',FileObject.biz_id==str(identifier)))==1
        assert owner.get('/api/v1/files/download/'+file_id).status_code==200
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='SYSTEM_ADMIN')
        ok(owner.put(ROOT+'/roles/'+role['id'],json={'name':role['name'],'baseRole':'SYSTEM_ADMIN',
            'permissions':[p for p in role['permissions'] if p!='system.person.export'],
            'expectedVersion':role['version'],'sensitiveConfirmed':True,'reason':'合成任务生成后撤回导出权限'}))
        assert owner.get(ROOT+f'/tasks/{identifier}/result').status_code==403
        assert owner.get('/api/v1/files/download/'+file_id).status_code in (403,404)
        assert owner.get('/api/v1/files/preview/'+file_id).status_code in (403,404)
    finally:owner.close()


def test_queued_directory_export_checks_changed_scope_at_execution_and_every_file_entry(school):
    from app.models import ExcelImportJob
    from app.services.system_import_runner import process_next
    owner=login(school)[0];actor=None
    try:
        _,credentials=import_rows(owner,'staff',[['T001','原院导出合成教师','C01','原院','','','','','']])
        import_rows(owner,'students',[
            ['S001','待调整合成学生','C01','原院','M01','原专业','B01','原班','2024'],
            ['S002','新院合成学生','C02','新院','M02','新专业','B02','新班','2024']])
        account=ok(owner.get(ROOT+'/accounts',params={'keyword':'T001'}))['items'][0]
        students={r['number']:r for r in ok(owner.get(ROOT+'/people/student'))['items']}
        role=next(r for r in ok(owner.get(ROOT+'/roles'))['items'] if r['code']=='COLLEGE_ADMIN')
        ok(owner.post(ROOT+'/roles/'+role['id']+'/members',json={'userId':account['id'],'enabled':True,'expectedVersion':role['version'],'reason':'合成院系导出前置'}))
        ok(owner.post(ROOT+'/scopes',json={'userId':account['id'],'roleCode':'COLLEGE_ADMIN','scopeType':'COLLEGE','refId':students['S001']['collegeId'],'reason':'原院合成范围'}))
        actor=reset_for_test(owner,school,'T001',credentials['T001']);activate(actor,'COLLEGE_ADMIN')
        done=ok(actor.post(ROOT+'/exports/student',json={}));assert process_next(int(done['id']))
        assert actor.get(ROOT+'/tasks/'+done['id']+'/result').status_code==200
        assert owner.get(ROOT+'/tasks/'+done['id']+'/result').status_code==403
        pending=ok(actor.post(ROOT+'/exports/student',json={}))
        with get_sessionmaker()() as db:file_id=db.get(ExcelImportJob,int(done['id'])).preview_json['fileId']
        student=students['S001']
        ok(owner.put(ROOT+'/people/student/'+student['id'],json={'name':student['name'],'collegeId':students['S002']['collegeId'],
            'classId':students['S002']['classId'],'grade':'2024','expectedVersion':student['version'],'confirmImpact':True,'reason':'生成与领取之间跨院调整'}))
        assert process_next(int(pending['id']))
        failed=next(r for r in ok(actor.get(ROOT+'/tasks'))['items'] if r['id']==pending['id'])
        assert failed['status']=='FAILED' and not failed['canDownload']
        assert actor.post(ROOT+'/tasks/'+pending['id']+'/retry-export',json={'expectedVersion':failed['version']}).status_code==403
        assert actor.get(ROOT+'/tasks/'+done['id']+'/result').status_code==403
        for suffix in [file_id,'preview/'+file_id,'download/'+file_id,file_id+'/url']:
            assert actor.get('/api/v1/files/'+suffix).status_code in (403,404)
        assert ok(actor.get(ROOT+'/people/student'))['total']==0
    finally:
        if actor:actor.close()
        owner.close()


def test_integration_reads_actual_receipt_time_and_isolates_school_secrets(school):
    from datetime import datetime
    from app.models.file import FileScanRecord
    from app.models.message import MessageChannelDelivery, UnifiedMessage
    stamp=datetime(2026,1,1)
    with get_sessionmaker()() as db:
        db.add(UnifiedMessage(tenant_id=school['id'],receiver_id=school['owner'],receiver_user_id=school['owner'],
            title='合成投递记录',delivery_status='DELIVERED',delivered_at=stamp))
        db.add(FileScanRecord(tenant_id=school['id'],file_id=1,attempt=1,result='ERROR',
            started_at=stamp,completed_at=stamp,error_code='ClamAVUnavailable',error_message='private-token'))
        db.add(MessageChannelDelivery(tenant_id=school['id'],campaign_id=1,channel='SMS',receiver_user_id=school['owner'],
            status='DEAD',attempt_count=1,last_error_code='SMS_PROVIDER_NOT_CONFIGURED',last_error_message_safe='private-token'))
        other=Tenant(tenant_code='PROBE-'+uuid4().hex[:12],school_name='另一合成学校',status='ACTIVE')
        db.add(other);db.flush()
        db.add(FileScanRecord(tenant_id=other.id,file_id=1,attempt=1,result='CLEAN',started_at=stamp,completed_at=datetime(2026,2,1)))
        db.commit()
    with login(school)[0] as client:
        data=ok(client.get(ROOT+'/integrations'))
        rows={row['name']:row for row in data['items']}
        assert rows['文件安全扫描']['lastCheckedAt'].startswith('2026-01-01')
        assert rows['文件安全扫描']['failureReason']=='扫描服务不可用或连接超时'
        assert rows['短信服务']['failureReason']=='短信服务商尚未配置'
        assert rows['站内消息']['statusLabel']=='最近站内消息已送达'
        assert rows['站内消息']['lastCheckedAt'].startswith('2026-01-01')
        assert rows['地图与位置']['lastCheckedAt'] is None
        assert 'private-token' not in str(data)
