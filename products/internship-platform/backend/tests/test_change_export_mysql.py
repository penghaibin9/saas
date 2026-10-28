"""F09: normal-role export against new synthetic rows; no retained data cleanup."""
import base64
from datetime import datetime
from io import BytesIO
import os

import pytest
from openpyxl import load_workbook
from sqlalchemy import select, text
from sqlalchemy.engine import make_url

from app.config import settings
from app.db.session import get_engine, get_sessionmaker
from app.models import (College, InternshipAuditTrail, InternshipChangeRequest, InternshipRecord,
                        Role, StudentProfile, TeacherStudentScope, User, UserRole)
from app.services.system_authorization import bootstrap_school
from test_yiyang_plan_export_options_mysql import login, seed_workspace

ROOT = '/api/v1/internship/change-requests'


def data(response):
    assert response.status_code == 200 and response.json()['code'] == 0, response.text
    return response.json()['data']


def ledger(response):
    receipt = data(response)
    sheet = load_workbook(BytesIO(base64.b64decode(receipt['contentBase64']))).active
    rows = list(sheet.iter_rows(min_row=3, values_only=True))
    assert receipt['rowCount'] == len(rows)
    return [dict(zip([cell.value for cell in sheet[2]], row)) for row in rows]


@pytest.fixture(scope='module')
def workspace():
    if os.environ.get('GAP09_MYSQL_ACCEPTANCE') != '1':
        pytest.skip('requires explicit isolated MySQL opt-in')
    url = make_url(settings.DATABASE_URL)
    assert settings.APP_ENV == 'test' and url.get_backend_name() == 'mysql'
    assert (url.host, url.port, url.database) in {
        ('127.0.0.1', 3318, 'internship_pr275_test_20261003'),
        ('127.0.0.1', 3306, 'internship_employment_test_ci'),
    }
    with get_engine().connect() as db:
        assert db.scalar(text('SELECT DATABASE()')) == url.database
    def retained():
        with get_sessionmaker()() as db:
            change = db.get(InternshipChangeRequest, 1)
            return None if not change else (change.tenant_id, change.internship_id,
                change.status, change.version, change.review_comment, change.reviewed_at)
    before = retained()
    ws = seed_workspace(2)
    tid, bid = int(ws['tenantId']), int(ws['batches'][0]['id'])
    with get_sessionmaker()() as db:
        bootstrap_school(db, tid)
        colleges = [College(tenant_id=tid, college_name='同名验收院系', code=f'F09-{index}', status='ACTIVE') for index in range(2)]
        db.add_all(colleges); db.flush()
        role = db.scalar(select(Role).where(Role.tenant_id == tid, Role.role_code == 'COLLEGE_ADMIN', Role.is_deleted.is_(False)))
        reader = User(tenant_id=tid, login_name='college_admin', real_name='变更台账院系管理员',
            password_hash=db.get(User, ws['users']['INTERN_MENTOR']).password_hash,
            user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
        db.add(reader); db.flush()
        db.add(UserRole(tenant_id=tid, user_id=reader.id, role_id=role.id, status='ACTIVE'))
        db.add(TeacherStudentScope(tenant_id=tid, teacher_key='college_admin', user_id=reader.id,
            role_code='COLLEGE_ADMIN', scope_type='COLLEGE', ref_id=colleges[0].id, status='ACTIVE'))
        ws['users']['COLLEGE_ADMIN'] = reader.id
        records = db.scalars(select(InternshipRecord).where(InternshipRecord.tenant_id == tid).order_by(InternshipRecord.id)).all()
        student = db.get(StudentProfile, records[0].student_id)
        student.real_name, student.student_no, student.college_id = '=合成%筛选', '000000000000000000123', colleges[0].id
        outsider = StudentProfile(tenant_id=tid, real_name='=合成%筛选', student_no='000000000000000000124',
            college_id=colleges[1].id, current_stage='ENROLLED', student_status='NORMAL', status='ACTIVE')
        db.add(outsider); db.flush()
        outside_record = InternshipRecord(tenant_id=tid, student_id=outsider.id, batch_id=bid,
            advisor_user_id=ws['users']['SCHOOL_ADMIN'], status='READY')
        db.add(outside_record); db.flush()
        records[0].enterprise_name, records[0].position_name = '后来当前单位', '后来当前岗位'
        own = []
        for index, status in enumerate(['REJECTED'] * 12 + ['APPROVED', 'PENDING']):
            change = InternshipChangeRequest(tenant_id=tid, internship_id=records[0].id, student_id=student.id,
                change_type='SELF_ARRANGED', status=status, reason='@仅用于合成导出台账',
                target_enterprise_name='申请目标单位', target_position_name='申请目标岗位', record_version_snapshot=1,
                created_at=datetime(2026, 10, 12, 23, 35), reviewed_at=None if status == 'PENDING' else datetime(2026, 10, 13, 0, 1),
                reviewed_by_name=None if status == 'PENDING' else '合成审核人', review_comment='+合成意见')
            db.add(change); db.flush(); own.append(str(change.id))
            if status == 'APPROVED':
                db.add(InternshipAuditTrail(tenant_id=tid, target_id=change.id, target_type='CHANGE_REQ',
                    action='APPROVE', operator_name='合成审核人', occurred_at=datetime(2026, 10, 13, 0, 1),
                    detail_json={'atomic': True, 'recordBefore': {'recordVersion': 1, 'enterpriseId': '',
                        'positionId': '', 'enterpriseName': '审批前单位', 'positionName': '审批前岗位'},
                        'recordAfter': {'recordVersion': 2, 'enterpriseId': '', 'positionId': '',
                        'enterpriseName': '本次生效单位', 'positionName': '本次生效岗位'}}))
        excluded = []
        for record in [outside_record, records[1]]:
            change = InternshipChangeRequest(tenant_id=tid, internship_id=record.id, student_id=record.student_id,
                change_type='WITHDRAW_POST', reason='同名院系或其他批次合成前置', status='PENDING')
            db.add(change); db.flush(); excluded.append(str(change.id))
        db.commit()
        ws.update(batchId=str(bid), own=own, excluded=excluded)
    yield ws
    if before is not None:
        assert retained() == before


def test_change_export_matches_all_filtered_pages_and_safe_chinese_cells(workspace):
    params = {'batchId': workspace['batchId'], 'status': 'REVIEWED', 'keyword': '%筛选'}
    with login(workspace) as school:
        first = data(school.get(ROOT, params={**params, 'pageSize': 10}))
        second = data(school.get(ROOT, params={**params, 'pageSize': 10, 'page': 2}))
        assert first['total'] == second['total'] == 13
        expected = {row['id'] for row in first['items'] + second['items']}
        rows = ledger(school.post(ROOT+'/export', params=params))
        assert {row['申请编号'] for row in rows} == expected == set(workspace['own'][:-1])
        for row in rows:
            assert row['学号'] == '000000000000000000123' and row['姓名'] == "'=合成%筛选"
            assert row['申请原因'] == "'@仅用于合成导出台账" and row['审核意见'] == "'+合成意见"
            assert row['申请时间'] == '2026-10-13 07:35' and row['审核时间'] == '2026-10-13 08:01'
            assert row['变更类型'] == '转自主实习' and row['状态'] in ('已通过', '已驳回')
        approved = next(row for row in rows if row['状态'] == '已通过')
        assert approved['本次变更前单位'] == '审批前单位' and approved['本次变更后单位'] == '本次生效单位'
        assert approved['当前实习单位'] == '后来当前单位' and approved['申请目标单位'] == '申请目标单位'
        assert not ledger(school.post(ROOT+'/export', params={**params, 'keyword': '不存在的学生'}))
    with get_sessionmaker()() as db:
        audits = db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id == int(workspace['tenantId']),
            InternshipAuditTrail.target_type == 'EXPORT', InternshipAuditTrail.action == 'EXPORT').order_by(InternshipAuditTrail.id)).all()
        detail = audits[0].detail_json
        assert detail['ledger'] == '实习变更台账' and detail['totalRows'] == 13
        assert detail['filters'] == params and detail['actorUserId'] == 'db-'+str(workspace['users']['SCHOOL_ADMIN'])


@pytest.mark.parametrize('role', ['INTERN_MENTOR', 'COLLEGE_ADMIN'])
def test_change_export_scope_matches_normal_teacher_and_same_name_college(workspace, role):
    params = {'batchId': workspace['batchId'], 'status': 'ALL'}
    with login(workspace, role) as reader:
        listed = data(reader.get(ROOT, params={**params, 'pageSize': 200}))
        rows = ledger(reader.post(ROOT+'/export', params=params))
        assert {row['申请编号'] for row in rows} == {row['id'] for row in listed['items']} == set(workspace['own'])
        assert not (set(workspace['excluded']) & {row['申请编号'] for row in rows})
        assert not ledger(reader.post(ROOT+'/export', params={**params, 'keyword': '000000000000000000124'}))
        assert not ledger(reader.post(ROOT+'/export', params={**params, 'batchId': workspace['batches'][1]['id']}))


def test_change_export_rejects_student_foreign_batch_and_invalid_filters(workspace):
    params = {'batchId': workspace['batchId']}
    foreign = seed_workspace(1)
    with login(workspace, 'STUDENT') as student:
        denied = student.post(ROOT+'/export', params=params)
        assert denied.status_code == 403 and 'contentBase64' not in denied.text
    with login(foreign) as outsider:
        denied = outsider.post(ROOT+'/export', params=params)
        assert denied.status_code in (403, 404) and 'contentBase64' not in denied.text
    with login(workspace) as school:
        for method, suffix in [('get', ''), ('post', '/export')]:
            rejected = getattr(school, method)(ROOT+suffix, params={**params, 'status': 'UNKNOWN'})
            assert rejected.status_code == 400 and 'contentBase64' not in rejected.text
        assert school.post(ROOT+'/export').status_code == 422
