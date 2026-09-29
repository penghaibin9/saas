from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.exceptions import AppException
from app.modules.internship.services import internship_regulatory_reporting_service as svc


def test_c07_status_contract_is_exact():
    assert svc.TASK_STATUSES == (
        "GENERATED",
        "VALIDATED",
        "EXPORTED",
        "SUBMITTED_EXTERNAL",
        "RECEIPT_PENDING",
        "ACCEPTED",
        "REJECTED",
    )


def test_rp01_required_enum_and_cross_field_validation():
    row = {field["key"]: "X" for field in svc.baseline_definition("RP01")["fields"]}
    row.update({
        "gender": "男",
        "insured": "否",
        "agreementSigned": "否",
        "majorMatch": "是",
        "highRisk": "是",
        "nightOrOvertime": "否",
        "holidayInternship": "否",
        "filingStatus": "待备案",
        "missingDocumentExplanation": "",
        "highRiskType": "",
        "internshipStartDate": "2026-09-01",
        "internshipEndDate": "2026-08-31",
    })
    errors = svc.validate_row("RP01", row)
    codes = {(e["field"], e["code"]) for e in errors}
    assert ("missingDocumentExplanation", "CROSS_FIELD") in codes
    assert ("highRiskType", "CROSS_FIELD") in codes
    assert ("internshipEndDate", "CROSS_FIELD") in codes


def test_rp02_overseas_requires_country():
    row = {field["key"]: "X" for field in svc.baseline_definition("RP02")["fields"]}
    row.update({
        "majorMatch": "是",
        "agreementSigned": "是",
        "overseas": "是",
        "crossProvince": "未知",
        "workCountry": "",
        "internshipStartDate": "2026-09-01",
        "internshipEndDate": "2026-12-01",
    })
    errors = svc.validate_row("RP02", row)
    assert any(
        e["field"] == "workCountry" and e["code"] == "CROSS_FIELD"
        for e in errors
    )


def test_formal_xlsx_keeps_credit_code_and_student_no_as_text():
    definition = svc.baseline_definition("RP01")
    row = {field["key"]: "" for field in definition["fields"]}
    row["studentNo"] = "001234"
    row["companyCreditCode"] = "012345678901234567"
    payload = svc.build_workbook("RP01", [row], definition)
    wb = load_workbook(BytesIO(payload))
    ws = wb["RP01"]
    labels = [c.value for c in ws[1]]
    credit_col = labels.index("统一社会信用代码") + 1
    student_col = labels.index("学号") + 1
    assert ws.cell(2, credit_col).value == "012345678901234567"
    assert ws.cell(2, credit_col).number_format == "@"
    assert ws.cell(2, student_col).value == "001234"
    assert ws.cell(2, student_col).number_format == "@"


def test_procurement_baseline_never_claims_official_template():
    for code in svc.REPORT_CODES:
        definition = svc.baseline_definition(code)
        assert definition["officialVerified"] is False
        assert "采购" in definition["templateName"]


def test_receipt_is_blocked_without_real_adapter(monkeypatch):
    monkeypatch.setattr(
        svc.settings,
        "REGULATORY_RECEIPT_ADAPTER_ENABLED",
        False,
        raising=False,
    )
    with pytest.raises(AppException) as exc:
        svc.record_external_receipt(
            1,
            {"status": "ACCEPTED", "receiptCode": "fake"},
        )
    assert "回执" in str(exc.value)


def test_g18_authoritative_teacher_term_and_rights_sources():
    class Batch:
        academic_year = "2025-2026"
        term = "第二学期"

    class SafePosition:
        night_shift = False
        overtime_allowed = False

    class NightPosition:
        night_shift = True
        overtime_allowed = False

    class UnknownPosition:
        night_shift = False
        overtime_allowed = None

    assert svc._batch_start_term(Batch()) == "2025-2026 第二学期"
    assert svc._night_or_overtime(SafePosition()) == "否"
    assert svc._night_or_overtime(NightPosition()) == "是"
    assert svc._night_or_overtime(UnknownPosition()) == "未知"
    assert svc._night_or_overtime(None) == "未知"

    rp01 = {field["key"]: field for field in svc.baseline_definition("RP01")["fields"]}
    rp02 = {field["key"]: field for field in svc.baseline_definition("RP02")["fields"]}
    assert "User.login_name" in rp01["advisorEmployeeNo"]["source"]
    assert "InternshipPosition.night_shift" in rp01["nightOrOvertime"]["source"]
    assert "InternshipBatch.academic_year" in rp02["startTerm"]["source"]


def test_g18_file_evidence_model_contract():
    from app.models import InternshipRegulatoryTask, InternshipRegulatoryTemplateVersion

    template_columns = set(InternshipRegulatoryTemplateVersion.__table__.c.keys())
    task_columns = set(InternshipRegulatoryTask.__table__.c.keys())
    assert {"source_file_id", "source_file_name", "source_file_sha256"} <= template_columns
    assert {"output_file_id", "output_sha256", "error_file_id", "error_sha256"} <= task_columns


def test_g18_school_confirmed_template_requires_real_source_file(monkeypatch):
    from app.services import file_service

    monkeypatch.setattr(
        file_service,
        "get_file_meta",
        lambda file_id, user=None: {
            "fileId": str(file_id),
            "fileName": "监管平台导入模板2026.xlsx",
            "ext": "xlsx",
            "sha256": "b" * 64,
            "readyForBusiness": True,
        },
    )
    meta = svc._regulatory_source_file_meta("901", user={"userId": "1"})
    assert meta["fileId"] == "901"
    assert meta["fileName"].endswith(".xlsx")
    assert meta["sha256"] == "b" * 64

    with pytest.raises(AppException):
        svc._regulatory_source_file_meta("", user={"userId": "1"})

    monkeypatch.setattr(
        file_service,
        "get_file_meta",
        lambda file_id, user=None: {
            "fileId": str(file_id),
            "fileName": "不是监管模板.pdf",
            "ext": "pdf",
            "sha256": "c" * 64,
            "readyForBusiness": True,
        },
    )
    with pytest.raises(AppException):
        svc._regulatory_source_file_meta("902", user={"userId": "1"})
