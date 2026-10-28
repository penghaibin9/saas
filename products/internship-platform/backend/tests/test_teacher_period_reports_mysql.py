"""教师本人周期报告：正常账号、同一正式事实、附件归属和真实导出。"""
from base64 import b64decode
from datetime import date
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import get_sessionmaker
from app.models import InternshipBatch, InternshipTeacherPeriodReport
from app.models.file import FileObject, FileBinding
from test_yiyang_write_flows_mysql import flow, login, seed_flow, upload, png_bytes

ROOT = '/api/v1/teacher-mobile/internship/activity/period-reports'
ADMIN = '/api/v1/internship/teacher-management/reports'


def payload(data, kind='WEEKLY', **extra):
    year, week, _ = date.today().isocalendar()
    key = f'{year}-W{week:02d}' if kind == 'WEEKLY' else date.today().strftime('%Y-%m')
    return {'batchId': data['batchId'], 'reportType': kind,
            'periodKey': 'SUMMARY' if kind == 'SUMMARY' else key,
            'content': '合成教师指导工作：巡访企业、核对安全记录、跟进学生问题。' * 14,
            'issueContent': '合成问题记录', 'nextPlan': '合成下期计划', 'studentCount': 1,
            'attachmentFileIds': [], **extra}


def rows(client, path, data):
    response = client.get(path, params={'batchId': data['batchId']})
    assert response.status_code == 200 and response.json()['code'] == 0, response.text
    return response.json()['data']


def test_teacher_cannot_take_student_private_file_as_own_report(flow):
    with login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        file = upload(student, biz='INTERNSHIP_REPORT')
        response = teacher.post(ROOT, json=payload(flow, attachmentFileIds=[file['fileId']]))
        assert response.status_code in (403, 404), response.text
        assert rows(teacher, ROOT, flow)['total'] == 0
        with get_sessionmaker()() as db:
            assert db.scalar(select(InternshipTeacherPeriodReport.id).where(
                InternshipTeacherPeriodReport.batch_id == int(flow['batchId']))) is None


@pytest.mark.parametrize('kind', ['WEEKLY', 'MONTHLY', 'SUMMARY'])
def test_teacher_submission_update_same_fact_admin_read_and_own_export(flow, kind):
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school, login(flow, 'STUDENT') as student:
        file = upload(teacher, name='teacher-evidence.png', biz='INTERNSHIP_REPORT')
        first = teacher.post(ROOT, json=payload(flow, kind, attachmentFileIds=[file['fileId']]))
        assert first.status_code == 200 and first.json()['code'] == 0, first.text
        row = first.json()['data']
        assert rows(teacher, ROOT, flow)['items'][0]['id'] == row['id']
        assert rows(school, ADMIN, flow)['items'][0]['content'] == row['content']
        assert teacher.get('/api/v1/files/download/'+file['fileId']).content == png_bytes()
        assert school.get('/api/v1/files/download/'+file['fileId']).content == png_bytes()
        assert student.get('/api/v1/files/download/'+file['fileId']).status_code == 404
        changed = payload(flow, kind, expectedVersion=row['version'],
                          content=row['content']+'补充实际办理结果。', attachmentFileIds=[file['fileId']])
        updated = teacher.post(ROOT, json=changed)
        assert updated.status_code == 200 and updated.json()['code'] == 0, updated.text
        assert updated.json()['data']['id'] == row['id']
        assert updated.json()['data']['version'] == row['version']+1
        stale = teacher.post(ROOT, json=changed)
        assert stale.status_code == 400 and stale.json()['bizCode'] == 'DATA_CONFLICT' and stale.json()['data'] is None, stale.text
        result = teacher.post(ROOT+'/export', params={'batchId': flow['batchId']})
        assert result.status_code == 200 and result.json()['code'] == 0, result.text
        exported = result.json()['data']
        workbook = load_workbook(BytesIO(b64decode(exported['contentBase64'])))
        cells = list(workbook.active.values)
        assert any(changed['content'] in values for values in cells)
        assert not any(kind in values for values in cells)


def test_cross_school_student_and_out_of_scope_batch_rejected(flow):
    foreign = seed_flow()
    with login(foreign, 'INTERN_MENTOR') as outsider, login(flow, 'STUDENT') as student, login(flow, 'INTERN_MENTOR') as teacher:
        assert outsider.post(ROOT, json=payload(flow)).status_code in (403, 404)
        assert student.post(ROOT, json=payload(flow)).status_code == 403
        assert teacher.post(ROOT, json=payload(flow, batchId=flow['batches'][1]['id'])).status_code == 403


def test_teacher_report_maximum_is_from_current_batch_and_enforced(flow):
    with get_sessionmaker()() as db:
        batch = db.get(InternshipBatch, int(flow['batchId']))
        batch.rules_config = {**(batch.rules_config or {}), 'teacherReport': {
            'weeklyMaxWords': 40, 'monthlyMaxWords': 120, 'summaryMaxWords': 320}}
        db.commit()
    with login(flow, 'INTERN_MENTOR') as teacher:
        response = teacher.post(ROOT, json=payload(flow, content='合成指导记录。'*10))
        assert response.status_code == 400 and response.json()['code'] != 0, response.text
        data = rows(teacher, ROOT, flow)
        assert data['rules']['weeklyMaxWords'] == 40 and data['total'] == 0


@pytest.mark.parametrize('condition', ['unsafe', 'deleted', 'another-business'])
def test_attachment_invalid_state_never_creates_report(flow, condition):
    with login(flow, 'INTERN_MENTOR') as teacher:
        file = upload(teacher, biz='INTERNSHIP_REPORT')
        with get_sessionmaker()() as db:
            row = db.get(FileObject, int(file['fileId']))
            if condition == 'unsafe': row.scan_status = 'PENDING'
            elif condition == 'deleted': row.is_deleted = True
            else: row.biz_id, row.biz_type = '999999', 'INTERNSHIP_INSURANCE_POLICY'
            db.commit()
        response = teacher.post(ROOT, json=payload(flow, attachmentFileIds=[file['fileId']]))
        assert response.status_code in (400, 404, 409), response.text
        assert rows(teacher, ROOT, flow)['total'] == 0


def test_removed_attachment_and_revoked_binding_do_not_use_owner_fallback(flow):
    with login(flow, 'INTERN_MENTOR') as teacher:
        file = upload(teacher, biz='INTERNSHIP_REPORT')
        row = teacher.post(ROOT, json=payload(flow, attachmentFileIds=[file['fileId']])).json()['data']
        assert teacher.get('/api/v1/files/download/'+file['fileId']).status_code == 200
        with get_sessionmaker()() as db:
            binding = db.scalar(select(FileBinding).where(FileBinding.file_id == int(file['fileId'])))
            binding.status = 'REVOKED'; db.commit()
        assert teacher.get('/api/v1/files/download/'+file['fileId']).status_code == 404
        changed = teacher.post(ROOT, json=payload(flow, expectedVersion=row['version'], attachmentFileIds=[]))
        assert changed.status_code == 200, changed.text
        assert teacher.get('/api/v1/files/download/'+file['fileId']).status_code == 404


def test_source_quarantine_keeps_report_visible_but_file_unavailable(flow):
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        file = upload(teacher, biz='INTERNSHIP_REPORT')
        row = teacher.post(ROOT, json=payload(flow, attachmentFileIds=[file['fileId']])).json()['data']
        with get_sessionmaker()() as db:
            stored = db.get(FileObject, int(file['fileId'])); stored.scan_status='INFECTED'; stored.status='QUARANTINED'; db.commit()
        for client, path in ((teacher, ROOT), (school, ADMIN)):
            read = rows(client, path, flow)['items'][0]
            assert read['id'] == row['id'] and read['attachments'][0]['canPreview'] is False
            assert client.get('/api/v1/files/download/'+file['fileId']).status_code == 404


@pytest.mark.parametrize('existing', [False, True])
def test_simultaneous_create_or_update_has_one_business_winner(flow, existing):
    from concurrent.futures import ThreadPoolExecutor
    with login(flow, 'INTERN_MENTOR') as first, login(flow, 'INTERN_MENTOR') as second:
        body = payload(flow)
        if existing:
            created = first.post(ROOT, json=body).json()['data']
            body.update(expectedVersion=created['version'], content=body['content']+'并发修改')
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures=[pool.submit(client.post, ROOT, json=body) for client in (first,second)]
            results=[f.result() for f in futures]
        assert sorted(r.status_code for r in results) == [200,400], [r.text for r in results]
        assert next(r for r in results if r.status_code == 400).json()['bizCode'] == 'DATA_CONFLICT'
        data=rows(first, ROOT, flow)
        assert data['total'] == 1 and data['items'][0]['version'] == (1 if existing else 0)


def test_report_export_contains_all_pages_and_current_chinese_type(flow):
    with login(flow, 'INTERN_MENTOR') as teacher, login(flow) as school:
        for week in range(1,25):
            response=teacher.post(ROOT, json=payload(flow, periodKey=f'2026-W{week:02d}'))
            assert response.status_code == 200, response.text
        data=rows(teacher, ROOT, flow)
        assert data['total'] == 24 and len(data['items']) == 20
        for client,path in ((teacher,ROOT),(school,ADMIN)):
            result=client.post(path+'/export',params={'batchId':flow['batchId']})
            assert result.status_code == 200 and result.json()['code'] == 0, result.text
            exported=result.json()['data']
            assert exported['rowCount'] == 24
            cells=list(load_workbook(BytesIO(b64decode(exported['contentBase64']))).active.values)
            assert sum('周报' in values for values in cells) == 24
            assert any('2026年第24周' in values for values in cells)


def test_teacher_cannot_read_or_export_schoolwide_reports_or_choose_another_actor(flow):
    with login(flow, 'INTERN_MENTOR') as teacher:
        assert teacher.get(ADMIN,params={'batchId':flow['batchId']}).status_code == 403
        assert teacher.post(ADMIN+'/export',params={'batchId':flow['batchId']}).status_code == 403
        assert teacher.post(ROOT,json=payload(flow,teacherUserId=flow['users']['SCHOOL_ADMIN'])).status_code == 422
