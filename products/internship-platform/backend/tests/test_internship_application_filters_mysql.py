"""Application filters: real role routes, isolated synthetic school, no fixture deletion."""
from io import BytesIO
import os
from urllib.parse import unquote

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import get_sessionmaker
from app.services.system_authorization import bootstrap_school
from app.models import College, Major, SchoolClass, StudentProfile, InternshipRecord, InternshipBatch, InternshipApplication, Role, User, UserRole, TeacherStudentScope
from test_yiyang_plan_export_options_mysql import seed_workspace, login

ROOT = '/api/v1/internship/applications'


def test_org_labels_keep_codes_and_distinguish_duplicate_or_missing_codes():
    from app.modules.internship.services.internship_application_service import _application_org_label
    first = _application_org_label('同名班级', 'CL01', '9007199254740993001')
    second = _application_org_label('同名班级', 'CL01', '9007199254740993002')
    assert first != second and 'CL01' in first and '同名班级' in first
    assert _application_org_label('同名班级', None, '9007199254740993001') == '同名班级（编号9007199254740993001）'


def data(response):
    assert response.status_code == 200, response.text
    assert response.json()['code'] == 0, response.text
    return response.json()['data']


@pytest.fixture(scope='module')
def workspace():
    if os.environ.get('GAP09_MYSQL_ACCEPTANCE') != '1':
        pytest.skip('requires isolated MySQL test opt-in')
    ws = seed_workspace(4)
    tid, bid = int(ws['tenantId']), int(ws['batches'][0]['id'])
    with get_sessionmaker()() as db:
        bootstrap_school(db, tid)
        db.get(InternshipBatch, bid).academic_year = '2026-2027'
        orgs = []
        for index in range(2):
            college = College(tenant_id=tid, college_name='同名院系', code=f'C{index}', status='ACTIVE')
            db.add(college); db.flush()
            major = Major(tenant_id=tid, college_id=college.id, major_name='同名专业', code=f'M{index}', status='ACTIVE')
            db.add(major); db.flush()
            school_class = SchoolClass(tenant_id=tid, major_id=major.id, class_name='同名班级', class_code=f'CL{index}', grade='2024', status='ACTIVE', class_status='NORMAL')
            db.add(school_class); db.flush()
            orgs.append(dict(collegeId=str(college.id), majorId=str(major.id), classId=str(school_class.id)))
        role = db.scalar(select(Role).where(Role.tenant_id == tid, Role.role_code == 'COLLEGE_ADMIN', Role.is_deleted.is_(False)))
        assert role is not None
        college_user = User(tenant_id=tid, login_name='college_admin', real_name='合成院系管理员',
            password_hash=db.get(User, ws['users']['INTERN_MENTOR']).password_hash,
            user_type='TEACHER', status='ACTIVE', credential_version=0, must_change_password=False)
        db.add(college_user); db.flush()
        db.add(UserRole(tenant_id=tid, user_id=college_user.id, role_id=role.id, status='ACTIVE'))
        db.add(TeacherStudentScope(tenant_id=tid, teacher_key='college_admin', user_id=college_user.id,
            role_code='COLLEGE_ADMIN', scope_type='COLLEGE', ref_id=int(orgs[0]['collegeId']), status='ACTIVE'))
        ws['users']['COLLEGE_ADMIN'] = college_user.id
        old_college = College(tenant_id=tid, college_name='已删除直挂院系', is_deleted=True)
        db.add(old_college); db.flush()
        old_major = Major(tenant_id=tid, college_id=old_college.id, major_name='已删除直挂专业', is_deleted=True)
        db.add(old_major); db.flush()
        records = db.scalars(select(InternshipRecord).where(InternshipRecord.tenant_id == tid).order_by(InternshipRecord.id)).all()
        apps = []
        for index, record in enumerate(records):
            record.batch_id = bid
            record.advisor_user_id = ws['users']['INTERN_MENTOR'] if index < 2 else ws['users']['SCHOOL_ADMIN']
            student = db.get(StudentProfile, record.student_id)
            student.class_id = int(orgs[1 if index == 2 else 0]['classId'])
            if index == 0:
                student.college_id, student.major_id = old_college.id, old_major.id
            student.grade = '2023' if index == 3 else None
            student.real_name = f'筛选学生{index}'
            student.student_no = f'0000000000000000000{index}'
            application = InternshipApplication(tenant_id=tid, record_id=record.id, student_id=student.id,
                batch_id=None if index == 0 else bid, application_type='EXEMPTION', volunteer_no=-1,
                exemption_type='MILITARY' if index == 2 else 'HEALTH', exemption_reason='合成只读筛选前置', status='PENDING_REVIEW')
            db.add(application); db.flush(); apps.append(str(application.id))
        db.commit()
    ws.update(batch=str(bid), orgs=orgs, apps=apps)
    return ws


def test_filter_combination_pagination_and_export_share_exact_scope(workspace):
    params = {'batchId': workspace['batch'], 'academicYear': '2026-2027', 'grade': '2024',
              **workspace['orgs'][0], 'studentName': '筛选学生', 'studentNo': '0000000000000000000',
              'applicationType': 'EXEMPTION', 'exemptionType': 'HEALTH', 'status': 'PENDING_REVIEW'}
    with login(workspace) as client:
        first = data(client.get(ROOT, params={**params, 'pageSize': 1}))
        second = data(client.get(ROOT, params={**params, 'pageSize': 1, 'page': 2}))
        assert first['total'] == second['total'] == 2
        assert {item['id'] for item in first['items'] + second['items']} == set(workspace['apps'][:2])
        assert all(item['collegeName'] == '同名院系' and item['majorName'] == '同名专业'
                   for item in first['items'] + second['items'])
        exported = client.get(ROOT + '/export.xlsx', params=params)
        assert exported.status_code == 200, exported.text
        disposition = exported.headers['content-disposition']
        assert disposition.isascii() and disposition.startswith("attachment; filename*=UTF-8''")
        assert unquote(disposition.split("UTF-8''", 1)[1]) == '实习申请审核台账.xlsx'
        assert exported.headers['cache-control'] == 'no-store'
        book = load_workbook(BytesIO(exported.content), read_only=True, data_only=True)
        exported_numbers = {str(value) for row in book.active.iter_rows(values_only=True) for value in row if str(value).startswith('0000000000000000000')}
        assert exported_numbers == {'00000000000000000000', '00000000000000000001'}
        for key, wrong in [('academicYear', '2025-2026'), ('grade', '2022'), ('collegeId', workspace['orgs'][1]['collegeId']),
                           ('majorId', workspace['orgs'][1]['majorId']), ('classId', workspace['orgs'][1]['classId']),
                           ('studentName', '其他'), ('studentNo', '%'), ('exemptionType', 'MILITARY')]:
            assert data(client.get(ROOT, params={**params, key: wrong}))['total'] == 0, key


def test_mobile_original_filter_keeps_teacher_scope_and_exact_application(workspace):
    path='/api/v1/mobile/teacher/internship/context/applications'
    params={'batchId':workspace['batch'],'status':'ALL'}
    with login(workspace,'INTERN_MENTOR') as teacher:
        for identifier in workspace['apps'][:2]:
            found=data(teacher.get(path,params={**params,'applicationId':identifier}))
            assert found['total']==1 and [row['id'] for row in found['items']]==[identifier]
        assert data(teacher.get(path,params={**params,'applicationId':workspace['apps'][2]}))['total']==0
        assert teacher.get(path,params={**params,'applicationId':'wrong'}).status_code==422
    with login(workspace) as school:
        assert data(school.get(path,params={**params,'applicationId':workspace['apps'][2]}))['total']==1


def test_options_and_queries_do_not_expand_teacher_or_school_scope(workspace):
    params = {'batchId': workspace['batch']}
    with login(workspace, 'INTERN_MENTOR') as teacher:
        options = data(teacher.get(ROOT + '/filter-options', params=params))
        assert options['academicYears'] == [{'value': '2026-2027', 'label': '2026-2027'}]
        assert options['grades'] == [{'value': '2024', 'label': '2024'}]
        for key, field in [('colleges', 'collegeId'), ('majors', 'majorId'), ('classes', 'classId')]:
            assert [item['value'] for item in options[key]] == [workspace['orgs'][0][field]]
        assert data(teacher.get(ROOT, params=params))['total'] == 2
        assert data(teacher.get(ROOT, params={**params, **workspace['orgs'][1]}))['total'] == 0
    with login(workspace) as school:
        options = data(school.get(ROOT + '/filter-options', params=params))
        assert len(options['colleges']) == len(options['majors']) == len(options['classes']) == 2
        for key, code in [('colleges', 'C'), ('majors', 'M'), ('classes', 'CL')]:
            assert len({item['label'] for item in options[key]}) == 2
            assert all(code in item['label'] and '已删除' not in item['label'] for item in options[key])
    with login(workspace, 'COLLEGE_ADMIN') as college:
        options = data(college.get(ROOT + '/filter-options', params=params))
        assert [item['value'] for item in options['colleges']] == [workspace['orgs'][0]['collegeId']]
        assert data(college.get(ROOT, params=params))['total'] == 3
        assert data(college.get(ROOT, params={**params, **workspace['orgs'][1]}))['total'] == 0
        for app_id in (workspace['apps'][0], workspace['apps'][1], workspace['apps'][3]):
            detail = data(college.get(ROOT + '/' + app_id))
            assert detail['id'] == app_id and detail['collegeName'] == '同名院系'
        assert college.get(ROOT + '/' + workspace['apps'][2]).status_code == 403
        body = {'action': 'REJECT', 'comment': '合成核对拒绝范围与旧版本', 'batchId': workspace['batch'], 'expectedVersion': 9999}
        rejected = college.post(ROOT + '/' + workspace['apps'][2] + '/review', json=body)
        assert rejected.status_code == 403, rejected.text
        conflict = college.post(ROOT + '/' + workspace['apps'][0] + '/review', json=body)
        assert conflict.status_code == 400 and conflict.json()['bizCode'] == 'DATA_CONFLICT', conflict.text
        with get_sessionmaker()() as db:
            for app_id in (workspace['apps'][0], workspace['apps'][2]):
                application = db.get(InternshipApplication, int(app_id))
                assert application.status == 'PENDING_REVIEW' and application.version == 0
    foreign = seed_workspace(1)
    with get_sessionmaker()() as db:
        bootstrap_school(db, int(foreign['tenantId']))
        db.commit()
    with login(foreign) as outsider, login(workspace, 'STUDENT') as student:
        assert data(outsider.get(ROOT + '/filter-options', params={'batchId': foreign['batches'][0]['id']}))['academicYears'] == []
        for suffix in ('', '/filter-options', '/export.xlsx'):
            assert outsider.get(ROOT + suffix, params=params).status_code in (403, 404)
            assert student.get(ROOT + suffix, params=params).status_code == 403


def test_filter_validation_is_identical_for_list_and_export(workspace):
    with login(workspace) as client:
        for suffix in ('', '/export.xlsx'):
            for invalid in ({'collegeId': '-1'}, {'majorId': '1.5'}, {'classId': 'text'}, {'grade': 'x' * 21}):
                assert client.get(ROOT + suffix, params={'batchId': workspace['batch'], **invalid}).status_code == 422
            assert client.get(ROOT + suffix, params={'batchId': workspace['batch'], 'exemptionType': 'UNKNOWN'}).status_code == 400
