"""Real normal-role MySQL/API employment declarations; synthetic records only."""
from concurrent.futures import ThreadPoolExecutor
from datetime import datetime, timedelta
import json
import pytest
from sqlalchemy import func, select

from app.core.context import set_tenant
from app.core.field_crypto import decrypt_field, encrypt_field
from app.db.session import get_sessionmaker
from app.models import (College, Major, SchoolClass, StudentContact, EmpCompany, InternshipRecord, StudentProfile, StudentAccountLink, UnifiedTodo, UnifiedMessage,
    InternshipAuditTrail, InternshipArchive, InternshipFinalScore, AuditOutbox, SecurityAuditLog, InternshipEmploymentRegistration as Registration,
    InternshipEmploymentSubmission as Submission, InternshipEmploymentReview as Review)
from app.models.file import FileObject, FileBinding
from app.adapters.employment_gateway import employment_gateway
from app.services.system_authorization import bootstrap_school
from app.modules.internship.services import internship_employment_service as service, internship_audit_service as audit
from test_yiyang_write_flows_mysql import flow as base_flow, context, login, upload, png_bytes, seed_flow
from test_student_weekly_detail_mysql import teammate

MOBILE = '/api/v1/mobile/internship/employment'
STAFF = '/api/v1/teacher-mobile/internship/employment'

@pytest.mark.parametrize('params,expected',[({'revisionNo':2},2),({'revisionNo':True},None),({'revisionNo':'2'},None),(None,None),([],None)])
def test_employment_notification_version_uses_its_own_valid_stored_metadata(params,expected):
    from types import SimpleNamespace
    message=SimpleNamespace(id=1,title='就业结果',content='原核验回执',status='UNREAD',created_at=None,read_at=None,action_params_json=params)
    assert service._notification_view(message)['revisionNo']==expected

@pytest.fixture
def flow(base_flow):
    with get_sessionmaker()() as db:
        tenant = int(base_flow['tenantId'])
        college = College(tenant_id=tenant,college_name='合成信息工程学院'); db.add(college); db.flush()
        major = Major(tenant_id=tenant,college_id=college.id,major_name='合成软件技术'); db.add(major); db.flush()
        school_class = SchoolClass(tenant_id=tenant,major_id=major.id,class_name='合成软件一班',grade='2024'); db.add(school_class); db.flush()
        student = db.get(StudentProfile,int(base_flow['studentId']))
        student.college_id, student.major_id, student.class_id, student.grade = college.id, major.id, school_class.id, '2024'
        student.enroll_date = datetime(2024,9,1)
        db.add(StudentContact(tenant_id=tenant,student_id=student.id,contact_type='PHONE',
            contact_value_encrypted=encrypt_field('13900000000'),is_primary=True))
        db.commit()
    return base_flow

def fields(**extra):
    return {'studentPhone':'13800000000','enrollmentYear':'2024','employmentType':'签订劳动合同就业',
        'companyName':'合成就业测试企业','creditCode':'91430100MA4L0ABC1D','companyAddress':'合成测试地址',
        'companyScale':'小型企业','companyPhone':'073100000000','companyEmail':'synthetic@example.invalid',
        'hrContact':'合成测试联系人','companyNature':'民营企业','industry':'软件服务',
        'jobCategory':'软件开发','positionName':'助理工程师','majorMatch':False,'salary':'0','agreementNo':'SYNTHETIC-AGREEMENT-01',**extra}

def body(data, **extra):
    return {**context(data),'registrationMode':'EMPLOYMENT','fields':fields(),'expectedVersion':0,**extra}

def ok(response):
    assert response.status_code == 200 and response.json()['code'] == 0, response.text
    return response.json()['data']

def decide(client, row, action, comment='合成登记信息已核对'):
    return client.post(STAFF+'/'+row['id']+'/review',json={'action':action,'comment':comment,
        'expectedVersion':row['version'],'submissionId':row['current']['id']})

def own_head(db, data):
    return db.scalar(select(Registration).where(Registration.internship_id == int(data['recordId'])))

def test_other_todo_group_includes_employment_and_unknown_types_with_scope_and_paging(flow):
    foreign = seed_flow()
    path = '/api/v1/teacher-mobile/internship/workbench/todos'
    with get_sessionmaker()() as db:
        bootstrap_school(db,int(flow['tenantId']))
        db.commit()
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher:
        submitted = ok(student.post(MOBILE,json=body(flow)))
        with get_sessionmaker()() as db:
            defaults = dict(tenant_id=int(flow['tenantId']),source_module='internship',
                assignee_id=flow['users']['INTERN_MENTOR'],title='合成分类待办',status='PENDING')
            rows = [
                UnifiedTodo(**{**defaults,'source_biz_id':1,'todo_type':'SYNTHETIC_UNKNOWN'}),
                UnifiedTodo(**{**defaults,'source_biz_id':2,'todo_type':'INTERN_WEEKLY_REVIEW'}),
                UnifiedTodo(**{**defaults,'source_biz_id':3,'todo_type':service.TODO_TYPE,'status':'DONE'}),
                UnifiedTodo(**{**defaults,'source_biz_id':4,'todo_type':service.TODO_TYPE,'is_deleted':True}),
                UnifiedTodo(**{**defaults,'source_biz_id':5,'todo_type':service.TODO_TYPE,'assignee_id':flow['users']['SCHOOL_ADMIN']}),
                UnifiedTodo(**{**defaults,'source_biz_id':6,'todo_type':service.TODO_TYPE,'tenant_id':int(foreign['tenantId'])}),
            ]
            db.add_all(rows); db.commit()
        first = ok(teacher.get(path,params={'group':'OTHER','pageSize':1}))
        second = ok(teacher.get(path,params={'group':'OTHER','pageSize':1,'page':2}))
        assert first['groupCounts']['ALL'] == 3 and first['groupCounts']['OTHER'] == 2
        assert first['total'] == second['total'] == 2 and first['hasMore'] and not second['hasMore']
        items = first['items'] + second['items']
        assert {x['todoType'] for x in items} == {service.TODO_TYPE,'SYNTHETIC_UNKNOWN'}
        employment = next(x for x in items if x['todoType'] == service.TODO_TYPE)
        assert employment['sourceBizId'] == submitted['id']
        assert employment['actionPath'].endswith('?registrationId='+submitted['id'])
        report = ok(teacher.get(path,params={'group':'REPORT'}))
        assert report['total'] == 1 and report['items'][0]['todoType'] == 'INTERN_WEEKLY_REVIEW'
        completed = ok(teacher.get(path,params={'group':'OTHER','status':'DONE'}))
        assert completed['total'] == completed['groupCounts']['OTHER'] == 1


def test_employment_ledger_exports_only_latest_filtered_masked_version_and_preserves_text(flow):
    import base64
    from io import BytesIO
    from openpyxl import load_workbook
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher, login(flow) as school:
        first = ok(student.post(MOBILE,json=body(flow)))
        returned = ok(decide(teacher,first,'RETURN','请补正本次岗位名称'))
        latest = ok(student.post(MOBILE,json=body(flow,expectedVersion=returned['version'],
            correctionReason='合成补正完整岗位名称',fields=fields(companyName='=合成就业单位',positionName='质量工程师'))))
        ok(decide(teacher,latest,'APPROVE'))
        params={**context(flow),'status':'APPROVED'}
        receipt=ok(school.post(STAFF+'/export',params=params))
        sheet=load_workbook(BytesIO(base64.b64decode(receipt['contentBase64']))).active
        assert receipt['rowCount']==1 and sheet.max_row==3
        values=dict(zip([c.value for c in sheet[2]],[c.value for c in sheet[3]]))
        assert values['登记编号']==first['id'] and values['实习记录编号']==flow['recordId']
        assert values['学号']==first['current']['fields']['studentNo'] and sheet.cell(3,8).data_type=='s'
        assert values['提交版本']==2 and values['状态']=='已核验' and values['专业是否对口']=='否'
        assert values['单位名称']=="'=合成就业单位" and values['岗位名称']=='质量工程师'
        assert all(values[service.FIELD_LABELS[key]]=='已隐藏，请说明核验用途后查看' for key in service.SENSITIVE_FIELDS)
        assert '13800000000' not in str(values) and '合成测试地址' not in str(values)
        assert ok(school.post(STAFF+'/export',params={**params,'status':'RETURNED'}))['rowCount']==0
        assert ok(school.post(STAFF+'/export',params={**params,'internshipId':str(int(flow['recordId'])+99999)}))['rowCount']==0
        assert student.post(STAFF+'/export',params=params).status_code==403
        assert school.post(STAFF+'/export',params={**params,'batchId':'99999999999999'}).status_code==404
        with get_sessionmaker()() as db:
            events=db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id==int(flow['tenantId']),
                InternshipAuditTrail.target_type=='EXPORT')).all()
            assert len(events)==3 and all(e.detail_json['filters']['masked'] for e in events)


def test_employment_export_is_scoped_to_tenant_and_current_advisor(flow):
    foreign=seed_flow()
    with login(flow,'STUDENT') as student,login(flow,'INTERN_MENTOR') as teacher,login(foreign) as outsider:
        row=ok(student.post(MOBILE,json=body(flow)))
        assert outsider.post(STAFF+'/export',params=context(flow)).status_code==404
        assert ok(teacher.post(STAFF+'/export',params=context(flow)))['rowCount']==1
        with get_sessionmaker()() as db:
            db.get(InternshipRecord,int(flow['recordId'])).advisor_user_id=None;db.commit()
        assert ok(teacher.post(STAFF+'/export',params=context(flow)))['rowCount']==0
        assert teacher.get(STAFF+'/'+row['id']).status_code==403


def test_approved_employment_correction_requires_new_review_and_does_not_mix_history_or_exports(flow):
    import base64
    from io import BytesIO
    from openpyxl import load_workbook
    def exported(client,status):
        receipt=ok(client.post(STAFF+'/export',params={**context(flow),'status':status}))
        sheet=load_workbook(BytesIO(base64.b64decode(receipt['contentBase64']))).active
        return [dict(zip([c.value for c in sheet[2]],values)) for values in sheet.iter_rows(min_row=3,values_only=True)]
    with login(flow,'STUDENT') as student,login(flow,'INTERN_MENTOR') as teacher,login(flow) as school:
        old_photo=upload(student,biz='INTERNSHIP_EMPLOYMENT')
        initial=ok(student.post(MOBILE,json=body(flow,attachmentFileIds=[old_photo['fileId']])))
        approved=ok(decide(teacher,initial,'APPROVE'))
        old=ok(student.get(MOBILE,params=context(flow)))['registration']['current']
        assert exported(school,'APPROVED')[0]['提交版本']==1
        new_photo=upload(student,biz='INTERNSHIP_EMPLOYMENT',name='correction.png')
        correction=ok(student.post(MOBILE,json=body(flow,expectedVersion=approved['version'],
            correctionReason='核验后更正岗位薪酬与协议',fields=fields(positionName='更正后质量工程师',salary='4200',agreementNo='SYN-CORRECTION-02'),
            attachmentFileIds=[new_photo['fileId']])))
        assert correction['id']==initial['id'] and correction['status']=='SUBMITTED'
        assert correction['current']['review'] is None and correction['current']['revisionNo']==2
        assert correction['history'][1]==old
        assert ok(student.get(MOBILE,params=context(flow)))['resultNotifications']['items'][0]['revisionNo']==1
        assert exported(school,'APPROVED')==[]
        pending=exported(school,'SUBMITTED');assert len(pending)==1 and pending[0]['岗位名称']=='更正后质量工程师' and pending[0]['提交版本']==2
        assert decide(teacher,initial,'APPROVE').status_code==409
        assert decide(teacher,approved,'APPROVE').status_code==409
        assert teacher.get('/api/v1/files/download/'+new_photo['fileId']).status_code==404
        ok(teacher.post(STAFF+'/'+correction['id']+'/disclose',json={'purpose':'核对更正后的就业资料及照片'}))
        assert teacher.get('/api/v1/files/download/'+new_photo['fileId']).content==png_bytes()
        verified=ok(decide(teacher,correction,'APPROVE'))
        reread=ok(student.get(MOBILE,params=context(flow)))['registration']
        assert reread['status']=='APPROVED' and reread['version']==verified['version']==4
        assert reread['current']['fields']['salary']=='4200.00' and reread['history'][1]==old
        assert [x['revisionNo'] for x in ok(student.get(MOBILE,params=context(flow)))['resultNotifications']['items']]==[2,1]
        ledger=exported(school,'APPROVED');assert len(ledger)==1 and ledger[0]['提交版本']==2 and ledger[0]['岗位名称']=='更正后质量工程师'
        assert exported(school,'SUBMITTED')==[]
        with get_sessionmaker()() as db:
            todo=db.scalar(select(UnifiedTodo).where(UnifiedTodo.tenant_id==int(flow['tenantId']),UnifiedTodo.todo_type==service.TODO_TYPE))
            assert todo.status=='DONE'
            assert db.scalar(select(func.count()).select_from(Review).where(Review.registration_id==int(initial['id'])))==2


def test_approved_correction_return_and_resubmit_keep_notices_files_and_old_review_separate(flow):
    with login(flow,'STUDENT') as student,login(flow,'INTERN_MENTOR') as teacher,login(flow) as school:
        old_file=upload(student,biz='INTERNSHIP_EMPLOYMENT')
        initial=ok(student.post(MOBILE,json=body(flow,attachmentFileIds=[old_file['fileId']])))
        approved=ok(decide(teacher,initial,'APPROVE'))
        original=ok(student.get(MOBILE,params=context(flow)))['registration']['current']
        correction=ok(student.post(MOBILE,json=body(flow,expectedVersion=approved['version'],fields=fields(salary='4500'),
            correctionReason='申请更正最新薪酬',attachmentFileIds=[old_file['fileId']])))
        returned=ok(decide(teacher,correction,'RETURN','请补4500元对应的薪酬佐证'))
        before=ok(student.get(MOBILE,params=context(flow)))
        rejected=before['registration']['current'];assert rejected['review']['action']=='RETURN'
        assert [x['revisionNo'] for x in before['resultNotifications']['items']]==[2,1]
        assert ok(school.post(STAFF+'/export',params={**context(flow),'status':'APPROVED'}))['rowCount']==0
        new_file=upload(student,biz='INTERNSHIP_EMPLOYMENT',name='corrected-salary.png')
        resubmitted=ok(student.post(MOBILE,json=body(flow,expectedVersion=returned['version'],fields=fields(salary='4500'),
            correctionReason='已补充一致的新薪酬佐证',attachmentFileIds=[new_file['fileId']])))
        assert resubmitted['id']==initial['id'] and resubmitted['revisionCount']==3 and resubmitted['current']['review'] is None
        assert resubmitted['history'][1]==rejected and resubmitted['history'][2]==original
        assert [x['fileId'] for x in resubmitted['current']['attachments']]==[new_file['fileId']]
        assert [x['fileId'] for x in resubmitted['history'][1]['attachments']]==[old_file['fileId']]
        assert decide(teacher,correction,'APPROVE').status_code==409
        assert teacher.get('/api/v1/files/download/'+new_file['fileId']).status_code==404
        ok(teacher.post(STAFF+'/'+resubmitted['id']+'/disclose',json={'purpose':'核对补正后的薪酬佐证'}))
        assert teacher.get('/api/v1/files/download/'+new_file['fileId']).content==png_bytes()
        ok(decide(teacher,resubmitted,'APPROVE'))
        final=ok(student.get(MOBILE,params=context(flow)))
        assert final['registration']['version']==6 and final['registration']['status']=='APPROVED'
        assert final['registration']['history'][1]==rejected and final['registration']['history'][2]==original
        assert [x['revisionNo'] for x in final['resultNotifications']['items']]==[3,2,1]
        assert ok(school.post(STAFF+'/export',params={**context(flow),'status':'RETURNED'}))['rowCount']==0
        assert ok(school.post(STAFF+'/export',params={**context(flow),'status':'APPROVED'}))['rowCount']==1


def test_employment_export_rejects_over_limit_before_projection(flow,monkeypatch):
    calls=[]
    def oversized(user,**filters):
        calls.append(filters['page_size']);return {'total':10001,'items':[]}
    monkeypatch.setattr(service,'list_for_teacher',oversized)
    with login(flow) as school:
        assert school.post(STAFF+'/export',params=context(flow)).json()['code']!=0
    assert calls==[0]


def test_full_23_fields_photo_todo_return_version_approval_and_formal_audits(flow):
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher:
        initial = ok(student.get(MOBILE,params=context(flow)))
        assert initial['registration'] is None and len(initial['prefill']['fieldLabels']) == 23
        photo = upload(student,biz='INTERNSHIP_EMPLOYMENT')
        assert teacher.get('/api/v1/files/download/'+photo['fileId']).status_code == 404
        submitted = ok(student.post(MOBILE,json=body(flow,attachmentFileIds=[photo['fileId']])))
        old = submitted['current']
        assert len(old['fields']) == 23 and old['fields']['salary'] == '0.00' and old['fields']['majorMatch'] is False
        assert old['fields']['studentName'] == initial['prefill']['identity']['studentName']
        assert all(value != '' for value in old['fields'].values()) and old['sources']['enrollmentYear'] == 'SCHOOL_PROFILE'
        assert old['sources']['studentPhone'] == 'STUDENT_DECLARED'
        assert teacher.get('/api/v1/files/download/'+photo['fileId']).status_code == 404
        ok(teacher.post(STAFF+'/'+submitted['id']+'/disclose',json={'purpose':'核对就业登记与佐证照片'}))
        assert teacher.get('/api/v1/files/download/'+photo['fileId']).content == png_bytes()
        queue = ok(teacher.get(STAFF,params={'batchId':flow['batchId'],'status':'SUBMITTED','pageSize':1}))
        assert queue['total'] == 1 and queue['items'][0]['id'] == submitted['id']
        todos = ok(teacher.get('/api/v1/teacher-mobile/internship/workbench/todos'))
        item = next(x for x in todos['items'] if x['todoType'] == service.TODO_TYPE)
        assert item['sourceBizId'] == submitted['id'] and 'registrationId='+submitted['id'] in item['actionPath']
        returned = ok(decide(teacher,submitted,'RETURN','请补充真实岗位职责与单位材料'))
        assert returned['status'] == 'RETURNED' and returned['history'][0]['review']['action'] == 'RETURN'
        payload = body(flow,expectedVersion=returned['version'],correctionReason='已经补充岗位职责与单位材料',
            fields=fields(positionName='软件测试工程师',salary='1234.50'),attachmentFileIds=[photo['fileId']])
        current = ok(student.post(MOBILE,json=payload))
        approved = ok(decide(teacher,current,'APPROVE'))
        reread = ok(student.get(MOBILE,params=context(flow)))['registration']
        assert reread['id'] == approved['id'] and reread['status'] == 'APPROVED' and reread['version'] == 4
        assert reread['sensitiveVisible'] and not approved['sensitiveVisible']
        assert [x['revisionNo'] for x in reread['history']] == [2,1]
        assert reread['history'][1]['fields'] == old['fields']
        assert reread['history'][1]['review']['action'] == 'RETURN' and reread['current']['review']['action'] == 'APPROVE'
    with get_sessionmaker()() as db:
        snapshots = db.scalars(select(Submission).where(Submission.registration_id == int(approved['id'])).order_by(Submission.revision_no)).all()
        assert len(snapshots) == 2 and all('13800000000' not in x.details_encrypted and 'example.invalid' not in x.details_encrypted for x in snapshots)
        contact = db.scalar(select(StudentContact).where(StudentContact.student_id == int(flow['studentId'])))
        assert decrypt_field(contact.contact_value_encrypted) == '13900000000'
        assert json.loads(decrypt_field(snapshots[0].details_encrypted,allow_legacy_plaintext=False))['fields'] == old['fields']
        assert db.scalar(select(func.count()).select_from(Review).where(Review.registration_id == int(approved['id']))) == 2
        todo = db.scalar(select(UnifiedTodo).where(UnifiedTodo.tenant_id == int(flow['tenantId']),UnifiedTodo.todo_type == service.TODO_TYPE))
        assert todo.status == 'DONE'
        messages = db.scalars(select(UnifiedMessage).where(UnifiedMessage.tenant_id == int(flow['tenantId']),UnifiedMessage.source_biz_id == int(approved['id']))).all()
        assert len(messages) == 2 and all(m.receiver_user_id == flow['users']['STUDENT'] and m.action_key == 'student.internship.employment-registration' for m in messages)
        trails = db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id == int(flow['tenantId']),InternshipAuditTrail.target_type == 'EMPLOYMENT_REGISTRATION')).all()
        assert len(trails) == 5 and all('13800000000' not in json.dumps(t.detail_json) for t in trails)
        events = db.scalars(select(AuditOutbox).where(AuditOutbox.tenant_id == int(flow['tenantId']),AuditOutbox.event_type.like('INTERNSHIP_EMPLOYMENT_%'))).all()
        assert len(events) == 5
        event_ids = [x.event_id for x in events]
        # Preserved synthetic test history can precede these events in the FIFO queue.
        preceding_due = db.scalar(select(func.count()).select_from(AuditOutbox).where(
            AuditOutbox.id <= max(x.id for x in events),
            AuditOutbox.status.in_(("PENDING", "RETRY_WAIT")),
            AuditOutbox.next_retry_at.is_(None) | (AuditOutbox.next_retry_at <= datetime.utcnow()),
        )) or 0
    for _ in range(max(1, (preceding_due + 999) // 1000)):
        audit.process_pending(limit=1000)
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(SecurityAuditLog).where(SecurityAuditLog.source_event_id.in_(event_ids))) == 5
    set_tenant(flow['tenantId'])
    try:
        destination = employment_gateway.get_student_destination(int(flow['studentId']))
        assert destination['recordId'] == approved['id'] and destination['externalVerified'] is False and destination['revisionNo'] == 2
        assert employment_gateway.get_student_destination(int(flow['studentId']),internship_id=int(flow['recordId'])) == destination
        assert employment_gateway.get_student_destination(int(flow['studentId']),internship_id=int(flow['recordId'])+1) is None
    finally:
        set_tenant(None)

def test_transfer_uses_current_canonical_company_without_mutating_placement(flow):
    with get_sessionmaker()() as db:
        company = EmpCompany(tenant_id=int(flow['tenantId']),name='合成实习企业',credit_code='91430100MA4L0ABC1D',status='ACTIVE',blacklist=False)
        db.add(company); db.flush()
        record = db.get(InternshipRecord,int(flow['recordId'])); record.enterprise_id = company.id; record.enterprise_name = company.name
        original = (record.enterprise_id,record.enterprise_name,record.position_name,record.status)
        db.commit()
    with login(flow,'STUDENT') as student:
        prefill = ok(student.get(MOBILE,params=context(flow)))['prefill']
        assert prefill['canTransfer'] and prefill['transfer']['companyName'] == '合成实习企业'
        bad = student.post(MOBILE,json=body(flow,registrationMode='INTERNSHIP_TRANSFER'))
        assert bad.json()['code'] != 0
        good = ok(student.post(MOBILE,json=body(flow,registrationMode='INTERNSHIP_TRANSFER',fields=fields(companyName='合成实习企业'))))
        assert good['current']['sourceCompanyId'] == str(original[0])
    with get_sessionmaker()() as db:
        record = db.get(InternshipRecord,int(flow['recordId']))
        assert (record.enterprise_id,record.enterprise_name,record.position_name,record.status) == original

def test_transfer_rejects_unlinked_redundant_company_name(flow):
    with login(flow,'STUDENT') as student:
        assert not ok(student.get(MOBILE,params=context(flow)))['prefill']['canTransfer']
        assert student.post(MOBILE,json=body(flow,registrationMode='INTERNSHIP_TRANSFER')).json()['code'] != 0
    with get_sessionmaker()() as db:
        assert own_head(db,flow) is None

@pytest.mark.parametrize('changed', [{'studentName':'另一个学生'},{'studentNo':'伪造学号'},{'collegeName':'伪造院系'},{'enrollmentYear':'2023'}])
def test_authoritative_school_identity_cannot_be_overridden(flow,changed):
    with get_sessionmaker()() as db:
        db.get(StudentProfile,int(flow['studentId'])).enroll_date = datetime(2024,9,1); db.commit()
    with login(flow,'STUDENT') as student:
        response = student.post(MOBILE,json=body(flow,fields=fields(**changed)))
        assert response.status_code == 403 and response.json()['code'] != 0
    with get_sessionmaker()() as db:
        assert own_head(db,flow) is None

@pytest.mark.parametrize('changed', [{'salary':'NaN'},{'salary':'-1'},{'salary':'0.001'},{'majorMatch':'false'},{'companyEmail':'bad-email'},{'companyName':'x'*201},{'employmentType':''}])
def test_invalid_business_fields_reject_without_persistent_facts(flow,changed):
    with login(flow,'STUDENT') as student:
        assert student.post(MOBILE,json=body(flow,fields=fields(**changed))).json()['code'] != 0
    with get_sessionmaker()() as db:
        assert own_head(db,flow) is None

def test_foreign_other_student_wrong_batch_revoked_link_and_staff_denied(flow):
    foreign = seed_flow(); other, other_context = teammate(flow)
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher, login(foreign,'INTERN_MENTOR') as outsider:
        photo = upload(student,biz='INTERNSHIP_EMPLOYMENT')
        row = ok(student.post(MOBILE,json=body(flow,attachmentFileIds=[photo['fileId']])))
        assert other.get(MOBILE,params=context(flow)).json()['code'] != 0
        assert other.get('/api/v1/files/download/'+photo['fileId']).status_code == 404
        assert student.get(MOBILE,params={**context(flow),'batchId':other_context['batchId']}).json()['code'] != 0
        assert teacher.get(MOBILE,params=context(flow)).status_code == 403
        assert outsider.get(STAFF+'/'+row['id']).status_code == 404
        assert outsider.post(STAFF+'/'+row['id']+'/review',json={'action':'APPROVE','submissionId':row['current']['id'],'expectedVersion':row['version']}).status_code == 404
        with get_sessionmaker()() as db:
            record = db.get(InternshipRecord,int(flow['recordId'])); record.advisor_user_id = None; record.advisor_name = None
            link = db.scalar(select(StudentAccountLink).where(StudentAccountLink.user_id == flow['users']['STUDENT'])); link.link_status = 'REVOKED'; db.commit()
        assert student.get(MOBILE,params=context(flow)).status_code == 403
        assert student.get('/api/v1/files/download/'+photo['fileId']).status_code == 403
        assert teacher.get(STAFF+'/'+row['id']).status_code == 403
        assert teacher.get('/api/v1/files/download/'+photo['fileId']).status_code == 404
    other.close()

def test_duplicate_concurrent_first_submit_and_old_review_have_one_fact(flow):
    with login(flow,'STUDENT') as one, login(flow,'STUDENT') as two, login(flow,'INTERN_MENTOR') as teacher:
        with ThreadPoolExecutor(max_workers=2) as pool:
            replies = [future.result(timeout=20) for future in [pool.submit(c.post,MOBILE,json=body(flow)) for c in (one,two)]]
        assert sum(r.status_code == 200 and r.json()['code'] == 0 for r in replies) == 1, [r.text for r in replies]
        row = next(r.json()['data'] for r in replies if r.json()['code'] == 0)
        approved = ok(decide(teacher,row,'APPROVE'))
        assert decide(teacher,row,'RETURN').json()['code'] != 0
        assert one.post(MOBILE,json=body(flow)).json()['code'] != 0
        reread = ok(one.get(MOBILE,params=context(flow)))['registration']
        assert (reread['id'],reread['status'],reread['version']) == (approved['id'],approved['status'],approved['version'])
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(Submission).where(Submission.registration_id == int(row['id']))) == 1
        assert db.scalar(select(func.count()).select_from(Review).where(Review.registration_id == int(row['id']))) == 1

def test_private_pending_and_already_bound_other_business_photos_are_not_rebound(flow):
    with login(flow,'STUDENT') as student, login(flow) as school:
        private = upload(school)
        assert student.post(MOBILE,json=body(flow,attachmentFileIds=[private['fileId']])).json()['code'] != 0
        pending = upload(student,biz='INTERNSHIP_EMPLOYMENT')
        with get_sessionmaker()() as db:
            file_obj = db.get(FileObject,int(pending['fileId'])); file_obj.scan_status='PENDING'; db.commit()
        assert student.post(MOBILE,json=body(flow,attachmentFileIds=[pending['fileId']])).json()['code'] != 0
        other = upload(student,biz='INTERNSHIP_INSURANCE_POLICY')
        from test_yiyang_write_flows_mysql import insurance_body, PORTAL
        ok(student.post(PORTAL+'/insurance',json=insurance_body(flow,other['fileId'])))
        assert student.post(MOBILE,json=body(flow,attachmentFileIds=[other['fileId']])).json()['code'] != 0
    with get_sessionmaker()() as db:
        assert own_head(db,flow) is None
        assert db.get(FileObject,int(other['fileId'])).biz_type == 'INTERNSHIP_INSURANCE'

def test_transaction_failure_rolls_back_declaration_photo_todo_and_audit(flow,monkeypatch):
    def fail(*args,**kwargs):
        raise RuntimeError('synthetic employment audit failure')
    with login(flow,'STUDENT') as student:
        photo = upload(student,biz='INTERNSHIP_EMPLOYMENT')
        monkeypatch.setattr(audit,'add_audit',fail)
        with pytest.raises(RuntimeError,match='synthetic employment audit failure'):
            student.post(MOBILE,json=body(flow,attachmentFileIds=[photo['fileId']]))
    with get_sessionmaker()() as db:
        assert own_head(db,flow) is None
        assert db.get(FileObject,int(photo['fileId'])).biz_type == 'TEMP_PRIVATE'
        assert db.scalar(select(func.count()).select_from(FileBinding).where(FileBinding.file_id == int(photo['fileId']))) == 0
        assert db.scalar(select(func.count()).select_from(UnifiedTodo).where(UnifiedTodo.tenant_id == int(flow['tenantId']),UnifiedTodo.todo_type == service.TODO_TYPE)) == 0
        assert db.scalar(select(func.count()).select_from(AuditOutbox).where(AuditOutbox.tenant_id == int(flow['tenantId']),AuditOutbox.event_type.like('INTERNSHIP_EMPLOYMENT_%'))) == 0

def test_masked_teacher_listing_and_explicit_scoped_disclosure_audit(flow):
    foreign = seed_flow()
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher, login(foreign,'INTERN_MENTOR') as outsider:
        row = ok(student.post(MOBILE,json=body(flow)))
        listed = ok(teacher.get(STAFF,params={'batchId':flow['batchId']}))['items'][0]
        detail = ok(teacher.get(STAFF+'/'+row['id']))
        for result in (listed,detail):
            assert not result['sensitiveVisible'] and '13800000000' not in json.dumps(result)
            assert 'example.invalid' not in json.dumps(result) and '合成测试地址' not in json.dumps(result)
        path = STAFF+'/'+row['id']+'/disclose'
        assert teacher.post(path,json={'purpose':'看'}).json()['code'] != 0
        assert outsider.post(path,json={'purpose':'核对本校就业登记材料'}).status_code == 404
        revealed = ok(teacher.post(path,json={'purpose':'核对学生就业单位与协议材料'}))
        assert revealed['sensitiveVisible'] and revealed['current']['fields'] == row['current']['fields']
    with get_sessionmaker()() as db:
        reads = db.scalars(select(InternshipAuditTrail).where(InternshipAuditTrail.tenant_id == int(flow['tenantId']),
            InternshipAuditTrail.action == 'EMPLOYMENT_SENSITIVE_VIEW')).all()
        assert len(reads) == 1 and '13800000000' not in json.dumps(reads[0].detail_json)

def test_teacher_photo_requires_own_disclosure_of_current_registration_version(flow):
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher, login(flow) as school:
        photo = upload(student,biz='INTERNSHIP_EMPLOYMENT')
        row = ok(student.post(MOBILE,json=body(flow,attachmentFileIds=[photo['fileId']])))
        path = '/api/v1/files/download/'+photo['fileId']
        assert teacher.get(path).status_code == 404
        ok(teacher.post(STAFF+'/'+row['id']+'/disclose',json={'purpose':'核对就业上报当前佐证材料'}))
        assert teacher.get(path).content == png_bytes() and school.get(path).status_code == 404
        returned = ok(decide(teacher,row,'RETURN','请补充就业岗位及佐证材料'))
        assert teacher.get(path).status_code == 404
        resubmitted = ok(student.post(MOBILE,json=body(flow,expectedVersion=returned['version'],correctionReason='已核对并补充就业岗位材料',attachmentFileIds=[photo['fileId']])))
        assert teacher.get(path).status_code == 404
        ok(teacher.post(STAFF+'/'+resubmitted['id']+'/disclose',json={'purpose':'核对补正后的就业佐证材料'}))
        assert teacher.get(path).content == png_bytes()
        assert student.get(path).content == png_bytes()

def test_school_identity_missing_blocks_and_missing_enrollment_can_be_declared(flow):
    with get_sessionmaker()() as db:
        student = db.get(StudentProfile,int(flow['studentId'])); student.class_id = None; student.enroll_date = None; db.commit()
    with login(flow,'STUDENT') as student:
        assert student.post(MOBILE,json=body(flow)).json()['code'] != 0
        with get_sessionmaker()() as db:
            profile = db.get(StudentProfile,int(flow['studentId']))
            profile.class_id = db.scalar(select(SchoolClass.id).where(SchoolClass.tenant_id == int(flow['tenantId']))); db.commit()
        row = ok(student.post(MOBILE,json=body(flow)))
        assert row['current']['sources']['enrollmentYear'] == 'STUDENT_DECLARED'
    with get_sessionmaker()() as db:
        assert db.get(StudentProfile,int(flow['studentId'])).enroll_date is None

def test_concurrent_review_records_one_immutable_decision(flow):
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as one, login(flow,'INTERN_MENTOR') as two:
        row = ok(student.post(MOBILE,json=body(flow)))
        with ThreadPoolExecutor(max_workers=2) as pool:
            futures = [pool.submit(decide,client,row,action) for client,action in ((one,'APPROVE'),(two,'RETURN'))]
            replies = [future.result(timeout=20) for future in futures]
        assert sum(r.status_code == 200 and r.json()['code'] == 0 for r in replies) == 1, [r.text for r in replies]
        accepted = next(r.json()['data'] for r in replies if r.json()['code'] == 0)
        assert ok(student.get(MOBILE,params=context(flow)))['registration']['status'] == accepted['status']
    with get_sessionmaker()() as db:
        assert db.scalar(select(func.count()).select_from(Review).where(Review.registration_id == int(row['id']))) == 1

def test_student_result_notifications_are_visible_without_auto_read_and_ack_is_idempotent(flow):
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher:
        row = ok(student.post(MOBILE,json=body(flow)))
        returned = ok(decide(teacher,row,'RETURN','请补充就业岗位职责与材料'))
        first = ok(student.get(MOBILE,params=context(flow)))['resultNotifications']
        assert first['total'] == first['unread'] == 1 and first['items'][0]['status'] == 'UNREAD'
        message = first['items'][0]
        assert '已退回补正' in message['content']
        with get_sessionmaker()() as db:
            original_version = db.get(UnifiedMessage,int(message['id'])).version
    with login(flow,'STUDENT') as one, login(flow,'STUDENT') as two:
        reread = ok(one.get(MOBILE,params=context(flow)))
        assert reread['resultNotifications']['unread'] == 1 and reread['registration']['version'] == returned['version']
        path = MOBILE+'/messages/'+message['id']+'/read'
        with ThreadPoolExecutor(max_workers=2) as pool:
            results = [future.result(timeout=20) for future in [pool.submit(c.post,path,params=context(flow)) for c in (one,two)]]
        assert all(ok(r)['status'] == 'READ' for r in results)
        assert ok(one.get(MOBILE,params=context(flow)))['resultNotifications']['unread'] == 0
    with get_sessionmaker()() as db:
        saved = db.get(UnifiedMessage,int(message['id']))
        assert saved.read_at is not None and saved.version == original_version+1
        assert own_head(db,flow).version == returned['version']

def test_result_notification_read_rejects_wrong_identity_context_and_revoked_binding(flow):
    foreign = seed_flow(); other, other_context = teammate(flow)
    try:
        with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher, login(foreign,'STUDENT') as outsider:
            row = ok(student.post(MOBILE,json=body(flow)))
            ok(decide(teacher,row,'APPROVE'))
            message = ok(student.get(MOBILE,params=context(flow)))['resultNotifications']['items'][0]
            path = MOBILE+'/messages/'+message['id']+'/read'
            for client, ctx in ((teacher,context(flow)),(outsider,context(flow)),(other,context(flow)),(other,other_context),
                (student,{**context(flow),'batchId':other_context['batchId']})):
                assert client.post(path,params=ctx).json()['code'] != 0
            with get_sessionmaker()() as db:
                link = db.scalar(select(StudentAccountLink).where(StudentAccountLink.user_id == flow['users']['STUDENT']))
                link.link_status = 'REVOKED';db.commit()
            assert student.post(path,params=context(flow)).status_code == 403
        with get_sessionmaker()() as db:
            assert db.get(UnifiedMessage,int(message['id'])).status == 'UNREAD'
    finally:
        other.close()

def test_result_notifications_are_sql_paged_and_hide_other_business_withdrawn_and_failed_messages(flow):
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher:
        row = ok(student.post(MOBILE,json=body(flow)));ok(decide(teacher,row,'APPROVE'))
        with get_sessionmaker()() as db:
            for index in range(26):
                values = dict(tenant_id=int(flow['tenantId']),receiver_id=flow['users']['STUDENT'],receiver_user_id=flow['users']['STUDENT'],
                    receiver_type='STUDENT',receiver_context_key='GLOBAL',source_module='internship',source_biz_id=int(row['id']),
                    title='合成分页结果通知',content='合成结果',status='UNREAD',action_key=service.RESULT_ACTION,delivery_status='DELIVERED')
                if index == 22:values['action_key']='other.synthetic.action'
                if index == 23:values['withdrawn_at']=datetime.utcnow()
                if index == 24:values['delivery_status']='FAILED'
                if index == 25:values['expire_at']=datetime.utcnow()-timedelta(seconds=1)
                db.add(UnifiedMessage(**values))
            db.commit()
        first = ok(student.get(MOBILE,params=context(flow)))['resultNotifications']
        second = ok(student.get(MOBILE,params={**context(flow),'notificationPage':2}))['resultNotifications']
        assert first['total'] == first['unread'] == second['total'] == 23
        assert len(first['items']) == 20 and first['hasMore'] and len(second['items']) == 3 and not second['hasMore']
        assert not ({x['id'] for x in first['items']} & {x['id'] for x in second['items']})
        assert student.get(MOBILE,params={**context(flow),'notificationPage':0}).status_code == 422

def test_archive_handoff_uses_exact_placement_and_preserves_published_score_freeze_gates(flow):
    """Prepared isolated archive facts test this read projection, not an archive UI journey."""
    from app.modules.internship.services.internship_score_archive_guard import _score_freeze
    path = '/api/v1/internship/archive/'+flow['recordId']+'/employment-transition'
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher:
        row = ok(student.post(MOBILE,json=body(flow)))
        assert teacher.get(path).json()['code'] != 0
        set_tenant(flow['tenantId'])
        try:
            with get_sessionmaker()() as db:
                score = InternshipFinalScore(tenant_id=int(flow['tenantId']),internship_id=int(flow['recordId']),
                    student_id=int(flow['studentId']),batch_id=int(flow['batchId']),status='PUBLISHED',version=2,total_score=80.0,is_pass=True,incomplete=False)
                db.add(score);db.flush();db.refresh(score);frozen, fingerprint = _score_freeze(db,score)
                db.add(InternshipArchive(tenant_id=int(flow['tenantId']),internship_id=int(flow['recordId']),student_id=int(flow['studentId']),
                    batch_id=int(flow['batchId']),status='ARCHIVED',material_snapshot={'finalScoreFreeze':frozen,'finalScoreFreezeHash':fingerprint}))
                score_id=score.id;db.commit()
        finally:
            set_tenant(None)
        pending = ok(teacher.get(path))
        assert not pending['employmentRecordExists'] and pending['employmentGatewayStatus'] == 'NOT_SUBMITTED'
        assert 'internshipId='+flow['recordId'] in pending['employmentPath'] and 'batchId='+flow['batchId'] in pending['employmentPath']
        ok(decide(teacher,row,'APPROVE'))
        linked = ok(teacher.get(path))
        assert linked['employmentRecordId'] == row['id'] and 'registrationId='+row['id'] in linked['employmentPath']
        assert linked['employmentGatewayStatus'] == 'SCHOOL_VERIFIED' and linked['finalScoreFreezeHash'] == fingerprint
        assert student.get(path).status_code == 403
        with get_sessionmaker()() as db:
            score = db.get(InternshipFinalScore,score_id);score.version += 1;db.commit()
        assert teacher.get(path).json()['code'] != 0
