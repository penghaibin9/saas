"""Change-ledger projection, cap and concurrent-count checks without a database."""
import base64
from contextlib import nullcontext
from io import BytesIO

import pytest
from openpyxl import load_workbook

from app.core.exceptions import AppException
from app.modules.internship.services import internship_change_service as svc
from app.modules.internship.services import internship_export_util as exports


def test_change_export_preserves_history_text_dates_and_formula_safety(monkeypatch):
    calls, audits = [], []
    row = dict(id='9007199254740993001', studentNo='000000000000000000123',
        studentName='=1+1', changeType='CHANGE_ENTERPRISE', changeTypeLabel='换实习单位',
        status='APPROVED', statusLabel='已通过', reason='@合成原因',
        createdAt='2026-10-12T23:35:00Z', reviewedAt='2026-10-13T00:01:00Z',
        reviewedBy='合成导师', reviewComment='+合成意见', targetEnterpriseName='申请单位',
        targetPositionName='申请岗位', currentEnterprise='后续当前单位', currentPosition='后续当前岗位',
        approvedHistory={'before': {'enterpriseName': '原单位', 'positionName': '原岗位'},
                         'after': {'enterpriseName': '本次新单位', 'positionName': '本次新岗位'},
                         'applicationBasisMatched': False})
    def listing(page, size, **filters):
        calls.append((page, size, filters))
        return ([row] if size else []), 1
    monkeypatch.setattr(svc, 'list_changes', listing)
    monkeypatch.setattr(svc, 'session', lambda: nullcontext(type('DB', (), {'commit': lambda self: None})()))
    monkeypatch.setattr(exports, 'write_export_audit', lambda *args, **kw: audits.append((args, kw)))
    result = svc.export_changes(batch_id='9007199254740993011', status='APPROVED', keyword='姓名', user={'realName':'导师'})
    assert [call[:2] for call in calls] == [(1, 0), (1, 1)]
    assert calls[0][2] == calls[1][2]
    sheet = load_workbook(BytesIO(base64.b64decode(result['contentBase64']))).active
    actual = dict(zip([cell.value for cell in sheet[2]], [cell.value for cell in sheet[3]]))
    assert actual['申请编号'] == row['id'] and actual['学号'] == row['studentNo']
    assert actual['姓名'] == "'=1+1" and actual['申请原因'] == "'@合成原因"
    assert actual['审核意见'] == "'+合成意见"
    assert actual['申请时间'] == '2026-10-13 07:35'
    assert actual['审核时间'] == '2026-10-13 08:01'
    assert actual['状态'] == '已通过' and actual['变更类型'] == '换实习单位'
    assert actual['本次变更前单位'] == '原单位' and actual['本次变更后单位'] == '本次新单位'
    assert actual['当前实习单位'] == '后续当前单位' and actual['申请目标单位'] == '申请单位'
    assert actual['历史依据说明'] == '仅记录本次审批前后去向，不能确认申请时去向'
    assert result['rowCount'] == 1 and audits[0][1]['filters']['batchId'] == '9007199254740993011'
    row['approvedHistory'] = None
    result = svc.export_changes(batch_id='18')
    sheet = load_workbook(BytesIO(base64.b64decode(result['contentBase64']))).active
    actual = dict(zip([cell.value for cell in sheet[2]], [cell.value for cell in sheet[3]]))
    assert actual['历史依据说明'] == '历史去向未记录，无法核实本次审批前后去向'
    assert actual['本次变更前单位'] is None and actual['本次变更后单位'] is None


@pytest.mark.parametrize('counts,code', [([10001], 'VALIDATION_ERROR'), ([1, 2], 'DATA_CONFLICT')])
def test_change_export_rejects_limit_or_count_change_before_file_and_audit(monkeypatch, counts, code):
    calls = []
    def listing(page, size, **filters):
        calls.append(size)
        return [], counts[len(calls)-1]
    monkeypatch.setattr(svc, 'list_changes', listing)
    monkeypatch.setattr(svc, 'session', lambda: pytest.fail('Rejected exports must not write an audit'))
    with pytest.raises(AppException) as exc:
        svc.export_changes(batch_id='18')
    assert exc.value.code == code
    assert calls == ([0] if len(counts) == 1 else [0, 1])
