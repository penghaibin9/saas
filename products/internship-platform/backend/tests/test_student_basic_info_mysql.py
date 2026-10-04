"""Normal roles, current master facts, atomic versions, import and identical ledger export."""
from base64 import b64decode
from concurrent.futures import ThreadPoolExecutor
from io import BytesIO

import pytest
from openpyxl import load_workbook
from sqlalchemy import select

from app.db.session import get_sessionmaker
from app.models import AuditOutbox, InternshipRecord, StudentAccountLink, StudentProfile
from app.services import xlsx_util
from test_yiyang_write_flows_mysql import flow, login, seed_flow

MOBILE = '/api/v1/mobile/internship/context/basic-info'
BASE = '/api/v1/internship/intern-students'


def params(data):
    return {'batchId': data['batchId'], 'internshipId': data['recordId']}


def fields(version=0, **extra):
    return {'sourceRegion': '湖南省 益阳市', 'nation': '汉族', 'gender': '男', 'expectedVersion': version, **extra}


def read(client, data):
    result = client.get(MOBILE, params=params(data))
    assert result.status_code == 200 and result.json()['code'] == 0, result.text
    return result.json()['data']


def test_uncollected_is_real_missing_task_not_default_completion(flow):
    with login(flow, 'STUDENT') as student:
        profile = read(student, flow)
        assert profile['sourceRegion'] == '' and profile['nation'] == ''
        assert not profile['complete'] and profile['todoCount'] == 1
        assert '生源地' in profile['missingFields'] and '民族' in profile['missingFields']
        dashboard = student.get('/api/v1/mobile/internship/context/my', params={'batchId': flow['batchId']})
        assert dashboard.status_code == 200, dashboard.text
        assert dashboard.json()['data']['basicInfo']['todoCount'] == 1


def test_student_school_ledger_export_share_master_and_clear_task(flow):
    with login(flow, 'STUDENT') as student, login(flow) as school:
        original = read(student, flow)
        saved = student.put(MOBILE, params=params(flow), json=fields(original['version']))
        assert saved.status_code == 200, saved.text
        current = saved.json()['data']
        assert current['complete'] and current['todoCount'] == 0
        url = BASE+'/'+flow['recordId']+'/basic-info'
        back = school.get(url, params={'batchId': flow['batchId']})
        assert back.status_code == 200 and back.json()['data']['version'] == current['version'], back.text
        corrected = school.put(url, params={'batchId': flow['batchId']}, json=fields(current['version'], sourceRegion='湖北省 武汉市'))
        assert corrected.status_code == 200, corrected.text
        assert read(student, flow)['sourceRegion'] == '湖北省 武汉市'
        listing = school.get(BASE, params={'batchId': flow['batchId'], 'sourceRegion': '湖北省', 'pageSize': 1})
        assert listing.status_code == 200 and listing.json()['data']['total'] == 1, listing.text
        row = listing.json()['data']['items'][0]
        assert row['sourceRegion'] == '湖北省 武汉市'
        assert school.get(BASE, params={'batchId': flow['batchId'], 'sourceRegion': '湖南省'}).json()['data']['total'] == 0
        export = school.post(BASE+'/export', params={'batchId': flow['batchId'], 'sourceRegion': '湖北省'})
        assert export.status_code == 200, export.text
        values = list(load_workbook(BytesIO(b64decode(export.json()['data']['contentBase64']))).active.values)
        assert export.json()['data']['rowCount'] == 1 and any('湖北省 武汉市' in row for row in values)
        assert read(student, flow)['todoCount'] == 0
        with get_sessionmaker()() as db:
            record = db.get(InternshipRecord, int(flow['recordId']))
            master = db.get(StudentProfile, record.student_id)
            assert (master.source_region, master.nation, master.gender, master.version) == ('湖北省 武汉市', '汉族', '男', original['version']+2)
            events = db.scalars(select(AuditOutbox).where(AuditOutbox.tenant_id == int(flow['tenantId']),
                                  AuditOutbox.event_type.in_(['INTERNSHIP_BASIC_INFO_SAVE_SELF', 'INTERNSHIP_BASIC_INFO_SAVE_SCHOOL']))).all()
            assert len(events) == 2
            assert all('湖南省' not in str(event.payload_json) and '汉族' not in str(event.payload_json) for event in events)


@pytest.mark.parametrize('extra', [{'sourceRegion': '未采集'}, {'gender': 'UNKNOWN'}, {'nation': '*族'},
                                  {'sourceRegion': 'x'*101}, {'studentId': '9999'}, {'expectedVersion': True}])
def test_invalid_values_and_client_identity_never_write(flow, extra):
    with login(flow, 'STUDENT') as student:
        original = read(student, flow)
        denied = student.put(MOBILE, params=params(flow), json=fields(original['version'], **extra))
        assert denied.status_code in (400,422), denied.text
        assert read(student, flow) == original


def test_empty_values_restore_truthful_todo(flow):
    with login(flow, 'STUDENT') as student:
        version = read(student, flow)['version']
        first = student.put(MOBILE, params=params(flow), json=fields(version)).json()['data']
        cleared = student.put(MOBILE, params=params(flow), json=fields(first['version'], sourceRegion='', nation='', gender=''))
        assert cleared.status_code == 200, cleared.text
        assert cleared.json()['data']['missingFields'] == ['生源地','民族','性别']
        assert cleared.json()['data']['todoCount'] == 1


def test_stale_and_simultaneous_normal_roles_have_one_winner(flow):
    with login(flow,'STUDENT') as student, login(flow) as school:
        version = read(student, flow)['version']
        with ThreadPoolExecutor(max_workers=2) as pool:
            a = pool.submit(student.put, MOBILE, params=params(flow), json=fields(version))
            b = pool.submit(school.put, BASE+'/'+flow['recordId']+'/basic-info', params={'batchId': flow['batchId']}, json=fields(version, sourceRegion='湖北省 武汉市'))
            results = [a.result(),b.result()]
        assert sorted(r.status_code for r in results) == [200,400], [r.text for r in results]
        assert next(r for r in results if r.status_code == 400).json()['bizCode'] == 'DATA_CONFLICT'
        assert read(student,flow)['version'] == version+1


def test_staff_foreign_school_wrong_context_and_revoked_binding_denied(flow):
    foreign = seed_flow()
    with login(flow,'STUDENT') as student, login(flow,'INTERN_MENTOR') as teacher, login(foreign) as outsider:
        url = BASE+'/'+flow['recordId']+'/basic-info'
        assert teacher.put(url, params={'batchId':flow['batchId']}, json=fields()).status_code == 403
        assert teacher.get(MOBILE,params=params(flow)).status_code == 403
        assert outsider.get(url,params={'batchId':flow['batchId']}).status_code in (403,404)
        assert student.get(MOBILE,params={**params(flow),'internshipId':'9999999'}).status_code in (403,404)
        assert student.get(MOBILE,params={**params(flow),'batchId':flow['batches'][1]['id']}).status_code in (400,403,404)
        with get_sessionmaker()() as db:
            link=db.scalar(select(StudentAccountLink).where(StudentAccountLink.user_id==int(flow['users']['STUDENT'])))
            link.link_status='REVOKED'; db.commit()
        assert student.put(MOBILE,params=params(flow),json=fields()).status_code == 403


def test_import_preview_preserves_versions_and_atomic_conflict(flow):
    with login(flow) as school, login(flow,'STUDENT') as student:
        row = read(student,flow)
        body = {'batchId': flow['batchId'], 'rows': [{'studentNo':row['studentNo'], **fields(row['version'])}]}
        pre = school.post(BASE+'/basic-info-import/dry-run',json=body)
        assert pre.status_code == 200 and pre.json()['data']['invalidRows'] == 0, pre.text
        body['rows'] = pre.json()['data']['rows']
        assert read(student,flow) == row
        saved = student.put(MOBILE,params=params(flow),json=fields(row['version'],sourceRegion='湖北省 武汉市'))
        assert saved.status_code == 200, saved.text
        denied = school.post(BASE+'/basic-info-import/confirm',json=body)
        assert denied.status_code == 400 and denied.json()['bizCode'] == 'DATA_CONFLICT', denied.text
        assert read(student,flow)['sourceRegion'] == '湖北省 武汉市'
        fresh = school.post(BASE+'/basic-info-import/dry-run',json={'batchId':flow['batchId'],'rows':body['rows']}).json()['data']
        accepted = school.post(BASE+'/basic-info-import/confirm',json={'batchId':flow['batchId'],'rows':fresh['rows']})
        assert accepted.status_code == 200 and accepted.json()['data']['updated'] == 1, accepted.text
        assert read(student,flow)['sourceRegion'] == '湖南省 益阳市'
        duplicated = school.post(BASE+'/basic-info-import/dry-run',json={'batchId':flow['batchId'],'rows':fresh['rows']*2})
        assert duplicated.status_code == 200 and duplicated.json()['data']['invalidRows'] == 1
        unknown = school.post(BASE+'/basic-info-import/dry-run',json={'batchId':flow['batchId'],'rows':[{'studentNo':'outside',**fields()}]})
        assert unknown.json()['data']['invalidRows'] == 1


def test_two_row_import_second_conflict_rolls_back_first_and_its_audit(flow):
    with get_sessionmaker()() as db:
        second = db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id == int(flow['batches'][1]['id'])))
        created = InternshipRecord(tenant_id=int(flow['tenantId']), student_id=second.student_id,
                     batch_id=int(flow['batchId']), status='PREPARING', eligibility_status='PENDING',
                     destination_type='NONE', risk_level='NONE')
        db.add(created); db.flush(); second_id=str(created.id)
        second_no=db.get(StudentProfile,second.student_id).student_no; db.commit()
    with login(flow) as school, login(flow,'STUDENT') as student:
        original=read(student,flow)
        rows=[{'studentNo':original['studentNo'],**fields(original['version'])},
              {'studentNo':second_no,**fields()}]
        pre=school.post(BASE+'/basic-info-import/dry-run',json={'batchId':flow['batchId'],'rows':rows})
        assert pre.status_code == 200 and pre.json()['data']['invalidRows']==0, pre.text
        change=school.put(BASE+'/'+second_id+'/basic-info',params={'batchId':flow['batchId']},json=fields(sourceRegion='湖北省 武汉市'))
        assert change.status_code==200,change.text
        confirmed=school.post(BASE+'/basic-info-import/confirm',json={'batchId':flow['batchId'],'rows':pre.json()['data']['rows']})
        assert confirmed.status_code==400 and confirmed.json()['bizCode']=='DATA_CONFLICT',confirmed.text
        assert read(student,flow)==original
        with get_sessionmaker()() as db:
            assert not db.scalars(select(AuditOutbox).where(AuditOutbox.tenant_id==int(flow['tenantId']),
                AuditOutbox.event_type=='INTERNSHIP_BASIC_INFO_IMPORT')).all()


@pytest.mark.parametrize('dimension', ['present', 'missing', 'understated'])
def test_xlsx_template_parse_preview_and_error_workbook_are_real_files(flow, dimension):
    from openpyxl import Workbook
    with login(flow) as school,login(flow,'STUDENT') as student:
        info=read(student,flow)
        template=school.get(BASE+'/basic-info-import/template')
        assert template.status_code==200,template.text
        assert list(load_workbook(BytesIO(template.content)).worksheets[0].values)[0] == ('学号 *','生源地','民族','性别')
        book=Workbook();sheet=book.active;sheet.title='导入模板'
        sheet.append(['学号','生源地','民族','性别']);sheet.append([info['studentNo'],'湖南省 益阳市','汉族','男'])
        buffer=BytesIO();book.save(buffer)
        content=buffer.getvalue()
        if dimension != 'present':
            import re, zipfile
            changed=BytesIO()
            with zipfile.ZipFile(BytesIO(content)) as source, zipfile.ZipFile(changed,'w',zipfile.ZIP_DEFLATED) as target:
                for item in source.infolist():
                    data=source.read(item.filename)
                    if item.filename=='xl/worksheets/sheet1.xml':
                        replacement=b'' if dimension=='missing' else b'<dimension ref="A1"/>'
                        data=re.sub(rb'<dimension[^>]*/>',replacement,data,count=1)
                    target.writestr(item,data)
            content=changed.getvalue()
        pre=school.post(BASE+'/basic-info-import/xlsx',params={'batchId':flow['batchId']},files={'file':('基础资料.xlsx',content,'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')})
        assert pre.status_code==200 and pre.json()['data']['invalidRows']==0,pre.text
        preview=pre.json()['data'];assert preview['rows'][0]['expectedVersion']==info['version']
        assert read(student,flow)==info
        errors=school.post(BASE+'/basic-info-import/errors-xlsx',json={'rows':preview['rows'],'errors':[{'rowNo':1,'field':'sourceRegion','message':'请核对生源地'}]})
        assert errors.status_code==200,errors.text
        error_values=list(load_workbook(BytesIO(errors.content)).active.values)
        assert any('请核对生源地' in str(row) for row in error_values)
        assert error_values[-1][-2]=='生源地'
        parsed=xlsx_util.read_xlsx(errors.content,{'学号':'studentNo','生源地':'sourceRegion','民族':'nation','性别':'gender'})
        assert parsed==[{'studentNo':info['studentNo'],'sourceRegion':'湖南省 益阳市','nation':'汉族','gender':'男'}]


@pytest.mark.parametrize('field,value', [('sourceRegion','未采集'),('nation','x'),('gender','无效性别')])
def test_import_invalid_field_points_to_actual_column_without_writing(flow, field, value):
    with login(flow) as school, login(flow,'STUDENT') as student:
        before=read(student,flow)
        row={'studentNo':before['studentNo'],**fields(),field:value}
        pre=school.post(BASE+'/basic-info-import/dry-run',json={'batchId':flow['batchId'],'rows':[row]})
        assert pre.status_code==200,pre.text
        data=pre.json()['data']
        assert data['invalidRows']==1 and data['errors'][0]['field']==field,data
        assert 'recordId' not in data['rows'][0] and 'expectedVersion' not in data['rows'][0]
        assert read(student,flow)==before
