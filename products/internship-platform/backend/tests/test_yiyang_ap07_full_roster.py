"""AP07: normal role requests, real MySQL, shared filters and actual xlsx bytes."""
import base64
import os
from datetime import date, datetime
from io import BytesIO
from decimal import Decimal
import pytest
from openpyxl import load_workbook
from sqlalchemy import event, select
from app.core.context import get_tenant, set_tenant
from app.db.session import get_engine, get_sessionmaker
from app.models import (College, EmpCompany, InternshipAgreement, InternshipApplication, InternshipBatchPlan,
    InternshipPlanAssignment, InternshipPosition, InternshipRecord, InternshipRotation, InternshipPlacementSnapshot,
    InternshipPayrollStatement, InternshipPayrollVersion, Major,
    Role, SchoolClass, StudentAccountLink, StudentProfile, TeacherStudentScope, User, UserRole)
from app.modules.internship.services.internship_placement_snapshot_service import capture_placement_snapshot_in_tx, snapshot_sha256
from test_yiyang_plan_export_options_mysql import seed_workspace, login

ROOT='/api/v1/internship/intern-students'


def data(response):
    assert response.status_code == 200, response.text
    payload=response.json()
    assert payload['code'] == 0, payload
    return payload['data']


def workbook(result):
    book=load_workbook(BytesIO(base64.b64decode(result['contentBase64'])),data_only=False)
    rows=list(book.active.iter_rows(values_only=True))
    header=next(index for index,row in enumerate(rows) if row[0] == '学号')
    return book, [dict(zip(rows[header],row)) for row in rows[header+1:] if row[0] is not None]


@pytest.fixture(scope='module')
def roster():
    if os.environ.get('GAP09_MYSQL_ACCEPTANCE') != '1':
        pytest.skip('explicit isolated MySQL opt-in required')
    ws=seed_workspace(2)
    tid=int(ws['tenantId']); batch=int(ws['batches'][0]['id'])
    with get_sessionmaker()() as db:
        teacher=db.get(User,ws['users']['INTERN_MENTOR']); teacher.real_name='校内甲教师'
        db.get(User,ws['users']['SCHOOL_ADMIN']).real_name='一览表学校管理员'
        college=College(tenant_id=tid,college_name='信息工程学院',status='ACTIVE'); db.add(college); db.flush()
        major=Major(tenant_id=tid,college_id=college.id,major_name='软件技术',status='ACTIVE'); db.add(major); db.flush()
        cls=SchoolClass(tenant_id=tid,major_id=major.id,class_name='软件2401班',grade='2024',head_teacher_id=teacher.id,status='ACTIVE'); db.add(cls); db.flush()
        first=db.scalar(select(InternshipRecord).where(InternshipRecord.batch_id == batch))
        first.advisor_name='历史导师称呼'
        student=db.get(StudentProfile,first.student_id)
        student.real_name='合成甲学生'; student.student_no='0009007199254740993111'
        student.class_id=cls.id; student.grade=student.college_id=student.major_id=None
        student.source_region='湖南省 益阳市'; student.gender='女'
        db.add(StudentAccountLink(tenant_id=tid,student_id=student.id,user_id=ws['users']['STUDENT'],link_status='ACTIVE',
            source='IDENTITY_IMPORT',bound_student_no=student.student_no,bound_login_name='student'))
        app=InternshipApplication(tenant_id=tid,record_id=first.id,student_id=student.id,batch_id=batch,
            application_type='SELF_ARRANGED',volunteer_no=0,status='PENDING_REVIEW',company_name='合成自主智造单位',position_name='合成研发%岗位',
            company_credit_code='000000000000000001',company_nature='民营',company_industry='制造业',
            company_registered_address='合成注册地址',work_country='中国',work_province='湖南省',work_city='益阳市',work_district='赫山区',
            work_address='合成研发中心',internship_department='研发部',enterprise_mentor_name='企业甲教师',
            position_category='技术岗',internship_mode='自主实习',major_match=True,agreed_salary=Decimal('3500.25'),
            internship_start_date=datetime(2026,9,1),internship_end_date=datetime(2026,12,31),registry_verification_status='UNVERIFIED')
        db.add(app); db.flush()
        ws['firstId']=str(first.id); ws['firstStudentId']=str(student.id); ws['appId']=str(app.id)
        ws['recordVersion']=first.version; ws['appVersion']=app.version; ws['collegeName']=college.college_name
        primary=db.get(InternshipBatchPlan,int(ws['batches'][0]['planId']))
        extra=db.get(InternshipBatchPlan,int(ws['batches'][1]['planId'])); extra.title='附加认知方案'; extra.internship_type='COGNITIVE'
        for index,plan in enumerate([primary,extra]):
            db.add(InternshipPlanAssignment(tenant_id=tid,internship_id=first.id,student_id=student.id,plan_id=plan.id,
                plan_batch_id=plan.batch_id,is_primary=index==0,status='ACTIVE'))
        db.add(InternshipAgreement(tenant_id=tid,internship_id=first.id,student_id=student.id,batch_id=batch,status='EFFECTIVE'))
        for index,department,score in [(1,'装配部',85),(2,'检验部',50)]:
            db.add(InternshipRotation(tenant_id=tid,internship_id=first.id,student_id=student.id,batch_id=batch,
                rotation_seq=index,department_name=department,mentor_name='轮岗甲教师',start_date=date(2026,9,1),
                end_date=date(2026,9,30),status='COMPLETED',total_score=Decimal(score)))
        for index in range(1,4):
            stu=StudentProfile(tenant_id=tid,student_no=f'AP07-{index}',real_name=f'合成对照学生{index}',class_id=cls.id,
                current_stage='ENROLLED',student_status='NORMAL',status='ACTIVE')
            db.add(stu); db.flush()
            rec=InternshipRecord(tenant_id=tid,student_id=stu.id,batch_id=batch,status='READY',destination_type='NONE',
                advisor_user_id=ws['users']['SCHOOL_ADMIN'],advisor_name='学校管理员')
            db.add(rec); db.flush()
            if index in (1,2):
                company=EmpCompany(tenant_id=tid,name=f'合成当前单位{index}',credit_code=f'91300000000000000{index}',nature='国有',industry='软件服务',status='ACTIVE')
                db.add(company); db.flush()
                position=InternshipPosition(tenant_id=tid,company_id=company.id,title=f'当前岗位{index}',category='工程岗',
                    work_location='长沙市',work_address=f'合成岗位地址{index}',remuneration_amount=Decimal('4200') if index==1 else Decimal('0'),
                    remuneration_cycle='MONTHLY',mentor_name=f'当前企业教师{index}',status='PUBLISHED')
                db.add(position); db.flush()
                rec.enterprise_id=company.id; rec.position_id=position.id; rec.destination_type='ASSIGNED'
                rec.enterprise_name=company.name; rec.position_name=position.title
                approved=InternshipApplication(tenant_id=tid,record_id=rec.id,student_id=stu.id,batch_id=batch,
                    campaign_id=987654321 if index==1 else None,application_type='POSITION',status='APPROVED',position_id=position.id,
                    company_name='错误历史单位',position_name='历史岗位',company_credit_code='历史信用代码',agreed_salary=Decimal('9999'),
                    company_nature='历史性质',company_industry='历史行业',internship_mode='集中实习',major_match=False)
                db.add(approved); db.flush()
                previous_tenant=get_tenant()
                try:
                    set_tenant({'tenantId':str(tid)})
                    capture_placement_snapshot_in_tx(db,record=rec,position=position,company=company,rights={'passed':True},
                        source_application_id=approved.id if index==1 else None)
                finally:
                    set_tenant(previous_tenant)
                ws['recruitId' if index==1 else 'directId']=str(rec.id)
            else:
                rec.destination_type='SELF_ARRANGED'; rec.enterprise_name='历史自主单位'; rec.position_name='历史自主岗位'
                db.add(InternshipApplication(tenant_id=tid,record_id=rec.id,student_id=stu.id,batch_id=batch,
                    application_type='SELF_ARRANGED',volunteer_no=0,status='APPROVED',company_name=rec.enterprise_name,
                    position_name=rec.position_name,company_credit_code='不应猜测的旧信用代码',agreed_salary=Decimal('9999')))
                ws['unattributedId']=str(rec.id)
        mentor_role=db.scalar(select(Role).where(Role.tenant_id==tid,Role.role_code=='INTERN_MENTOR'))
        for username,role_code in [('same_name','INTERN_MENTOR'),('scope_manager','COLLEGE_ADMIN')]:
            role=mentor_role
            if role_code=='COLLEGE_ADMIN':
                role=Role(tenant_id=tid,role_code=role_code,role_name='学院管理员',role_type='SYSTEM',status='ACTIVE'); db.add(role); db.flush()
            user=User(tenant_id=tid,login_name=username,real_name=teacher.real_name if username=='same_name' else '专业范围管理人',
                password_hash=teacher.password_hash,user_type='TEACHER',status='ACTIVE',credential_version=0,must_change_password=False)
            db.add(user); db.flush(); db.add(UserRole(tenant_id=tid,user_id=user.id,role_id=role.id,status='ACTIVE'))
        db.add(TeacherStudentScope(tenant_id=tid,teacher_key='scope_manager',role_code='COLLEGE_ADMIN',scope_type='MAJOR',ref_value='软件技术',status='ACTIVE'))
        db.commit()
    with login(ws) as client:
        data(client.post('/api/v1/internship/applications/'+ws['appId']+'/review',json={'action':'APPROVE',
            'comment':'合成资料已核验','expectedVersion':ws['appVersion'],'recordExpectedVersion':ws['recordVersion'],'batchId':str(batch)}))
    return ws


@pytest.fixture(scope='module')
def admin(roster):
    with login(roster) as client:
        yield client


def test_complete_fields_current_sources_and_chinese_xlsx(roster,admin):
    params={'batchId':roster['batches'][0]['id']}
    listed=data(admin.get(ROOT,params=params)); rows={row['id']:row for row in listed['items']}
    assert listed['total']==4 and len(rows)==4
    first=rows[roster['firstId']]
    expected={'grade':'2024','collegeName':'信息工程学院','majorName':'软件技术','className':'软件2401班',
        'headTeacherName':'校内甲教师','advisorName':'校内甲教师','mentorName':'企业甲教师','internshipMode':'自主实习',
        'sourceApplicationId':roster['appId'],'companyCreditCode':'000000000000000001','companyNature':'民营','companyIndustry':'制造业',
        'applicationStatusLabel':'已通过','agreementStatusLabel':'已生效','majorMatchLabel':'对口','registryVerificationStatusLabel':'未核验'}
    for key,value in expected.items(): assert first[key]==value,key
    assert set(first['planIds'])=={item['planId'] for item in roster['batches']}
    assert '附加认知方案' in first['planNames'] and '认知实习' in first['internshipTypeLabel']
    assert first['agreedSalary']==3500.25
    assert all(value in first['rotationInfo'] for value in ['装配部','检验部','轮岗甲教师','2026-09-01','已完成'])
    assert all(value in first['workLocation'] for value in ['中国','湖南省','益阳市','赫山区','合成研发中心'])
    recruit=rows[roster['recruitId']]; direct=rows[roster['directId']]; orphan=rows[roster['unattributedId']]
    assert recruit['companyCreditCode']=='913000000000000001' and recruit['agreedSalary']==4200 and recruit['internshipMode']=='集中实习'
    assert direct['companyNature']=='国有' and direct['agreedSalary']==0 and direct['sourceApplicationId'] is None
    assert orphan['enterpriseName']=='历史自主单位' and orphan['companyCreditCode'] is None and orphan['agreedSalary'] is None
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(roster['recruitId']))
        snapshot=db.get(InternshipPlacementSnapshot,record.current_placement_snapshot_id)
        assert Decimal(snapshot.snapshot_json['remunerationAmount'])==Decimal('4200')
        assert snapshot_sha256(snapshot.snapshot_json)==snapshot.snapshot_sha256
    result=data(admin.post(ROOT+'/export',params=params)); book,exported=workbook(result)
    assert result['totalRows']==result['exportedRows']==4 and result['truncated'] is False
    assert {row['学号'] for row in exported}=={row['studentNo'] for row in rows.values()}
    line=next(row for row in exported if row['学号']==first['studentNo'])
    assert line['实习计划']==first['planNames'] and line['企业老师']=='企业甲教师' and line['班主任']=='校内甲教师'
    assert line['三方协议状态']=='已生效' and line['约定报酬币种']=='人民币' and line['工商核验状态']=='未核验'
    assert line['部门轮岗信息']==first['rotationInfo'] and line['实习地点']==first['workLocation']
    codes=[cell for row in book.active for cell in row if cell.value=='000000000000000001']
    assert codes and all(cell.data_type=='s' for cell in codes)


def test_every_required_filter_and_export_use_same_predicates(roster,admin):
    conditions={'studentName':'合成甲','studentNo':'0009007199254740993111','grade':'2024','collegeName':'信息工程',
        'majorName':'软件','className':'2401','planName':'附加认知','internshipType':'COGNITIVE','internshipMode':'自主',
        'enterpriseName':'合成自主','companyCreditCode':'000000000000000001','positionName':'研发%',
        'salaryMin':'3500.25','salaryMax':'3500.25','workLocation':'中国 湖南省 益阳市 赫山区 合成研发中心','headTeacherName':'校内甲',
        'advisorName':'校内甲','mentorName':'企业甲','applicationStatus':'APPROVED','companyNature':'民营',
        'companyIndustry':'制造','majorMatch':'true','sourceRegion':'益阳','rotationInfo':'装配','rotationScoreMin':'80',
        'rotationScoreMax':'90','agreementStatus':'EFFECTIVE','status':'READY'}
    base={'batchId':roster['batches'][0]['id'],'studentNo':'0009007199254740993111'}
    for key,value in conditions.items():
        params={**base,key:value}
        found=data(admin.get(ROOT,params=params))
        assert found['total']==1 and found['items'][0]['id']==roster['firstId'],key
        _,exported=workbook(data(admin.post(ROOT+'/export',params=params)))
        assert len(exported)==1 and exported[0]['学号']==base['studentNo'],key
    for params in [{'planName':'附加认知','internshipType':'POST'}, {'rotationInfo':'装配','rotationScoreMax':'60'},
                   {'companyCreditCode':'不存在'}, {'majorMatch':'false'}, {'salaryMin':'3500.26'}, {'salaryMax':'3500.24'}]:
        assert data(admin.get(ROOT,params={**base,**params}))['total']==0,params
    percent=data(admin.get(ROOT,params={'batchId':base['batchId'],'positionName':'%'}))
    assert percent['total']==1 and percent['items'][0]['id']==roster['firstId']
    zero=data(admin.get(ROOT,params={'batchId':base['batchId'],'salaryMin':'0','salaryMax':'0','majorMatch':'false'}))
    assert zero['total']==0  # Direct placement has no inferred professional-match claim.
    zero=data(admin.get(ROOT,params={'batchId':base['batchId'],'salaryMin':'0','salaryMax':'0'}))
    assert zero['total']==1 and zero['items'][0]['id']==roster['directId']


def test_database_pagination_query_count_and_cross_batch_aggregation(roster,admin):
    base={'batchId':roster['batches'][0]['id']}; ids=[]
    for page in range(1,5):
        result=data(admin.get(ROOT,params={**base,'page':page,'pageSize':1}))
        assert result['total']==4 and len(result['items'])==1
        ids.append(result['items'][0]['id'])
    assert len(set(ids))==4
    assert data(admin.get(ROOT,params={**base,'page':5,'pageSize':1}))['items']==[]
    counts=[]
    for size in [1,4]:
        queries=[]
        def observe(conn,cursor,statement,parameters,context,many):
            if statement.lstrip().upper().startswith('SELECT'): queries.append(statement)
        event.listen(get_engine(),'before_cursor_execute',observe)
        try: data(admin.get(ROOT,params={**base,'pageSize':size}))
        finally: event.remove(get_engine(),'before_cursor_execute',observe)
        counts.append(len(queries))
        assert any('LIMIT' in statement.upper() and 't_internship_record' in statement for statement in queries)
    assert counts[1]<=counts[0]+1,counts
    result=data(admin.get(ROOT,params={**base,'allBatches':'true'}))
    assert result['total']==5 and len({row['id'] for row in result['items']})==5
    exported=data(admin.post(ROOT+'/export',params={'allBatches':'true'}))
    _,rows=workbook(exported)
    assert len(rows)==5 and {row['实习批次'] for row in rows}=={item['batchName'] for item in result['items']}


def test_normal_teacher_same_name_denial_and_major_scope_detail_parity(roster):
    params={'batchId':roster['batches'][0]['id']}
    with login(roster,'INTERN_MENTOR') as teacher:
        own=data(teacher.get(ROOT,params={**params,'allBatches':'true','pageSize':1}))
        assert own['total']==1 and own['items'][0]['id']==roster['firstId']
        assert teacher.get(ROOT+'/'+roster['directId']).status_code==403
    with login(roster,'SAME_NAME') as stranger:
        assert data(stranger.get(ROOT,params={**params,'allBatches':'true'}))['total']==0
        assert stranger.get(ROOT+'/'+roster['firstId']).status_code==403
    with login(roster,'SCOPE_MANAGER') as scoped:
        found=data(scoped.get(ROOT,params=params))
        assert found['total']==4
        assert data(scoped.get(ROOT+'/'+roster['firstId']))['id']==roster['firstId']
    with login(roster,'STUDENT') as student:
        assert student.get(ROOT,params=params).status_code==403
        assert student.post(ROOT+'/export',params=params).status_code==403
    other=seed_workspace(1)
    with login(other) as foreign:
        assert foreign.get(ROOT,params=params).status_code==404
        assert foreign.get(ROOT+'/'+roster['firstId']).status_code==404
        assert foreign.post(ROOT+'/export',params=params).status_code==404


def test_filter_validation_is_shared_and_rejects_unknown_or_invalid_conditions(roster,admin):
    base={'batchId':roster['batches'][0]['id']}
    for wrong in [{'internshipType':'NOT_A_TYPE'},{'salaryMin':'-1'},{'salaryMin':'nan'},
                  {'salaryMin':'100','salaryMax':'99'},{'rotationScoreMin':'101'},{'majorMatch':'maybe'},{'unknownField':'x'}]:
        for method,path in [(admin.get,ROOT),(admin.post,ROOT+'/export')]:
            assert method(path,params={**base,**wrong}).status_code==422,wrong


def test_current_source_survives_same_destination_and_is_cleared_by_real_change(roster,admin):
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(roster['firstId'])); version=record.version
    body={'destination':'SELF_ARRANGED','reason':'保持当前已批准去向','expectedVersion':version}
    data(admin.post(ROOT+'/'+roster['firstId']+'/destination',json=body))
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(roster['firstId']))
        assert str(record.current_application_id)==roster['appId']
        body.update(destination='NONE',reason='合成测试解除去向',expectedVersion=record.version)
    data(admin.post(ROOT+'/'+roster['firstId']+'/destination',json=body))
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(roster['firstId']))
        assert record.current_application_id is None and record.current_placement_snapshot_id is None
        assert record.enterprise_name is None and record.position_name is None and record.enterprise_mentor_name is None


def test_pending_self_application_cannot_overwrite_an_already_placed_record(roster,admin):
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(roster['directId']))
        app=InternshipApplication(tenant_id=int(roster['tenantId']),record_id=record.id,student_id=record.student_id,
            batch_id=record.batch_id,application_type='SELF_ARRANGED',volunteer_no=0,status='PENDING_REVIEW',company_name='不得覆盖的自主单位')
        db.add(app); db.commit()
        app_id=str(app.id); app_version=app.version; record_version=record.version
        snapshot_id=record.current_placement_snapshot_id; position_id=record.position_id
    rejected=admin.post('/api/v1/internship/applications/'+app_id+'/review',json={'action':'APPROVE',
        'comment':'不应覆盖已落岗','expectedVersion':app_version,'recordExpectedVersion':record_version,'batchId':roster['batches'][0]['id']})
    assert rejected.status_code==409,rejected.text
    with get_sessionmaker()() as db:
        record=db.get(InternshipRecord,int(roster['directId']))
        assert record.current_placement_snapshot_id==snapshot_id and record.position_id==position_id
        assert record.current_application_id is None and db.get(InternshipApplication,int(app_id)).status=='PENDING_REVIEW'


def test_malformed_relationships_do_not_change_roster_facts_or_leak_detail():
    if os.environ.get('GAP09_MYSQL_ACCEPTANCE') != '1':
        pytest.skip('explicit isolated MySQL opt-in required')
    ws=seed_workspace(2); foreign=seed_workspace(1)
    tid=int(ws['tenantId']); batch=int(ws['batches'][0]['id']); other_batch=int(ws['batches'][1]['id'])
    with get_sessionmaker()() as db:
        rec=db.scalar(select(InternshipRecord).where(InternshipRecord.tenant_id==tid,InternshipRecord.batch_id==batch))
        other=db.scalar(select(InternshipRecord).where(InternshipRecord.tenant_id==tid,InternshipRecord.batch_id==other_batch))
        for index,plan_info in enumerate(ws['batches']):
            db.add(InternshipPlanAssignment(tenant_id=tid,internship_id=rec.id,student_id=rec.student_id,
                plan_id=int(plan_info['planId']),plan_batch_id=batch,is_primary=index==0,status='ACTIVE'))
        db.add(InternshipAgreement(tenant_id=tid,internship_id=rec.id,student_id=rec.student_id,batch_id=batch,status='EFFECTIVE'))
        db.add(InternshipAgreement(tenant_id=tid,internship_id=rec.id,student_id=rec.student_id,batch_id=other_batch,status='DRAFT'))
        for index,month,owner,amount in [(1,'2026-09',rec,3200),(2,'2026-10',other,9999)]:
            statement=InternshipPayrollStatement(tenant_id=tid,internship_id=rec.id,student_id=owner.student_id,
                batch_id=owner.batch_id,pay_month=month)
            db.add(statement); db.flush()
            version=InternshipPayrollVersion(tenant_id=tid,statement_id=statement.id,revision_no=1,
                actual_amount=Decimal(amount),status='APPROVED',is_current=True)
            db.add(version); db.flush(); statement.current_version_id=version.id
        db.add(InternshipRotation(tenant_id=tid,internship_id=rec.id,student_id=other.student_id,batch_id=other_batch,
            rotation_seq=1,department_name='错误轮岗部门',mentor_name='错误教师',start_date=date(2026,9,1),
            end_date=date(2026,9,30),total_score=Decimal(99)))
        foreign_record=db.scalar(select(InternshipRecord).where(InternshipRecord.tenant_id==int(foreign['tenantId'])))
        foreign_company=EmpCompany(tenant_id=int(foreign['tenantId']),name='外校单位',status='ACTIVE')
        db.add(foreign_company); db.flush()
        foreign_position=InternshipPosition(tenant_id=int(foreign['tenantId']),company_id=foreign_company.id,title='外校保密岗位',status='PUBLISHED')
        db.add(foreign_position); db.flush()
        rec.position_id=foreign_position.id
        record_id=str(rec.id); student_id=rec.student_id; foreign_student_id=foreign_record.student_id
        db.commit()
    params={'batchId':str(batch)}
    with login(ws) as client:
        row=data(client.get(ROOT,params=params))['items'][0]
        assert row['planIds']==[ws['batches'][0]['planId']]
        assert row['agreementStatus']=='EFFECTIVE' and row['latestActualSalary']==3200 and row['latestPayrollMonth']=='2026-09'
        assert row['rotationCount']==0 and data(client.get(ROOT,params={**params,'rotationInfo':'错误轮岗'}))['total']==0
        assert data(client.get(ROOT+'/'+record_id))['position'] is None
        assert data(client.get(ROOT,params={**params,'planName':ws['batches'][1]['title']}))['total']==0
        assert data(client.get(ROOT,params={**params,'agreementStatus':'DRAFT'}))['total']==0
        with get_sessionmaker()() as db:
            assignment=db.scalar(select(InternshipPlanAssignment).where(InternshipPlanAssignment.internship_id==int(record_id),
                InternshipPlanAssignment.plan_id==int(ws['batches'][1]['planId'])))
            assignment.plan_batch_id=other_batch
            from app.models import InternshipBatch
            db.get(InternshipBatch,other_batch).is_deleted=True; db.commit()
        assert data(client.get(ROOT,params={**params,'planName':ws['batches'][1]['title']}))['total']==0
        with get_sessionmaker()() as db:
            db.get(InternshipRecord,int(record_id)).student_id=foreign_student_id; db.commit()
        assert data(client.get(ROOT,params=params))['total']==0
        assert client.get(ROOT+'/'+record_id).status_code==404
        with get_sessionmaker()() as db:
            db.get(InternshipRecord,int(record_id)).student_id=student_id
            db.get(StudentProfile,student_id).is_deleted=True; db.commit()
        assert data(client.get(ROOT,params=params))['total']==0
        assert client.get(ROOT+'/'+record_id).status_code==404
        with get_sessionmaker()() as db:
            db.get(StudentProfile,student_id).is_deleted=False
            db.get(InternshipRecord,int(record_id)).batch_id=other_batch; db.commit()
        assert client.get(ROOT+'/'+record_id).status_code==404
        with get_sessionmaker()() as db:
            db.get(InternshipRecord,int(record_id)).batch_id=int(foreign['batches'][0]['id']); db.commit()
        assert client.get(ROOT+'/'+record_id).status_code==404


def test_preloaded_scope_does_not_grant_non_advisors_by_account_id():
    from types import SimpleNamespace
    from app.modules.internship.services.internship_advisor_identity_guard import _domain_rec_in_scope_pre
    student=SimpleNamespace(id=1,student_no='S1',class_id=None,college_id=None)
    record=SimpleNamespace(advisor_user_id=42)
    scope={'mode':'SCOPED','roleCode':'ACADEMIC_TEACHER','advisorUserIds':{42},'advisorNames':{'同名教师'}}
    assert not _domain_rec_in_scope_pre(scope,record,student,{},{})
    assert _domain_rec_in_scope_pre({**scope,'studentNos':{'S1'}},record,student,{},{})
    for role in ('INTERN_MENTOR','GD_MENTOR'):
        assert _domain_rec_in_scope_pre({**scope,'roleCode':role},record,student,{},{})
        assert not _domain_rec_in_scope_pre({**scope,'roleCode':role,'advisorUserIds':{43}},record,student,{},{})
