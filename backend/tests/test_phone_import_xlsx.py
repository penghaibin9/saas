"""Actual XLSX parser/staging regressions; spreadsheet bytes contain synthetic data only."""
import io
import secrets
import pytest
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from tests.test_phone_login_mysql import phone_db, phone_identity


def workbook_bytes(kind, phone='13800138000', owner='SELF', contact=''):
    from app.services.identity_import_file_service import STUDENT_HEADERS, TEACHER_HEADERS
    wb = Workbook()
    ws = wb.active
    ws.title = '导入模板'
    headers = STUDENT_HEADERS if kind == 'STUDENT' else TEACHER_HEADERS
    ws.append(headers)
    row = {'学号': '000012345678901234', '工号': 'T0000001', '姓名': '虚构测试',
        '班级名称': '测试班', '预设角色编码': 'ACADEMIC_TEACHER',
        '本人手机号': phone, '号码归属': owner, '联系手机号': contact}
    ws.append([row.get(h, '') for h in headers])
    out = io.BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


@pytest.mark.parametrize('kind', ['STUDENT', 'TEACHER'])
def test_both_parsers_preserve_phone_columns_and_account_precision(kind, tmp_path):
    from app.services import identity_import_file_service as files
    from app.services.identity_import_path_parser import parse_identity_xlsx_path
    content = workbook_bytes(kind, '+8613800138000')
    path = tmp_path / 'synthetic.xlsx'
    path.write_bytes(content)
    first = (files.parse_student_xlsx if kind == 'STUDENT' else files.parse_teacher_xlsx)(content, 'synthetic.xlsx')
    second = parse_identity_xlsx_path(path, 'synthetic.xlsx', kind)
    key = 'students' if kind == 'STUDENT' else 'teachers'
    for parsed in (first, second):
        assert parsed['errors'] == []
        assert parsed[key][0]['selfPhone'] == '+8613800138000'
        assert parsed[key][0]['phoneOwnerType'] == 'SELF'
    assert first[key] == second[key]
    if kind == 'STUDENT':
        assert first[key][0]['studentNo'] == '000012345678901234'


@pytest.mark.parametrize('value', ['=CONCAT("138","00138000")', '1.3800138E10', '138****8000', True, 13800138000.5])
def test_phone_formula_mask_scientific_text_boolean_fraction_rejected(value):
    from app.services.identity_import_file_service import parse_teacher_xlsx
    parsed = parse_teacher_xlsx(workbook_bytes('TEACHER', value), 'synthetic.xlsx')
    assert any(e['field'] == '本人手机号' for e in parsed['errors'])
    assert str(value) not in str(parsed['errors'])


def test_exact_numeric_phone_is_warning_not_verified():
    from app.services.identity_import_file_service import parse_teacher_xlsx
    parsed = parse_teacher_xlsx(workbook_bytes('TEACHER', 13800138000), 'synthetic.xlsx')
    assert not parsed['errors']
    assert parsed['teachers'][0]['_phoneWarnings']
    assert parsed['teachers'][0]['selfPhone'] == '+8613800138000'
    assert 'phoneVerified' not in parsed['teachers'][0]


def test_staging_encrypts_phone_and_roundtrips_without_changing_digest(phone_identity, tmp_path):
    from app.services.identity_import_staging_service import stage_identity_xlsx, StagingRowSequence, staging_fingerprint
    from app.db.session import get_sessionmaker
    from app.models.data_exchange import IdentityImportStagingRow
    path = tmp_path / 'synthetic.xlsx'
    path.write_bytes(workbook_bytes('TEACHER'))
    job_id = secrets.randbelow(10**10) + 10**10
    result = stage_identity_xlsx(path=path, filename=path.name, kind='TEACHER',
        tenant_id=phone_identity['tenant_id'], job_id=job_id, actor_id=phone_identity['user_id'])
    with get_sessionmaker()() as db:
        row = db.scalar(select(IdentityImportStagingRow).where(IdentityImportStagingRow.import_job_id == job_id))
        assert '13800138000' not in str(row.payload_json)
        assert 'selfPhone' not in row.payload_json
        assert row.payload_json['selfPhoneEncrypted']
    decoded = list(StagingRowSequence(phone_identity['tenant_id'], job_id, 'TEACHER'))
    assert decoded[0]['selfPhone'] == '+8613800138000'
    assert staging_fingerprint(phone_identity['tenant_id'], job_id) == (1, result['stagingDigest'])


def test_import_rejects_shared_login_number_and_client_verified_flags():
    from app.services.school_onboarding_service import _validate_rows
    common = {'name': '测试', 'roleCodes': 'ACADEMIC_TEACHER', 'selfPhone': '13800138000', 'phoneOwnerType': 'SELF'}
    errors = _validate_rows({'teachers': [{'loginName': 'a', **common}, {'loginName': 'b', **common}]})
    assert any(e.get('reasonCode') == 'PHONE_DUPLICATE_IN_FILE' for e in errors)
    assert '13800138000' not in str(errors)
    for field in ('phoneVerified', 'verifiedAt', 'credentialVersion', 'bindingStatus'):
        assert _validate_rows({'teachers': [{'loginName': 'a', **common, field: True}]})


def test_identity_audit_failure_cannot_commit_new_account(phone_identity, monkeypatch):
    from app.services import identity_import_service, audit_log
    from app.core.context import get_tenant, set_tenant
    from app.db.session import get_sessionmaker
    from app.models import User
    operator = {'userId': f"db-{phone_identity['user_id']}", 'tenantId': str(phone_identity['tenant_id']),
        'userType': 'ADMIN', 'currentRoleCode': 'SCHOOL_ADMIN'}
    body = {'tenantId': str(phone_identity['tenant_id']), 'teachers': [{'loginName': 'audit-rollback',
        'name': '虚构教师', 'roleCodes': 'ACADEMIC_TEACHER', 'selfPhone': '13700137000', 'phoneOwnerType': 'SELF'}]}
    def broken(action, *args, **kwargs):
        if action == 'ONBOARD_IMPORT':
            raise RuntimeError('injected import audit failure')
    original = audit_log.record_critical_in_session
    def critical(db, action, *args, **kwargs):
        broken(action)
        return original(db, action, *args, **kwargs)
    monkeypatch.setattr(audit_log, 'record', broken)
    monkeypatch.setattr(audit_log, 'record_critical_in_session', critical)
    previous = get_tenant()
    set_tenant(phone_identity['tenant_id'])
    try:
        with pytest.raises(RuntimeError, match='injected import audit failure'):
            identity_import_service.run_identity_import(operator, body, dry_run=False)
        with get_sessionmaker()() as db:
            assert db.scalar(select(User.id).where(User.tenant_id == phone_identity['tenant_id'], User.login_name == 'audit-rollback')) is None
    finally:
        set_tenant(previous)


def test_preview_rejects_verified_occupancy_and_different_existing_candidate(phone_identity):
    from app.services.phone_login_service import preview_import_phones, create_pending_candidate_in_session
    from app.db.session import get_sessionmaker
    report = {'errors': [], 'summary': {}}
    body = {'teachers': [{'loginName': 'new-teacher', 'name': '测试', 'selfPhone': '13800138000', 'phoneOwnerType': 'SELF'}]}
    with get_sessionmaker()() as db:
        preview_import_phones(db, phone_identity['tenant_id'], body, report)
        assert any(e.get('reasonCode') == 'PHONE_OCCUPIED' for e in report['errors'])
        create_pending_candidate_in_session(db, tenant_id=phone_identity['tenant_id'], user_id=phone_identity['user_id'],
            phone='13900139000', source_kind='SELF_SERVICE', expected_version=0)
        db.commit()
    body['teachers'][0].update(loginName=phone_identity['login'], selfPhone='13700137000')
    with get_sessionmaker()() as db:
        report = {'errors': [], 'summary': {}}
        preview_import_phones(db, phone_identity['tenant_id'], body, report)
        assert report['errors']


def test_original_teacher_writer_creates_only_candidate_and_preserves_password_on_retry(phone_identity):
    from app.services.identity_import_service import run_identity_import
    from app.core.context import get_tenant, get_current_user_ctx, set_tenant, set_current_user
    from app.db.session import get_sessionmaker
    from app.models import User, PhoneLoginCandidate, PhoneLoginBinding, PasswordResetSmsJob
    from sqlalchemy import func
    previous = get_tenant(), get_current_user_ctx()
    operator = {'userId': f"db-{phone_identity['user_id']}", 'tenantId': str(phone_identity['tenant_id']),
                'userType': 'ADMIN', 'currentRoleCode': 'SCHOOL_ADMIN'}
    set_tenant(phone_identity['tenant_id'])
    set_current_user(operator)
    body = {'tenantId': str(phone_identity['tenant_id']), 'teachers': [{'loginName': 'phone-import-teacher',
        'name': '虚构导入教师', 'roleCodes': 'ACADEMIC_TEACHER', 'selfPhone': '+8613700137000',
        'phoneOwnerType': 'SELF', 'contactPhone': '+8613600136000', '_rowNo': 2}]}
    try:
        preview = run_identity_import(operator, body, dry_run=True)
        assert not preview['errors']
        assert preview['phoneSummary']['phonePending'] == 1
        result = run_identity_import(operator, body, dry_run=False)
        assert result['phoneWriteSummary']['CREATED'] == 1
        with get_sessionmaker()() as db:
            user = db.scalar(select(User).where(User.tenant_id == phone_identity['tenant_id'], User.login_name == 'phone-import-teacher'))
            original_id, original_hash = user.id, user.password_hash
            candidate = db.scalar(select(PhoneLoginCandidate).where(PhoneLoginCandidate.user_id == original_id))
            assert candidate.state == 'PENDING' and candidate.source_row_no == 2
            assert db.scalar(select(func.count()).select_from(PhoneLoginBinding).where(PhoneLoginBinding.user_id == original_id)) == 0
            assert db.scalar(select(func.count()).select_from(PasswordResetSmsJob).where(PasswordResetSmsJob.user_id == original_id)) == 0
        again = run_identity_import(operator, body, dry_run=False)
        assert again['teacherCredentials'] == []
        with get_sessionmaker()() as db:
            user = db.get(User, original_id)
            assert user.password_hash == original_hash and user.login_name == 'phone-import-teacher'
    finally:
        set_tenant(previous[0])
        set_current_user(previous[1])


def test_real_20k_staging_remains_encrypted_keyset_and_masked_error_workbook(phone_identity, tmp_path):
    from app.services.identity_import_file_service import TEACHER_HEADERS
    from app.services import identity_import_staging_service as svc
    from app.models.data_exchange import IdentityImportStagingRow, ImportRowError
    from app.db.session import get_sessionmaker
    wb = Workbook(write_only=True)
    ws = wb.create_sheet('导入模板')
    ws.append(TEACHER_HEADERS)
    for i in range(20_000):
        ws.append([f'T{i:018d}', '虚构测试教师', '', '', 'ACADEMIC_TEACHER', '', '', f'139{i:08d}', 'SELF', ''])
    path = tmp_path / 'synthetic-20k.xlsx'
    wb.save(path)
    wb.close()
    job = secrets.randbelow(10**10) + 10**10
    tid = phone_identity['tenant_id']
    result = svc.stage_identity_xlsx(path=path, filename=path.name, kind='TEACHER', tenant_id=tid, job_id=job, actor_id=phone_identity['user_id'])
    assert result['totalRows'] == 20_000 and not result['parserErrors']
    assert sum(1 for _ in svc.StagingRowSequence(tid, job, 'TEACHER')) == 20_000
    with get_sessionmaker()() as db:
        row = db.scalar(select(IdentityImportStagingRow).where(IdentityImportStagingRow.tenant_id == tid,
            IdentityImportStagingRow.import_job_id == job).order_by(IdentityImportStagingRow.row_no).limit(1))
        assert '13900000000' not in str(row.payload_json)
        db.add(ImportRowError(tenant_id=tid, import_job_id=job, row_no=2, field_code='selfPhone',
            error_code='PHONE_CONFLICT', error_message='=not-a-formula', raw_snapshot_json={'entity': 'teacher'}))
        db.commit()
    receipt = svc.build_staging_error_workbook(tenant_id=tid, job_id=job)
    read = load_workbook(io.BytesIO(receipt), data_only=False)
    values = list(read.active.values)
    assert values[1][-1] and '****' in values[1][-1]
    assert '13900000000' not in str(values)
    assert read.active.cell(2, 7).data_type != 'f'
    read.close()


def test_original_student_writer_keeps_link_and_never_merges_by_contact(phone_identity):
    from app.services.identity_import_service import run_identity_import
    from app.services.school_onboarding_service import run_onboarding
    from app.core.context import get_tenant, set_tenant
    from app.db.session import get_sessionmaker
    from app.models import User, StudentAccountLink, PhoneLoginCandidate, StudentProfile
    operator = {'userId': f"db-{phone_identity['user_id']}", 'tenantId': str(phone_identity['tenant_id']),
        'userType': 'ADMIN', 'currentRoleCode': 'SCHOOL_ADMIN'}
    previous = get_tenant()
    set_tenant(phone_identity['tenant_id'])
    try:
        run_onboarding(operator, {'tenantId': str(phone_identity['tenant_id']),
            'colleges': [{'name': '测试学院'}], 'majors': [{'name': '测试专业', 'collegeName': '测试学院'}],
            'classes': [{'name': '测试班', 'majorName': '测试专业', 'grade': '2026'}]}, dry_run=False)
        rows = [{'studentNo': f'0000000000000000{i}', 'name': f'虚构学生{i}', 'className': '测试班',
            'selfPhone': f'1370000000{i}', 'phoneOwnerType': 'SELF', 'contactPhone': '13600136000'} for i in (1, 2)]
        result = run_identity_import(operator, {'tenantId': str(phone_identity['tenant_id']), 'students': rows}, dry_run=False)
        assert result['phoneWriteSummary']['CREATED'] == 2
        with get_sessionmaker()() as db:
            profiles = list(db.scalars(select(StudentProfile).where(StudentProfile.tenant_id == phone_identity['tenant_id'])))
            assert len(profiles) == 2 and profiles[0].id != profiles[1].id
            for profile in profiles:
                link = db.scalar(select(StudentAccountLink).where(StudentAccountLink.tenant_id == phone_identity['tenant_id'], StudentAccountLink.student_id == profile.id))
                user = db.get(User, link.user_id)
                assert user.login_name == profile.student_no and user.credential_version == 0
                assert db.scalar(select(PhoneLoginCandidate.state).where(PhoneLoginCandidate.user_id == user.id)) == 'PENDING'
    finally:
        set_tenant(previous)


def test_receipt_generation_failure_rolls_back_original_import_job(phone_identity, monkeypatch):
    """MySQL transaction test; scan readiness is isolated, not a scanner acceptance test."""
    from app.services import identity_import_file_service as files, data_exchange_job_service as jobs
    from app.services.identity_import_service import preview_identity_import
    from app.services import file_scan_service
    from app.core.context import get_tenant, set_tenant
    from app.db.session import get_sessionmaker
    from app.models import User, IdentityImportBatch
    from app.models.data_exchange import ImportJob
    previous = get_tenant()
    set_tenant(phone_identity['tenant_id'])
    operator = {'userId': f"db-{phone_identity['user_id']}", 'tenantId': str(phone_identity['tenant_id']),
        'userType': 'ADMIN', 'currentRoleCode': 'SCHOOL_ADMIN'}
    try:
        parsed = files.parse_teacher_xlsx(workbook_bytes('TEACHER', '13700137000'), 'synthetic.xlsx')
        report = preview_identity_import(operator, {'tenantId': str(phone_identity['tenant_id']), 'teachers': parsed['teachers']})
        batch = files.create_batch(operator, parsed, report)
        job = jobs.create_identity_import_job(kind='TEACHER', source_file_id='123456', parsed=parsed, batch_result=batch, user=operator)
        monkeypatch.setattr(file_scan_service, 'assert_file_ready_for_business', lambda *a, **k: None)
        def broken(*a, **k):
            raise RuntimeError('injected receipt generation failure')
        monkeypatch.setattr(files, 'build_credential_receipt', broken)
        with pytest.raises(RuntimeError, match='injected receipt'):
            jobs.confirm_identity_import_job(job['id'], expected_version=job['version'], user=operator,
                idempotency_key='synthetic-idempotency-key')
        with get_sessionmaker()() as db:
            assert db.scalar(select(User.id).where(User.tenant_id == phone_identity['tenant_id'], User.login_name == 'T0000001')) is None
            assert db.get(ImportJob, int(job['id'])).status == 'VALIDATED'
            assert db.scalar(select(IdentityImportBatch.status).where(IdentityImportBatch.batch_no == batch['batchNo'])) == 'VALIDATED'
    finally:
        set_tenant(previous)


def test_import_receipt_and_business_commit_once_in_mysql(phone_identity, monkeypatch, tmp_path):
    from app.services import identity_import_file_service as files, data_exchange_job_service as jobs
    from app.services.identity_import_service import preview_identity_import
    from app.services import file_scan_service
    from app.services.storage.local import LocalStorageBackend
    from app.core.config import settings
    from app.core.context import get_tenant, set_tenant
    from app.db.session import get_sessionmaker
    from app.models import User, IdentityImportBatch, FileObject
    from app.models.data_exchange import ExportJob
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    monkeypatch.setattr(jobs, 'get_backend', LocalStorageBackend)
    # Only the readiness dependency is isolated; receipt bytes and MySQL are real.
    monkeypatch.setattr(file_scan_service, 'assert_file_ready_for_business', lambda *a, **k: None)
    previous = get_tenant()
    set_tenant(phone_identity['tenant_id'])
    operator = {'userId': f"db-{phone_identity['user_id']}", 'tenantId': str(phone_identity['tenant_id']),
        'userType': 'ADMIN', 'currentRoleCode': 'SCHOOL_ADMIN'}
    try:
        parsed = files.parse_teacher_xlsx(workbook_bytes('TEACHER', '13700137000'), 'synthetic.xlsx')
        report = preview_identity_import(operator, {'tenantId': str(phone_identity['tenant_id']), 'teachers': parsed['teachers']})
        batch = files.create_batch(operator, parsed, report)
        job = jobs.create_identity_import_job(kind='TEACHER', source_file_id='123456', parsed=parsed, batch_result=batch, user=operator)
        result = jobs.confirm_identity_import_job(job['id'], expected_version=job['version'], user=operator,
            idempotency_key='synthetic-idempotency-key')
        assert result['status'] == 'SUCCEEDED' and result['credentialReceiptFileId']
        with get_sessionmaker()() as db:
            file = db.get(FileObject, int(result['credentialReceiptFileId']))
            assert file.visibility == 'PRIVATE' and file.security_level == 'HIGHLY_SENSITIVE'
            wb = load_workbook(tmp_path / file.file_key)
            assert wb.active.max_row >= 2
            wb.close()
            account = db.scalar(select(User).where(User.tenant_id == phone_identity['tenant_id'], User.login_name == 'T0000001'))
            original_hash = account.password_hash
            assert db.get(ExportJob, int(result['result']['credentialExportJobId'])).file_object_id == file.id
            assert db.scalar(select(IdentityImportBatch.status).where(IdentityImportBatch.batch_no == batch['batchNo'])) == 'IDENTITY_CONFIRMED'
        again = jobs.confirm_identity_import_job(job['id'], expected_version=job['version'], user=operator,
            idempotency_key='synthetic-idempotency-key')
        assert again['credentialReceiptFileId'] == result['credentialReceiptFileId']
        with get_sessionmaker()() as db:
            assert db.get(User, account.id).password_hash == original_hash
        assert 'teacherCredentials' not in str(result)
    finally:
        set_tenant(previous)
