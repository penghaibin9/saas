"""Real orientation workbooks retain batch boundaries and readable business labels."""
import io

from openpyxl import load_workbook
from sqlalchemy import select

from test_orientation_o3_self_service import _seed_o3, TID


def test_orientation_export_workbooks_keep_batch_and_chinese_business_states(client, db_mode, auth_headers, monkeypatch):
    from app.api.v1 import import_export as export_api
    from app.db.session import get_sessionmaker
    from app.models import (OrientationBatch, OrientationStudent, OrientationMaterial,
                            OrientationStudentStep, OrientationFlowStep, OrientationPaymentAccount,
                            GreenChannelApplication, OrientationException)

    # Repeated workbook variants share one isolated fixture; quota behavior has its own tests.
    monkeypatch.setattr(export_api, '_limit_operation', lambda *args, **kwargs: None)
    ids = _seed_o3(db_mode, include_payment=True)
    with get_sessionmaker()() as db:
        student = db.get(OrientationStudent, ids['orientationId'])
        student.name = '=公式注入测试'
        batch = db.get(OrientationBatch, student.batch_id)
        batch_id, batch_no = batch.id, batch.batch_no
        other_batch = OrientationBatch(tenant_id=TID, batch_no='EXPORT-OTHER', batch_name='其他批次', year='2026', status='ACTIVE', flow_version_id=batch.flow_version_id)
        foreign_batch = OrientationBatch(tenant_id=TID + 1, batch_no='EXPORT-FOREIGN', batch_name='其他学校批次', year='2026', status='DRAFT')
        db.add_all([other_batch, foreign_batch]); db.flush()
        foreign_id = foreign_batch.id
        other = OrientationStudent(tenant_id=TID, batch_id=other_batch.id, student_id=ids['otherProfileId'], name='不应出现在目标批次', admission_no='EXPORT-OTHER-STUDENT', source_type='MANUAL', source_record_id='EXPORT-OTHER-STUDENT', identity_status='LINKED', record_status='ACTIVE')
        db.add(other)
        db.add(OrientationMaterial(tenant_id=TID, ori_student_id=student.id, material_type='ID_CARD', file_name='fictional.png', status='RETURNED', is_current=True, return_reason='请补充清晰证明'))
        db.add(OrientationPaymentAccount(tenant_id=TID, orientation_student_id=student.id, student_id=student.student_id, status='DEFERRED', source_type='MANUAL_VERIFIED', source_biz_id='export-fixture-payment'))
        db.add(GreenChannelApplication(tenant_id=TID, ori_student_id=student.id, student_id=student.student_id, apply_type='TUITION_DEFERMENT', status='RETURNED', reject_reason='请补充预计缴费日期'))
        db.add(OrientationException(tenant_id=TID, ori_student_id=student.id, exception_type='MATERIAL', status='OPEN', risk_level='HIGH'))
        step = db.scalar(select(OrientationStudentStep).where(OrientationStudentStep.orientation_student_id == student.id, OrientationStudentStep.step_key == 'MATERIAL'))
        step.status = 'BLOCKED'; step.blocked_reason = '材料待补充'
        db.get(OrientationFlowStep, step.flow_step_id).step_name = '本批次入学材料核验'
        db.commit()

    expected = {
        'students': {'缴费事实': '已批准缓缴', '绿色通道审批': '暂无通过记录'},
        'progress': {'受阻环节': '本批次入学材料核验：材料待补充'},
        'materials': {'材料类型': '身份证明', '审核状态': '已退回'},
        'payment': {'缴费状态': '已批准缓缴', '事实来源': '人工核验'},
        'green-channel': {'申请类型': '学费缓缴', '状态': '已退回', '退回或驳回原因': '请补充预计缴费日期'},
        'exceptions': {'异常类型': '材料异常', '风险等级': '高风险', '状态': '待处理'},
        'dorm': {'住宿状态': '未形成住宿事实'},
        'no-show': {'报到状态': '未报到'},
        'checkin': {},
    }
    for report, labels in expected.items():
        purpose = f'迎新批次台账验收用途：{report}'
        created = client.post('/api/v1/export/domain/orientation', headers=auth_headers,
                              json={'batchId': batch_id, 'reportType': report, 'purpose': purpose})
        assert created.status_code == 200, created.text
        task = created.json()['data']
        assert task['status'] == 'SUCCESS'
        assert task['rowCount'] == (0 if report == 'checkin' else 1)
        downloaded = client.get(f"/api/v1/export/tasks/{task['taskId']}/download", headers=auth_headers)
        assert downloaded.status_code == 200, downloaded.text
        assert downloaded.headers['content-type'].startswith('application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
        workbook = load_workbook(io.BytesIO(downloaded.content), read_only=True)
        values = list(workbook.active.values)
        assert batch_no in values[0][0] and purpose in values[0][0]
        assert not any('不应出现在目标批次' in str(value) for row in values for value in row)
        if labels:
            row = dict(zip(values[1], values[2]))
            assert row['迎新批次编号'] == batch_no
            assert row['姓名'] == "'=公式注入测试"
            for column, value in labels.items():
                assert row[column] == value, (report, column, row[column])
        workbook.close()

    # Exercise all supported client values through the real workbook endpoint.
    for apply_type, expected_label in (("POVERTY", "家庭经济困难"), ("DISASTER", "突发灾害")):
        with get_sessionmaker()() as db:
            account = db.scalar(select(OrientationPaymentAccount).where(OrientationPaymentAccount.orientation_student_id == ids['orientationId']))
            account.status = "WAIVED"
            application = db.scalar(select(GreenChannelApplication).where(GreenChannelApplication.ori_student_id == ids['orientationId']))
            application.apply_type = apply_type
            db.commit()
        for report, column, expected_label_value in (("green-channel", "申请类型", expected_label), ("payment", "缴费状态", "已减免"), ("students", "缴费事实", "已减免")):
            created = client.post('/api/v1/export/domain/orientation', headers=auth_headers,
                                  json={'batchId': batch_id, 'reportType': report, 'purpose': '受支持业务标签台账核验用途'})
            assert created.status_code == 200, created.text
            downloaded = client.get(f"/api/v1/export/tasks/{created.json()['data']['taskId']}/download", headers=auth_headers)
            assert downloaded.status_code == 200, downloaded.text
            workbook = load_workbook(io.BytesIO(downloaded.content), read_only=True)
            values = list(workbook.active.values)
            row = dict(zip(values[1], values[2]))
            assert row[column] == expected_label_value, (report, column, row[column])
            workbook.close()

    # Absence of a finance fact is not unpaid. Approved green-channel evidence is
    # independent of both finance and the student's compatibility projection.
    with get_sessionmaker()() as db:
        account = db.scalar(select(OrientationPaymentAccount).where(OrientationPaymentAccount.orientation_student_id == ids['orientationId']))
        account.is_deleted = True
        application = db.scalar(select(GreenChannelApplication).where(GreenChannelApplication.ori_student_id == ids['orientationId']))
        application.status = 'APPROVED'
        db.get(OrientationStudent, ids['orientationId']).green_channel_status = 'NOT_APPLIED'
        db.commit()
    created = client.post('/api/v1/export/domain/orientation', headers=auth_headers,
                          json={'batchId': batch_id, 'reportType': 'students', 'purpose': '缺少财务记录的绿色通道事实核对'})
    assert created.status_code == 200, created.text
    downloaded = client.get(f"/api/v1/export/tasks/{created.json()['data']['taskId']}/download", headers=auth_headers)
    assert downloaded.status_code == 200
    workbook = load_workbook(io.BytesIO(downloaded.content), read_only=True)
    values = list(workbook.active.values)
    row = dict(zip(values[1], values[2]))
    assert row['缴费事实'] == '未同步缴费事实'
    assert row['绿色通道审批'] == '已通过'
    workbook.close()

    for invalid_batch in (None, foreign_id, 99999999):
        denied = client.post('/api/v1/export/domain/orientation', headers=auth_headers,
                             json={'batchId': invalid_batch, 'purpose': '无效批次导出拒绝验收'})
        assert denied.status_code >= 400, denied.text
