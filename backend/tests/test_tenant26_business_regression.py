"""Tenant26: real MySQL services, synthetic schools, isolated file/SMS effects."""
from datetime import datetime, timedelta
from pathlib import Path
import time

import pytest

from app.core.context import get_tenant, set_tenant
from app.db.session import get_sessionmaker

A = 1000000000000000001
B = 1000000000000000002


@pytest.fixture()
def school_context(db_mode):
    previous = get_tenant()
    set_tenant({'tenantId': str(A), 'tenantCode': 'demo'})
    yield
    set_tenant(previous)


@pytest.mark.usefixtures('school_context')
@pytest.mark.parametrize('global_worker', [False, True])
def test_export_cleanup_preserves_foreign_file_and_cleans_own(tmp_path, monkeypatch, global_worker):
    from app.models.file import FileObject
    from app.models.data_exchange import ExportJob
    from app.services import data_exchange_job_service as jobs
    own, foreign = tmp_path / 'own.xlsx', tmp_path / 'foreign.xlsx'
    own.write_bytes(b'own synthetic export')
    foreign.write_bytes(b'foreign synthetic export')
    deleted = []

    class IsolatedStorage:
        def delete(self, key):
            target = Path(key)
            assert target.parent == tmp_path
            deleted.append(key)
            target.unlink()

    monkeypatch.setattr(jobs, 'get_backend', lambda: IsolatedStorage())
    with get_sessionmaker()() as db:
        files = [FileObject(tenant_id=tid, file_key=str(p), file_name=p.name)
                 for tid, p in [(A, own), (B, foreign)]]
        db.add_all(files); db.flush()
        ids = [f.id for f in files]
        db.add_all([ExportJob(tenant_id=A, module_code='SYSTEM', export_type='TEST',
                             status='SUCCEEDED', file_object_id=f.id,
                             expires_at=datetime.utcnow() - timedelta(days=1)) for f in files])
        db.commit()
    if global_worker:
        set_tenant(None)
    result = jobs.cleanup_expired_jobs()
    assert result == {'expiredImports': 0, 'expiredExports': 2, 'deletedFiles': 1}
    assert deleted == [str(own)] and not own.exists()
    assert foreign.read_bytes() == b'foreign synthetic export'
    with get_sessionmaker()() as db:
        assert db.get(FileObject, ids[0]).is_deleted
        row = db.get(FileObject, ids[1])
        assert not row.is_deleted and row.status == 'AVAILABLE'


@pytest.mark.usefixtures('school_context')
@pytest.mark.parametrize('link_kind', ['valid', 'foreign', 'missing', 'legacy'])
def test_account_projection_respects_stable_link_without_name_fallback(link_kind):
    from app.models import College, Major, SchoolClass, StudentProfile, StudentAccountLink, User
    from app.modules.system_admin.routers.system_bundle import _student_account_meta
    with get_sessionmaker()() as db:
        college = College(tenant_id=B, college_name='FOREIGN_COLLEGE_SECRET')
        db.add(college); db.flush()
        major = Major(tenant_id=B, college_id=college.id, major_name='FOREIGN_MAJOR_SECRET')
        db.add(major); db.flush()
        cls = SchoolClass(tenant_id=B, major_id=major.id, class_name='FOREIGN_CLASS_SECRET')
        db.add(cls); db.flush()
        own = StudentProfile(tenant_id=A, student_no='T26-ACCOUNT', real_name='本校学生',
                             college_id=college.id, major_id=major.id, class_id=cls.id)
        other = StudentProfile(tenant_id=B, student_no='T26-FOREIGN', real_name='FOREIGN_STUDENT_SECRET')
        account = User(tenant_id=A, login_name='T26-ACCOUNT', real_name='测试账号',
                       password_hash='unused', user_type='STUDENT', status='ACTIVE')
        db.add_all([own, other, account]); db.flush()
        if link_kind != 'legacy':
            target = own.id if link_kind == 'valid' else other.id if link_kind == 'foreign' else 999999999
            db.add(StudentAccountLink(tenant_id=A, user_id=account.id, student_id=target,
                                      link_status='ACTIVE'))
        db.commit()
        value = _student_account_meta(db, account)
        assert 'SECRET' not in str(value)
        if link_kind in ('foreign', 'missing'):
            assert value['profileBound'] is False and value['studentId'] == ''
        else:
            assert value['profileBound'] is True and value['studentId'] == str(own.id)
            assert value['collegeName'] == value['majorName'] == value['className'] == ''


@pytest.mark.usefixtures('school_context')
def test_organization_and_counselor_chain_rejects_each_foreign_hop():
    from app.models import College, Major, SchoolClass, StudentProfile, AcademicStudent
    from app.modules.academic_affairs.services import academic_affairs_org_service as org
    from app.modules.academic_affairs.services import academic_affairs_warning_service as warning
    from app.modules.academic_affairs.services import academic_affairs_service as affairs
    with get_sessionmaker()() as db:
        colleges = [College(tenant_id=t, college_name=f't26-{t}') for t in (A, B)]
        db.add_all(colleges); db.flush()
        majors = [Major(tenant_id=t, college_id=c.id, major_name='测试专业') for t, c in zip((A, B), colleges)]
        db.add_all(majors); db.flush()
        classes = [SchoolClass(tenant_id=t, major_id=m.id, class_name='测试班', counselor_id=123 if t == A else 456)
                   for t, m in zip((A, B), majors)]
        db.add_all(classes); db.flush()
        students = [StudentProfile(tenant_id=t, class_id=c.id, student_no=f't26-{t}', real_name='测试学生')
                    for t, c in zip((A, B), classes)]
        db.add_all(students); db.flush()
        acad = AcademicStudent(tenant_id=A, student_id=students[0].id, student_no='T26', name='测试学生')
        db.add(acad); db.commit()
        assert org._class_college_id(db, classes[0].id) == colleges[0].id
        assert org._class_college_id(db, classes[1].id) is None
        assert warning._counselor_of(db, acad.id) == (123, students[0].id)
        classes[0].major_id = majors[1].id; db.flush()
        assert org._class_college_id(db, classes[0].id) is None
        students[0].class_id = classes[1].id; db.flush()
        assert warning._counselor_of(db, acad.id)[0] == 0
        assert affairs._counselor_of(db, students[0].id) == 0
        acad.student_id = students[1].id; db.flush()
        assert warning._counselor_of(db, acad.id) == (0, None)
        acad.tenant_id = B; db.flush()
        assert warning._counselor_of(db, acad.id) == (0, None)


@pytest.mark.usefixtures('school_context')
def test_resource_labels_only_resolve_own_school():
    from app.models import AaClassroom, AaLabResource
    from app.modules.academic_affairs.services.academic_affairs_resource_service import _resolve_owner_label
    with get_sessionmaker()() as db:
        rooms = [AaClassroom(tenant_id=t, building_code='T26', building_name=f'building-{t}', room_code='101') for t in (A, B)]
        labs = [AaLabResource(tenant_id=t, lab_code='T26', lab_name=f'lab-{t}') for t in (A, B)]
        db.add_all(rooms + labs); db.commit()
        assert _resolve_owner_label(db, 'CLASSROOM', rooms[0].id) == f'building-{A}101'
        assert _resolve_owner_label(db, 'LAB', labs[0].id) == f'lab-{A}'
        assert _resolve_owner_label(db, 'CLASSROOM', rooms[1].id) is None
        assert _resolve_owner_label(db, 'LAB', labs[1].id) is None
        assert _resolve_owner_label(db, 'CLASSROOM', 999999999) is None


@pytest.mark.usefixtures('school_context')
def test_legacy_transcript_bytes_never_contain_foreign_student():
    from io import BytesIO
    from openpyxl import load_workbook
    from app.models import StudentProfile
    from app.modules.academic_affairs.services import academic_affairs_grade_core_service as legacy
    user = {'userId': 'u_school_admin01', 'loginName': 'school_admin01', 'realName': '测试管理员',
            'tenantId': str(A), 'currentRoleCode': 'SCHOOL_ADMIN'}
    with get_sessionmaker()() as db:
        students = [StudentProfile(tenant_id=t, student_no=f'T26-EXPORT-{t}', real_name=f'NAME-{t}') for t in (A, B)]
        db.add_all(students); db.commit()
        identities = [(s.id, s.real_name, s.student_no) for s in students]
    for index, (sid, name, no) in enumerate(identities):
        blob = legacy.export_transcript_xlsx(user, sid, purpose='租户隔离测试用途')
        book = load_workbook(BytesIO(blob), read_only=True)
        cells = str([cell for sheet in book for row in sheet.values for cell in row])
        book.close()
        if index == 0:
            assert name in cells and no in cells
        else:
            assert name not in cells and no not in cells


def test_formal_grade_routes_keep_authoritative_command():
    from app.main import app
    from app.core.route_introspection import iter_effective_api_routes
    from app.modules.academic_affairs.services import academic_affairs_grade_service as grade
    from app.modules.academic_affairs.services import academic_affairs_grade_correction_command as command
    from app.modules.academic_affairs.services import academic_affairs_grade_core_service as legacy
    assert grade.change_academic_review is command.change_academic_review
    assert grade.change_college_review is command.change_college_review
    assert grade.export_transcript_xlsx is not legacy.export_transcript_xlsx
    endpoints = [r.endpoint for r in iter_effective_api_routes(app.routes) if '/grade-change/' in r.path]
    assert len(endpoints) == 2
    assert all(endpoint.__module__.endswith('grade_change_recheck_router') for endpoint in endpoints)


@pytest.mark.usefixtures('school_context')
def test_graduation_roster_wrong_profile_never_leaks_name_or_number():
    from app.models import StudentProfile, AaGraduationAuditResult
    from app.modules.academic_affairs.services.academic_affairs_graduation_service import rosters
    with get_sessionmaker()() as db:
        students = [StudentProfile(tenant_id=t, student_no=f'NO-{t}', real_name=f'NAME-{t}') for t in (A, B)]
        db.add_all(students); db.flush()
        db.add_all([AaGraduationAuditResult(tenant_id=A, batch_id=1234, student_id=s.id, conclusion='GRADUATED')
                    for s in students])
        db.commit()
    result = rosters(1234, {'currentRoleCode': 'ACADEMIC_ADMIN', 'tenantId': str(A)})
    assert f'NAME-{A}' in str(result) and f'NO-{A}' in str(result)
    assert f'NAME-{B}' not in str(result) and f'NO-{B}' not in str(result)


@pytest.mark.usefixtures('school_context')
def test_graduation_grade_scope_excludes_foreign_student_reference():
    from app.core.context import get_current_user_ctx, set_current_user
    from app.models import GraduationStudent, GraduationGrade
    from app.modules.graduation.services.graduation_grade_service import list_grades
    previous = get_current_user_ctx()
    set_current_user({'currentRoleCode': 'SCHOOL_ADMIN', 'tenantId': str(A), 'realName': '测试管理员'})
    try:
        with get_sessionmaker()() as db:
            students = [GraduationStudent(tenant_id=t, name=f'NAME-{t}', student_no=f'NO-{t}') for t in (A, B)]
            db.add_all(students); db.flush()
            db.add_all([GraduationGrade(tenant_id=A, gd_student_id=s.id, status='DRAFT') for s in students])
            db.commit()
        rows, total = list_grades(1, 20)
        assert total == 1 and rows[0]['studentName'] == f'NAME-{A}'
        assert f'NAME-{B}' not in str(rows) and f'NO-{B}' not in str(rows)
    finally:
        set_current_user(previous)


@pytest.mark.usefixtures('school_context')
def test_identity_import_error_receipt_is_real_xlsx_bound_to_own_job(tmp_path, monkeypatch):
    from io import BytesIO
    from openpyxl import Workbook, load_workbook
    from app.core.config import settings
    from app.core.exceptions import AppException
    from app.models import FileObject
    from app.models.data_exchange import ImportJob, ImportRowError
    from app.services import identity_import_file_service as files, data_exchange_job_service as jobs
    from app.services.storage.local import LocalStorageBackend
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    monkeypatch.setattr(jobs, 'get_backend', LocalStorageBackend)
    book = Workbook(); sheet = book.active
    sheet.append(files.TEACHER_HEADERS)
    account_no = '000012345678901234'
    sheet.append([account_no if header == '工号' else '' for header in files.TEACHER_HEADERS])
    stream = BytesIO(); book.save(stream); book.close()
    parsed = files.parse_teacher_xlsx(stream.getvalue(), 'tenant26-invalid.xlsx')
    assert parsed['errors']
    user = {'userId': '101', 'tenantId': str(A), 'realName': '测试导入员', 'currentRoleCode': 'SCHOOL_ADMIN'}
    batch = files.create_batch(user, parsed, {'tenantId': str(A), 'errors': parsed['errors']})
    result = jobs.create_identity_import_job(kind='TEACHER', source_file_id='0', parsed=parsed, batch_result=batch, user=user)
    assert result['status'] == 'VALIDATION_FAILED' and result['errorReceiptFileId']
    with get_sessionmaker()() as db:
        row = db.get(ImportJob, int(result['id']))
        receipt = db.get(FileObject, row.error_receipt_file_id)
        assert row.tenant_id == receipt.tenant_id == A
        assert receipt.visibility == 'PRIVATE' and receipt.biz_id == f'IMPORT:{row.id}:ERRORS'
        assert db.query(ImportRowError).filter_by(tenant_id=A, import_job_id=row.id).count() > 0
        output = load_workbook(tmp_path / receipt.file_key, read_only=True)
        assert account_no in str([cell for sheet in output for values in sheet.values for cell in values])
        output.close()
    replay = jobs.create_identity_import_job(kind='TEACHER', source_file_id='0', parsed=parsed, batch_result=batch, user=user)
    assert replay['id'] == result['id'] and replay['errorReceiptFileId'] == result['errorReceiptFileId']
    set_tenant(B)
    with pytest.raises(AppException):
        files.get_batch({**user, 'tenantId': str(B)}, B, batch['batchNo'])


def test_recheck_does_not_update_foreign_academic_aggregate(client, db_mode):
    from tests import test_aa_grade_recheck as scenario
    from app.models import AcademicStudent, AcademicGrade, AaGradeRecheck, User
    gid = scenario._seed_grade('T26-RECHECK', '测试复查学生', '测试课程', 58)
    student = scenario._stu_token('测试复查学生', 'T26-RECHECK')
    created = client.post(f'{scenario.BASE}/grade-recheck/submit', headers=student,
                          json={'acadGradeId': str(gid), 'reason': '申请核对课程成绩'})
    assert created.status_code == 200, created.text
    rid = int(created.json()['data']['recheckId'])
    admin = scenario._admin(client)
    with get_sessionmaker()() as db:
        if not db.query(User).filter_by(tenant_id=A, login_name='school_admin01').first():
            db.add(User(tenant_id=A, login_name='school_admin01', real_name='测试审核员',
                        password_hash='unused', user_type='SCHOOL_ADMIN', status='ACTIVE'))
        acad = db.get(AcademicStudent, db.get(AcademicGrade, gid).acad_student_id)
        acad.tenant_id, acad.gpa = B, 3.75
        db.commit()
        aid = acad.id
    response = client.post(f'{scenario.ADMIN_BASE}/grade-rechecks/{rid}/review', headers=admin,
                           json={'action': 'ADJUST', 'newScore': 72, 'note': '重新核对原始试卷成绩'})
    assert response.status_code in (200, 403, 404, 409), response.text
    with get_sessionmaker()() as db:
        assert float(db.get(AcademicStudent, aid).gpa) == 3.75
        if response.status_code != 200:
            assert db.get(AaGradeRecheck, rid).status == 'SUBMITTED'
            assert db.get(AcademicGrade, gid).record_status == 'ACTIVE'
        else:
            assert db.get(AaGradeRecheck, rid).status == 'ADJUSTED'


def test_makeup_does_not_update_foreign_academic_aggregate(client, db_mode):
    from tests import test_aa_makeup as scenario
    from app.models import AcademicStudent, AaMakeupBatch
    ids = scenario._seed(db_mode)
    admin = scenario._hdr(client, 'school_admin01')
    base = scenario.BASE
    batch = client.post(f'{base}/makeup/batches', headers=admin, json={'batchName': '租户隔离补考测试'})
    assert batch.status_code == 200, batch.text
    bid = batch.json()['data']['batchId']
    enrolled = client.post(f'{base}/makeup/batches/{bid}/enroll', headers=admin,
                           json={'gradeId': str(ids['failGrade']), 'acadStudentId': str(ids['acad'])})
    assert enrolled.status_code == 200, enrolled.text
    mid = enrolled.json()['data']['makeupId']
    for suffix, body in [(f'batches/{bid}/publish', None), (f'records/{mid}/score', {'score': 72}),
                         (f'batches/{bid}/college-review', None)]:
        response = client.post(f'{base}/makeup/{suffix}', headers=admin, json=body)
        assert response.status_code == 200, response.text
    with get_sessionmaker()() as db:
        acad = db.get(AcademicStudent, ids['acad'])
        acad.tenant_id, acad.gpa = B, 3.75
        db.commit()
    result = client.post(f'{base}/makeup/batches/{bid}/finish', headers=admin)
    assert result.status_code in (200, 403, 404, 409), result.text
    with get_sessionmaker()() as db:
        assert float(db.get(AcademicStudent, ids['acad']).gpa) == 3.75
        assert db.get(AaMakeupBatch, int(bid)).status == ('FINISHED' if result.status_code == 200 else 'REVIEWED')


@pytest.mark.usefixtures('db_mode')
def test_dynamic_correction_rejects_foreign_component_without_partial_commit():
    from tests import test_aa_grade_correction_command as scenario
    from app.core.exceptions import AppException
    from app.models.academic_affairs_r10 import AaGradeComponentScore
    from app.models.academic_affairs_effective_grade import AaGradeChangeRequest
    ids = scenario._seed_published_grade(dynamic=True)
    scenario._apply(ids)
    scenario._college(ids)
    with get_sessionmaker()() as db:
        component = db.query(AaGradeComponentScore).filter_by(tenant_id=A, grade_task_id=ids['taskId']).first()
        component.tenant_id = B
        db.commit()
        component_id, score, version = component.id, component.score, component.version
    before = scenario._state(ids)
    with pytest.raises(AppException) as failure:
        scenario._final(ids, command_key='tenant26-foreign-component')
    assert failure.value.http_status == 409
    assert scenario._state(ids) == before
    with get_sessionmaker()() as db:
        component = db.get(AaGradeComponentScore, component_id)
        assert (component.tenant_id, component.score, component.version) == (B, score, version)
        request = db.query(AaGradeChangeRequest).filter_by(tenant_id=A, grade_record_id=ids['recordId']).one()
        assert request.status == 'PENDING'


@pytest.mark.usefixtures('school_context')
@pytest.mark.parametrize('delivery_status', ['SENT', 'RETRY_WAIT', 'EXPIRED'])
def test_global_sms_worker_keeps_two_school_identities_and_snapshot(delivery_status, monkeypatch):
    from app.core.config import settings
    from app.core.field_crypto import encrypt_field
    from app.models import Tenant, User, Role, UserRole, PhoneLoginBinding, SysConfig, PasswordResetSmsJob
    from app.services import password_reset_service as reset
    from app.services.phone_login_service import phone_lookup
    monkeypatch.setattr(settings, 'SENSITIVE_SEARCH_HMAC_KEY', 'ab' * 32)
    reset.reset_for_tests()
    requests = []
    with get_sessionmaker()() as db:
        for index in (1, 2):
            tenant = Tenant(tenant_code=f't26-sms-{index}', school_name='短信测试学校', status='ACTIVE')
            db.add(tenant); db.flush()
            user = User(tenant_id=tenant.id, login_name=f't26-sms-{index}', real_name='测试学生',
                        user_type='STUDENT', status='ACTIVE', password_hash='unused', must_change_password=False)
            role = Role(tenant_id=tenant.id, role_code='STUDENT', role_name='学生')
            db.add_all([user, role]); db.flush()
            phone = f'1380000000{index}'
            db.add_all([UserRole(tenant_id=tenant.id, user_id=user.id, role_id=role.id, status='ACTIVE'),
                        PhoneLoginBinding(tenant_id=tenant.id, user_id=user.id, state='VERIFIED',
                                          phone_ciphertext=encrypt_field(phone), active_phone_lookup=phone_lookup(tenant.id, phone), version=1),
                        SysConfig(tenant_id=tenant.id, config_key='SEC_PHONE_RECOVERY_ENABLED', value_text='1')])
            db.commit()
            snapshot = reset._find_reset_account(user.login_name, tenant.tenant_code, 'PC')
            assert snapshot is not None
            request_id = f't26-sms-{index}'
            snapshot.update(purpose='RESET_PASSWORD', expiresAt=time.time() + (-1 if delivery_status == 'EXPIRED' else 300))
            reset._set('reset-operation', request_id, snapshot, 300)
            db.add(PasswordResetSmsJob(tenant_id=tenant.id, user_id=user.id, request_id=request_id,
                                      phone_encrypted=encrypt_field(phone), code_encrypted=encrypt_field('123456'),
                                      expires_at=datetime.utcnow() + timedelta(minutes=5)))
            requests.append((int(tenant.id), phone))
        db.commit()
    delivered = []
    def transport(tenant_id, phone, code):
        delivered.append((int(tenant_id), phone))
        return {'status': 'SENT'} if delivery_status == 'SENT' else {'status': 'FAILED', 'retryable': True}
    monkeypatch.setattr('app.services.notification.sms_service.notify_password_reset', transport)
    set_tenant(None)
    assert reset.process_delivery_jobs(worker_id='tenant26-test') == (2 if delivery_status == 'SENT' else 0)
    assert delivered == ([] if delivery_status == 'EXPIRED' else requests)
    with get_sessionmaker()() as db:
        rows = db.query(PasswordResetSmsJob).order_by(PasswordResetSmsJob.id).all()
        assert len(rows) == 2 and all(r.status == delivery_status and r.locked_by is None for r in rows)
        assert all(r.attempt_count == 1 for r in rows)
        if delivery_status == 'RETRY_WAIT':
            assert all(r.next_retry_at is not None and r.code_encrypted for r in rows)
    reset.reset_for_tests()
