"""Real MySQL: bulk preview, atomic writes/receipts, Excel, tenant and permission boundaries."""
from io import BytesIO
from openpyxl import load_workbook
from tests.test_aa_resource import _hdr, _mk, BASE


def building(client, hdr, code="BULK"):
    r = client.post(BASE + "/classroom-buildings", headers=hdr, json={
        "buildingCode": code, "buildingName": code + "教学楼", "campusCode": "本部", "floorCount": 10})
    assert r.status_code == 200, r.text
    return r.json()["data"]


def preview(client, hdr, b, **changes):
    body = {"buildingId": b["buildingId"], "startFloor": 1, "endFloor": 10,
        "roomsPerFloor": 20, "capacity": 60, "examSeats": 30, "roomType": "MULTIMEDIA"}
    body.update(changes)
    r = client.post(BASE + "/classroom-batches/generate-preview", headers=hdr, json=body)
    assert r.status_code == 200, r.text
    return r.json()["data"]


def test_bulk_200_roundtrip_idempotency_and_keep(client, db_mode):
    hdr = _hdr(client, "school_admin01")
    b = building(client, hdr)
    p = preview(client, hdr, b)
    assert p["createCount"] == 200
    path = BASE + f'/classroom-batches/{p["batchNo"]}/confirm'
    r = client.post(path, headers=hdr)
    assert r.status_code == 200, r.text
    result = r.json()["data"]
    assert result["createdCount"] == 200 and len(set(result["classroomIds"])) == 200
    assert client.post(path, headers=hdr).json()["data"] == result
    rows = client.get(BASE + "/classrooms", headers=hdr, params={"buildingId": b["buildingId"], "floorNo": 2}).json()["data"]
    assert rows["total"] == 20 and all(r["floorNo"] == 2 and r["examSeats"] == 30 for r in rows["items"])
    again = preview(client, hdr, b)
    assert again["createCount"] == 0 and again["keepCount"] == 200
    x = client.get(BASE + f'/classroom-batches/{p["batchNo"]}/result.xlsx', headers=hdr)
    sheet = load_workbook(BytesIO(x.content)).active
    assert sheet.max_row == 201 and sheet["J2"].value == "已创建"


def test_preview_stale_conflict_rolls_back_entire_batch(client, db_mode):
    hdr = _hdr(client, "school_admin01")
    b = building(client, hdr, "RACE")
    p = preview(client, hdr, b, endFloor=1, roomsPerFloor=2)
    r = _mk(client, hdr, "RACE", "102", buildingName="RACE教学楼", campusCode="本部")
    assert r.status_code == 200
    r = client.post(BASE + f'/classroom-batches/{p["batchNo"]}/confirm', headers=hdr)
    assert r.status_code == 409, r.text
    rows = client.get(BASE + "/classrooms", headers=hdr, params={"buildingCode": "RACE"}).json()["data"]
    assert rows["total"] == 1 and rows["items"][0]["roomCode"] == "102"


def test_xlsx_errors_duplicates_and_leading_zero(client, db_mode):
    from app.modules.academic_affairs.services.academic_affairs_classroom_catalog_service import workbook
    hdr = _hdr(client, "school_admin01")
    building(client, hdr, "XLS")
    content = workbook([["XLS", 1, "0101", "", "普通教室", 60, 30, "否", ""],
                        ["XLS", 1, "0101", "", "普通教室", 60, 30, "否", ""],
                        ["XLS", 22, "0102", "", "普通教室", 60, 30, "否", ""]])
    r = client.post(BASE + "/classroom-batches/import-preview", headers=hdr,
                    files={"file": ("rooms.xlsx", content)})
    assert r.status_code == 200, r.text
    p = r.json()["data"]
    assert p["errorCount"] == 2 and p["items"][0]["row"]["roomCode"] == "0101"
    assert client.post(BASE + f'/classroom-batches/{p["batchNo"]}/confirm', headers=hdr).status_code == 400
    corrected = [p["items"][0]["row"]]
    p2 = client.post(BASE + "/classroom-batches/preview", headers=hdr, json={"rows": corrected}).json()["data"]
    assert client.post(BASE + f'/classroom-batches/{p2["batchNo"]}/confirm', headers=hdr).status_code == 200
    rows = client.get(BASE + "/classrooms", headers=hdr, params={"buildingCode": "XLS"}).json()["data"]
    assert rows["items"][0]["roomCode"] == "0101"


def test_limits_floor_version_and_permission(client, db_mode):
    hdr = _hdr(client, "school_admin01")
    b = building(client, hdr, "LIMIT")
    p = preview(client, hdr, b, endFloor=1, roomsPerFloor=1)
    r = client.put(BASE + f'/classroom-buildings/{b["buildingId"]}', headers=hdr,
        json={**b, "floorCount": 11, "expectedVersion": b["version"]})
    assert r.status_code == 200
    assert client.post(BASE + f'/classroom-batches/{p["batchNo"]}/confirm', headers=hdr).status_code == 409
    student = _hdr(client, "student01")
    assert client.get(BASE + "/classroom-buildings", headers=student).status_code == 403
    assert client.post(BASE + f'/classroom-batches/{p["batchNo"]}/confirm', headers=student).status_code == 403
    assert client.get(BASE + "/classroom-batches/unknown", headers=hdr).status_code == 404
    bad = client.post(BASE + "/classroom-batches/generate-preview", headers=hdr, json={
        "buildingId": b["buildingId"], "startFloor": 1, "endFloor": 11, "roomsPerFloor": 100, "capacity": 60})
    assert bad.status_code == 400


def test_preview_owner_and_tenant_isolation(client, db_mode):
    import pytest
    from sqlalchemy import select
    from app.core.context import set_tenant, set_current_user, get_tenant, get_current_user_ctx
    from app.core.exceptions import AppException
    from app.models import SharedImportBatch
    from app.services.db_service import session
    from app.modules.academic_affairs.services import academic_affairs_classroom_catalog_service as svc
    hdr = _hdr(client, "school_admin01")
    b = building(client, hdr, "SCOPE")
    p = preview(client, hdr, b, endFloor=1, roomsPerFloor=1)
    with session() as db:
        batch = db.scalar(select(SharedImportBatch).where(SharedImportBatch.batch_no == p["batchNo"]))
        tid, uid = batch.tenant_id, batch.operator_key
    old_tenant, old_user = get_tenant(), get_current_user_ctx()
    try:
        set_tenant(tid)
        set_current_user({"userId": "different-operator"})
        for action in [svc.get_batch, svc.confirm_batch, svc.batch_workbook]:
            with pytest.raises(AppException) as exc:
                action(p["batchNo"])
            assert exc.value.http_status == 404
        set_tenant(tid + 999999)
        set_current_user({"userId": uid})
        assert svc.list_buildings()["total"] == 0
        for action in [svc.get_batch, svc.confirm_batch, svc.batch_workbook]:
            with pytest.raises(AppException) as exc:
                action(p["batchNo"])
            assert exc.value.http_status == 404
    finally:
        set_tenant(old_tenant)
        set_current_user(old_user)


def test_legacy_identity_and_individual_room_version(client, db_mode):
    hdr = _hdr(client, "school_admin01")
    old = _mk(client, hdr, "OLD", "0101", buildingName="OLD教学楼", campusCode="本部").json()["data"]
    b = building(client, hdr, "OLD")
    reread = client.get(BASE + f'/classrooms/{old["classroomId"]}', headers=hdr).json()["data"]
    assert reread["buildingId"] == b["buildingId"] and reread["floorNo"] is None
    assert reread["classroomId"] == old["classroomId"] and reread["roomCode"] == "0101"
    body = {"floorNo": 1, "examSeats": 25, "isExclusive": True, "expectedVersion": reread["version"]}
    updated = client.put(BASE + f'/classrooms/{old["classroomId"]}', headers=hdr, json=body)
    assert updated.status_code == 200, updated.text
    assert updated.json()["data"]["examSeats"] == 25 and updated.json()["data"]["isExclusive"] is True
    assert client.put(BASE + f'/classrooms/{old["classroomId"]}', headers=hdr, json=body).status_code == 409


def test_xlsx_formula_and_empty_rejected(client, db_mode):
    from openpyxl import Workbook
    from app.modules.academic_affairs.services.academic_affairs_classroom_catalog_service import HEADERS
    hdr = _hdr(client, "school_admin01")
    wb = Workbook(); ws = wb.active; ws.append(list(HEADERS))
    empty = BytesIO(); wb.save(empty)
    assert client.post(BASE + "/classroom-batches/import-preview", headers=hdr,
        files={"file": ("empty.xlsx", empty.getvalue())}).status_code == 400
    ws.append(["A", 1, "=1+1", "", "普通教室", 60, 30, "否", ""])
    formula = BytesIO(); wb.save(formula)
    r = client.post(BASE + "/classroom-batches/import-preview", headers=hdr,
        files={"file": ("formula.xlsx", formula.getvalue())})
    assert r.status_code == 400 and "公式" in r.json()["message"]
