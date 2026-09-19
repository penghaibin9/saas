"""Real local browser plus original backend; only synthetic isolated rows.

Run the two existing Vite apps at 15310/15311 with VITE_PROXY_TARGET=18310.
No mock HTTP, schema cleanup, production data or external SMS.
"""
import json
import os
from pathlib import Path
import secrets
import subprocess
import sys
import threading
import time
import urllib.request

import pytest

from tests.test_phone_login_mysql import phone_db, phone_identity


@pytest.fixture
def phone_surface_identities(phone_db, phone_identity):
    """Prepare a synthetic student through the existing identity entities.

    The browser story still performs the real first-login password mutation.  A
    VERIFIED binding is fixture state for testing PHONE login; it is not evidence
    that an xlsx candidate became verified.
    """
    from datetime import datetime
    from app.core.field_crypto import encrypt_sensitive
    from app.core.security import hash_password
    from app.models import PhoneLoginBinding, Role, StudentAccountLink, StudentProfile, User, UserRole
    from app.services.phone_login_service import phone_lookup

    suffix = secrets.token_hex(5)
    account = 'student-' + suffix
    initial_account = 'student-new-' + suffix
    phone = '137' + str(secrets.randbelow(10**8)).zfill(8)
    teacher_account = 'mini-teacher-' + suffix
    teacher_phone = '139' + str(secrets.randbelow(10**8)).zfill(8)
    role = Role(tenant_id=phone_identity['tenant_id'], role_code='STUDENT', role_name='学生')
    teacher_role = phone_db.query(Role).filter(Role.tenant_id == phone_identity['tenant_id'],
        Role.role_code == 'ACADEMIC_TEACHER').one()
    profile = StudentProfile(tenant_id=phone_identity['tenant_id'], student_no=account,
        real_name='虚构学生移动验收', grade='2026')
    initial_profile = StudentProfile(tenant_id=phone_identity['tenant_id'], student_no=initial_account,
        real_name='虚构学生首登验收', grade='2026')
    user = User(tenant_id=phone_identity['tenant_id'], login_name=account,
        real_name=profile.real_name, password_hash=hash_password('Student-local-Password2!'),
        user_type='STUDENT', must_change_password=False)
    initial_user = User(tenant_id=phone_identity['tenant_id'], login_name=initial_account,
        real_name=initial_profile.real_name, password_hash=hash_password('Student-initial-Password1!'),
        user_type='STUDENT', must_change_password=True)
    teacher = User(tenant_id=phone_identity['tenant_id'], login_name=teacher_account,
        real_name='虚构教师移动验收', password_hash=hash_password('Local-test-Password1!'), user_type='TEACHER')
    phone_db.add_all([role, profile, initial_profile, user, initial_user, teacher])
    phone_db.flush()
    phone_db.add_all([
        UserRole(tenant_id=phone_identity['tenant_id'], user_id=user.id, role_id=role.id, status='ACTIVE'),
        UserRole(tenant_id=phone_identity['tenant_id'], user_id=initial_user.id, role_id=role.id, status='ACTIVE'),
        UserRole(tenant_id=phone_identity['tenant_id'], user_id=teacher.id, role_id=teacher_role.id, status='ACTIVE'),
        StudentAccountLink(tenant_id=phone_identity['tenant_id'], student_id=profile.id, user_id=user.id,
            link_status='ACTIVE', bound_login_name=account, bound_student_no=account,
            source='IDENTITY_IMPORT', bound_at=datetime.utcnow()),
        StudentAccountLink(tenant_id=phone_identity['tenant_id'], student_id=initial_profile.id, user_id=initial_user.id,
            link_status='ACTIVE', bound_login_name=initial_account, bound_student_no=initial_account,
            source='IDENTITY_IMPORT', bound_at=datetime.utcnow()),
        PhoneLoginBinding(tenant_id=phone_identity['tenant_id'], user_id=user.id, state='VERIFIED',
            active_phone_lookup=phone_lookup(phone_identity['tenant_id'], phone),
            phone_ciphertext=encrypt_sensitive(phone), version=1),
        PhoneLoginBinding(tenant_id=phone_identity['tenant_id'], user_id=teacher.id, state='VERIFIED',
            active_phone_lookup=phone_lookup(phone_identity['tenant_id'], teacher_phone),
            phone_ciphertext=encrypt_sensitive(teacher_phone), version=1),
    ])
    phone_db.commit()
    return {'student_account': account, 'student_phone': phone, 'initial_student_account': initial_account,
        'teacher_account': teacher_account, 'teacher_phone': teacher_phone}


@pytest.fixture
def imported_phone_surface_identities(phone_identity, monkeypatch, tmp_path):
    """Prepare supplied-template copies and an authorized importer; the browser creates subjects."""
    from datetime import datetime, timedelta

    from openpyxl import load_workbook
    from sqlalchemy import select

    from app.core.config import settings
    from app.core.context import get_current_user_ctx, get_tenant, set_current_user, set_tenant
    from app.db.session import get_sessionmaker
    from app.models import Permission, PlatformConfig, Role, RolePermission, UserRole
    from app.services import storage
    from app.services.file_scan_config import get_file_scan_config
    from app.services.school_onboarding_service import run_onboarding

    if os.getenv('PHONE_REAL_IMPORT_BROWSER') != '1':
        pytest.skip('set PHONE_REAL_IMPORT_BROWSER=1 for the real XLSX-to-browser acceptance')
    template_paths = {
        'TEACHER': Path(os.environ['PHONE_TEACHER_TEMPLATE_XLSX']),
        'STUDENT': Path(os.environ['PHONE_STUDENT_TEMPLATE_XLSX']),
    }
    assert all(path.is_file() for path in template_paths.values())

    suffix = secrets.token_hex(5)
    teacher_account = 'TE2E' + suffix
    student_account = '2026' + str(secrets.randbelow(10**14)).zfill(14)
    teacher_phone = '139' + str(secrets.randbelow(10**8)).zfill(8)
    student_phone = '137' + str(secrets.randbelow(10**8)).zfill(8)
    changed_student_phone = '136' + str(secrets.randbelow(10**8)).zfill(8)
    rows = {
        'TEACHER': {'工号': teacher_account, '姓名': '虚构导入教师浏览器验收', '所属部门': '教务处',
                    '岗位名称': '任课教师', '预设角色编码': 'ACADEMIC_TEACHER',
                    '本人手机号': teacher_phone, '号码归属': 'SELF'},
        'STUDENT': {'学号': student_account, '姓名': '虚构导入学生浏览器验收', '所属学院': '测试学院',
                    '所属专业': '测试专业', '班级名称': '测试班', '年级': '2026',
                    '本人手机号': student_phone, '号码归属': 'SELF'},
    }
    tid, uid = phone_identity['tenant_id'], phone_identity['user_id']
    actor = {'userId': f'db-{uid}', 'tenantId': str(tid), 'realName': '手机号导入经办人',
             'userType': 'ADMIN', 'currentRoleCode': 'SCHOOL_ADMIN', 'permissions': ['*']}
    previous = get_tenant(), get_current_user_ctx()
    monkeypatch.setattr(settings, 'UPLOAD_DIR', str(tmp_path))
    monkeypatch.setattr(settings, 'FILE_STORAGE_BACKEND', 'local')
    monkeypatch.setenv('CLAMAV_ENABLED', 'true')
    monkeypatch.setenv('FILE_SCAN_REQUIRED', 'true')
    assert get_file_scan_config().enabled is True and get_file_scan_config().required is True
    storage.reset_backend()
    set_tenant({'tenantId': str(tid), 'tenantCode': phone_identity['tenant']})
    set_current_user(actor)

    def filled_template(kind):
        source = template_paths[kind]
        workbook = load_workbook(source)
        sheet = workbook[workbook.sheetnames[0]]
        headers = [str(cell.value or '').strip() for cell in sheet[1]]
        assert set(rows[kind]).issubset(headers)
        for column, header in enumerate(headers, 1):
            sheet.cell(2, column).value = rows[kind].get(header, '')
            if header in {'工号', '学号', '本人手机号', '联系手机号'}:
                sheet.cell(2, column).number_format = '@'
        target = tmp_path / f'{kind.lower()}-real-template.xlsx'
        workbook.save(target)
        workbook.close()
        return target

    try:
        with get_sessionmaker()() as db:
            db.add(PlatformConfig(tenant_id=tid, config_type='TENANT_META', config_key='-',
                config_json={'status': 'trial', 'packageCode': 'trial', 'environment': 'test',
                    'trialEndAt': (datetime.now() + timedelta(days=1)).isoformat(timespec='seconds')},
                enabled=True, status='ACTIVE'))
            teacher_role = db.scalar(select(Role).where(Role.tenant_id == tid,
                Role.role_code == 'ACADEMIC_TEACHER'))
            role = Role(tenant_id=tid, role_code='SCHOOL_ADMIN', role_name='学校管理员',
                        role_type='SYSTEM', status='ACTIVE')
            db.add(role)
            db.flush()
            user_role = db.scalar(select(UserRole).where(
                UserRole.tenant_id == tid, UserRole.user_id == uid,
                UserRole.role_id == teacher_role.id))
            user_role.role_id = role.id
            for code in ('systemAdmin.user.import', 'systemAdmin.dataExchange.confirm',
                         'systemAdmin.dataExchange.viewOwn'):
                permission = db.scalar(select(Permission).where(Permission.permission_code == code))
                if permission is None:
                    permission = Permission(permission_code=code, permission_name=code,
                                            module_code='systemAdmin', action='import')
                    db.add(permission)
                    db.flush()
                if db.scalar(select(RolePermission.id).where(
                        RolePermission.tenant_id == tid, RolePermission.role_id == role.id,
                        RolePermission.permission_id == permission.id)) is None:
                    db.add(RolePermission(tenant_id=tid, role_id=role.id,
                                          permission_id=permission.id, status='ACTIVE'))
            db.commit()
        run_onboarding(actor, {'tenantId': str(tid), 'colleges': [{'name': '测试学院'}],
            'majors': [{'name': '测试专业', 'collegeName': '测试学院'}],
            'classes': [{'name': '测试班', 'majorName': '测试专业', 'grade': '2026'}]}, dry_run=False)
        yield {
            'teacher': {'account': teacher_account, 'phone': teacher_phone,
                        'template': str(filled_template('TEACHER'))},
            'student': {'account': student_account, 'phone': student_phone,
                        'template': str(filled_template('STUDENT'))},
            'student_changed_phone': changed_student_phone,
            'importer_account': phone_identity['login'],
            'importer_password': 'Local-test-Password1!',
        }
    finally:
        storage.reset_backend()
        set_current_user(previous[1])
        set_tenant(previous[0])


def test_local_four_surface_phone_and_account_browser(phone_identity, phone_surface_identities):
    from app.core.config import settings
    root = Path(__file__).resolve().parents[2]
    env = {**os.environ, 'APP_ENV': 'test', 'DEPLOYMENT_MODE': 'development', 'DB_ENABLED': 'true',
        'DATABASE_URL': os.environ['TEST_DATABASE_URL'], 'SCHEDULER_MODE': 'external',
        'SMS_ENABLED': 'false', 'SMS_PROVIDER': 'mock',
        'SENSITIVE_SEARCH_HMAC_KEY': settings.SENSITIVE_SEARCH_HMAC_KEY,
        'DEFAULT_TENANT_CODE': phone_identity['tenant'],
        'CORS_ORIGINS': 'http://127.0.0.1:15310,http://127.0.0.1:15311,http://localhost:5188',
        'E2E_ALLOW_DESTRUCTIVE_TESTS': 'true', 'E2E_STAFF_BASE_URL': 'http://127.0.0.1:15310',
        'E2E_STUDENT_BASE_URL': 'http://127.0.0.1:15311/portal', 'E2E_API_BASE_URL': 'http://127.0.0.1:18310/api/v1',
        'PHONE_TEST_TENANT': phone_identity['tenant'], 'PHONE_TEST_ACCOUNT': phone_identity['login'],
        'PHONE_TEST_STUDENT_ACCOUNT': phone_surface_identities['student_account'],
        'PHONE_TEST_STUDENT_PHONE': phone_surface_identities['student_phone'],
        'PHONE_TEST_INITIAL_STUDENT_ACCOUNT': phone_surface_identities['initial_student_account'],
        'PHONE_TEST_MINI_TEACHER_ACCOUNT': phone_surface_identities['teacher_account'],
        'PHONE_TEST_MINI_TEACHER_PHONE': phone_surface_identities['teacher_phone'],
        'PHONE_TEST_MINI_BASE_URL': 'http://localhost:5188'}
    phone_log_dir = Path('/tmp/phone-browser')
    phone_log_dir.mkdir(parents=True, exist_ok=True)
    backend_log_path = phone_log_dir / 'backend.log'
    backend_log = backend_log_path.open('w', encoding='utf-8')
    server = subprocess.Popen([sys.executable, '-m', 'uvicorn', 'app.main:app', '--host', '127.0.0.1',
        '--port', '18310', '--no-access-log'], cwd=root / 'backend', env=env,
        stdout=backend_log, stderr=subprocess.STDOUT)
    try:
        ready = False
        deadline = time.monotonic() + 45
        while time.monotonic() < deadline and server.poll() is None:
            try:
                with urllib.request.urlopen('http://127.0.0.1:18310/health', timeout=1) as response:
                    ready = response.status == 200
                if ready:
                    break
            except OSError:
                time.sleep(.2)
        assert ready, 'Dedicated local backend did not become healthy'
        command = ['node', 'node_modules/@playwright/test/cli.js', 'test', '--config=playwright.phone.config.mjs']
        if os.environ.get('PHONE_E2E_GREP'):
            command.extend(['--grep', os.environ['PHONE_E2E_GREP']])
        result = subprocess.run(command,
            cwd=root / 'e2e', env=env, timeout=240, capture_output=True, text=True, encoding='utf-8', errors='replace')
        if result.returncode != 0:
            backend_log.flush()
            backend_text = backend_log_path.read_text(encoding='utf-8', errors='replace')
            raise AssertionError(
                result.stdout + result.stderr
                + '\n--- dedicated backend log ---\n'
                + backend_text[-40000:]
            )
    finally:
        server.terminate()
        server.wait(timeout=15)
        backend_log.close()


def test_real_xlsx_import_to_four_surface_phone_browser(
        phone_identity, imported_phone_surface_identities, monkeypatch, tmp_path):
    """Join Staff-PC import to self-verification and all four UI entries.

    Staff PC uploads and confirms the supplied production-format XLSX copies.
    The provider boundary is a local capture sink, but the durable SMS job, worker,
    encryption, Redis proof, binding transaction and browser requests are real.
    """
    from openpyxl import load_workbook
    from sqlalchemy import func, select

    from app.core.config import settings
    from app.core.redis_client import cache_delete_pattern
    from app.db.session import get_sessionmaker
    from app.models import (FileObject, PhoneLoginBinding, PhoneLoginCandidate, SecurityAuditLog,
                            StudentAccountLink, StudentProfile, User)
    from app.models.data_exchange import ImportJob
    from app.services import password_reset_service
    from app.services.clamav_client import ClamAVClient
    from app.services.file_scan_config import get_file_scan_config
    from app.services.file_scan_service import process_next_scan_job
    from app.services.notification import sms_service
    from app.services.notification.sms_provider import SmsProvider, SmsResult
    from app.services.phone_login_service import phone_lookup
    from app.services.storage.finalize import finalize_scan_storage
    from app.workers.identity_import_worker import process_next_identity_import

    root = Path(__file__).resolve().parents[2]
    identities = imported_phone_surface_identities
    teacher, student = identities['teacher'], identities['student']
    changed_phone = identities['student_changed_phone']
    mailbox = tmp_path / 'synthetic-sms-mailbox.jsonl'
    mailbox_lock = threading.Lock()
    worker_stop = threading.Event()
    worker_errors = []

    monkeypatch.setattr(settings, 'SMS_ENABLED', True)
    monkeypatch.setattr(settings, 'SMS_PROVIDER', 'mock')
    monkeypatch.setattr(settings, 'PHONE_BINDING_ENABLED', True)
    monkeypatch.setattr(settings, 'PHONE_SMS_CONSUMERS_READY', True)
    monkeypatch.setattr(settings, 'SMS_PHONE_DAILY_TENANT_BUDGET', 100)
    monkeypatch.setattr(settings, 'SMS_PHONE_DAILY_PLATFORM_BUDGET', 100)
    monkeypatch.setattr(settings, 'SMS_RATE_LIMIT_PER_MINUTE', 100)

    class CaptureProvider(SmsProvider):
        name = 'local-capture'

        def send(self, phone, template_id, params):
            local_phone = str(phone)[-11:]
            assert local_phone in {teacher['phone'], student['phone'], changed_phone}
            payload = {'phone': local_phone, 'template': template_id, 'code': str(params['code'])}
            with mailbox_lock:
                with mailbox.open('a', encoding='utf-8') as stream:
                    stream.write(json.dumps(payload, ensure_ascii=False) + '\n')
            # This acceptance performs bind then change without sleeping for the
            # production 60-second resend window. Only isolated test Redis rate
            # keys are advanced; durable jobs, proofs and business limits remain.
            cache_delete_pattern(
                f"rate:phone-send-cooldown:{phone_identity['tenant_id']}:*"
            )
            cache_delete_pattern('rate:login:*')
            return SmsResult(success=True, request_id='local-capture-' + secrets.token_hex(8),
                             provider=self.name, retryable=False)

    capture_provider = CaptureProvider()
    monkeypatch.setattr(sms_service, 'get_provider', lambda name=None: capture_provider)
    cache_delete_pattern('rate:login:*')
    # The dedicated phone-test Redis persists between local runs. Clear only its
    # phone acceptance counters so yesterday's synthetic platform budget cannot
    # make today's first challenge look like a product failure.
    cache_delete_pattern('rate:phone-*')

    def run_local_workers():
        while not worker_stop.wait(.1):
            try:
                finalize_scan_storage(process_next_scan_job(
                    'phone-browser-clamav', client=ClamAVClient(get_file_scan_config())))
                process_next_identity_import('phone-browser-identity-import')
                password_reset_service.process_delivery_jobs(
                    limit=10, worker_id='phone-browser-local-capture',
                    tenant_id=phone_identity['tenant_id'], purposes=('BIND_PHONE', 'CHANGE_PHONE'))
            except Exception as exc:  # surfaced after Playwright, without leaking OTP values
                worker_errors.append(f'{type(exc).__name__}: {exc}')
                worker_stop.set()

    env = {
        **os.environ,
        'APP_ENV': 'test', 'DEPLOYMENT_MODE': 'development', 'DB_ENABLED': 'true',
        'DATABASE_URL': os.environ['TEST_DATABASE_URL'], 'SCHEDULER_MODE': 'external',
        'SMS_ENABLED': 'true', 'SMS_PROVIDER': 'mock',
        'PHONE_BINDING_ENABLED': 'true', 'PHONE_SMS_CONSUMERS_READY': 'true',
        'SMS_PHONE_DAILY_TENANT_BUDGET': '100', 'SMS_PHONE_DAILY_PLATFORM_BUDGET': '100',
        'SMS_RATE_LIMIT_PER_MINUTE': '100',
        'UPLOAD_DIR': str(tmp_path), 'FILE_STORAGE_BACKEND': 'local',
        'CLAMAV_ENABLED': 'true', 'FILE_SCAN_REQUIRED': 'true',
        'CLAMAV_HOST': os.environ.get('CLAMAV_HOST', '127.0.0.1'),
        'CLAMAV_PORT': os.environ.get('CLAMAV_PORT', '3310'),
        'SENSITIVE_SEARCH_HMAC_KEY': settings.SENSITIVE_SEARCH_HMAC_KEY,
        'DEFAULT_TENANT_CODE': phone_identity['tenant'],
        'CORS_ORIGINS': 'http://127.0.0.1:15310,http://127.0.0.1:15311,http://localhost:5188',
        'E2E_ALLOW_DESTRUCTIVE_TESTS': 'true',
        'E2E_STAFF_BASE_URL': 'http://127.0.0.1:15310',
        'E2E_STUDENT_BASE_URL': 'http://127.0.0.1:15311/portal',
        'E2E_API_BASE_URL': 'http://127.0.0.1:18310/api/v1',
        'PHONE_REAL_IMPORT_BROWSER': '1', 'PHONE_TEST_TENANT': phone_identity['tenant'],
        'PHONE_IMPORTER_ACCOUNT': identities['importer_account'],
        'PHONE_IMPORTER_PASSWORD': identities['importer_password'],
        'PHONE_IMPORT_TEACHER_ACCOUNT': teacher['account'],
        'PHONE_IMPORT_TEACHER_PHONE': teacher['phone'],
        'PHONE_IMPORT_TEACHER_TEMPLATE': teacher['template'],
        'PHONE_IMPORT_STUDENT_ACCOUNT': student['account'],
        'PHONE_IMPORT_STUDENT_PHONE': student['phone'],
        'PHONE_IMPORT_STUDENT_CHANGED_PHONE': changed_phone,
        'PHONE_IMPORT_STUDENT_TEMPLATE': student['template'],
        'PHONE_TEST_MINI_BASE_URL': 'http://localhost:5188',
        'PHONE_SMS_MAILBOX': str(mailbox),
    }
    server = subprocess.Popen([
        sys.executable, '-m', 'uvicorn', 'app.main:app', '--host', '127.0.0.1',
        '--port', '18310', '--no-access-log'], cwd=root / 'backend', env=env,
        stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    worker = threading.Thread(target=run_local_workers, name='phone-browser-local-workers', daemon=True)
    worker.start()
    try:
        ready = False
        deadline = time.monotonic() + 45
        while time.monotonic() < deadline and server.poll() is None:
            try:
                with urllib.request.urlopen('http://127.0.0.1:18310/health', timeout=1) as response:
                    ready = response.status == 200
                if ready:
                    break
            except OSError:
                time.sleep(.2)
        assert ready, 'Dedicated local backend did not become healthy'
        import_command = ['node', 'node_modules/@playwright/test/cli.js', 'test',
                          '--config=playwright.phone.config.mjs', '--grep',
                          'real xlsx import creates accounts through Staff PC']
        result = subprocess.run(import_command, cwd=root / 'e2e', env=env, timeout=300,
                                capture_output=True, text=True, encoding='utf-8', errors='replace')
        assert result.returncode == 0, result.stdout + result.stderr
        assert not worker_errors, '; '.join(worker_errors)

        def imported_subject(db, kind, expected):
            user = db.scalar(select(User).where(
                User.tenant_id == phone_identity['tenant_id'], User.login_name == expected['account']))
            assert user is not None and user.must_change_password is True
            candidate = db.scalar(select(PhoneLoginCandidate).where(
                PhoneLoginCandidate.tenant_id == phone_identity['tenant_id'],
                PhoneLoginCandidate.user_id == user.id))
            assert candidate is not None and candidate.state == 'PENDING'
            assert db.scalar(select(PhoneLoginBinding.id).where(
                PhoneLoginBinding.tenant_id == phone_identity['tenant_id'],
                PhoneLoginBinding.user_id == user.id)) is None
            job = db.get(ImportJob, int(candidate.source_job_id))
            assert job is not None and job.status == 'SUCCEEDED'
            source = db.get(FileObject, int(job.source_file_id))
            assert source.scan_status == 'CLEAN' and source.status == 'AVAILABLE'
            receipt = db.get(FileObject, int(job.result_json['credentialReceiptFileId']))
            workbook = load_workbook(tmp_path / receipt.file_key, read_only=True, data_only=True)
            values = list(workbook.active.values)
            workbook.close()
            header_index, header = next(
                (index, [str(value or '') for value in row])
                for index, row in enumerate(values) if '初始密码' in row
            )
            record = next(row for row in values[header_index + 1:]
                          if str(row[1]) == expected['account'])
            initial_password = str(record[header.index('初始密码')])
            assert initial_password and initial_password not in str(job.result_json)
            if kind == 'STUDENT':
                link = db.scalar(select(StudentAccountLink).where(
                    StudentAccountLink.tenant_id == phone_identity['tenant_id'],
                    StudentAccountLink.user_id == user.id))
                assert link and db.get(StudentProfile, link.student_id).student_no == expected['account']
            return {'user_id': user.id, 'job_id': str(job.id), 'initial_password': initial_password}

        with get_sessionmaker()() as db:
            teacher.update(imported_subject(db, 'TEACHER', teacher))
            student.update(imported_subject(db, 'STUDENT', student))

        env.update(
            PHONE_IMPORT_TEACHER_INITIAL_PASSWORD=teacher['initial_password'],
            PHONE_IMPORT_STUDENT_INITIAL_PASSWORD=student['initial_password'],
        )
        flow_command = ['node', 'node_modules/@playwright/test/cli.js', 'test',
                        '--config=playwright.phone.config.mjs', '--grep',
                        'imported accounts complete verified bind']
        result = subprocess.run(flow_command, cwd=root / 'e2e', env=env, timeout=300,
                                capture_output=True, text=True, encoding='utf-8', errors='replace')
        assert result.returncode == 0, result.stdout + result.stderr
        assert not worker_errors, '; '.join(worker_errors)

        with get_sessionmaker()() as db:
            teacher_user = db.get(User, teacher['user_id'])
            student_user = db.get(User, student['user_id'])
            teacher_binding = db.scalar(select(PhoneLoginBinding).where(
                PhoneLoginBinding.tenant_id == phone_identity['tenant_id'],
                PhoneLoginBinding.user_id == teacher_user.id))
            student_binding = db.scalar(select(PhoneLoginBinding).where(
                PhoneLoginBinding.tenant_id == phone_identity['tenant_id'],
                PhoneLoginBinding.user_id == student_user.id))
            teacher_candidate = db.scalar(select(PhoneLoginCandidate).where(
                PhoneLoginCandidate.tenant_id == phone_identity['tenant_id'],
                PhoneLoginCandidate.user_id == teacher_user.id))
            student_candidate = db.scalar(select(PhoneLoginCandidate).where(
                PhoneLoginCandidate.tenant_id == phone_identity['tenant_id'],
                PhoneLoginCandidate.user_id == student_user.id))
            assert teacher_user.login_name == teacher['account'] and not teacher_user.must_change_password
            assert student_user.login_name == student['account'] and not student_user.must_change_password
            assert teacher_binding.state == 'VERIFIED'
            assert teacher_binding.active_phone_lookup == phone_lookup(phone_identity['tenant_id'], teacher['phone'])
            assert teacher_candidate.state == 'APPLIED'
            assert teacher_binding.source_candidate_id == teacher_candidate.id
            assert teacher_binding.source_job_id == int(teacher['job_id'])
            assert student_binding.state == 'VERIFIED' and student_binding.version == 2
            assert student_binding.active_phone_lookup == phone_lookup(phone_identity['tenant_id'], changed_phone)
            assert student_binding.active_phone_lookup != phone_lookup(
                phone_identity['tenant_id'], student['phone'])
            assert student_candidate.state == 'APPLIED'
            assert student_binding.source_candidate_id == student_candidate.id
            assert db.scalar(select(func.count(SecurityAuditLog.id)).where(
                SecurityAuditLog.tenant_id == phone_identity['tenant_id'],
                SecurityAuditLog.action == 'PHONE_BINDING_CHANGE',
                SecurityAuditLog.resource_id.in_((str(teacher_user.id), str(student_user.id))))) == 3
    finally:
        worker_stop.set()
        worker.join(timeout=5)
        server.terminate()
        server.wait(timeout=15)
        mailbox.unlink(missing_ok=True)
