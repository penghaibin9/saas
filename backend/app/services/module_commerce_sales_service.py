"""M3 module sales workspace: live catalogue, priced drafts and scoped ledger.

Only the existing itemized writer persists an unpaid order. Preview is not a
reservation or entitlement. The writer invokes our context check in its own
transaction, after durable replay detection and before any new order insert.
"""
from __future__ import annotations

import base64
from datetime import datetime, timezone
from decimal import Decimal, localcontext
from io import BytesIO

from sqlalchemy import func, or_, select

from app.core.exceptions import AppException
from app.db.session import db_enabled, get_sessionmaker
from app.services.commercial_catalog_contract import (
    ContractError, MODULE_FEATURES, Snapshot, canonical, code,
    compile_order_draft, identifier, instant, integer, money,
)
from app.services import commercial_catalog_service as catalog

MODULE_LABELS = {"internship": "岗位实习中心", "graduationDesign": "毕业设计中心",
                 "studentAffairs": "学工中心", "academicAffairs": "教务中心"}
ORDER_STATUSES = {"unpaid", "paid", "cancelled"}
MAX_EXPORT_ROWS = 1000


def _db_required():
    if not db_enabled():
        raise AppException("SERVER_ERROR", "商业销售工作区需要数据库", http_status=503)


def _invalid(message):
    return AppException("VALIDATION_ERROR", message, http_status=422)


def _tid(value):
    try:
        return int(identifier(str(value) if type(value) is int else value, "tenantId"))
    except ContractError as exc:
        raise _invalid(str(exc)) from exc


def _page(page, page_size):
    try:
        integer(page, "page", minimum=1, maximum=10000)
        integer(page_size, "pageSize", minimum=1, maximum=100)
    except ContractError as exc:
        raise _invalid(str(exc)) from exc
    return (page - 1) * page_size


def _iso(value):
    if value is None:
        return None
    aware = value.replace(tzinfo=timezone.utc) if value.tzinfo is None else value
    return aware.astimezone(timezone.utc).isoformat(timespec="auto").replace("+00:00", "Z")


def _tenant(db, tenant_id):
    from app.models import Tenant
    row = db.get(Tenant, _tid(tenant_id))
    if row is None or row.is_deleted:
        raise AppException("DATA_NOT_FOUND", "学校不存在", http_status=404)
    return row


def list_sales_tenants(*, keyword="", page=1, page_size=30):
    _db_required()
    from app.models import Tenant
    offset = _page(page, page_size)
    text = str(keyword or "").strip()
    if len(text) > 100:
        raise _invalid("学校检索词最多100字符")
    filters = [Tenant.is_deleted.is_(False), Tenant.id > 0]
    if text:
        filters.append(or_(Tenant.school_name.contains(text, autoescape=True),
                           Tenant.tenant_code.contains(text, autoescape=True)))
    with get_sessionmaker()() as db:
        total = db.scalar(select(func.count()).select_from(Tenant).where(*filters))
        rows = db.scalars(select(Tenant).where(*filters).order_by(Tenant.school_name, Tenant.id)
                          .offset(offset).limit(page_size)).all()
        return {"items": [{"tenantId": str(r.id), "tenantName": r.school_name,
                           "tenantCode": r.tenant_code, "status": r.status} for r in rows],
                "total": int(total or 0), "page": page, "pageSize": page_size}


def list_sale_skus(*, module_key="", keyword="", page=1, page_size=30):
    _db_required()
    from app.models import CommercialSkuVersion
    offset = _page(page, page_size)
    if module_key and module_key not in MODULE_FEATURES:
        raise _invalid("请选择四个可售核心模块之一")
    text = str(keyword or "").strip()
    if len(text) > 100:
        raise _invalid("商品检索词最多100字符")
    filters = [CommercialSkuVersion.tenant_id == 0, CommercialSkuVersion.is_deleted.is_(False),
               CommercialSkuVersion.publish_status == "PUBLISHED",
               CommercialSkuVersion.product_type == "MODULE",
               CommercialSkuVersion.module_key.in_(tuple(MODULE_FEATURES))]
    if module_key:
        filters.append(CommercialSkuVersion.module_key == module_key)
    if text:
        filters.append(or_(CommercialSkuVersion.sku_code.contains(text, autoescape=True),
                           CommercialSkuVersion.snapshot_json["name"].as_string().contains(text, autoescape=True)))
    with get_sessionmaker()() as db:
        total = db.scalar(select(func.count()).select_from(CommercialSkuVersion).where(*filters))
        rows = db.scalars(select(CommercialSkuVersion).where(*filters)
                          .order_by(CommercialSkuVersion.sku_code, CommercialSkuVersion.sku_revision.desc())
                          .offset(offset).limit(page_size)).all()
        items = []
        for row in rows:
            snapshot = Snapshot(canonical(row.snapshot_json))
            data = snapshot.as_dict()
            if (snapshot.content_hash != row.content_hash or data.get("skuCode") != row.sku_code
                    or data.get("revision") != row.sku_revision or data.get("moduleKey") != row.module_key
                    or data.get("productType") != "MODULE"):
                raise AppException("DATA_CONFLICT", "商品快照不一致，停止展示可售商品", http_status=409)
            items.append({**data, "contentHash": snapshot.content_hash, "publishStatus": row.publish_status})
        return {"items": items, "total": int(total or 0), "page": page, "pageSize": page_size,
                "modules": [{"moduleKey": k, "label": MODULE_LABELS[k],
                             "approvedFeatures": list(catalog.APPROVED_FEATURE_SCOPES[k])}
                            for k in MODULE_FEATURES]}


def _sales_context(db, tenant, *, lock=False):
    """Bounded state lookup and paid-source aggregate; never use a UI entitlement flag."""
    from app.models import (CommercialOrderItem, PlatformOrder, TenantCommercialProfile,
                            TenantModuleState, TenantModuleSubscriptionSource)
    query = select(TenantModuleState).where(TenantModuleState.tenant_id == tenant.id,
                TenantModuleState.is_deleted.is_(False)).order_by(TenantModuleState.module_key)
    if lock:
        query = query.with_for_update().execution_options(populate_existing=True)
    states = {r.module_key: {"generation": int(r.generation), "dataState": r.data_state}
              for r in db.scalars(query).all()}
    # Join all owners, generations and paid headers. A cancelled/unpaid/foreign
    # source cannot provide the renewal boundary for this school.
    paid = db.execute(select(TenantModuleSubscriptionSource.module_key,
                             func.max(TenantModuleSubscriptionSource.ends_at))
        .join(TenantModuleState, (TenantModuleState.tenant_id == TenantModuleSubscriptionSource.tenant_id)
              & (TenantModuleState.module_key == TenantModuleSubscriptionSource.module_key)
              & (TenantModuleState.generation == TenantModuleSubscriptionSource.module_generation))
        .join(CommercialOrderItem, (CommercialOrderItem.id == TenantModuleSubscriptionSource.order_item_id)
              & (CommercialOrderItem.tenant_id == TenantModuleSubscriptionSource.tenant_id)
              & (CommercialOrderItem.module_key == TenantModuleSubscriptionSource.module_key)
              & (CommercialOrderItem.module_generation == TenantModuleSubscriptionSource.module_generation))
        .join(PlatformOrder, (PlatformOrder.id == CommercialOrderItem.order_id)
              & (PlatformOrder.tenant_id == CommercialOrderItem.tenant_id))
        .where(TenantModuleSubscriptionSource.tenant_id == tenant.id,
               TenantModuleSubscriptionSource.source_type == "PAID_ORDER_ITEM",
               TenantModuleSubscriptionSource.status.in_(("ACTIVE", "SCHEDULED")),
               TenantModuleSubscriptionSource.is_deleted.is_(False),
               TenantModuleState.is_deleted.is_(False), CommercialOrderItem.is_deleted.is_(False),
               PlatformOrder.is_deleted.is_(False), PlatformOrder.status == "paid")
        .group_by(TenantModuleSubscriptionSource.module_key)).all()
    profile = db.scalars(select(TenantCommercialProfile).where(
        TenantCommercialProfile.tenant_id == tenant.id, TenantCommercialProfile.is_deleted.is_(False))).first()
    return {"tenantId": str(tenant.id), "tenantStatus": tenant.status, "tenantName": tenant.school_name,
            "states": states, "paidThrough": {k: _iso(v) for k, v in paid},
            "readerVersion": profile.reader_version if profile else "LEGACY"}


def validate_sales_context(compiled, order_type, context):
    """Pure sale guard, reused at preview and under the order writer's parent lock."""
    if context["tenantStatus"] != "ACTIVE":
        raise AppException("DATA_CONFLICT", "学校当前不可销售，请先处理学校状态", http_status=409)
    if order_type not in {"NEW", "RENEW"}:
        raise _invalid("此工作区仅办理新购与续费，不自动执行升级或退款")
    if compiled["tenantId"] != context["tenantId"]:
        raise AppException("DATA_CONFLICT", "订单学校与办理上下文不一致", http_status=409)
    seen = set()
    for line in compiled["items"]:
        module = line["skuSnapshot"].get("moduleKey")
        if line["skuSnapshot"].get("productType") != "MODULE" or module not in MODULE_FEATURES or module in seen:
            raise _invalid("一张销售单每个核心模块只列一行；组合商品和附加能力另行办理")
        seen.add(module)
        state = context["states"].get(module)
        generation = state["generation"] if state else 1
        if state and state["dataState"] != "AVAILABLE":
            raise AppException("DATA_CONFLICT", f"{MODULE_LABELS[module]}已冻结、保留或销毁，不能在此恢复历史数据", http_status=409)
        if line["requestedGeneration"] != generation:
            raise AppException("DATA_CONFLICT", "模块代次已变化，请重新预检；未创建新订单", http_status=409)
        if order_type == "RENEW":
            end = context["paidThrough"].get(module)
            if not end:
                raise AppException("DATA_CONFLICT", f"{MODULE_LABELS[module]}没有当前代次的已付分项来源，不能冒充续费", http_status=409)
            if instant(line["startAt"], "startAt") < instant(end, "paidThrough"):
                raise AppException("DATA_CONFLICT", "续费起点早于已付服务截止，请避免重复计费区间", http_status=409)


def compile_sales_preview(body, context, *, resolve_sku):
    """Derive money/hash/generation server-side. Prices are explicit agreed inputs."""
    if not isinstance(body, dict) or set(body) - {"tenantId", "orderType", "remark", "items"}:
        raise _invalid("销售预检参数不完整或包含未知字段")
    raw_items = body.get("items")
    if not isinstance(raw_items, list) or not 1 <= len(raw_items) <= 4:
        raise _invalid("请选择1至4个核心模块")
    order_type = str(body.get("orderType") or "NEW")
    remark = str(body.get("remark") or "").strip()
    if not 5 <= len(remark) <= 500:
        raise _invalid("请填写5至500字符的合同/报价说明")
    items, snapshots, currencies = [], {}, set()
    try:
        tid = identifier(body.get("tenantId"), "tenantId")
        total = Decimal("0.00")
        for index, raw in enumerate(raw_items, 1):
            allowed = {"skuCode", "skuRevision", "skuContentHash", "quantity", "unitPrice", "discountAmount", "startAt", "endAt"}
            if not isinstance(raw, dict) or set(raw) != allowed:
                raise ContractError(f"第{index}行参数不完整或包含未知字段")
            key = (code(raw["skuCode"], "skuCode"), integer(raw["skuRevision"], "skuRevision", minimum=1))
            snapshot = resolve_sku(*key)
            snapshots[key] = snapshot
            sku = snapshot.as_dict()
            if snapshot.content_hash != raw["skuContentHash"]:
                raise ContractError("商品版本已变化，请重新选择")
            currencies.add(sku["pricePolicy"]["currency"])
            quantity = integer(raw["quantity"], "quantity", minimum=1, maximum=1000000)
            unit, discount = money(raw["unitPrice"], "unitPrice"), money(raw["discountAmount"], "discountAmount")
            with localcontext() as ctx:
                ctx.prec = 40
                net = unit * quantity - discount
                total += net
            state = context["states"].get(sku.get("moduleKey"))
            items.append({**raw, "lineNo": index, "netAmount": format(net, ".2f"),
                          "requestedGeneration": state["generation"] if state else 1})
        if len(currencies) != 1:
            raise ContractError("同一订单不能混用币种")
        contract = {"tenantId": tid, "currency": next(iter(currencies)),
                    "totalAmount": format(total, ".2f"), "items": items}
        compiled = compile_order_draft(contract, resolve_sku=lambda c, r: snapshots[c, r]).as_dict()
        validate_sales_context(compiled, order_type, context)
    except ContractError as exc:
        raise _invalid(str(exc)) from exc
    warnings = []
    for item in compiled["items"]:
        end = context["paidThrough"].get(item["moduleKey"])
        if end:
            warnings.append({"moduleKey": item["moduleKey"], "paidThrough": end,
                             "code": "PAID_COVERAGE_EXISTS" if order_type == "NEW" else "RENEWAL_BOUNDARY",
                             "message": "已有已付截止时间；新订单不会覆盖旧合同，请核对是否应选择续费。" if order_type == "NEW" else "续费为新增付费来源，不修改原合同；起点晚于此时间会形成服务空档。"})
    return {"order": {**contract, "orderType": order_type, "remark": remark},
            "items": compiled["items"], "totalAmount": compiled["totalAmount"], "currency": compiled["currency"],
            "tenantId": tid, "tenantName": context["tenantName"], "readerVersion": context["readerVersion"],
            "warnings": warnings, "validationOnly": True, "paymentRecorded": False, "rightsMaterialized": False}


def get_sales_context(tenant_id):
    _db_required()
    with get_sessionmaker()() as db:
        return _sales_context(db, _tenant(db, tenant_id))


def preview_sales_order(body):
    _db_required()
    with get_sessionmaker()() as db:
        tenant = _tenant(db, body.get("tenantId") if isinstance(body, dict) else None)
        return compile_sales_preview(body, _sales_context(db, tenant),
                                     resolve_sku=lambda c, r: catalog.get_sku_snapshot(c, r, db_session=db))


def _validate_new_sales_order(db, tenant, compiled, order_type):
    validate_sales_context(compiled, order_type, _sales_context(db, tenant, lock=True))


def create_sales_order(body, *, idempotency_key, actor_id):
    # Do not accept UI-only metadata or silently drop a changed quotation field.
    if not isinstance(body, dict) or set(body) != {"tenantId", "currency", "totalAmount", "items", "orderType", "remark"}:
        raise _invalid("请先完成销售预检，提交其完整订单合同")
    if not isinstance(body["remark"], str) or not 5 <= len(body["remark"].strip()) <= 500:
        raise _invalid("请填写5至500字符的合同/报价说明")
    from app.services import commercial_order_item_service as orders
    return orders.create_itemized_order(body, idempotency_key=idempotency_key,
                                       actor_id=actor_id, context_validator=_validate_new_sales_order)


def _order_filters(tenant_id, status):
    from app.models import PlatformOrder
    if status and status not in ORDER_STATUSES:
        raise _invalid("订单状态必须为未支付、已支付或已取消")
    filters = [PlatformOrder.tenant_id == _tid(tenant_id), PlatformOrder.is_deleted.is_(False),
               PlatformOrder.package_code == "MODULE_V2"]
    if status:
        filters.append(PlatformOrder.status == status)
    return filters


def _order_dict(row):
    return {"orderId": str(row.id), "tenantId": str(row.tenant_id), "orderNo": row.order_no,
            "orderType": row.order_type, "status": row.status, "totalAmount": format(row.amount, ".2f"),
            "paidAmount": format(row.paid_amount, ".2f"), "version": int(row.version or 0),
            "startAt": _iso(row.start_at), "endAt": _iso(row.end_at),
            "createdAt": _iso(row.created_at), "remark": row.remark or ""}


def _item_dict(row):
    return {"lineNo": int(row.line_no), "moduleKey": row.module_key, "skuCode": row.sku_code,
            "skuRevision": row.sku_revision, "skuContentHash": row.sku_content_hash,
            "generation": row.module_generation, "quantity": int(row.quantity),
            "unitPrice": format(row.unit_price, ".2f"), "discountAmount": format(row.discount_amount, ".2f"),
            "netAmount": format(row.net_amount, ".2f"), "currency": row.currency,
            "startAt": _iso(row.service_start_at), "endAt": _iso(row.service_end_at),
            "fulfillmentStatus": row.fulfillment_status,
            "features": dict(row.feature_snapshot_json or {}), "quotas": dict(row.quota_snapshot_json or {})}


def list_sales_orders(tenant_id, *, status="", page=1, page_size=20):
    _db_required()
    from app.models import CommercialOrderItem, PlatformOrder
    offset = _page(page, page_size)
    filters = _order_filters(tenant_id, status)
    with get_sessionmaker()() as db:
        _tenant(db, tenant_id)
        total = db.scalar(select(func.count()).select_from(PlatformOrder).where(*filters))
        rows = db.scalars(select(PlatformOrder).where(*filters).order_by(PlatformOrder.id.desc())
                          .offset(offset).limit(page_size)).all()
        ids = [r.id for r in rows]
        counts = {oid: (count, low if low == high else "MIXED") for oid, count, low, high in db.execute(select(CommercialOrderItem.order_id, func.count(),
            func.min(CommercialOrderItem.currency), func.max(CommercialOrderItem.currency))
            .where(CommercialOrderItem.tenant_id == _tid(tenant_id), CommercialOrderItem.order_id.in_(ids),
                   CommercialOrderItem.is_deleted.is_(False)).group_by(CommercialOrderItem.order_id)).all()} if ids else {}
        return {"items": [{**_order_dict(r), "itemCount": counts.get(r.id, (0, "UNKNOWN"))[0],
                           "currency": counts.get(r.id, (0, "UNKNOWN"))[1]} for r in rows],
                "total": int(total or 0), "page": page, "pageSize": page_size}


def get_sales_order(tenant_id, order_id):
    _db_required()
    from app.models import CommercialOrderItem, PlatformOrder
    with get_sessionmaker()() as db:
        row = db.scalars(select(PlatformOrder).where(*_order_filters(tenant_id, ""),
                                                    PlatformOrder.id == _tid(order_id))).first()
        if row is None:
            raise AppException("DATA_NOT_FOUND", "此学校的分项订单不存在", http_status=404)
        items = db.scalars(select(CommercialOrderItem).where(
            CommercialOrderItem.tenant_id == row.tenant_id, CommercialOrderItem.order_id == row.id,
            CommercialOrderItem.is_deleted.is_(False)).order_by(CommercialOrderItem.line_no).limit(101)).all()
        if len(items) > 100:
            raise AppException("DATA_CONFLICT", "订单分项超过合同上限，需核对数据", http_status=409)
        return {**_order_dict(row), "items": [_item_dict(i) for i in items],
                "paymentRecorded": row.status == "paid", "snapshotSource": "ORDER_ITEM_NOT_LIVE_CATALOGUE"}


def export_sales_orders(body):
    _db_required()
    from openpyxl import Workbook
    from openpyxl.cell import WriteOnlyCell
    from app.models import CommercialOrderItem, PlatformOrder
    from app.services import audit_log
    if not isinstance(body, dict) or set(body) - {"tenantId", "status", "reason"}:
        raise _invalid("导出参数无效")
    reason = str(body.get("reason") or "").strip()
    if not 5 <= len(reason) <= 500:
        raise _invalid("导出原因须为5至500字符")
    tid = _tid(body.get("tenantId"))
    with get_sessionmaker()() as db:
        tenant = _tenant(db, tid)
        rows = db.execute(select(PlatformOrder, CommercialOrderItem)
            .join(CommercialOrderItem, (CommercialOrderItem.order_id == PlatformOrder.id)
                  & (CommercialOrderItem.tenant_id == PlatformOrder.tenant_id))
            .where(*_order_filters(tid, body.get("status") or ""), CommercialOrderItem.is_deleted.is_(False))
            .order_by(PlatformOrder.id.desc(), CommercialOrderItem.line_no).limit(MAX_EXPORT_ROWS + 1)).all()
        if len(rows) > MAX_EXPORT_ROWS:
            raise AppException("EXPORT_LIMIT_EXCEEDED", "超过1000条分项，请按订单状态缩小范围；未截断导出", http_status=413)
        workbook = Workbook(write_only=True)
        sheet = workbook.create_sheet("模块分项订单")
        headers = ["学校ID", "学校", "订单号", "订单类型", "支付状态", "订单总额（勿逐行累加）", "已付金额（勿逐行累加）", "行号", "模块", "SKU", "商品版本", "代次", "数量", "成交单价", "优惠金额", "分项净额", "币种", "服务开始UTC", "服务截止UTC", "履约状态", "商品指纹"]
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = f"A1:U{len(rows) + 1}"
        from openpyxl.styles import Font, PatternFill
        from openpyxl.utils import get_column_letter
        for index in range(1, 22):
            sheet.column_dimensions[get_column_letter(index)].width = 24 if index != 21 else 68
        title_cells = []
        for title in headers:
            cell = WriteOnlyCell(sheet, value=title)
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="163D64")
            title_cells.append(cell)
        sheet.append(title_cells)
        for order, item in rows:
            values = [str(tid), tenant.school_name, order.order_no, order.order_type, order.status,
                      format(order.amount, ".2f"), format(order.paid_amount, ".2f"), item.line_no,
                      MODULE_LABELS.get(item.module_key, item.module_key), item.sku_code, str(item.sku_revision),
                      str(item.module_generation), item.quantity, format(item.unit_price, ".2f"),
                      format(item.discount_amount, ".2f"), format(item.net_amount, ".2f"), item.currency,
                      _iso(item.service_start_at), _iso(item.service_end_at), item.fulfillment_status, item.sku_content_hash]
            cells = []
            for value in values:
                cell = WriteOnlyCell(sheet, value=value)
                if isinstance(value, str):
                    cell.data_type = "s"  # No formula execution or BIGINT/scientific-notation loss.
                cells.append(cell)
            sheet.append(cells)
        output = BytesIO()
        workbook.save(output)
        audit_log.record_critical_in_session(db, "COMMERCIAL_SALES_LEDGER_EXPORT", f"tenant:{tid}",
            detail={"reason": reason, "rowCount": len(rows), "status": body.get("status") or "ALL"},
            tenant_id=tid, resource_id=str(tid))
        db.commit()  # An audit failure must prevent the response/download.
        return {"filename": f"module-orders-{tid}.xlsx", "rowCount": len(rows),
                "mediaType": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                "contentBase64": base64.b64encode(output.getvalue()).decode("ascii"), "truncated": False}
