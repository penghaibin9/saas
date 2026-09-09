"""Bounded classroom catalogue construction with persisted previews and atomic receipts."""
from __future__ import annotations

from datetime import datetime, timedelta
from io import BytesIO
from uuid import uuid4
from zipfile import ZipFile, BadZipFile

from openpyxl import Workbook, load_workbook
from sqlalchemy import select, func
from sqlalchemy.exc import IntegrityError

from app.core.exceptions import AppException, not_found
from app.models import AaClassroom, AaTeachingBuilding
from app.services.db_service import _tid, session
from app.services import shared_import_batch_service as batches
from app.services.import_export_service import _excel_safe
from app.modules.academic_affairs.services import academic_affairs_resource_service as rooms

NAMESPACE = "AA_CLASSROOM_BUILD"
MAX_ROWS = 1000
HEADERS = {"楼栋编码": "buildingCode", "楼层": "floorNo", "教室编号": "roomCode",
           "教室名称": "roomName", "教室类型": "roomType", "教学座位": "capacity",
           "考试座位": "examSeats", "专用教室": "isExclusive", "备注": "remark"}


def _bad(message):
    raise AppException("VALIDATION_ERROR", message)


def _text(value, label, size, required=False):
    text = str(value if value is not None else "").strip()
    if (required and not text) or len(text) > size:
        _bad(f"{label}{'必填且' if required else ''}最多 {size} 字")
    return text


def _integer(value, label, low, high):
    if isinstance(value, bool):
        _bad(f"{label}须为 {low}–{high} 的整数")
    try:
        n = int(value)
        if str(n) != str(value).strip() and not (isinstance(value, float) and value == n):
            raise ValueError()
    except (ValueError, TypeError):
        _bad(f"{label}须为 {low}–{high} 的整数")
    if not low <= n <= high:
        _bad(f"{label}须为 {low}–{high} 的整数")
    return n


def _building(db, building_id, lock=False):
    q = select(AaTeachingBuilding).where(AaTeachingBuilding.id == int(building_id),
        AaTeachingBuilding.tenant_id == _tid(), AaTeachingBuilding.is_deleted.is_(False))
    b = db.scalar(q.with_for_update() if lock else q)
    if not b:
        raise not_found("教学楼不存在")
    return b


def _building_row(b):
    return {"buildingId": str(b.id), "buildingCode": b.building_code, "buildingName": b.building_name,
            "campusCode": b.campus_code or "", "floorCount": b.floor_count, "version": b.version}


def list_buildings(keyword="", page=1):
    with session() as db:
        conditions = [AaTeachingBuilding.tenant_id == _tid(), AaTeachingBuilding.is_deleted.is_(False)]
        if keyword:
            conditions.append((AaTeachingBuilding.building_name.contains(keyword)) |
                              (AaTeachingBuilding.building_code.contains(keyword)) |
                              (AaTeachingBuilding.campus_code.contains(keyword)))
        total = db.scalar(select(func.count()).select_from(AaTeachingBuilding).where(*conditions))
        buildings = db.scalars(select(AaTeachingBuilding).where(*conditions)
            .order_by(AaTeachingBuilding.building_code).offset((page - 1) * 50).limit(50)).all()
        counts = db.execute(select(AaClassroom.building_id, func.count(), func.sum(AaClassroom.capacity))
            .where(AaClassroom.tenant_id == _tid(), AaClassroom.is_deleted.is_(False),
                   AaClassroom.building_id.in_([b.id for b in buildings]))
            .group_by(AaClassroom.building_id)).all()
        stats = {bid: (count, capacity) for bid, count, capacity in counts}
        return {"items": [{**_building_row(b), "roomCount": stats.get(b.id, (0, 0))[0],
                "capacity": int(stats.get(b.id, (0, 0))[1])} for b in buildings], "total": total, "page": page}


def save_building(body, building_id=None):
    code = _text(body.buildingCode, "楼栋编码", 50, True)
    name = _text(body.buildingName, "楼栋名称", 100, True)
    campus = _text(body.campusCode, "校区", 50)
    try:
        with session() as db:
            if building_id:
                b = _building(db, building_id, True)
                if body.expectedVersion != b.version:
                    raise AppException("DATA_CONFLICT", "楼栋资料已变化，请刷新后核对")
                # Location identity is stable; changing labels would rewrite existing schedule snapshots.
                if (code, name, campus) != (b.building_code, b.building_name, b.campus_code or ""):
                    _bad("已建教学楼的编码、名称和校区保持不变，可扩展楼层")
                occupied = db.scalar(select(func.max(AaClassroom.floor_no)).where(
                    AaClassroom.tenant_id == _tid(), AaClassroom.building_id == b.id,
                    AaClassroom.is_deleted.is_(False))) or 0
                if body.floorCount < occupied:
                    _bad(f"第 {occupied} 层已有教室，不能缩减到该楼层以下")
                b.floor_count = body.floorCount
                b.version += 1
            else:
                # Never silently absorb historical rooms into a similarly named building.
                legacy = db.scalar(select(AaClassroom.id).where(AaClassroom.tenant_id == _tid(),
                    AaClassroom.building_code == code,
                    ((AaClassroom.building_name != name) |
                     (func.coalesce(AaClassroom.campus_code, "") != campus))).limit(1))
                if legacy:
                    raise AppException("DATA_CONFLICT", "该编码已有不同楼栋名称或校区的教室，请核对旧资料或使用全校唯一编码")
                b = AaTeachingBuilding(tenant_id=_tid(), building_code=code, building_name=name,
                    campus_code=campus or None, floor_count=body.floorCount)
                db.add(b)
                db.flush()
                # Exact identity match only; keep IDs, capacities, statuses and unknown floors untouched.
                db.query(AaClassroom).filter(AaClassroom.tenant_id == _tid(),
                    AaClassroom.building_code == code, AaClassroom.building_id.is_(None)).update(
                    {AaClassroom.building_id: b.id}, synchronize_session=False)
            rooms._audit(db, b.id, "UPDATE" if building_id else "CREATE", name, "AA_TEACHING_BUILDING")
            db.commit()
            return _building_row(b)
    except IntegrityError:
        raise AppException("DATA_CONFLICT", "楼栋编码在本校已存在，请刷新核对") from None


def bind_location(db, body, existing=None):
    code = getattr(body, "buildingCode", None) or (existing.building_code if existing else "")
    name = getattr(body, "buildingName", None) or (existing.building_name if existing else "")
    _text(code, "楼栋编码", 50, True)
    _text(name, "楼栋名称", 100, True)
    _text(getattr(body, "roomCode", None) or (existing.room_code if existing else ""), "教室编号", 50, True)
    bid = getattr(body, "buildingId", None)
    b = _building(db, bid) if bid else db.scalar(select(AaTeachingBuilding).where(
        AaTeachingBuilding.tenant_id == _tid(), AaTeachingBuilding.building_code == code,
        AaTeachingBuilding.is_deleted.is_(False)))
    floor = getattr(body, "floorNo", None)
    if existing and "floorNo" not in body.model_fields_set:
        floor = existing.floor_no if existing.building_code == code else None
    if b:
        campus = getattr(body, "campusCode", None)
        if campus is None and existing:
            campus = existing.campus_code
        if (code, name, campus or "") != (b.building_code, b.building_name, b.campus_code or ""):
            _bad("教室所在楼栋名称、编码与校区须与教学楼资料一致")
        if floor is not None and not 1 <= floor <= b.floor_count:
            _bad(f"楼层须在 1–{b.floor_count} 层")
    elif floor is not None:
        _bad("请先建立教学楼，再维护楼层")
    return (b.id if b else None), floor


def generate_preview(body):
    with session() as db:
        b = _building(db, body.buildingId)
        if body.startFloor > body.endFloor or body.endFloor > b.floor_count:
            _bad("楼层范围超出教学楼，或起止楼层颠倒")
        if (body.endFloor - body.startFloor + 1) * body.roomsPerFloor > MAX_ROWS:
            _bad(f"每批最多 {MAX_ROWS} 间教室，请分批建设")
        if body.startSequence + body.roomsPerFloor - 1 >= 10 ** body.digits:
            _bad("流水号位数不足，请增加位数或减少每层数量")
        excluded = set(body.excludeCodes)
        generated = [{"buildingCode": b.building_code, "floorNo": floor,
            "roomCode": f"{body.prefix}{floor}{seq:0{body.digits}d}", "capacity": body.capacity,
            "examSeats": body.examSeats, "roomType": body.roomType, "isExclusive": body.isExclusive}
            for floor in range(body.startFloor, body.endFloor + 1)
            for seq in range(body.startSequence, body.startSequence + body.roomsPerFloor)]
        known = {r["roomCode"] for r in generated}
        if excluded - known:
            _bad("跳过编号不在本次生成范围内，请核对")
        return preview_rows([r for r in generated if r["roomCode"] not in excluded])


def preview_rows(raw_rows):
    if not raw_rows or len(raw_rows) > MAX_ROWS:
        _bad(f"请提供 1–{MAX_ROWS} 行教室数据")
    items, seen, building_versions = [], set(), {}
    with session() as db:
        codes = {str(r.get("buildingCode", "")).strip() for r in raw_rows}
        catalog = {b.building_code: b for b in db.scalars(select(AaTeachingBuilding).where(
            AaTeachingBuilding.tenant_id == _tid(), AaTeachingBuilding.is_deleted.is_(False),
            AaTeachingBuilding.building_code.in_(codes))).all()}
        for number, raw in enumerate(raw_rows, 2):
            item = {"line": number, "input": raw, "action": "ERROR", "message": ""}
            try:
                code = _text(raw.get("buildingCode"), "楼栋编码", 50, True)
                b = catalog.get(code)
                if not b:
                    _bad("楼栋尚未建立，请先新建教学楼")
                building_versions[str(b.id)] = b.version
                row = {"buildingId": str(b.id), "buildingCode": code, "buildingName": b.building_name,
                    "campusCode": b.campus_code, "floorNo": _integer(raw.get("floorNo"), "楼层", 1, b.floor_count),
                    "roomCode": _text(raw.get("roomCode"), "教室编号", 50, True),
                    "roomName": _text(raw.get("roomName"), "教室名称", 100),
                    "capacity": _integer(raw.get("capacity"), "教学座位", 0, 1000),
                    "examSeats": None if raw.get("examSeats") in (None, "") else _integer(raw["examSeats"], "考试座位", 0, 1000),
                    "roomType": rooms._norm_type({v: k for k, v in rooms.ROOM_TYPE_LABEL.items()}.get(raw.get("roomType"), raw.get("roomType"))),
                    "remark": _text(raw.get("remark"), "备注", 500)}
                exclusive = raw.get("isExclusive", False)
                if exclusive not in (True, False, "是", "否", ""):
                    _bad("专用教室请填写“是”或“否”")
                row["isExclusive"] = exclusive in (True, "是")
                key = (code, row["roomCode"])
                if key in seen:
                    _bad("本批次编号重复")
                seen.add(key)
                item.update(row=row, action="CREATE", message="待创建")
            except AppException as e:
                item["message"] = str(e.message)
            items.append(item)
        # One bounded query, including deleted IDs: bulk creation must never resurrect old linked rooms.
        existing = {(c.building_code, c.room_code): c for c in db.scalars(select(AaClassroom).where(
            AaClassroom.tenant_id == _tid(), AaClassroom.building_code.in_(codes),
            AaClassroom.room_code.in_([i["row"]["roomCode"] for i in items if "row" in i]))).all()}
        for item in items:
            if item["action"] != "CREATE":
                continue
            row = item["row"]
            old = existing.get((row["buildingCode"], row["roomCode"]))
            if old:
                item.update(action="KEEP", message="历史编号已保留" if old.is_deleted else "已存在，保留原资料",
                            existingId=str(old.id), existing=rooms._row(old))
    batch_no = uuid4().hex
    payload = {"items": items, "buildingVersions": building_versions}
    batches.create(_tid(), NAMESPACE, batch_no, "PREVIEW", payload,
        errors=[i for i in items if i["action"] == "ERROR"], operator_key=rooms._op()[2])
    return _preview_dto(batch_no, payload)


def _preview_dto(batch_no, payload):
    items = payload["items"]
    return {"batchNo": batch_no, "items": items, "total": len(items),
        "createCount": sum(i["action"] == "CREATE" for i in items),
        "keepCount": sum(i["action"] == "KEEP" for i in items),
        "errorCount": sum(i["action"] == "ERROR" for i in items)}


def _owned_batch(db, batch_no, lock=False):
    b = batches._row(db, _tid(), NAMESPACE, batch_no, lock=lock)
    if b.operator_key != rooms._op()[2]:
        raise not_found("该建库批次不属于当前操作人")
    return b


def get_batch(batch_no):
    with session() as db:
        b = _owned_batch(db, batch_no)
        return {**_preview_dto(batch_no, b.payload_json), "status": b.status, "result": b.public_result_json}


def confirm_batch(batch_no):
    try:
        with session() as db:
            batch = _owned_batch(db, batch_no, True)
            if batch.status == "SUCCESS":
                return batch.public_result_json
            payload = batch.payload_json
            if batch.status != "PREVIEW" or any(i["action"] == "ERROR" for i in payload["items"]):
                _bad("本批次存在错误，请修正后重新预检")
            for bid, version in sorted(payload["buildingVersions"].items(), key=lambda x: int(x[0])):
                if _building(db, bid, True).version != version:
                    raise AppException("DATA_CONFLICT", "楼栋资料已变化，请重新预检")
            created = []
            for item in payload["items"]:
                if item["action"] != "CREATE":
                    continue
                r = item["row"]
                c = AaClassroom(tenant_id=_tid(), building_id=int(r["buildingId"]), floor_no=r["floorNo"],
                    building_code=r["buildingCode"], building_name=r["buildingName"], campus_code=r["campusCode"],
                    room_code=r["roomCode"], room_name=r["roomName"] or None, capacity=r["capacity"],
                    exam_seats=r["examSeats"], is_exclusive=r["isExclusive"], room_type=r["roomType"],
                    remark=r["remark"] or None, status="AVAILABLE")
                db.add(c)
                created.append(c)
            if not created:
                _bad("没有待创建的教室，已有资料均已保留")
            db.flush()  # Unique constraint also catches per-room writers between preview and confirm.
            for c in created:
                rooms._audit(db, c.id, "BATCH_CREATE", f"批次 {batch_no} · {c.building_name}{c.room_code}")
            result = {"batchNo": batch_no, "createdCount": len(created),
                "keptCount": sum(i["action"] == "KEEP" for i in payload["items"]),
                "classroomIds": [str(c.id) for c in created], "completedAt": datetime.utcnow().isoformat()}
            batch.public_result_json = result
            batch.status = "SUCCESS"
            batch.confirmed_at = datetime.utcnow()
            batch.expires_at = datetime.utcnow() + timedelta(days=7)
            db.commit()  # Rooms, audit and retry receipt share one transaction.
            return result
    except IntegrityError:
        raise AppException("DATA_CONFLICT", "预检后教室编号已被使用，本批未写入。请重新预检核对") from None


def workbook(rows, headers=None):
    wb = Workbook()
    ws = wb.active
    ws.title = "教室台账"
    headers = headers or list(HEADERS)
    ws.append(headers)
    for row in rows:
        ws.append([_excel_safe(v) for v in row])
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column in ws.columns:
        ws.column_dimensions[column[0].column_letter].width = 22
    for cells in ws.iter_rows(min_row=2):
        for cell in cells:
            if isinstance(cell.value, str):
                cell.number_format = "@"
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def import_preview(content):
    try:
        with ZipFile(BytesIO(content)) as z:
            if len(z.infolist()) > 200 or sum(i.file_size for i in z.infolist()) > 30 * 1024 * 1024:
                _bad("表格解压后过大，请按每批最多1000行拆分")
        wb = load_workbook(BytesIO(content), read_only=True, data_only=False)
    except (BadZipFile, ValueError, KeyError, OSError):
        _bad("无法读取 xlsx 文件，请使用下载的模板")
    try:
        ws = wb.active
        ws.reset_dimensions()  # Do not trust uploaded dimension metadata.
        iterator = ws.iter_rows(values_only=True)
        headers = [str(x or "").strip() for x in next(iterator, ())]
        if headers != list(HEADERS):
            _bad("模板列不匹配，请下载教室台账模板，保留列名和顺序")
        raw = []
        for line, values in enumerate(iterator, 2):
            if line > MAX_ROWS + 1:
                _bad("每批最多1000行，请拆分文件")
            if len(values) > len(headers) and any(v is not None for v in values[len(headers):]):
                _bad(f"第{line}行存在模板之外的列")
            if any(isinstance(v, str) and v.lstrip().startswith(("=", "+", "-", "@")) for v in values):
                _bad(f"第{line}行包含公式或危险文本，请改成普通文本")
            raw.append({key: values[i] if i < len(values) else None for i, key in enumerate(HEADERS.values())})
        return preview_rows(raw)
    finally:
        wb.close()


def batch_workbook(batch_no):
    data = get_batch(batch_no)
    return workbook([[i.get("existing", i.get("row", i["input"])).get(key, "") for key in HEADERS.values()] +
        ["已创建" if data["status"] == "SUCCESS" and i["action"] == "CREATE" else i["message"]]
        for i in data["items"]], list(HEADERS) + ["处理结果"])
