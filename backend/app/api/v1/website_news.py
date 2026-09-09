"""The website editorial plane: platform roots only; no school wildcard bypass."""
from __future__ import annotations
from contextlib import contextmanager
from io import BytesIO
from pathlib import PurePath
from typing import Literal
from fastapi import APIRouter, Depends, File, UploadFile, Query
from pydantic import BaseModel, ConfigDict, Field, StrictBool, StrictInt
from sqlalchemy import select, func, exists, and_
from sqlalchemy.exc import SQLAlchemyError
from fastapi.responses import Response, HTMLResponse
from app.core.exceptions import AppException
from app.core.response import success
from app.core.platform_principal import require_platform_root
from app.db.session import get_sessionmaker, DBNotEnabledError
from app.models.website_news import NewsPackage, NewsArticle, NewsMedia, NewsMediaLink, NewsAudit
from app.services.website_news import service, render
from app.services.website_news.package import parse_package, MAX_ARCHIVE, PackageError, CATEGORIES

router = APIRouter(prefix="/platform/website-news", tags=["官网内容运营"])
public_api = APIRouter(prefix="/website-news", tags=["官网公开资讯"])
public_pages = APIRouter(prefix="/news", tags=["官网资讯页面"])

def news_db():
    db = None
    try:
        db = get_sessionmaker()()
        if db.get_bind().dialect.name != "mysql":
            raise AppException("SERVER_ERROR", "新闻发布要求连接 MySQL", http_status=503)
        yield db  # Writers explicitly commit before reporting HTTP success.
    except service.NewsError as exc:
        if db: db.rollback()
        raise AppException("DATA_CONFLICT" if exc.status == 409 else "VALIDATION_ERROR", str(exc), http_status=exc.status) from exc
    except (DBNotEnabledError, SQLAlchemyError) as exc:
        if db: db.rollback()
        raise AppException("SERVER_ERROR", "新闻服务暂不可用，请稍后重试；不会将失败报告为发布成功", http_status=503) from exc
    except BaseException:
        if db: db.rollback()
        raise
    finally:
        if db: db.close()

class Confirm(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expectedVersion: StrictInt = Field(ge=0)
    reviewDigest: str = Field(pattern=r"^[a-f0-9]{64}$")
    confirmed: StrictBool
    excludedIds: list[str] = Field(default_factory=list, max_length=2000)
    intervalMinutes: StrictInt = Field(default=3, ge=1, le=60)

class Control(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expectedVersion: StrictInt = Field(ge=0)
    action: Literal["pause", "resume", "cancel"]

@router.get("/package-format")
def package_format(user=Depends(require_platform_root)):
    from app.services.website_news.format import GPT_PACKAGE_PROMPT
    return success({"prompt": GPT_PACKAGE_PROMPT})

@router.post("/packages")
def upload_package(user=Depends(require_platform_root), file: UploadFile=File(...), db=Depends(news_db)):
    try:
        if not (file.filename or "").lower().endswith(".zip"):
            raise PackageError("请上传 ZIP 新闻资源包")
        blob = file.file.read(MAX_ARCHIVE + 1)
        parsed = parse_package(blob)
    except PackageError as exc:
        raise AppException("VALIDATION_ERROR", str(exc), http_status=400) from exc
    finally:
        file.file.close()
    name = (file.filename or "news.zip").replace("\\", "/").split("/")[-1]
    p, reused = service.import_package(db, parsed, name, user.get("userId", "platform"))
    result={"package": service.summary(db,p), "reused": reused}
    db.commit()
    return success(result, "已解析，等待你整包审核确认；尚未发布")

@router.get("/packages")
def list_packages(page: int=Query(1,ge=1,le=10000), user=Depends(require_platform_root), db=Depends(news_db)):
    rows = db.scalars(select(NewsPackage).order_by(NewsPackage.created_at.desc(), NewsPackage.id.desc()).offset((page-1)*20).limit(20))
    return success({"items":[service.summary(db,p) for p in rows], "total":db.scalar(select(func.count()).select_from(NewsPackage)), "health":service.get_health(db)})

@router.get("/packages/{pid}")
def get_package(pid: str, page: int=Query(1,ge=1,le=10000), user=Depends(require_platform_root), db=Depends(news_db)):
    p = service.package(db,pid)
    rows = db.scalars(select(NewsArticle).where(NewsArticle.package_id==pid).order_by(NewsArticle.ordinal).offset((page-1)*100).limit(100))
    return success({"package":service.summary(db,p), "articles":[service.article_data(a) for a in rows], "page":page, "pageSize":100, "health":service.get_health(db)})

@router.get("/articles/{aid}")
def review_article(aid: str, user=Depends(require_platform_root), db=Depends(news_db)):
    a = db.get(NewsArticle,aid)
    if not a: raise service.NewsError("文章不存在",404)
    return success(service.article_data(a,full=True))

@router.post("/packages/{pid}/confirm")
def confirm_package(pid: str, body: Confirm, user=Depends(require_platform_root), db=Depends(news_db)):
    if not body.confirmed: raise AppException("VALIDATION_ERROR","请明确确认审核结果",http_status=400)
    p,reused=service.approve(db,pid,user.get("userId","platform"),body.expectedVersion,body.reviewDigest,body.excludedIds,body.intervalMinutes)
    result={"package":service.summary(db,p), "reused":reused, "health":service.get_health(db)}
    db.commit()
    return success(result, "整包已确认并排期；将在服务器持续发布全部合格新闻")

@router.post("/packages/{pid}/control")
def control_package(pid: str, body: Control, user=Depends(require_platform_root), db=Depends(news_db)):
    p=service.control(db,pid,user.get("userId","platform"),body.expectedVersion,body.action)
    result=service.summary(db,p)
    db.commit()
    return success(result)

@router.post("/articles/{aid}/withdraw")
def withdraw(aid: str, user=Depends(require_platform_root), db=Depends(news_db)):
    result=service.article_data(service.withdraw(db,aid,user.get("userId","platform")))
    db.commit()
    return success(result)

@router.get("/packages/{pid}/ledger.xlsx")
def export_ledger(pid: str, user=Depends(require_platform_root), db=Depends(news_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    p=service.package(db,pid)
    rows=list(db.scalars(select(NewsArticle).where(NewsArticle.package_id==pid).order_by(NewsArticle.ordinal).limit(2001)))
    if len(rows)>2000: raise AppException("VALIDATION_ERROR","单包台账过大")
    wb=Workbook();ws=wb.active;ws.title="新闻发布台账"
    ws.append(["序号","标题","分类","状态","问题说明","预计发布时间（UTC）","实际发布时间（UTC）","公开地址"])
    def cell(v):
        text=str(v or "")
        return "'"+text if text.lstrip().startswith(("=","+","-","@")) else text
    for a in rows:
        ws.append([a.ordinal,cell(a.title),cell(CATEGORIES.get(a.category)),a.state,cell(a.issue),service.iso(a.scheduled_at),service.iso(a.published_at),render.origin()+"/news/"+a.slug])
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="193A62")
    for col,width in zip("ABCDEFGH",[8,50,18,18,52,27,27,70]): ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    log=wb.create_sheet("操作审计");log.append(["时间（UTC）","操作人","动作","文章ID","记录"])
    audit_rows = list(db.scalars(select(NewsAudit).where(NewsAudit.package_id==pid).order_by(NewsAudit.created_at).limit(10001)))
    if len(audit_rows) > 10000:
        raise AppException("VALIDATION_ERROR", "本包审计记录超过单次导出安全上限，请联系运维分页导出，未生成不完整台账")
    for a in audit_rows:
        log.append([service.iso(a.created_at),cell(a.actor),a.action,a.article_id,cell(str(a.detail))])
    log.freeze_panes="A2"
    out=BytesIO();wb.save(out)
    service.audit(db,p,user.get("userId","platform"),"EXPORT_LEDGER",{"rows":len(rows)})
    db.commit()
    return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="news-{pid}.xlsx"',"Cache-Control":"no-store"})

def published(db,category=None):
    query=select(NewsArticle).where(NewsArticle.state=="PUBLISHED")
    if category in CATEGORIES: query=query.where(NewsArticle.category==category)
    return query

@public_api.get("/latest")
def latest(db=Depends(news_db)):
    rows=db.scalars(published(db).order_by(NewsArticle.published_at.desc(),NewsArticle.id.desc()).limit(3))
    return success({"items":[service.article_data(a) for a in rows]})

@public_pages.get("/media/{mid}")
def media(mid: str, db=Depends(news_db)):
    # An uploaded asset is not public until a published article references it.
    import re
    if not re.fullmatch(r"[a-f0-9]{64}",mid): return Response(status_code=404)
    visible=db.scalar(select(NewsArticle.id).join(NewsMediaLink, NewsMediaLink.article_id==NewsArticle.id).where(
        NewsArticle.state=="PUBLISHED", NewsMediaLink.media_id==mid).limit(1))
    if not visible: return Response(status_code=404,headers={"Cache-Control":"no-store"})
    m=db.get(NewsMedia,mid)
    if not m: return Response(status_code=404)
    return Response(m.content,media_type=m.mime,headers={"X-Content-Type-Options":"nosniff","Cache-Control":"no-cache"})

@public_pages.get("/sitemap.xml")
def sitemap(page: int=Query(1,ge=1,le=10000), db=Depends(news_db)):
    from xml.etree.ElementTree import Element,SubElement,tostring
    total=db.scalar(select(func.count()).select_from(NewsArticle).where(NewsArticle.state=="PUBLISHED"))
    namespace="http://www.sitemaps.org/schemas/sitemap/0.9"
    if page==1 and total>1000:
        root=Element("sitemapindex",xmlns=namespace)
        for n in range((total+999)//1000):
            item=SubElement(root,"sitemap");SubElement(item,"loc").text=render.origin()+f"/news/sitemap-part/{n+1}.xml"
    else:
        return sitemap_part(page,db)
    return Response(tostring(root,encoding="utf-8",xml_declaration=True),media_type="application/xml")

@public_pages.get("/sitemap-part/{number}.xml")
def sitemap_part(number: int, db=Depends(news_db)):
    from xml.etree.ElementTree import Element,SubElement,tostring
    if number<1 or number>10000: return Response(status_code=404)
    root=Element("urlset",xmlns="http://www.sitemaps.org/schemas/sitemap/0.9")
    if number==1:
        u=SubElement(root,"url");SubElement(u,"loc").text=render.origin()+"/news"
    for a in db.scalars(published(db).order_by(NewsArticle.id).offset((number-1)*1000).limit(1000)):
        u=SubElement(root,"url");SubElement(u,"loc").text=render.origin()+"/news/"+a.slug
        SubElement(u,"lastmod").text=a.updated_at.isoformat()+"Z"
    return Response(tostring(root,encoding="utf-8",xml_declaration=True),media_type="application/xml")

@public_pages.get("/feed.xml")
def feed(db=Depends(news_db)):
    from xml.etree.ElementTree import Element,SubElement,tostring
    from email.utils import format_datetime
    from datetime import timezone
    root=Element("rss",version="2.0");channel=SubElement(root,"channel")
    for name,text in [("title","跃科 · 新闻与资讯"),("link",render.origin()+"/news"),("description","教育管理动态与实践")]: SubElement(channel,name).text=text
    for a in db.scalars(published(db).order_by(NewsArticle.published_at.desc(),NewsArticle.id.desc()).limit(30)):
        item=SubElement(channel,"item")
        for name,text in [("title",a.title),("link",render.origin()+"/news/"+a.slug),("guid",render.origin()+"/news/"+a.slug),("description",a.summary),("pubDate",format_datetime(a.published_at.replace(tzinfo=timezone.utc)))]: SubElement(item,name).text=text
    return Response(tostring(root,encoding="utf-8",xml_declaration=True),media_type="application/rss+xml")

@public_pages.get("",response_class=HTMLResponse)
@public_pages.get("/",response_class=HTMLResponse,include_in_schema=False)
def news_index(page: int=Query(1,ge=1,le=10000), category: str|None=None, db=Depends(news_db)):
    query=published(db,category)
    rows=list(db.scalars(query.order_by(NewsArticle.published_at.desc(),NewsArticle.id.desc()).offset((page-1)*12).limit(12)))
    total=db.scalar(select(func.count()).select_from(query.subquery()))
    return HTMLResponse(render.list_html(rows,page,total,category),headers={"Cache-Control":"no-cache","X-Content-Type-Options":"nosniff"})

@public_pages.get("/{slug}",response_class=HTMLResponse)
def news_detail(slug: str, db=Depends(news_db)):
    a=db.scalar(published(db).where(NewsArticle.slug==slug))
    if not a:
        return HTMLResponse(render.page("内容不存在","这篇内容尚未发布或已下架。",'<h1>内容暂不可用</h1><p><a href="/news">返回新闻与资讯</a></p>',noindex=True),status_code=404,headers={"Cache-Control":"no-store"})
    return HTMLResponse(render.article_html(a),headers={"Cache-Control":"no-cache","X-Content-Type-Options":"nosniff"})
