"""Independent website editorial and public news routes."""
from __future__ import annotations
import json
from datetime import date
from contextlib import contextmanager
from io import BytesIO
from pathlib import PurePath
from typing import Literal
from fastapi import APIRouter, Depends, File, UploadFile, Query
from pydantic import BaseModel, ConfigDict, Field, StrictBool, StrictInt
from sqlalchemy import select, func, exists, and_, or_
from sqlalchemy.exc import SQLAlchemyError
from fastapi.responses import Response, HTMLResponse
from website_app.errors import AppException
from website_app.errors import success
from website_app.auth import WebsiteUser, require_editor, require_publisher
from website_app.db import get_sessionmaker, DBNotEnabledError
from website_app.news.models import NewsPackage, NewsArticle, NewsRevision, NewsMedia, NewsMediaLink, NewsAudit
from website_app.news import service, render
from website_app.news import editor
from website_app.news.package import parse_package, MAX_ARCHIVE, PackageError, CATEGORIES

router = APIRouter(prefix="/manage/news", tags=["官网内容运营"])
public_api = APIRouter(prefix="/news", tags=["官网公开资讯"])
public_pages = APIRouter(prefix="/news", tags=["官网资讯页面"])

def news_db():
    db = None
    try:
        db = get_sessionmaker()()
        if db.get_bind().dialect.name != "mysql":
            raise AppException("SERVER_ERROR", "新闻发布要求连接 MySQL", http_status=503)
        yield db  # Writers explicitly commit before reporting HTTP success.
    except service.NewsError as exc:
        if db: db.rollback()
        raise AppException("DATA_CONFLICT" if exc.status == 409 else "VALIDATION_ERROR", str(exc), http_status=exc.status) from exc
    except (DBNotEnabledError, SQLAlchemyError) as exc:
        if db: db.rollback()
        raise AppException("SERVER_ERROR", "新闻服务暂不可用，请稍后重试；不会将失败报告为发布成功", http_status=503) from exc
    except BaseException:
        if db: db.rollback()
        raise
    finally:
        if db: db.close()

class Confirm(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expectedVersion: StrictInt = Field(ge=0)
    reviewDigest: str = Field(pattern=r"^[a-f0-9]{64}$")
    confirmed: StrictBool
    excludedIds: list[str] = Field(default_factory=list, max_length=2000)
    intervalMinutes: StrictInt = Field(default=3, ge=1, le=60)

class Control(BaseModel):
    model_config = ConfigDict(extra="forbid")
    expectedVersion: StrictInt = Field(ge=0)
    action: Literal["pause", "resume", "cancel"]

class EditorialCommand(BaseModel):
    model_config = ConfigDict(extra='forbid')
    expectedVersion: StrictInt = Field(ge=0)
    revisionId: str | None = Field(default=None,pattern=r'^[a-f0-9]{32}$')
    reviewDigest: str = Field(default='',max_length=64)
    reason: str = Field(default='',max_length=1000)
    publishAt: editor.datetime | None = None

    @editor.field_validator('publishAt')
    @classmethod
    def publish_time(cls,value):
        from datetime import timezone
        if value and not value.tzinfo: raise ValueError('发布时间须包含时区')
        return value.astimezone(timezone.utc).replace(tzinfo=None) if value else None

class CoverCommand(BaseModel):
    model_config = ConfigDict(extra='forbid')
    expectedVersion: StrictInt = Field(ge=0)
    revisionId: str | None = Field(default=None, pattern=r'^[a-f0-9]{32}$')
    mediaId: str | None = Field(pattern=r'^[a-f0-9]{64}$')

@router.get('/products')
def products(user=Depends(require_editor)):
    return success({'items':list(service.product_catalog().values())})

@router.get('/overview')
def overview(user=Depends(require_editor),db=Depends(news_db)):
    counts=dict(db.execute(select(NewsArticle.editorial_state,func.count()).group_by(NewsArticle.editorial_state)).all())
    for state,count in db.execute(select(NewsRevision.editorial_state,func.count()).where(NewsRevision.editorial_state.not_in(['BASELINE','APPLIED'])).group_by(NewsRevision.editorial_state)):
        counts[state]=counts.get(state,0)+count
    return success({'counts':counts,'health':service.get_health(db)})

@router.get('/articles')
def list_articles(page:int=Query(1,ge=1,le=10000),q:str=Query('',max_length=80),category:str=Query('',max_length=40),
    product:str=Query('',max_length=64),state:str=Query('',max_length=24),day:date|None=None,
    sort:Literal['updated','scheduled']='updated',user=Depends(require_editor),db=Depends(news_db)):
    import json
    def filters_for(model):
        predicates=[]
        if q.strip(): predicates.append(model.title.contains(q.strip(),autoescape=True))
        if category: predicates.append(model.category==category)
        if product: predicates.append(func.json_contains(model.related_products,json.dumps(product)))
        if state: predicates.append(model.editorial_state==state)
        return and_(*predicates) if predicates else True
    revision_match=exists(select(NewsRevision.id).where(NewsRevision.article_id==NewsArticle.id,
        NewsRevision.editorial_state.not_in(['BASELINE','APPLIED','CANCELLED']),filters_for(NewsRevision)))
    query=select(NewsArticle).where(or_(filters_for(NewsArticle),revision_match))
    pending_time=select(NewsRevision.scheduled_at).where(NewsRevision.article_id==NewsArticle.id,
        NewsRevision.editorial_state.not_in(['BASELINE','APPLIED','CANCELLED'])).limit(1).scalar_subquery()
    schedule=func.coalesce(pending_time,NewsArticle.scheduled_at)
    if day:
        from datetime import datetime,timedelta
        start=datetime.combine(day,datetime.min.time())-timedelta(hours=8)
        query=query.where(schedule>=start,schedule<start+timedelta(days=1))
    total=db.scalar(select(func.count()).select_from(query.subquery()))
    order=(schedule.is_(None),schedule,NewsArticle.id) if sort=='scheduled' else (NewsArticle.updated_at.desc(),NewsArticle.id.desc())
    rows=list(db.scalars(query.order_by(*order).offset((page-1)*20).limit(20)))
    packages={p.id:p for p in db.scalars(select(NewsPackage).where(NewsPackage.id.in_({a.package_id for a in rows})))}
    revisions={}
    for r in db.scalars(select(NewsRevision).where(NewsRevision.article_id.in_({a.id for a in rows}),NewsRevision.editorial_state.not_in(['BASELINE','APPLIED','CANCELLED']))): revisions[r.article_id]=r
    items=[]
    for a in rows:
        data=service.article_data(a);data.update(packageVersion=packages[a.package_id].version,packageId=a.package_id)
        if a.id in revisions: data['pendingRevision']=service.article_data(revisions[a.id])
        items.append(data)
    return success({'items':items,'total':total,'page':page})

def actor_names(db, rows):
    names={u.id:u.display_name for u in db.scalars(select(WebsiteUser).where(WebsiteUser.id.in_({r.actor for r in rows})))}
    return {r.actor:names.get(r.actor, '官网发布任务' if r.actor=='website-news-worker' else '新闻迁移人员' if r.actor=='website-transfer' else '原有操作人员') for r in rows}

@router.get('/audit')
def audits(page:int=Query(1,ge=1,le=10000),article:str=Query('',max_length=32),user=Depends(require_editor),db=Depends(news_db)):
    query=select(NewsAudit)
    if article: query=query.where(NewsAudit.article_id==article)
    count=db.scalar(select(func.count()).select_from(query.subquery()))
    rows=list(db.scalars(query.order_by(NewsAudit.created_at.desc(),NewsAudit.id.desc()).offset((page-1)*30).limit(30)))
    names=actor_names(db,rows)
    return success({'items':[{'id':r.id,'createdAt':service.iso(r.created_at),'actorName':names[r.actor],
        'actionName':service.ACTIONS.get(r.action,'其他操作'),'articleId':r.article_id,
        'reason':r.detail.get('reason',''),'version':r.detail.get('version'),'number':r.detail.get('number')} for r in rows],'total':count})

@router.post('/articles/{aid}/submit')
def submit(aid:str,body:EditorialCommand,user=Depends(require_editor),db=Depends(news_db)):
    p,a=service.editorial_action(db,aid,body,user['userId'],'submit')
    result={'package':service.summary(db,p),'article':service.article_data(a,full=True)};db.commit();return success(result)

@router.post('/articles/{aid}/actions/{action}')
def article_action(aid:str,action:Literal['return','approve','schedule','pause','resume','cancel'],body:EditorialCommand,user=Depends(require_publisher),db=Depends(news_db)):
    p,a=service.editorial_action(db,aid,body,user['userId'],action)
    result={'package':service.summary(db,p),'article':service.article_data(a,full=True)};db.commit();return success(result)

@router.post('/articles/{aid}/revisions')
def revision(aid:str,body:EditorialCommand,user=Depends(require_editor),db=Depends(news_db)):
    p,a=service.create_revision(db,aid,body.expectedVersion,user['userId'])
    result={'package':service.summary(db,p),'article':service.article_data(a,full=True)};db.commit();return success(result)

@router.post('/articles/{aid}/restore')
def restore_article(aid:str,body:EditorialCommand,user=Depends(require_publisher),db=Depends(news_db)):
    if len(body.reason.strip())<4: raise service.NewsError('请说明恢复发布的原因',400)
    p,a=service.restore(db,aid,body,user['userId'])
    result={'package':service.summary(db,p),'article':service.article_data(a,full=True)};db.commit();return success(result)

@router.post('/articles/{aid}/media')
def upload_media(aid:str,expectedVersion:int=Query(...,ge=0),revisionId:str|None=None,
    alt:str=Query(...,min_length=2,max_length=240),basis:str=Query(...,min_length=8,max_length=1000),
    file:UploadFile=File(...),user=Depends(require_editor),db=Depends(news_db)):
    from website_app.news.package import _picture,digest
    try:
        data=file.file.read(4_000_001)
        if len(data)>4_000_000: raise PackageError('单张图片不能超过4MB')
        picture,mime=_picture(data)
    except PackageError as exc: raise AppException('VALIDATION_ERROR',str(exc),http_status=400) from exc
    finally: file.file.close()
    p,a=service.locked_target(db,aid,expectedVersion,revisionId)
    if a.editorial_state not in ('DRAFT','RETURNED','APPROVED'): raise service.NewsError('请先返回草稿，再修改配图')
    mid=digest(picture)
    if not db.get(NewsMedia,mid): db.add(NewsMedia(id=mid,content=picture,mime=mime));db.flush()
    name='assets/'+mid+'.webp';a.media_map={**a.media_map,name:mid}
    addition='\n\n!['+alt.replace(']','').replace('\n',' ')+']('+name+')'
    if len(a.body)+len(addition)>50000 or len(a.media_map)>50: raise service.NewsError('正文或配图数量已到安全上限，请精简后再上传',400)
    a.body+=addition
    if not a.cover_id: a.cover_id=mid
    a.editorial_state='DRAFT';a.approved_digest=None;a.updated_at=service.now();p.version+=1
    if not isinstance(a,NewsRevision): service.sync_media_links(db,a)
    service.audit(db,p,user['userId'],'UPLOAD_MEDIA',{'mediaId':mid,'description':alt,'basis':basis,'revisionId':revisionId},article=aid)
    service.redigest(db,p)
    result={'package':service.summary(db,p),'article':service.article_data(a,full=True)};db.commit();return success(result)

@router.post('/articles/{aid}/cover')
def set_cover(aid: str, body: CoverCommand, user=Depends(require_editor), db=Depends(news_db)):
    p, a = service.set_cover(db, aid, body, user['userId'])
    result = {'package': service.summary(db, p), 'article': service.article_data(a, full=True)}
    db.commit()
    return success(result)

@router.get('/categories')
def categories(user=Depends(require_editor)):
    return success({'items': [{'value': key, 'label': label} for key, label in CATEGORIES.items()]})

@router.post('/drafts')
def create_draft(body: editor.CreateDraft, user=Depends(require_editor), db=Depends(news_db)):
    package, article = editor.create(db, body, user['userId'])
    result = {'package': service.summary(db, package), 'article': service.article_data(article, full=True)}
    db.commit()
    return success(result, '草稿已保存，尚未公开。请继续审核。')

@router.put('/articles/{aid}')
def edit_draft(aid: str, body: editor.UpdateDraft, user=Depends(require_editor), db=Depends(news_db)):
    package, article = editor.update(db, aid, body, user['userId'])
    result = {'package': service.summary(db, package), 'article': service.article_data(article, full=True)}
    db.commit()
    return success(result, '修改已保存，需要重新核对审核结果。')

@router.get("/package-format")
def package_format(user=Depends(require_editor)):
    from website_app.news.format import GPT_PACKAGE_PROMPT
    return success({"prompt": GPT_PACKAGE_PROMPT})

@router.post("/packages")
def upload_package(user=Depends(require_editor), file: UploadFile=File(...), db=Depends(news_db)):
    try:
        if not (file.filename or "").lower().endswith(".zip"):
            raise PackageError("请上传 ZIP 新闻资源包")
        blob = file.file.read(MAX_ARCHIVE + 1)
        parsed = parse_package(blob)
    except PackageError as exc:
        raise AppException("VALIDATION_ERROR", str(exc), http_status=400) from exc
    finally:
        file.file.close()
    name = (file.filename or "news.zip").replace("\\", "/").split("/")[-1]
    p, reused = service.import_package(db, parsed, name, user.get("userId", "platform"))
    result={"package": service.summary(db,p), "reused": reused}
    db.commit()
    return success(result, "已解析，等待你整包审核确认；尚未发布")

@router.get("/packages")
def list_packages(page: int=Query(1,ge=1,le=10000), user=Depends(require_editor), db=Depends(news_db)):
    rows = db.scalars(select(NewsPackage).order_by(NewsPackage.created_at.desc(), NewsPackage.id.desc()).offset((page-1)*20).limit(20))
    return success({"items":[service.summary(db,p) for p in rows], "total":db.scalar(select(func.count()).select_from(NewsPackage)), "health":service.get_health(db)})

@router.get("/packages/{pid}")
def get_package(pid: str, page: int=Query(1,ge=1,le=10000), user=Depends(require_editor), db=Depends(news_db)):
    p = service.package(db,pid)
    rows = db.scalars(select(NewsArticle).where(NewsArticle.package_id==pid).order_by(NewsArticle.ordinal).offset((page-1)*100).limit(100))
    return success({"package":service.summary(db,p), "articles":[service.article_data(a) for a in rows], "page":page, "pageSize":100, "health":service.get_health(db)})

@router.get("/articles/{aid}")
def review_article(aid: str, revisionId:str|None=None,user=Depends(require_editor), db=Depends(news_db)):
    a = db.get(NewsArticle,aid)
    if not a: raise service.NewsError("文章不存在",404)
    package = service.package(db, a.package_id)
    original=a
    if revisionId:
        a=db.get(NewsRevision,revisionId)
        if not a or a.article_id!=aid: raise service.NewsError('修订不存在',404)
    history=db.scalars(select(NewsRevision).where(NewsRevision.article_id==aid).order_by(NewsRevision.number.desc()).limit(50))
    uploads=list(db.scalars(select(NewsAudit).where(NewsAudit.article_id==aid,NewsAudit.action=='UPLOAD_MEDIA').order_by(NewsAudit.created_at.desc()).limit(1000)))
    names=actor_names(db,uploads)
    materials=[];seen=set()
    applied=select(NewsRevision.id).where(NewsRevision.article_id==aid,NewsRevision.editorial_state=='APPLIED')
    if isinstance(a,NewsRevision): applied=applied.where(NewsRevision.number<=a.number)
    permitted_revisions={None,*db.scalars(applied)}
    if isinstance(a,NewsRevision): permitted_revisions.add(a.id)
    for event in uploads:
        mid=event.detail.get('mediaId')
        if event.detail.get('revisionId') in permitted_revisions and mid in set(a.media_map.values()) and mid not in seen:
            materials.append({'description':event.detail.get('description',''),'basis':event.detail.get('basis',''),
                'actorName':names[event.actor],'createdAt':service.iso(event.created_at)})
            seen.add(mid)
    import re
    descriptions = {name: description for description, name in re.findall(r'!\[([^\]]*)\]\(([^)]+)\)', a.body)}
    stored = set(db.scalars(select(NewsMedia.id).where(NewsMedia.id.in_(set(a.media_map.values())))))
    images, seen_images = [], set()
    for name, mid in a.media_map.items():
        if mid not in stored or mid in seen_images: continue
        seen_images.add(mid)
        images.append({'id': mid, 'url': '/api/manage/news/media/'+mid, 'description': descriptions.get(name) or f'配图 {len(images)+1}', 'isCover': mid == a.cover_id})
    return success({**service.article_data(a,full=True),'url':'/news/'+original.slug,
        'packageVersion': package.version,'packageId':package.id,'packageState':package.state,
        'materials':materials,'images':images,'coverId':a.cover_id,
        'history':[{'id':r.id,'number':r.number,'stateName':service.STATES.get(r.editorial_state,'待核对'),'updatedAt':service.iso(r.updated_at),'appliedAt':service.iso(r.applied_at)} for r in history]})

@router.get("/articles/{aid}/preview", response_class=HTMLResponse)
def preview_article(aid: str, revisionId:str|None=None,user=Depends(require_editor), db=Depends(news_db)):
    a = db.get(NewsArticle, aid)
    if not a: raise service.NewsError("文章不存在", 404)
    if revisionId:
        original=a;a=db.get(NewsRevision,revisionId)
        if not a or a.article_id!=aid: raise service.NewsError('修订不存在',404)
        a.slug=original.slug;a.published_at=original.published_at
    # Only change the media URL prefix in this trusted renderer's output.
    html = render.article_html(a, preview=True).replace('src="/news/media/', 'src="/api/manage/news/media/')
    return HTMLResponse(html, headers={"Cache-Control": "no-store", "X-Robots-Tag": "noindex, nofollow"})

@router.get("/media/{mid}")
def preview_media(mid: str, user=Depends(require_editor), db=Depends(news_db)):
    media = db.get(NewsMedia, mid)
    if not media: raise service.NewsError("图片不存在", 404)
    return Response(media.content, media_type=media.mime, headers={"Cache-Control": "no-store", "X-Content-Type-Options": "nosniff"})

@router.post("/packages/{pid}/confirm")
def confirm_package(pid: str, body: Confirm, user=Depends(require_publisher), db=Depends(news_db)):
    if not body.confirmed: raise AppException("VALIDATION_ERROR","请明确确认审核结果",http_status=400)
    p,reused=service.approve(db,pid,user.get("userId","platform"),body.expectedVersion,body.reviewDigest,body.excludedIds,body.intervalMinutes)
    result={"package":service.summary(db,p), "reused":reused, "health":service.get_health(db)}
    db.commit()
    return success(result, "整包已确认并排期；将在服务器持续发布全部合格新闻")

@router.post("/packages/{pid}/control")
def control_package(pid: str, body: Control, user=Depends(require_publisher), db=Depends(news_db)):
    p=service.control(db,pid,user.get("userId","platform"),body.expectedVersion,body.action)
    result=service.summary(db,p)
    db.commit()
    return success(result)

@router.post("/articles/{aid}/withdraw")
def withdraw(aid: str,body:EditorialCommand,user=Depends(require_publisher), db=Depends(news_db)):
    result=service.article_data(service.withdraw(db,aid,user.get("userId","platform"),body.reason,body.expectedVersion))
    db.commit()
    return success(result)

@router.get("/packages/{pid}/ledger.xlsx")
def export_ledger(pid: str, user=Depends(require_editor), db=Depends(news_db)):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment
    p=service.package(db,pid)
    rows=list(db.scalars(select(NewsArticle).where(NewsArticle.package_id==pid).order_by(NewsArticle.ordinal).limit(2001)))
    if len(rows)>2000: raise AppException("VALIDATION_ERROR","单包台账过大")
    wb=Workbook();ws=wb.active;ws.title="新闻发布台账"
    ws.append(["序号","标题","分类","状态","问题说明","预计发布时间（北京时间）","实际发布时间（北京时间）","公开地址"])
    def cell(v):
        text=str(v or "")
        return "'"+text if text.lstrip().startswith(("=","+","-","@")) else text
    from datetime import timezone
    from zoneinfo import ZoneInfo
    def local_date(v): return v.replace(tzinfo=timezone.utc).astimezone(ZoneInfo('Asia/Shanghai')).strftime('%Y-%m-%d %H:%M:%S') if v else ''
    for a in rows:
        ws.append([a.ordinal,cell(a.title),cell(CATEGORIES.get(a.category)),service.STATES.get(a.state,'待核对'),cell(a.issue),local_date(a.scheduled_at),local_date(a.published_at),render.origin()+"/news/"+a.slug])
    for c in ws[1]: c.font=Font(bold=True,color="FFFFFF");c.fill=PatternFill("solid",fgColor="193A62")
    for col,width in zip("ABCDEFGH",[8,50,18,18,52,27,27,70]): ws.column_dimensions[col].width=width
    for row in ws.iter_rows(min_row=2):
        for c in row: c.alignment=Alignment(vertical="top",wrap_text=True)
    ws.freeze_panes="A2";ws.auto_filter.ref=ws.dimensions
    log=wb.create_sheet("操作审计");log.append(["时间（北京时间）","操作人","动作","文章标识","记录"])
    audit_rows = list(db.scalars(select(NewsAudit).where(NewsAudit.package_id==pid).order_by(NewsAudit.created_at).limit(10001)))
    if len(audit_rows) > 10000:
        raise AppException("VALIDATION_ERROR", "本包审计记录超过单次导出安全上限，请联系运维分页导出，未生成不完整台账")
    names=actor_names(db,audit_rows)
    for a in audit_rows:
        log.append([local_date(a.created_at),cell(names[a.actor]),service.ACTIONS.get(a.action,'其他操作'),a.article_id,cell(a.detail.get('reason') or a.detail.get('version') or '')])
    log.freeze_panes="A2"
    out=BytesIO();wb.save(out)
    service.audit(db,p,user.get("userId","platform"),"EXPORT_LEDGER",{"rows":len(rows)})
    db.commit()
    return Response(out.getvalue(),media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",headers={"Content-Disposition":f'attachment; filename="news-{pid}.xlsx"',"Cache-Control":"no-store"})

def published(db,category=None,product=None):
    query=select(NewsArticle).where(NewsArticle.state=="PUBLISHED")
    if category in CATEGORIES: query=query.where(NewsArticle.category==category)
    if product:
        if product not in service.product_catalog(): raise AppException('NOT_FOUND','未找到这个产品的公开资料。',404)
        query=query.where(func.json_contains(NewsArticle.related_products,json.dumps(product)))
    return query

@public_api.get("/latest")
def latest(product:str=Query('',max_length=64),db=Depends(news_db)):
    rows=db.scalars(published(db,product=product).order_by(NewsArticle.published_at.desc(),NewsArticle.id.desc()).limit(3))
    return success({"items":[service.article_data(a) for a in rows]})

@public_pages.get("/media/{mid}")
def media(mid: str, db=Depends(news_db)):
    # An uploaded asset is not public until a published article references it.
    import re
    if not re.fullmatch(r"[a-f0-9]{64}",mid): return Response(status_code=404)
    visible=db.scalar(select(NewsArticle.id).join(NewsMediaLink, NewsMediaLink.article_id==NewsArticle.id).where(
        NewsArticle.state=="PUBLISHED", NewsMediaLink.media_id==mid).limit(1))
    if not visible: return Response(status_code=404,headers={"Cache-Control":"no-store"})
    m=db.get(NewsMedia,mid)
    if not m: return Response(status_code=404)
    return Response(m.content,media_type=m.mime,headers={"X-Content-Type-Options":"nosniff","Cache-Control":"no-cache"})

@public_pages.get("/sitemap.xml")
def sitemap(page: int=Query(1,ge=1,le=10000), db=Depends(news_db)):
    from xml.etree.ElementTree import Element,SubElement,tostring
    total=db.scalar(select(func.count()).select_from(NewsArticle).where(NewsArticle.state=="PUBLISHED"))
    namespace="http://www.sitemaps.org/schemas/sitemap/0.9"
    if page==1 and total>1000:
        root=Element("sitemapindex",xmlns=namespace)
        for n in range((total+999)//1000):
            item=SubElement(root,"sitemap");SubElement(item,"loc").text=render.origin()+f"/news/sitemap-part/{n+1}.xml"
    else:
        return sitemap_part(page,db)
    return Response(tostring(root,encoding="utf-8",xml_declaration=True),media_type="application/xml")

@public_pages.get("/sitemap-part/{number}.xml")
def sitemap_part(number: int, db=Depends(news_db)):
    from xml.etree.ElementTree import Element,SubElement,tostring
    if number<1 or number>10000: return Response(status_code=404)
    root=Element("urlset",xmlns="http://www.sitemaps.org/schemas/sitemap/0.9")
    if number==1:
        u=SubElement(root,"url");SubElement(u,"loc").text=render.origin()+"/news"
    for a in db.scalars(published(db).order_by(NewsArticle.id).offset((number-1)*1000).limit(1000)):
        u=SubElement(root,"url");SubElement(u,"loc").text=render.origin()+"/news/"+a.slug
        SubElement(u,"lastmod").text=a.updated_at.isoformat()+"Z"
    return Response(tostring(root,encoding="utf-8",xml_declaration=True),media_type="application/xml")

@public_pages.get("/feed.xml")
def feed(db=Depends(news_db)):
    from xml.etree.ElementTree import Element,SubElement,tostring
    from email.utils import format_datetime
    from datetime import timezone
    root=Element("rss",version="2.0");channel=SubElement(root,"channel")
    for name,text in [("title","跃科 · 新闻与资讯"),("link",render.origin()+"/news"),("description","教育管理动态与实践")]: SubElement(channel,name).text=text
    for a in db.scalars(published(db).order_by(NewsArticle.published_at.desc(),NewsArticle.id.desc()).limit(30)):
        item=SubElement(channel,"item")
        for name,text in [("title",a.title),("link",render.origin()+"/news/"+a.slug),("guid",render.origin()+"/news/"+a.slug),("description",a.summary),("pubDate",format_datetime(a.published_at.replace(tzinfo=timezone.utc)))]: SubElement(item,name).text=text
    return Response(tostring(root,encoding="utf-8",xml_declaration=True),media_type="application/rss+xml")

@public_pages.get("",response_class=HTMLResponse)
@public_pages.get("/",response_class=HTMLResponse,include_in_schema=False)
def news_index(
    page: int=Query(1,ge=1,le=10000),
    category: str|None=None,
    q: str|None=Query(None,max_length=80),
    product: str=Query('',max_length=64),
    db=Depends(news_db),
):
    if category is not None and category not in CATEGORIES:
        return HTMLResponse(render.page('栏目不存在','未找到这个资讯栏目。','<h1>栏目不存在</h1><p><a href="/news">返回新闻资讯</a></p>',noindex=True),status_code=404)
    query=published(db,category,product)
    search=(q or "").strip()
    if search:
        query=query.where(or_(
            NewsArticle.title.contains(search,autoescape=True),
            NewsArticle.summary.contains(search,autoescape=True),
            NewsArticle.body.contains(search,autoescape=True),
        ))
    rows=list(db.scalars(query.order_by(NewsArticle.published_at.desc(),NewsArticle.id.desc()).offset((page-1)*12).limit(12)))
    total=db.scalar(select(func.count()).select_from(query.subquery()))
    if page>1 and not rows:
        return HTMLResponse(render.page('页面不存在','该资讯页码不存在。','<h1>页面不存在</h1><p><a href="/news">返回新闻资讯</a></p>',noindex=True),status_code=404)
    product_name=service.product_catalog()[product]['name'] if product else ''
    return HTMLResponse(render.list_html(rows,page,total,category,search,search_page=q is not None,product=product,product_name=product_name),headers={"Cache-Control":"no-cache","X-Content-Type-Options":"nosniff"})

@public_pages.get("/{slug}",response_class=HTMLResponse)
def news_detail(slug: str, db=Depends(news_db)):
    a=db.scalar(published(db).where(NewsArticle.slug==slug))
    if not a:
        return HTMLResponse(render.page("内容不存在","这篇内容尚未发布或已下架。",'<h1>内容暂不可用</h1><p><a href="/news">返回新闻与资讯</a></p>',noindex=True),status_code=404,headers={"Cache-Control":"no-store"})
    return HTMLResponse(render.article_html(a),headers={"Cache-Control":"no-cache","X-Content-Type-Options":"nosniff"})
